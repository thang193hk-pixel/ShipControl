import streamlit as st
import os
import json
import pandas as pd
import sqlite3
import hashlib
import hmac
import html
import secrets
from datetime import datetime, date, timedelta
import altair as alt
import extra_streamlit_components as stx
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart


MODERN_CSS_TEMPLATE = r'''
/* ================================================================
   🎨 GIAO DIỆN HIỆN ĐẠI (lớp phủ cuối cùng, ghi đè các kiểu cũ)
   ================================================================ */

html, body, .stApp, .stApp p, .stApp label, .stApp input, .stApp textarea, .stApp button,
.stApp h1, .stApp h2, .stApp h3, .stApp h4, .stApp li, .stApp span:not([data-testid="stIconMaterial"]) {
    font-family: 'Twemoji Country Flags', 'Be Vietnam Pro', system-ui, -apple-system, 'Segoe UI', Roboto, sans-serif !important;
}

/* ---------- 🏳️ Cờ các nước: Windows không tự hiện emoji cờ (chỉ hiện chữ "VN"),
   nên tải một phông chữ nhỏ CHỈ chứa hình cờ (các ký tự khác vẫn dùng phông bình thường) ---------- */
@font-face {
    font-family: 'Twemoji Country Flags';
    unicode-range: U+1F1E6-1F1FF, U+1F3F4, U+E0062-E0063, U+E0065, U+E0067, U+E006C, U+E006E, U+E0073-E0074, U+E0077, U+E007F;
    src: url('https://cdn.jsdelivr.net/npm/country-flag-emoji-polyfill@0.1/dist/TwemojiCountryFlags.woff2') format('woff2');
    font-display: swap;
}
[data-baseweb="popover"] li, [data-baseweb="popover"] li *, [role="listbox"] [role="option"], [role="listbox"] [role="option"] * {
    font-family: 'Twemoji Country Flags', 'Be Vietnam Pro', system-ui, -apple-system, 'Segoe UI', Roboto, sans-serif !important;
}

/* ---------- Nền: ảnh tàu, mỗi trang một ảnh ----------
   Ảnh nằm trên một lớp cố định riêng (nhẹ hơn nhiều so với background-attachment: fixed),
   mờ dần vào khi đổi trang, và tất cả ảnh được tải trước nên đổi trang không bị nháy. */
.stApp {
    background: __BASE_BG__ !important;
}
.stApp::after {
    content: __PRELOAD__;
    position: absolute;
    width: 0;
    height: 0;
    overflow: hidden;
    opacity: 0;
    z-index: -1;
    pointer-events: none;
}
div[data-testid="stAppViewContainer"] {
    /* Giữ nguyên kiểu định vị gốc của Streamlit (để thanh bên và trang cuộn được), chỉ đưa lên trên lớp ảnh nền */
    z-index: 1;
    background: transparent !important;
}
.main, .stMain, section[data-testid="stMain"], div[data-testid="stAppViewContainer"] > .main {
    background: transparent !important;
}
header[data-testid="stHeader"] {
    background: transparent !important;
}

/* ---------- Thanh bên ---------- */
section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #0369a1 0%, #075985 55%, #0c4a6e 100%) !important;
    border-right: 1px solid rgba(255, 255, 255, 0.08) !important;
    box-shadow: 4px 0 24px rgba(12, 74, 110, 0.18) !important;
}
.sidebar-header {
    font-size: 1.7rem !important;
    letter-spacing: -0.02em !important;
}
.user-card {
    background: rgba(255, 255, 255, 0.10) !important;
    border: 1px solid rgba(255, 255, 255, 0.16) !important;
    border-radius: 16px !important;
}
.made-by-minh {
    font-size: 1.25rem !important;
    border-radius: 999px !important;
    box-shadow: 0 6px 20px rgba(250, 204, 21, 0.35) !important;
    letter-spacing: -0.01em !important;
}

/* Menu: mục thường phẳng, mục đang chọn là viên thuốc màu vàng phát sáng nhẹ */
section[data-testid="stSidebar"] div.stButton > button[kind="secondary"] {
    background: transparent !important;
    border: 1px solid transparent !important;
    border-radius: 12px !important;
}
section[data-testid="stSidebar"] div.stButton > button[kind="secondary"]:hover:not(:disabled) {
    background: rgba(255, 255, 255, 0.12) !important;
    border-color: rgba(255, 255, 255, 0.12) !important;
    transform: translateX(3px) !important;
}
section[data-testid="stSidebar"] div.stButton > button[kind="primary"] {
    background: linear-gradient(135deg, #fde047 0%, #facc15 100%) !important;
    border: none !important;
    border-radius: 12px !important;
    box-shadow: 0 6px 18px rgba(250, 204, 21, 0.40) !important;
}
section[data-testid="stSidebar"] div.stButton > button[kind="primary"]:hover:not(:disabled) {
    box-shadow: 0 8px 22px rgba(250, 204, 21, 0.50) !important;
}
section[data-testid="stSidebar"] div.st-key-btn_logout_bottom div.stButton > button[kind] {
    background: rgba(255, 255, 255, 0.06) !important;
    border: 1px solid rgba(255, 255, 255, 0.35) !important;
    box-shadow: none !important;
}
section[data-testid="stSidebar"] div.st-key-btn_logout_bottom div.stButton > button[kind]:hover:not(:disabled) {
    background: rgba(239, 68, 68, 0.85) !important;
    border-color: transparent !important;
    transform: none !important;
}

/* ---------- Nút trong trang: phẳng, bo tròn, bóng mềm ---------- */
div.stButton > button,
div.stDownloadButton > button,
div[data-testid="stFormSubmitButton"] > button {
    border-radius: 12px !important;
    min-height: 48px !important;
    transition: transform 0.15s ease, box-shadow 0.2s ease, background 0.2s ease, border-color 0.2s ease !important;
}
div.stButton > button *,
div.stDownloadButton > button *,
div[data-testid="stFormSubmitButton"] > button * {
    font-size: 1.02rem !important;
    font-weight: 700 !important;
    letter-spacing: 0 !important;
}
.stMain div.stButton > button[kind="primary"],
.main div.stButton > button[kind="primary"],
div[data-testid="stFormSubmitButton"] > button,
div.stDownloadButton > button {
    background: linear-gradient(135deg, #22c55e 0%, #16a34a 100%) !important;
    border: none !important;
    box-shadow: 0 4px 14px rgba(22, 163, 74, 0.30) !important;
}
.stMain div.stButton > button[kind="primary"]:hover:not(:disabled),
.main div.stButton > button[kind="primary"]:hover:not(:disabled),
div[data-testid="stFormSubmitButton"] > button:hover:not(:disabled),
div.stDownloadButton > button:hover:not(:disabled) {
    background: linear-gradient(135deg, #16a34a 0%, #15803d 100%) !important;
    box-shadow: 0 8px 22px rgba(22, 163, 74, 0.38) !important;
    transform: translateY(-2px) !important;
}
.stMain div.stButton > button[kind="secondary"],
.main div.stButton > button[kind="secondary"] {
    background: __SURFACE__ !important;
    border: 1px solid __SURFACE_BORDER__ !important;
    box-shadow: 0 1px 3px rgba(15, 23, 42, 0.08) !important;
}
.stMain div.stButton > button[kind="secondary"]:hover:not(:disabled),
.main div.stButton > button[kind="secondary"]:hover:not(:disabled) {
    border-color: #0ea5e9 !important;
    box-shadow: 0 6px 16px rgba(14, 165, 233, 0.18) !important;
    transform: translateY(-2px) !important;
}
div.stButton > button:active:not(:disabled),
div.stDownloadButton > button:active:not(:disabled),
div[data-testid="stFormSubmitButton"] > button:active:not(:disabled) {
    transform: scale(0.97) !important;
}
/* Nút xóa: nền đỏ nhạt, chữ đỏ; rê chuột thì đỏ đậm */
__DANGER_SEL__ {
    background: __DANGER_SOFT__ !important;
    border: 1px solid rgba(220, 38, 38, 0.35) !important;
    box-shadow: none !important;
}
__DANGER_SEL_HOVER__ {
    background: linear-gradient(135deg, #ef4444 0%, #dc2626 100%) !important;
    border-color: transparent !important;
    box-shadow: 0 8px 20px rgba(220, 38, 38, 0.30) !important;
}

/* ---------- Tiêu đề ---------- */
.main-title {
    font-size: 2rem !important;
    font-weight: 800 !important;
    letter-spacing: -0.03em !important;
    border-bottom: none !important;
    padding-bottom: 14px !important;
    position: relative;
}
.main-title::after {
    content: "";
    position: absolute;
    left: 50%;
    bottom: 0;
    width: 120px;
    height: 4px;
    transform: translateX(-50%);
    border-radius: 999px;
    background: linear-gradient(90deg, #0ea5e9, #22c55e, #facc15);
}
.big-table-title {
    font-size: 1.55rem !important;
    font-weight: 800 !important;
    letter-spacing: -0.02em !important;
}
.stApp h3 {
    font-weight: 700 !important;
    letter-spacing: -0.01em !important;
}

/* ---------- Thẻ (form, khung mở rộng, bảng) ---------- */
div[data-testid="stForm"], div.st-key-sc_add_task_box {
    background: __SURFACE__ !important;
    border: 1px solid __SURFACE_BORDER__ !important;
    border-radius: 18px !important;
    padding: 22px !important;
    box-shadow: 0 1px 2px rgba(15, 23, 42, 0.04), 0 8px 24px rgba(15, 23, 42, 0.06) !important;
}
div[data-testid="stExpander"] details {
    background: __SURFACE__ !important;
    border: 1px solid __SURFACE_BORDER__ !important;
    border-radius: 16px !important;
    box-shadow: 0 4px 16px rgba(15, 23, 42, 0.05) !important;
    overflow: hidden;
}
div[data-testid="stExpander"] summary {
    background: transparent !important;
    padding: 14px 18px !important;
}
div[data-testid="stExpander"] summary:hover {
    background: __HOVER_TINT__ !important;
}
div[data-testid="stDataFrame"] {
    border-radius: 14px !important;
    overflow: hidden;
    border: 1px solid __SURFACE_BORDER__ !important;
    box-shadow: 0 4px 16px rgba(15, 23, 42, 0.05) !important;
}
div[data-testid="stDataFrame"] th,
div[data-testid="stTable"] th {
    font-size: 1rem !important;
    font-weight: 700 !important;
}
div[data-testid="stDataFrame"] td,
div[data-testid="stTable"] td,
div[data-testid="stDataFrame"] [role="gridcell"] {
    font-size: 1rem !important;
    font-weight: 500 !important;
}
hr {
    border-color: __SURFACE_BORDER__ !important;
    opacity: 0.8;
}

/* ---------- Ô nhập liệu ----------
   Khung ngoài (data-baseweb="input") là "hộp" hiển thị: nền, viền, bo góc.
   Bên trong (ô gõ chữ, nút con mắt) để trong suốt → không bị viền đôi, nút con mắt nằm gọn trong hộp. */
:root {
    color-scheme: __SCHEME__;
}
div[data-testid="stTextInputRootElement"],
div[data-testid="stTextAreaRootElement"],
div[data-testid="stNumberInputContainer"],
div[data-testid="stDateInputField"],
div[data-testid="stSelectbox"] div:has(> input),
div[data-testid="stMultiSelect"] div:has(> input),
div[data-baseweb="input"],
div[data-baseweb="textarea"] {
    background-color: __INPUT_BG__ !important;
    border: 1.5px solid __INPUT_BORDER__ !important;
    border-radius: 12px !important;
    transition: border-color 0.15s ease, box-shadow 0.15s ease !important;
}
div[data-testid="stTextInputRootElement"]:hover,
div[data-testid="stTextAreaRootElement"]:hover,
div[data-testid="stNumberInputContainer"]:hover,
div[data-testid="stDateInputField"]:hover,
div[data-testid="stSelectbox"] div:has(> input):hover,
div[data-testid="stMultiSelect"] div:has(> input):hover,
div[data-baseweb="input"]:hover,
div[data-baseweb="textarea"]:hover {
    border-color: __INPUT_BORDER_HOVER__ !important;
}
div[data-testid="stTextInputRootElement"]:focus-within,
div[data-testid="stTextAreaRootElement"]:focus-within,
div[data-testid="stNumberInputContainer"]:focus-within,
div[data-testid="stDateInputField"]:focus-within,
div[data-testid="stSelectbox"] div:has(> input):focus-within,
div[data-testid="stMultiSelect"] div:has(> input):focus-within,
div[data-baseweb="input"]:focus-within,
div[data-baseweb="textarea"]:focus-within {
    border-color: #0ea5e9 !important;
    box-shadow: 0 0 0 4px rgba(14, 165, 233, 0.18) !important;
}
div[data-testid="stTextInputRootElement"] *,
div[data-testid="stTextAreaRootElement"] *,
div[data-testid="stNumberInputContainer"] *,
div[data-testid="stDateInputField"] *,
div[data-testid="stSelectbox"] div:has(> input) *,
div[data-testid="stMultiSelect"] div:has(> input) *,
div[data-baseweb="input"] *,
div[data-baseweb="textarea"] * {
    background-color: transparent !important;
    border-color: transparent !important;
    box-shadow: none !important;
}
/* Ô nhập nằm trong ô khác (ví dụ ô số): không vẽ hộp lồng nhau */
div[data-testid="stTextInputRootElement"] div[data-baseweb="input"],
div[data-testid="stTextAreaRootElement"] div[data-baseweb="input"],
div[data-testid="stNumberInputContainer"] div[data-baseweb="input"],
div[data-testid="stDateInputField"] div[data-baseweb="input"] {
    border: none !important;
    background: transparent !important;
    box-shadow: none !important;
}
.stApp input, .stApp textarea {
    background: transparent !important;
    border: none !important;
    outline: none !important;
    box-shadow: none !important;
    color: __INPUT_TEXT__ !important;
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
    caret-color: #0ea5e9 !important;
    font-size: 1.05rem !important;
    font-weight: 500 !important;
}
.stApp input::placeholder, .stApp textarea::placeholder {
    color: __PLACEHOLDER__ !important;
    -webkit-text-fill-color: __PLACEHOLDER__ !important;
    opacity: 1 !important;
}
div[data-testid="stSelectbox"] div:has(> input) *,
div[data-baseweb="select"] * {
    color: __INPUT_TEXT__ !important;
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
}
/* Trình duyệt tự điền (autofill): giữ đúng màu nền và màu chữ của app */
.stApp input:-webkit-autofill,
.stApp input:-webkit-autofill:hover,
.stApp input:-webkit-autofill:focus {
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
    -webkit-box-shadow: 0 0 0 1000px __INPUT_BG__ inset !important;
    box-shadow: 0 0 0 1000px __INPUT_BG__ inset !important;
    transition: background-color 99999s ease-out 0s;
}

/* Biểu tượng trong ô nhập: con mắt (hiện mật khẩu), mũi tên chọn, nút +/- */
div[data-testid="stTextInputRootElement"] button,
div[data-baseweb="input"] button,
div[data-testid="stNumberInputContainer"] button {
    background: transparent !important;
    border: none !important;
    color: __ICON__ !important;
    opacity: 1 !important;
    padding: 0 12px !important;
    min-height: 0 !important;
}
div[data-testid="stTextInputRootElement"] button svg,
div[data-baseweb="input"] button svg,
div[data-testid="stNumberInputContainer"] button svg,
div[data-testid="stSelectbox"] svg,
div[data-testid="stMultiSelect"] svg,
div[data-baseweb="select"] svg {
    color: __ICON__ !important;
    fill: __ICON__ !important;
    width: 22px !important;
    height: 22px !important;
    opacity: 1 !important;
}
div[data-testid="stTextInputRootElement"] button:hover svg,
div[data-baseweb="input"] button:hover svg,
div[data-testid="stNumberInputContainer"] button:hover svg {
    color: #0ea5e9 !important;
    fill: #0ea5e9 !important;
}

div[data-testid="stTextInputRootElement"] button span,
div[data-testid="stTextInputRootElement"] button i {
    font-size: 22px !important;
    color: __ICON__ !important;
    opacity: 1 !important;
}
div[data-testid="stDateInputField"] *,
div[data-testid="stDateInputField"] input {
    color: __INPUT_TEXT__ !important;
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
    opacity: 1 !important;
}

/* ---------- 🪪 Thẻ tài khoản dưới menu ---------- */
.profile-card {
    margin: 14px 0 12px 0;
    padding: 14px 14px 10px 14px;
    border-radius: 16px;
    background: rgba(255, 255, 255, 0.10);
    border: 1px solid rgba(255, 255, 255, 0.18);
    color: #ffffff;
}
.profile-card * { color: #ffffff; }
.pc-top {
    display: flex;
    align-items: center;
    gap: 12px;
    padding-bottom: 10px;
    margin-bottom: 6px;
    border-bottom: 1px solid rgba(255, 255, 255, 0.15);
}
.pc-avatar {
    flex: 0 0 44px;
    width: 44px;
    height: 44px;
    border-radius: 50%;
    display: flex;
    align-items: center;
    justify-content: center;
    background: linear-gradient(135deg, #fde047, #facc15);
    color: #0c4a6e !important;
    font-weight: 800;
    font-size: 1.05rem;
    box-shadow: 0 4px 12px rgba(250, 204, 21, 0.35);
}
.pc-id { min-width: 0; }
.pc-name {
    font-weight: 800;
    font-size: 1.05rem;
    line-height: 1.25;
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
}
.pc-status { font-size: 0.8rem; color: #86efac !important; }
.pc-status span { color: rgba(255, 255, 255, 0.75) !important; }
.pc-row {
    display: flex;
    justify-content: space-between;
    align-items: center;
    gap: 10px;
    padding: 5px 0;
    font-size: 0.9rem;
}
.pc-label { color: rgba(255, 255, 255, 0.70) !important; white-space: nowrap; }
.pc-val {
    font-weight: 700;
    text-align: right;
    overflow-wrap: anywhere;
}
.pc-role {
    font-weight: 800;
    font-size: 0.85rem;
    padding: 3px 10px;
    border-radius: 999px;
    background: rgba(250, 204, 21, 0.18);
    border: 1px solid rgba(250, 204, 21, 0.55);
    color: #fde047 !important;
    white-space: nowrap;
}
/* Menu gọn hơn một chút để thẻ tài khoản hiện ra mà không phải cuộn */
section[data-testid="stSidebar"] div[data-testid="stVerticalBlock"] {
    gap: 0.4rem !important;
}

/* ---------- Công tắc bật/tắt (toggle): rõ ràng ở cả trạng thái Tắt và Bật ---------- */
label:has(> span > input[role="switch"]) > div:not([data-testid]) {
    background: __TOGGLE_OFF__ !important;
    border: 2px solid __TOGGLE_OFF_BORDER__ !important;
    box-shadow: inset 0 1px 3px rgba(15, 23, 42, 0.25) !important;
    opacity: 1 !important;
    transition: background 0.2s ease, border-color 0.2s ease !important;
}
label:has(> span > input[role="switch"]) > div:not([data-testid]) > div {
    background: #ffffff !important;
    box-shadow: 0 1px 4px rgba(15, 23, 42, 0.45) !important;
    opacity: 1 !important;
}
label:has(> span > input[role="switch"]:checked) > div:not([data-testid]) {
    background: #0ea5e9 !important;
    border-color: #0284c7 !important;
}
label:has(> span > input[role="switch"]:focus-visible) > div:not([data-testid]) {
    outline: 3px solid rgba(14, 165, 233, 0.45) !important;
    outline-offset: 2px !important;
}
label:has(> span > input[role="switch"]:disabled) {
    opacity: 0.5 !important;
    cursor: not-allowed !important;
}

/* ---------- 🌐 Chọn chủ đề + ngôn ngữ trong Cài đặt: nút xanh, chữ trắng, mũi tên trắng ---------- */
html body .st-key-w_ui_theme,
html body .st-key-w_ui_lang {
    width: auto !important;
}
html body .st-key-w_ui_theme div[data-testid="stSelectbox"] div:has(> input),
html body .st-key-w_ui_lang div[data-testid="stSelectbox"] div:has(> input),
html body .st-key-w_ui_theme div[data-baseweb="select"] > div,
html body .st-key-w_ui_lang div[data-baseweb="select"] > div {
    background: #4f8ef7 !important;
    background-color: #4f8ef7 !important;
    border: none !important;
    border-radius: 8px !important;
    min-height: 38px !important;
    width: fit-content !important;
    min-width: 170px !important;
    box-shadow: 0 2px 8px rgba(37, 99, 235, 0.35) !important;
    cursor: pointer !important;
}
html body .st-key-w_ui_theme div[data-testid="stSelectbox"] div:has(> input):hover,
html body .st-key-w_ui_lang div[data-testid="stSelectbox"] div:has(> input):hover {
    background: #3b7cf0 !important;
    background-color: #3b7cf0 !important;
}
html body .st-key-w_ui_theme div[data-testid="stSelectbox"] div:has(> input) *,
html body .st-key-w_ui_lang div[data-testid="stSelectbox"] div:has(> input) *,
html body .st-key-w_ui_theme div[data-baseweb="select"] *,
html body .st-key-w_ui_lang div[data-baseweb="select"] * {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    font-weight: 700 !important;
    caret-color: transparent !important;
}
html body .st-key-w_ui_theme div[data-testid="stSelectbox"] svg,
html body .st-key-w_ui_lang div[data-testid="stSelectbox"] svg {
    color: #ffffff !important;
    fill: #ffffff !important;
}

/* ---------- 🌐 Nút ngôn ngữ: ô vuông xanh, quả địa cầu trắng + mũi tên ---------- */
section[data-testid="stSidebar"] .st-key-pop_lang [data-testid="stPopover"] button,
section[data-testid="stSidebar"] .st-key-pop_lang button {
    background: #4f8ef7 !important;
    border: none !important;
    border-radius: 8px !important;
    min-height: 34px !important;
    height: 34px !important;
    min-width: 0 !important;
    width: auto !important;
    padding: 0 8px 0 10px !important;
    box-shadow: 0 2px 8px rgba(37, 99, 235, 0.35) !important;
    gap: 2px !important;
}
section[data-testid="stSidebar"] .st-key-pop_lang [data-testid="stPopover"] button:hover,
section[data-testid="stSidebar"] .st-key-pop_lang button:hover {
    background: #3b7cf0 !important;
}
section[data-testid="stSidebar"] .st-key-pop_lang button * {
    color: #ffffff !important;
    fill: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
}
section[data-testid="stSidebar"] .st-key-pop_lang [data-testid="stIconMaterial"] {
    font-size: 22px !important;
}
section[data-testid="stSidebar"] .st-key-pop_lang button::after {
    display: none !important;
}

/* ---------- 📲 Chọn thiết bị: 2 ô lớn thay cho nút tròn nhỏ ---------- */
.st-key-device_mode_login,
.st-key-w_device_mode {
    width: 100% !important;
}
.st-key-device_mode_login [data-testid="stRadioGroup"],
.st-key-w_device_mode [data-testid="stRadioGroup"] {
    display: grid !important;
    grid-template-columns: 1fr 1fr !important;
    gap: 12px !important;
    width: 100% !important;
}
.st-key-device_mode_login [data-testid="stRadioGroup"] > div,
.st-key-w_device_mode [data-testid="stRadioGroup"] > div {
    width: 100% !important;
    margin: 0 !important;
}
.st-key-device_mode_login label[data-testid="stRadioOption"],
.st-key-w_device_mode label[data-testid="stRadioOption"] {
    position: relative;
    display: flex !important;
    align-items: center !important;
    justify-content: center !important;
    width: 100% !important;
    min-height: 64px !important;
    padding: 12px 10px !important;
    margin: 0 !important;
    border-radius: 16px !important;
    border: 2px solid __INPUT_BORDER__ !important;
    background: __SURFACE__ !important;
    box-shadow: 0 2px 8px rgba(15, 23, 42, 0.06) !important;
    cursor: pointer !important;
    transition: transform 0.15s ease, border-color 0.15s ease, box-shadow 0.15s ease, background 0.15s ease !important;
}
.st-key-device_mode_login label[data-testid="stRadioOption"]:hover,
.st-key-w_device_mode label[data-testid="stRadioOption"]:hover {
    border-color: #0ea5e9 !important;
    transform: translateY(-2px);
    box-shadow: 0 6px 16px rgba(14, 165, 233, 0.18) !important;
}
.st-key-device_mode_login label[data-testid="stRadioOption"][data-selected="true"],
.st-key-w_device_mode label[data-testid="stRadioOption"][data-selected="true"] {
    border-color: #0ea5e9 !important;
    background: __SELECTED_BG__ !important;
    box-shadow: 0 0 0 4px rgba(14, 165, 233, 0.18), 0 6px 16px rgba(14, 165, 233, 0.20) !important;
}
.st-key-device_mode_login label[data-testid="stRadioOption"][data-selected="true"]::after,
.st-key-w_device_mode label[data-testid="stRadioOption"][data-selected="true"]::after {
    content: "✓";
    position: absolute;
    top: -9px;
    right: -9px;
    width: 24px;
    height: 24px;
    border-radius: 50%;
    background: #0ea5e9;
    color: #ffffff;
    font-size: 14px;
    font-weight: 900;
    display: flex;
    align-items: center;
    justify-content: center;
    box-shadow: 0 2px 6px rgba(2, 132, 199, 0.4);
}
/* Ẩn nút tròn nhỏ, chỉ giữ chữ */
.st-key-device_mode_login label[data-testid="stRadioOption"] > div > div:first-child:not([data-testid]),
.st-key-w_device_mode label[data-testid="stRadioOption"] > div > div:first-child:not([data-testid]) {
    display: none !important;
}
.st-key-device_mode_login label[data-testid="stRadioOption"] p,
.st-key-w_device_mode label[data-testid="stRadioOption"] p {
    font-size: 1.2rem !important;
    font-weight: 800 !important;
    white-space: nowrap !important;
    margin: 0 !important;
}
.st-key-device_mode_login label[data-testid="stRadioOption"]:has(input:focus-visible),
.st-key-w_device_mode label[data-testid="stRadioOption"]:has(input:focus-visible) {
    outline: 3px solid rgba(14, 165, 233, 0.5) !important;
    outline-offset: 2px !important;
}
/* Trong bảng Cài đặt: nhỏ gọn hơn một chút */
.st-key-w_device_mode label[data-testid="stRadioOption"] {
    min-height: 50px !important;
    border-radius: 12px !important;
}
.st-key-w_device_mode label[data-testid="stRadioOption"] p {
    font-size: 1.02rem !important;
}

/* ---------- ⚙️ Nút Cài đặt trên thanh bên ---------- */
section[data-testid="stSidebar"] [data-testid="stPopover"] button {
    background: rgba(255, 255, 255, 0.12) !important;
    border: 1px solid rgba(255, 255, 255, 0.25) !important;
    border-radius: 999px !important;
    min-height: 34px !important;
    height: 34px !important;
    padding: 0 14px !important;
    width: auto !important;
    box-shadow: none !important;
    transition: background 0.15s ease !important;
}
section[data-testid="stSidebar"] [data-testid="stPopover"] button:hover {
    background: rgba(255, 255, 255, 0.24) !important;
}
section[data-testid="stSidebar"] [data-testid="stPopover"] button * {
    color: #ffffff !important;
    fill: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    font-weight: 600 !important;
    font-size: 0.88rem !important;
}
/* Bảng cài đặt bật ra */
div[data-testid="stPopoverBody"] {
    background: __SURFACE__ !important;
    border: 1px solid __SURFACE_BORDER__ !important;
    border-radius: 16px !important;
    box-shadow: 0 16px 40px rgba(15, 23, 42, 0.25) !important;
    padding: 18px !important;
    min-width: 260px !important;
}
div[data-testid="stPopoverBody"] * {
    color: __INPUT_TEXT__ !important;
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
}
div[data-testid="stPopoverBody"] [data-testid="stWidgetLabel"] p {
    font-weight: 700 !important;
    font-size: 1rem !important;
}

/* ---------- Nút mở / đóng thanh bên (») và («) ---------- */
[data-testid="stExpandSidebarButton"],
[data-testid="stSidebarCollapsedControl"] button {
    background: linear-gradient(135deg, #0284c7 0%, #0369a1 100%) !important;
    border-radius: 12px !important;
    width: 44px !important;
    height: 44px !important;
    box-shadow: 0 6px 16px rgba(3, 105, 161, 0.35) !important;
    opacity: 1 !important;
}
[data-testid="stExpandSidebarButton"] *,
[data-testid="stSidebarCollapsedControl"] button * {
    color: #ffffff !important;
    fill: #ffffff !important;
    font-size: 26px !important;
    opacity: 1 !important;
}
[data-testid="stExpandSidebarButton"]:hover,
[data-testid="stSidebarCollapsedControl"] button:hover {
    background: linear-gradient(135deg, #0ea5e9 0%, #0284c7 100%) !important;
}
[data-testid="stSidebarCollapseButton"] button {
    background: rgba(255, 255, 255, 0.14) !important;
    border-radius: 10px !important;
    width: 40px !important;
    height: 40px !important;
    opacity: 1 !important;
    visibility: visible !important;
}
[data-testid="stSidebarCollapseButton"] {
    visibility: visible !important;
    opacity: 1 !important;
}
[data-testid="stSidebarCollapseButton"] button * {
    color: #ffffff !important;
    fill: #ffffff !important;
    font-size: 24px !important;
    opacity: 1 !important;
}
[data-testid="stSidebarCollapseButton"] button:hover {
    background: rgba(255, 255, 255, 0.28) !important;
}
/* Nút ⋮ và Deploy ở góc trên */
header[data-testid="stHeader"] button {
    color: __INPUT_TEXT__ !important;
}
/* Vệt sáng chỉ chạy trên mục menu đang chọn; mục thường nền trong suốt nên vệt sáng trông như vết lỗi */
section[data-testid="stSidebar"] div.stButton > button[kind="secondary"]::after,
section[data-testid="stSidebar"] [data-testid="stPopover"] button::after {
    display: none !important;
}

/* ---------- Thông báo: mỗi loại một màu, chữ vừa phải ---------- */
div[data-testid="stAlert"] {
    border-radius: 14px !important;
    border: 1px solid transparent !important;
}
div[data-testid="stAlert"] * {
    font-size: 1rem !important;
    font-weight: 600 !important;
}
div[data-testid="stAlert"]:has([data-testid="stAlertContentInfo"]) {
    background: __INFO_BG__ !important; border-color: rgba(14, 165, 233, 0.35) !important;
}
div[data-testid="stAlert"]:has([data-testid="stAlertContentInfo"]) * {
    color: __INFO_TX__ !important; -webkit-text-fill-color: __INFO_TX__ !important;
}
div[data-testid="stAlert"]:has([data-testid="stAlertContentSuccess"]) {
    background: __OK_BG__ !important; border-color: rgba(34, 197, 94, 0.35) !important;
}
div[data-testid="stAlert"]:has([data-testid="stAlertContentSuccess"]) * {
    color: __OK_TX__ !important; -webkit-text-fill-color: __OK_TX__ !important;
}
div[data-testid="stAlert"]:has([data-testid="stAlertContentWarning"]) {
    background: __WARN_BG__ !important; border-color: rgba(234, 179, 8, 0.45) !important;
}
div[data-testid="stAlert"]:has([data-testid="stAlertContentWarning"]) * {
    color: __WARN_TX__ !important; -webkit-text-fill-color: __WARN_TX__ !important;
}
div[data-testid="stAlert"]:has([data-testid="stAlertContentError"]) {
    background: __ERR_BG__ !important; border-color: rgba(239, 68, 68, 0.40) !important;
}
div[data-testid="stAlert"]:has([data-testid="stAlertContentError"]) * {
    color: __ERR_TX__ !important; -webkit-text-fill-color: __ERR_TX__ !important;
}

/* ---------- 💬 Khung tin nhắn ---------- */
.sc-chat {
    height: 420px;
    overflow-y: auto;
    display: flex;
    flex-direction: column-reverse;      /* tin mới nhất ở dưới, tự cuộn xuống cuối */
    gap: 10px;
    padding: 14px;
    border-radius: 16px;
    background: __SURFACE__;
    border: 1px solid __SURFACE_BORDER__;
    box-shadow: inset 0 1px 4px rgba(15, 23, 42, 0.05);
    margin-bottom: 10px;
}
/* Ô ẩn dùng để gửi lệnh từ menu chuột phải */
.st-key-chat_action {
    position: absolute !important;
    left: -10000px !important;
    width: 1px !important;
    height: 1px !important;
    overflow: hidden !important;
}
.sc-recalled { font-style: italic; opacity: 0.75; }
.sc-msg { cursor: context-menu; -webkit-touch-callout: none; }
.sc-pinned {
    margin-bottom: 8px;
    padding: 8px 12px;
    border-radius: 12px;
    background: __SELECTED_BG__;
    border-left: 4px solid #0ea5e9;
    font-size: 0.9rem;
    color: __INPUT_TEXT__;
}
.sc-pin-title { font-weight: 800; margin-bottom: 2px; }
.sc-pin-row { white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
.sc-chat-empty { margin: auto; color: __PLACEHOLDER__; font-weight: 600; text-align: center; }
.sc-msg { display: flex; flex-direction: column; max-width: 78%; }
.sc-mine { align-self: flex-end; align-items: flex-end; }
.sc-theirs { align-self: flex-start; align-items: flex-start; }
.sc-meta { font-size: 0.75rem; color: __PLACEHOLDER__; margin: 0 6px 3px 6px; }
.sc-bubble {
    padding: 9px 14px;
    border-radius: 18px;
    font-size: 1rem;
    line-height: 1.4;
    overflow-wrap: anywhere;
    animation: scFadeUp 0.2s ease-out both;
}
.sc-mine .sc-bubble {
    background: linear-gradient(135deg, #0ea5e9, #0284c7);
    color: #ffffff;
    border-bottom-right-radius: 6px;
}
.sc-theirs .sc-bubble {
    background: __BUBBLE_THEIRS__;
    color: __INPUT_TEXT__;
    border-bottom-left-radius: 6px;
}
[data-testid="stChatInput"] {
    border-radius: 16px !important;
}
/* Ô nhập tin nhắn: luôn nền sáng / chữ đậm (hoặc nền tối / chữ sáng ở chế độ Tối), không phụ thuộc giao diện máy */
html body [data-testid="stChatInput"] > div {
    background-color: __INPUT_BG__ !important;
    border: 1.5px solid __INPUT_BORDER__ !important;
    border-radius: 16px !important;
}
html body [data-testid="stChatInput"] > div:focus-within {
    border-color: #0ea5e9 !important;
    box-shadow: 0 0 0 4px rgba(14, 165, 233, 0.18) !important;
}
html body [data-testid="stChatInput"] > div div {
    background-color: transparent !important;
}
html body [data-testid="stChatInput"] textarea {
    color: __INPUT_TEXT__ !important;
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
    caret-color: #0ea5e9 !important;
    font-size: 1.05rem !important;
    font-weight: 500 !important;
}
html body [data-testid="stChatInput"] textarea::placeholder {
    color: __PLACEHOLDER__ !important;
    -webkit-text-fill-color: __PLACEHOLDER__ !important;
}
html body [data-testid="stChatInput"] button:not(.sc-mic) {
    background-color: #0ea5e9 !important;
    border-radius: 10px !important;
}
html body [data-testid="stChatInput"] button:not(.sc-mic) svg {
    fill: #ffffff !important;
    color: #ffffff !important;
}
/* Danh sách người trong Tin nhắn: chữ căn trái, gọn hơn */
div[class*="st-key-chat_open_"] div.stButton > button {
    justify-content: flex-start !important;
    min-height: 44px !important;
    margin-bottom: 2px !important;
}
div[class*="st-key-chat_open_"] div.stButton > button > div,
div[class*="st-key-chat_open_"] div.stButton > button p {
    justify-content: flex-start !important;
    text-align: left !important;
    white-space: normal !important;
    overflow: visible !important;
    text-overflow: clip !important;
}

/* ---------- Nút mở liên kết (Bắt đầu / Mở phòng họp): giống nút chính màu xanh ---------- */
div[data-testid="stLinkButton"] a {
    min-height: 48px !important;
    border-radius: 12px !important;
    font-weight: 700 !important;
    transition: transform 0.15s ease, box-shadow 0.2s ease !important;
}
div[data-testid="stLinkButton"] a[kind="primary"],
div[data-testid="stLinkButton"] a[data-testid="stBaseLinkButton-primary"] {
    background: linear-gradient(135deg, #22c55e 0%, #16a34a 100%) !important;
    border: none !important;
    color: #ffffff !important;
    box-shadow: 0 4px 14px rgba(22, 163, 74, 0.30) !important;
}
div[data-testid="stLinkButton"] a[kind="primary"] *,
div[data-testid="stLinkButton"] a[data-testid="stBaseLinkButton-primary"] * {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    font-weight: 700 !important;
}
div[data-testid="stLinkButton"] a:hover {
    transform: translateY(-2px);
    box-shadow: 0 8px 22px rgba(22, 163, 74, 0.38) !important;
}
div[data-testid="stTimeInput"] div:has(> input),
div[data-testid="stTimeInput"] div[data-baseweb="select"] > div {
    background-color: __INPUT_BG__ !important;
    border: 1.5px solid __INPUT_BORDER__ !important;
    border-radius: 12px !important;
}
div[data-testid="stTimeInput"] input, div[data-testid="stTimeInput"] div[data-baseweb="select"] * {
    color: __INPUT_TEXT__ !important;
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
}

/* ---------- Ô chọn nhiều người: tên đã chọn hiện chữ đen trên nền xanh nhạt ---------- */
/* Hộp ngoài là khung nhập; vùng chứa tên bên trong trong suốt (không bị khung lồng khung) */
html body div[data-testid="stMultiSelect"] div:has(> div[data-testid="stMultiSelectTagsContainer"]) {
    background-color: __INPUT_BG__ !important;
    border: 1.5px solid __INPUT_BORDER__ !important;
    border-radius: 12px !important;
}
html body div[data-testid="stMultiSelect"] div:has(> div[data-testid="stMultiSelectTagsContainer"]):focus-within {
    border-color: #0ea5e9 !important;
    box-shadow: 0 0 0 4px rgba(14, 165, 233, 0.18) !important;
}
html body div[data-testid="stMultiSelect"] div[data-testid="stMultiSelectTagsContainer"] {
    background-color: transparent !important;
    border: none !important;
    box-shadow: none !important;
}
/* Từng tên đã chọn */
html body div[data-testid="stMultiSelect"] div[data-testid="stMultiSelectTagsContainer"] > span > span,
html body div[data-testid="stMultiSelect"] [data-baseweb="tag"] {
    background-color: __SELECTED_BG__ !important;
    border: 1px solid #0ea5e9 !important;
    border-radius: 14px !important;
    padding: 4px 4px 4px 10px !important;
}
html body div[data-testid="stMultiSelect"] div[data-testid="stMultiSelectTagsContainer"] > span > span *,
html body div[data-testid="stMultiSelect"] [data-baseweb="tag"] * {
    color: __INPUT_TEXT__ !important;
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
    background-color: transparent !important;
    font-weight: 600 !important;
}
html body div[data-testid="stMultiSelect"] div[data-testid="stMultiSelectTagsContainer"] > span > span svg,
html body div[data-testid="stMultiSelect"] [data-baseweb="tag"] svg {
    fill: __INPUT_TEXT__ !important;
    color: __INPUT_TEXT__ !important;
    width: 14px !important;
    height: 14px !important;
}
/* Hiện đủ tên (xuống dòng nếu dài), không cắt thành "Minh ..." */
html body div[data-testid="stMultiSelect"] div[data-testid="stMultiSelectTagsContainer"] > span > span {
    max-width: 100% !important;
    height: auto !important;
}
html body div[data-testid="stMultiSelect"] div[data-testid="stMultiSelectTagsContainer"] > span > span > span {
    max-width: none !important;
    overflow: visible !important;
    text-overflow: clip !important;
    white-space: normal !important;
    line-height: 1.3 !important;
}
/* Danh sách xổ xuống: chữ đen */
[data-baseweb="popover"] li, [data-baseweb="popover"] li * {
    color: __INPUT_TEXT__ !important;
    -webkit-text-fill-color: __INPUT_TEXT__ !important;
}

/* ================================================================
   🧭 THANH MENU NGANG (thay thanh bên trái)
   ================================================================ */
section[data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"],
[data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapseButton"] {
    display: none !important;
}
.st-key-sc_js_slot {
    position: absolute !important;
    width: 0 !important;
    height: 0 !important;
    overflow: hidden !important;
}
/* 🔘 Thanh mini: tiêu đề + nút tròn ➕ ➖ — luôn nằm trên 1 hàng kể cả trên điện thoại */
[class*="st-key-sc_mini_"] { margin-top: 4px !important; }
[class*="st-key-sc_mini_"] h4 { margin: 0 !important; padding: 0 !important; font-size: 1.05rem !important; }
[class*="st-key-sc_mini_"] [data-testid="stColumn"]:first-child [data-testid="stElementContainer"],
[class*="st-key-sc_mini_"] [data-testid="stColumn"]:first-child [data-testid="stMarkdownContainer"],
[class*="st-key-sc_mini_"] [data-testid="stColumn"]:first-child .stMarkdown { margin: 0 !important; }
[class*="st-key-sc_mini_"] [data-testid="stHorizontalBlock"] { flex-wrap: nowrap !important; gap: 6px !important; align-items: center !important; }
[class*="st-key-sc_mini_"] [data-testid="stColumn"] { min-width: 0 !important; width: auto !important; flex: 0 0 auto !important; }
[class*="st-key-sc_mini_"] [data-testid="stColumn"]:first-child { flex: 1 1 auto !important; }
[class*="st-key-sc_mini_"] [data-testid="stButton"] { display: flex !important; justify-content: flex-end !important; }
[class*="st-key-sc_mini_"] [data-testid="stButton"] button {
    width: 30px !important; min-width: 30px !important; height: 30px !important; min-height: 30px !important;
    padding: 0 !important; border-radius: 50% !important; font-size: 14px !important; line-height: 1 !important;
    letter-spacing: 0 !important; box-shadow: none !important;
}
.st-key-sc_mini_cc [data-testid="stButton"], .st-key-sc_mini_block [data-testid="stButton"] { margin-bottom: 5px !important; }
.st-key-sc_mini_cc, .st-key-sc_mini_block { margin-top: 0 !important; }
[class*="st-key-sc_minip_"] { padding: 8px !important; gap: 6px !important; animation: scPop 0.15s ease-out both; }
[class*="st-key-sc_minip_"] [data-testid="stButton"] button { min-height: 34px !important; padding: 4px 10px !important; font-size: 13px !important; }
/* 🖥️ Khung slide kiểu PowerPoint (báo cáo hình ảnh theo tuần) */
[class*="st-key-sc_slide_"] {
    background: #ffffff !important; border-radius: 18px !important; overflow: hidden !important;
    border: 1px solid #dbe3ee !important; box-shadow: 0 18px 40px rgba(11, 37, 69, 0.18) !important;
    padding: 0 22px 14px !important; gap: 12px !important; margin: 6px 0 14px !important;
}
[class*="st-key-sc_slide_"] .sc-slide-head {
    margin: 0 -22px 6px; padding: 22px 28px; display: flex; justify-content: space-between; align-items: flex-end; gap: 12px;
    background: #0b2545; color: #fff;
}
[class*="st-key-sc_slide_"] .sc-slide-kicker { color: #f2a541 !important; font-size: 0.78rem; font-weight: 800; letter-spacing: 0.18em; }
[class*="st-key-sc_slide_"] .sc-slide-title { color: #fff !important; font-family: Cambria, Georgia, serif; font-size: 2.3rem; font-weight: 800; line-height: 1.1; }
[class*="st-key-sc_slide_"] .sc-slide-dates { color: #8da9c4 !important; font-size: 1rem; margin-top: 2px; }
[class*="st-key-sc_slide_"] .sc-slide-right { display: flex; gap: 8px; flex-wrap: wrap; justify-content: flex-end; }
[class*="st-key-sc_slide_"] .sc-slide-badge { background: #134074; color: #fff !important; border-radius: 999px; padding: 6px 14px; font-weight: 700; font-size: 0.9rem; white-space: nowrap; }
[class*="st-key-sc_slide_"] .sc-slide-badge.sc-done { background: #f2a541; color: #0b2545 !important; }
[class*="st-key-sc_slide_"] .sc-slide-empty { background: #eef3f8; color: #64748b !important; border-radius: 14px; padding: 90px 20px; text-align: center; font-size: 1.15rem; }
[class*="st-key-sc_slide_"] .sc-slide-cap { color: #64748b !important; font-size: 0.8rem; text-align: center; margin: -6px 0 2px; }
[class*="st-key-sc_slide_"] .sc-slide-foot { display: flex; justify-content: space-between; color: #94a3b8 !important; font-size: 0.75rem; border-top: 1px solid #eef2f7; padding-top: 8px; }
[class*="st-key-sc_slide_"] .sc-slide-foot span { color: #94a3b8 !important; }
[class*="st-key-sc_slide_"] img { border-radius: 10px; }
[class*="st-key-sc_slide_n1_"] img { max-height: 460px !important; width: 100% !important; object-fit: contain !important; }
[class*="st-key-sc_slide_n2_"] img { height: 300px !important; width: 100% !important; object-fit: cover !important; }
[class*="st-key-sc_slide_n3_"] img { height: 210px !important; width: 100% !important; object-fit: cover !important; }
.st-key-sc_topbar {
    position: sticky !important;
    top: 0;
    z-index: 999;
    background: __NAV_BG__ !important;
    border-radius: 0 0 14px 14px !important;
    padding: 4px 12px !important;
    margin: 0 0 18px 0 !important;
    box-shadow: 0 6px 20px rgba(15, 23, 42, 0.22) !important;
}
/* Các khối vô hình (chỉ chứa CSS / mã chạy ngầm) không được tạo khoảng trống trên trang */
[data-testid="stMainBlockContainer"] div[data-testid="stElementContainer"]:has(> div > div > style),
[data-testid="stMainBlockContainer"] div[data-testid="stElementContainer"]:has(style:only-child) {
    display: none !important;
}
div[data-testid="stLayoutWrapper"]:has(> .st-key-sc_js_slot),
div[data-testid="stLayoutWrapper"]:has(.st-key-sc_js_slot):not(:has(.st-key-sc_topbar)) {
    margin-top: -1rem !important;
    height: 0 !important;
}
/* Đã đăng nhập: logo nằm trên thanh menu rồi → ẩn tiêu đề lớn để nội dung lên cao, đỡ phải cuộn */
.stApp:has(.st-key-sc_nav button) .main-title,
.stApp:has(.st-key-sc_nav button) div[data-testid="stElementContainer"]:has(.main-title) { display: none !important; }
.sc-brand {
    font-weight: 900;
    font-size: 1.05rem;
    letter-spacing: 0.04em;
    color: #ffffff !important;
    white-space: nowrap;
    padding-right: 10px;
    margin-right: 4px;
    border-right: 1px solid rgba(255, 255, 255, 0.25);
}
.st-key-sc_topbar .stMarkdown, .st-key-sc_topbar [data-testid="stMarkdownContainer"] p { margin: 0 !important; }
/* Các mục menu: chữ trắng IN HOA, nền trong suốt; cuộn ngang nếu màn hình hẹp */
.st-key-sc_nav {
    overflow-x: auto !important;
    scrollbar-width: none;
    flex-wrap: nowrap !important;
}
.st-key-sc_nav::-webkit-scrollbar { display: none; }
/* Màn hình rộng: menu xuống dòng nếu nhiều mục (không bị che). Điện thoại: vuốt ngang. */
@media (min-width: 900px) {
    .st-key-sc_nav { flex-wrap: wrap !important; overflow-x: visible !important; row-gap: 2px !important; }
}
/* Máy tính: logo | menu | ⚙️ + tài khoản trên CÙNG một thanh (gọn, không phải cuộn) */
/* Điện thoại: hàng 1 = logo + ⚙️ + tài khoản, hàng 2 = menu (vuốt ngang) */
@media (max-width: 899px) {
    .st-key-sc_topbar { flex-wrap: wrap !important; row-gap: 4px !important; position: relative !important; }
    .st-key-sc_topbar > div:has(.sc-brand) { flex: 1 1 auto !important; }
    .st-key-sc_topbar > .st-key-sc_nav,
    .st-key-sc_topbar > div:has(.st-key-sc_nav) {
        order: 3 !important; flex: 1 1 100% !important; width: 100% !important; max-width: 100% !important; min-width: 0 !important;
    }
    .st-key-sc_nav { width: 100% !important; }
    .sc-brand { border-right: none !important; font-size: 1.1rem !important; }
}
/* Ẩn thanh tiêu đề mặc định của Streamlit (nó nằm đè lên thanh menu và chặn bấm) */
header[data-testid="stHeader"] { display: none !important; }
/* Bớt khoảng trống phía trên trang */
.block-container, [data-testid="stMainBlockContainer"] { padding-top: 1.6rem !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button,
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="secondary"],
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] {
    width: auto !important;
    min-height: 38px !important;
    height: 38px !important;
    padding: 0 8px !important;
    margin: 0 !important;
    background: transparent !important;
    border: none !important;
    border-radius: 8px !important;
    box-shadow: none !important;
    white-space: nowrap !important;
    transform: none !important;
}
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button * {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    text-transform: uppercase !important;
    font-size: 0.88rem !important;
    font-weight: 700 !important;
    letter-spacing: 0.01em !important;
    white-space: nowrap !important;
}
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button:hover:not(:disabled) {
    background: rgba(255, 255, 255, 0.14) !important;
}
/* Trang đang mở: gạch chân vàng + nền nhạt */
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] {
    background: rgba(255, 255, 255, 0.16) !important;
    box-shadow: inset 0 -3px 0 #facc15 !important;
}
html body .st-key-sc_tools [data-testid="stPopover"] button * { font-size: 0.92rem !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] * { font-weight: 800 !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button::after { display: none !important; }
/* Nút ⚙️ và nút tài khoản ở bên phải thanh */
html body .st-key-sc_tools [data-testid="stPopover"] button {
    min-height: 38px !important;
    height: 38px !important;
    min-width: 38px !important;
    padding: 0 10px !important;
    border-radius: 999px !important;
    background: rgba(255, 255, 255, 0.16) !important;
    border: 1px solid rgba(255, 255, 255, 0.3) !important;
    box-shadow: none !important;
}
html body .st-key-sc_tools [data-testid="stPopover"] button:hover { background: rgba(255, 255, 255, 0.28) !important; }
html body .st-key-sc_tools [data-testid="stPopover"] button * {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    fill: #ffffff !important;
    font-weight: 800 !important;
}
html body .st-key-pop_account button {
    background: linear-gradient(135deg, #fde047, #facc15) !important;
    border: none !important;
}
html body .st-key-pop_account button * { color: #0c4a6e !important; -webkit-text-fill-color: #0c4a6e !important; }
html body .st-key-sc_tools button::after { display: none !important; }
/* Thẻ tài khoản trong bảng bật ra: nền xanh để chữ trắng dễ đọc */
div[data-testid="stPopoverBody"] .profile-card {
    background: linear-gradient(135deg, #1d4ed8, #0c4a6e) !important;
    margin-top: 0 !important;
}
div[data-testid="stPopoverBody"] .profile-card,
div[data-testid="stPopoverBody"] .profile-card * {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
}
div[data-testid="stPopoverBody"] .profile-card .pc-avatar,
div[data-testid="stPopoverBody"] .profile-card .pc-avatar * { color: #0c4a6e !important; -webkit-text-fill-color: #0c4a6e !important; }
div[data-testid="stPopoverBody"] .profile-card .pc-role,
div[data-testid="stPopoverBody"] .profile-card .pc-role * { color: #fde047 !important; -webkit-text-fill-color: #fde047 !important; }
div[data-testid="stPopoverBody"] .profile-card .pc-label { color: rgba(255,255,255,0.75) !important; -webkit-text-fill-color: rgba(255,255,255,0.75) !important; }

/* ---------- 🏠 Trang "Về chúng tôi" ---------- */
.sc-about-hero {
    text-align: center;
    padding: 34px 24px 26px 24px;
    margin: 6px auto 18px auto;
    max-width: 900px;
    border-radius: 22px;
    background: __SURFACE__;
    border: 1px solid __SURFACE_BORDER__;
    box-shadow: 0 10px 30px rgba(15, 23, 42, 0.10);
    animation: scFadeUp 0.35s ease-out both;
}
.sc-about-kicker { font-weight: 900; letter-spacing: 0.12em; color: #0ea5e9; font-size: 0.95rem; }
.sc-about-hero h1 { font-size: 2rem !important; font-weight: 900 !important; margin: 8px 0 10px 0 !important; color: __INPUT_TEXT__ !important; line-height: 1.25 !important; }
.sc-about-hero p { font-size: 1.1rem; color: __PLACEHOLDER__; margin: 0 auto; max-width: 680px; }
.sc-about-grid {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
    gap: 14px;
    max-width: 1000px;
    margin: 0 auto 22px auto;
}
.sc-about-card {
    padding: 18px;
    border-radius: 18px;
    background: __SURFACE__;
    border: 1px solid __SURFACE_BORDER__;
    box-shadow: 0 6px 18px rgba(15, 23, 42, 0.07);
    transition: transform 0.2s ease, box-shadow 0.2s ease;
    animation: scFadeUp 0.45s ease-out both;
}
.sc-about-card:hover { transform: translateY(-3px); box-shadow: 0 12px 26px rgba(15, 23, 42, 0.12); }
.sc-about-card h3 { font-size: 1.1rem !important; margin: 0 0 6px 0 !important; color: __INPUT_TEXT__ !important; }
.sc-about-card p { margin: 0; color: __PLACEHOLDER__; font-size: 0.95rem; }

.sc-about-hero.sc-about-hero-left { text-align: left; margin: 0; padding: 30px 28px; max-width: none; }
.sc-about-hero.sc-about-hero-left h1 { font-size: 1.9rem !important; }
.sc-about-hero.sc-about-hero-left p { margin: 0; }
.sc-about-video-title { text-align: center; margin: 6px 0 10px 0; }
.sc-about-video-title h3 { font-size: 1.4rem !important; font-weight: 800 !important; margin: 0 0 4px 0 !important; color: __INPUT_TEXT__ !important; }
.sc-about-video-title p { margin: 0; color: __PLACEHOLDER__; }
.st-key-sc_about_video {
    border-radius: 18px !important;
    overflow: hidden !important;
    box-shadow: 0 14px 34px rgba(15, 23, 42, 0.22) !important;
    border: 1px solid __SURFACE_BORDER__ !important;
}
.st-key-sc_about_video iframe, .st-key-sc_about_video video { display: block; border-radius: 18px; }

/* ---------- 🎥 Thẻ ID + mật khẩu phòng họp ---------- */
.sc-meet-card {
    display: flex;
    flex-wrap: wrap;
    gap: 12px;
    margin: 6px 0 4px 0;
}
.sc-meet-card > div {
    flex: 1 1 160px;
    padding: 12px 16px;
    border-radius: 14px;
    background: __SELECTED_BG__;
    border: 2px dashed #0ea5e9;
}
.sc-meet-card span { display: block; font-size: 0.85rem; color: __PLACEHOLDER__; font-weight: 600; }
.sc-meet-card b { display: block; font-size: 1.7rem; letter-spacing: 0.08em; color: __INPUT_TEXT__; font-family: ui-monospace, 'SF Mono', Menlo, Consolas, monospace !important; }
.sc-meet-card.sc-compact > div { padding: 8px 12px; }
.sc-meet-card.sc-compact b { font-size: 1.25rem; }

/* ================================================================
   🌊 MƯỢT MÀ HƠN
   ================================================================ */
html { scroll-behavior: smooth; }
body, .stApp {
    -webkit-font-smoothing: antialiased;
    -moz-osx-font-smoothing: grayscale;
    text-rendering: optimizeLegibility;
}
/* Khi app đang xử lý, Streamlit làm mờ nội dung cũ (nhấp nháy) → giữ nguyên độ rõ */
[data-stale="true"], .stale-element {
    opacity: 1 !important;
    filter: none !important;
    transition: none !important;
}
/* Bỏ thanh màu nhấp nháy ở mép trên */
[data-testid="stDecoration"] { display: none !important; }
/* Thông báo, khung mở rộng, bảng cài đặt và danh sách chọn hiện ra nhẹ nhàng */
div[data-testid="stAlert"] { animation: scFadeUp 0.28s ease-out both; }
[data-testid="stExpanderDetails"] { animation: scFadeUp 0.25s ease-out both; }
div[data-testid="stPopoverBody"] { animation: scPop 0.18s cubic-bezier(0.2, 0.8, 0.2, 1) both; transform-origin: top left; }
[data-baseweb="popover"] ul[role="listbox"] { animation: scPop 0.15s ease-out both; transform-origin: top center; }
@keyframes scFadeUp {
    from { opacity: 0; transform: translate3d(0, 6px, 0); }
    to   { opacity: 1; transform: none; }
}
@keyframes scPop {
    from { opacity: 0; transform: scale(0.96) translate3d(0, -4px, 0); }
    to   { opacity: 1; transform: none; }
}
/* Đổi chủ đề / Sáng-Tối: màu chuyển dần thay vì đổi phụt */
section[data-testid="stSidebar"], div[data-testid="stForm"], div.st-key-sc_add_task_box, div[data-testid="stExpander"] details,
div[data-testid="stAlert"], .profile-card, .main-title, .big-table-title, .stApp p, .stApp label {
    transition: background-color 0.35s ease, color 0.35s ease, border-color 0.35s ease, box-shadow 0.35s ease;
}
/* Nút chọn thiết bị, bảng, liên kết: chuyển trạng thái mượt */
a, label[data-testid="stRadioOption"], div[data-testid="stDataFrame"] {
    transition: color 0.2s ease, background-color 0.2s ease, border-color 0.2s ease, box-shadow 0.2s ease, transform 0.2s ease;
}
/* Cuộn thanh bên mượt, không kéo theo cả trang */
section[data-testid="stSidebar"] [data-testid="stSidebarContent"] {
    scroll-behavior: smooth;
    overscroll-behavior: contain;
}
@media (prefers-reduced-motion: reduce) {
    html { scroll-behavior: auto; }
    div[data-testid="stAlert"], [data-testid="stExpanderDetails"], div[data-testid="stPopoverBody"],
    [data-baseweb="popover"] ul[role="listbox"] { animation: none !important; }
}

/* ---------- Thanh cuộn mảnh ---------- */
::-webkit-scrollbar { width: 10px; height: 10px; }
::-webkit-scrollbar-thumb { background: rgba(100, 116, 139, 0.35); border-radius: 999px; }
::-webkit-scrollbar-track { background: transparent; }
'''

# Phần CSS thay đổi theo từng trang (ảnh nền) được tách riêng và rất nhỏ,
# để khi đổi trang trình duyệt không phải xử lý lại toàn bộ CSS lớn.
BG_CSS_TEMPLATE = r'''
.stApp::before {
    content: "";
    position: fixed;
    inset: 0;
    z-index: 0;
    pointer-events: none;
    background: __APP_BG__;
    background-size: cover;
    background-position: center;
    animation: __BG_ANIM__ 0.45s ease-out both;
    will-change: opacity;
}
@keyframes __BG_ANIM__ {
    from { opacity: 0.35; }
    to   { opacity: 1; }
}
'''


# ==========================================
# 🌐 BỘ DỊCH GIAO DIỆN (Tiếng Việt → English)
# ==========================================
VI_EN = {
    # --- Tiêu đề & menu ---
    "🚢 SHIPCONTROL - QUẢN LÝ CÔNG VIỆC TÀU": "🚢 SHIPCONTROL - SHIP WORK MANAGEMENT",
    "🔐 Xác Thực": "🔐 Sign In",
    "Vui lòng đăng nhập hoặc đăng ký để tiếp tục.": "Please sign in or register to continue.",
    "🧰 Bảng Công Việc": "🧰 Task Board",
    "➕ Thêm Công Việc": "➕ Add Task",
    "📋 Giao Việc": "📋 Assign Tasks",
    "👥 Team Của Tôi": "👥 My Team",
    "✏️ Chỉnh Sửa/Xóa": "✏️ Edit / Delete",
    "👥 Quản Lý Phân Quyền": "👥 Roles & Permissions",
    "🗑️ Thùng Rác": "🗑️ Trash",
    "📊 Báo Cáo & Khai Báo": "📊 Reports",
    "🔑 Đổi Mật Khẩu": "🔑 Change Password",
    "🚪 Đăng Xuất": "🚪 Log Out",
    "⚙️ Cài đặt": "⚙️ Settings",
    "📲 Thiết bị": "📲 Device",
    "📲 Bạn đang dùng:": "📲 You are using:",
    "💻 Máy tính": "💻 Computer",
    "📱 Điện thoại": "📱 Phone",
    "🎤 Bấm nút micro trong ô nhập chữ để nói thay vì gõ.": "🎤 Tap the microphone in a text box to speak instead of typing.",
    "🎤 Chế độ điện thoại: bấm nút micro trong ô Tên tài khoản để nói thay vì gõ.":
        "🎤 Phone mode: tap the microphone in the Username box to speak instead of typing.",
    "Trình duyệt này chưa hỗ trợ nhập bằng giọng nói. Hãy dùng Chrome (Android) hoặc Safari (iPhone).":
        "This browser doesn't support voice input yet. Please use Chrome (Android) or Safari (iPhone).",
    "Chưa được phép dùng micro. Hãy cho phép micro cho trang web này trong cài đặt trình duyệt.":
        "Microphone access is blocked. Please allow the microphone for this site in your browser settings.",
    "🏭 Workshop: xem danh sách / thêm / xóa": "🏭 Workshops: view / add / delete",
    "Chưa có Workshop nào. Hãy thêm Workshop đầu tiên bên dưới.": "No workshops yet. Add the first one below.",
    "#### ➕ Thêm Workshop Mới": "#### ➕ Add New Workshop",
    "Mã Workshop (Ví dụ: WOS_07):": "Workshop code (e.g. WOS_07):",
    "➕ THÊM WORKSHOP": "➕ ADD WORKSHOP",
    "Mã tối đa 20 ký tự, tên tối đa 80 ký tự!": "Code max 20 characters, name max 80 characters!",
    "#### 🗑️ Xóa Workshop": "#### 🗑️ Delete Workshop",
    "Chọn Workshop để xóa:": "Workshop to delete:",
    "🗑️ XÓA WORKSHOP": "🗑️ DELETE WORKSHOP",
    "Tên Workshop": "Workshop name",
    "Số tài khoản": "Accounts",
    "Số công việc": "Tasks",
    "🎨 Chủ đề": "🎨 Theme",
    "✨ Hiện đại (Modern)": "✨ Modern",
    "🚀 Tương lai (Futuristic)": "🚀 Futuristic",
    "🎈 Vui nhộn (Playful)": "🎈 Playful",
    "📜 Cổ điển (Old times)": "📜 Old times",
    "Chủ đề Tương lai luôn dùng nền tối.": "The Futuristic theme is always dark.",
    ">Tài khoản<": ">Account<",
    "🔒 Tài khoản của bạn vừa được mở ở một tab hoặc thiết bị khác, nên bạn đã bị đăng xuất khỏi trang này. ":
        "🔒 Your account was just opened in another tab or on another device, so you have been signed out here. ",
    "Nếu đó không phải bạn, hãy đăng nhập lại và đổi mật khẩu ngay.":
        "If that wasn't you, sign in again and change your password right away.",
    ">Vai trò<": ">Role<",
    ">Đang đăng nhập<": ">Signed in<",
    "🌙 Chế độ Tối (Dark)": "🌙 Dark mode",
    "🌐 Ngôn ngữ": "🌐 Language",

    # --- Tiêu đề trang ---
    "📋 Bảng Quản Lý Tiến Độ Công Việc": "📋 Task Progress Board",
    "➕ Thêm Công Việc Mới": "➕ Add New Task",
    "✏️ Chỉnh Sửa & Xóa Quản Lý": "✏️ Edit & Delete",
    "👥 Quản Lý & Cấp Quyền Tài Khoản (Role List)": "👥 Accounts & Roles (Role List)",
    "📊 Báo Cáo & Thống Kê Tiến Độ": "📊 Progress Reports & Statistics",
    "🗑️ Thùng Rác & Khôi Phục Tổng Hợp": "🗑️ Trash & Restore",
    "### ⚡ Cập Nhật Tiến Độ Công Việc Nhanh": "### ⚡ Quick Progress Update",
    "### 🏭 Thống Kê Theo Workshop": "### 🏭 Statistics by Workshop",
    "### 📈 Phân Bố Tiến Độ Công Việc": "### 📈 Task Progress Distribution",
    "### 📌 Giao Việc Cho Worker": "### 📌 Assign Tasks to Workers",
    "### 📥 Tải Báo Cáo Dữ Liệu": "### 📥 Download Report Data",
    "### 🔄 Thay Đổi Quyền Hạn Cho Tài Khoản": "### 🔄 Change Account Role",
    "#### ➕ Thêm Worker": "#### ➕ Add Workers",
    "#### ➖ Xóa Khỏi Team": "#### ➖ Remove from Team",
    "#### 👷 Thành Viên": "#### 👷 Members",
    "#### 🚪 Rời Team": "#### 🚪 Leave Team",

    # --- Đăng nhập / đăng ký ---
    "🔑 Đăng Nhập": "🔑 Sign In",
    "📝 Đăng Ký Tài Khoản": "📝 Register",
    "Tên tài khoản (Username):": "Username:",
    "Mật khẩu (Password):": "Password:",
    "🚀 ĐĂNG NHẬP": "🚀 SIGN IN",
    "✨ ĐĂNG KÝ NGAY": "✨ REGISTER NOW",
    "Họ và Tên:": "Full name:",
    "Tên đăng nhập mới (Username):": "New username:",
    "Xác nhận lại mật khẩu:": "Confirm password:",
    "Vui lòng nhập đầy đủ Username và Mật khẩu!": "Please enter both username and password!",
    "Sai tên tài khoản, mật khẩu hoặc tài khoản đã bị khóa/xóa!": "Wrong username or password, or the account is locked/deleted!",
    "⏳ Tài khoản của bạn đang chờ Admin / WOS Manager cấp Role. Vui lòng quay lại sau!": "⏳ Your account is waiting for an Admin / WOS Manager to assign a role. Please come back later!",
    "Chào mừng ": "Welcome ",
    ") đã quay trở lại!": ") is back!",
    "Vui lòng điền đầy đủ các thông tin!": "Please fill in all fields!",
    "Mật khẩu phải có ít nhất 8 ký tự!": "Password must be at least 8 characters!",
    "Username hoặc Họ tên quá dài!": "Username or full name is too long!",
    "Mật khẩu xác nhận không trùng khớp!": "Passwords do not match!",
    "Tên đăng nhập này đã tồn tại, vui lòng chọn tên khác!": "This username already exists, please choose another!",
    "🎉 Đăng ký thành công! Tài khoản đang chờ WOS Manager phê duyệt Role.": "🎉 Registered! Your account is waiting for a WOS Manager to approve a role.",

    # --- Bảng công việc ---
    "👷 Đây là các công việc được giao cho bạn.": "👷 These are the tasks assigned to you.",
    "Chỉ xem công việc của team tôi": "Show only my team's tasks",
    "Chưa có công việc nào để hiển thị.": "No tasks to show yet.",
    "Bạn chưa có công việc nào để cập nhật tiến độ.": "You have no tasks to update.",
    "Chọn công việc cần cập nhật tiến độ:": "Choose a task to update:",
    "Mức Tiến Độ Mới (%)": "New progress (%)",
    "Ghi Chú Thi Công (Remark):": "Work note (Remark):",
    "🚀 CẬP NHẬT TIẾN ĐỘ": "🚀 UPDATE PROGRESS",
    "Đã cập nhật tiến độ công việc thành công!": "Task progress updated!",
    " (Hiện tại: ": " (Current: ",

    # --- Thêm công việc ---
    "⚙️ Thêm / Xóa trong danh sách Block và WS Cost Code": "⚙️ Add / Remove Blocks and WS Cost Codes",
    "Block mới (Ví dụ: 170150):": "New block (e.g. 170150):",
    "➕ THÊM BLOCK": "➕ ADD BLOCK",
    "Vui lòng nhập tên Block!": "Please enter a block name!",
    "Đã thêm Block **": "Added block **",
    "Chọn Block để xóa:": "Block to delete:",
    "🗑️ XÓA BLOCK": "🗑️ DELETE BLOCK",
    "Đã xóa Block **": "Deleted block **",
    "Danh sách Block đang trống.": "The block list is empty.",
    "Mã (Ví dụ: WOS_07):": "Code (e.g. WOS_07):",
    "Tên Workshop / Phòng ban:": "Workshop / department name:",
    "➕ THÊM COST CODE": "➕ ADD COST CODE",
    "Vui lòng nhập cả Mã và Tên!": "Please enter both code and name!",
    "Chọn Cost Code để xóa:": "Cost code to delete:",
    "🗑️ XÓA COST CODE": "🗑️ DELETE COST CODE",
    "Danh sách Cost Code đang trống.": "The cost code list is empty.",
    "Không thể xóa: còn **": "Can't delete: **",
    "** tài khoản thuộc workshop này. Hãy chuyển họ sang workshop khác ở 👥 Quản Lý Phân Quyền trước.":
        "** account(s) still belong to this workshop. Move them to another workshop in 👥 Roles & Permissions first.",
    "** khỏi danh sách. Công việc cũ vẫn giữ nguyên.": "** from the list. Existing tasks are kept.",
    "** đã có trong danh sách!": "** is already in the list!",
    "Mã **": "Code **",
    "Đã thêm **": "Added **",
    "Đã xóa **": "Deleted **",
    "Description (Mô tả công việc):": "Description:",
    "Initial By (Người khởi tạo):": "Initial By:",
    "Initial Date (Ngày tạo):": "Initial Date:",
    "In Charge By (Người phụ trách):": "In Charge By:",
    "Remark (Ghi chú):": "Remark:",
    "— Không chọn —": "— None —",
    "📌 Công việc này sẽ thuộc team của bạn. Vào 📋 Giao Việc để giao cho Worker.":
        "📌 This task will belong to your team. Go to 📋 Assign Tasks to give it to a worker.",
    "💾 LƯU CÔNG VIỆC MỚI": "💾 SAVE NEW TASK",
    "Vui lòng điền đầy đủ thông tin bắt buộc: Task ID và Task Name!": "Please fill in the required fields: Task ID and Task Name!",
    "Plan Finish Date không được trước Plan Start Date!": "Plan Finish Date can't be before Plan Start Date!",
    "** đã tồn tại (có thể đang trong Thùng Rác). Vui lòng dùng ID khác!": "** already exists (it may be in the Trash). Please use another ID!",
    "Đã lưu thành công công việc **": "Saved task **",
    "⚠️ CHƯA CÓ WS COST CODE: Foreman / WOS Manager hãy thêm Cost Code ở mục ⚙️ phía trên trước khi thêm công việc!":
        "⚠️ NO WS COST CODE YET: a Foreman / WOS Manager must add one in the ⚙️ section above before adding tasks!",

    # --- Giao việc ---
    "👑 Chọn công việc và giao cho một **Team Leader** hoặc **Foreman**. Họ sẽ giao tiếp cho Worker trong team.":
        "👑 Choose a task and assign it to a **Team Leader** or **Foreman**. They will pass it on to workers in their team.",
    "Chưa có công việc nào. Vào ➕ Thêm Công Việc để tạo.": "No tasks yet. Go to ➕ Add Task to create one.",
    "Chưa có Team Leader / Foreman nào. Vào 👥 Quản Lý Phân Quyền để cấp role trước.":
        "No Team Leaders / Foremen yet. Assign roles in 👥 Roles & Permissions first.",
    "Chỉ hiện công việc chưa giao": "Show only unassigned tasks",
    "Tất cả công việc đã được giao!": "All tasks have been assigned!",
    "Chọn công việc:": "Choose a task:",
    "Giao cho Team Leader / Foreman:": "Assign to Team Leader / Foreman:",
    "Giao cho Worker:": "Assign to worker:",
    "— Chưa giao —": "— Unassigned —",
    "Chưa giao": "Unassigned",
    "💾 LƯU GIAO VIỆC": "💾 SAVE ASSIGNMENT",
    "Đã giao công việc cho **": "Task assigned to **",
    "Team của bạn chưa có Worker nào. Vào 👥 Team Của Tôi để tạo team và thêm Worker.":
        "Your team has no workers yet. Go to 👥 My Team to create a team and add workers.",
    "Bạn chưa được giao công việc nào.": "You haven't been given any tasks yet.",
    "Cần có Worker trong team trước khi giao việc.": "You need workers in your team before assigning tasks.",
    "(chưa đặt tên)": "(unnamed)",

    # --- Team ---
    "⚠️ Tài khoản của bạn chưa được gán Workshop. Nhờ WOS Manager gán Workshop trước khi lập team.":
        "⚠️ Your account has no workshop yet. Ask a WOS Manager to assign one before creating a team.",
    "Bạn chưa có team. Hãy đặt tên để tạo team.": "You don't have a team yet. Give it a name to create one.",
    "Tên Team *": "Team name *",
    "Ví dụ: Team Hàn Block 170": "e.g. Welding Team Block 170",
    "➕ TẠO TEAM": "➕ CREATE TEAM",
    "Vui lòng nhập tên team!": "Please enter a team name!",
    "Tên team tối đa 60 ký tự!": "Team name can be at most 60 characters!",
    "Đã tạo team **": "Created team **",
    "✏️ Đổi tên team": "✏️ Rename team",
    "Tên team mới:": "New team name:",
    "💾 LƯU TÊN": "💾 SAVE NAME",
    "Tên team không hợp lệ (1–60 ký tự)!": "Invalid team name (1–60 characters)!",
    "Đã đổi tên team!": "Team renamed!",
    "🗑️ Xóa team": "🗑️ Delete team",
    "Khi xóa team: tất cả Worker sẽ rời team, và các công việc đang giao cho họ sẽ trở về trạng thái chưa giao. Công việc vẫn thuộc về bạn, bạn có thể tạo team mới sau.":
        "Deleting the team: all workers leave the team and their tasks become unassigned. The tasks stay with you, and you can create a new team later.",
    "Tôi chắc chắn muốn xóa team này": "I'm sure I want to delete this team",
    "🗑️ XÓA TEAM": "🗑️ DELETE TEAM",
    "Đã xóa team.": "Team deleted.",
    "Team chưa có Worker nào.": "The team has no workers yet.",
    "Không còn Worker nào trong workshop ": "No workers left in workshop ",
    " chưa có team.": " without a team.",
    "Chọn Worker (cùng workshop, chưa có team):": "Choose workers (same workshop, no team yet):",
    "➕ THÊM VÀO TEAM": "➕ ADD TO TEAM",
    "Hãy chọn ít nhất 1 Worker!": "Choose at least 1 worker!",
    " Worker vào team!": " worker(s) to the team!",
    "Đã thêm ": "Added ",
    "Chưa có thành viên để xóa.": "No members to remove.",
    "Chọn Worker:": "Choose a worker:",
    "➖ XÓA KHỎI TEAM": "➖ REMOVE FROM TEAM",
    "Đã xóa Worker khỏi team. Các việc của team giao cho người này đã được bỏ giao.":
        "Worker removed from the team. Their team tasks are now unassigned.",
    "Bạn chưa thuộc team nào. Team Leader / Foreman trong workshop của bạn có thể thêm bạn vào team.":
        "You're not in a team yet. A Team Leader / Foreman in your workshop can add you.",
    "(Team chưa đặt tên)": "(Unnamed team)",
    "**Trưởng team:** ": "**Team leader:** ",
    "Khi rời team, các công việc của team đang giao cho bạn sẽ trở về trạng thái chưa giao.":
        "When you leave, the team tasks assigned to you become unassigned.",
    "Tôi chắc chắn muốn rời team này": "I'm sure I want to leave this team",
    "🚪 RỜI TEAM": "🚪 LEAVE TEAM",
    "Bạn đã rời team.": "You left the team.",

    # --- Chỉnh sửa / xóa ---
    "🧰 Xóa Tạm Công Việc": "🧰 Move Tasks to Trash",
    "👤 Xóa Tạm Tài Khoản Người Dùng": "👤 Lock User Accounts",
    "Hiện không có công việc nào để chỉnh sửa hoặc xóa.": "There are no tasks to edit or delete.",
    "Chọn công việc cần chuyển vào Thùng Rác:": "Task to move to Trash:",
    "🗑️ Chuyển Công Việc Vào Thùng Rác": "🗑️ Move Task to Trash",
    "Đã chuyển công việc vào Thùng Rác thành công!": "Task moved to Trash!",
    "Không có tài khoản khác khả dụng để xóa.": "No other accounts available to delete.",
    "Chọn tài khoản muốn khóa/xóa tạm:": "Account to lock:",
    "🗑️ Khóa/Xóa Tạm Tài Khoản Này": "🗑️ Lock This Account",
    "Đã khóa/chuyển tài khoản vào Thùng Rác thành công!": "Account locked and moved to Trash!",

    # --- Phân quyền ---
    "🛡️ Admin chỉ cấp quyền **WOS Manager** và chọn **Workshop** mà Manager đó phụ trách. Các role Worker / Team Leader / Foreman do WOS Manager cấp.":
        "🛡️ The Admin only grants the **WOS Manager** role and chooses the **Workshop** that manager runs. Worker / Team Leader / Foreman roles are given by WOS Managers.",
    "👑 Bạn có thể cấp **Worker / Team Leader / Foreman**, chọn **Workshop**, và xếp Worker vào **team** của một Team Leader / Foreman.":
        "👑 You can give **Worker / Team Leader / Foreman** roles, choose the **Workshop**, and put workers into a Team Leader / Foreman's **team**.",
    "Không có tài khoản nào để cấp quyền.": "No accounts to assign roles to.",
    "Chọn tài khoản cần chuyển đổi Role:": "Account to change:",
    "Chọn thao tác:": "Action:",
    "👑 Cấp quyền WOS Manager": "👑 Make WOS Manager",
    "⛔ Thu hồi quyền (về Pending)": "⛔ Revoke role (back to Pending)",
    "Workshop mà Manager này phụ trách:": "Workshop this manager runs:",
    "⚠️ Chưa có Workshop nào. Vào ➕ Thêm Công Việc → ⚙️ Thêm / Xóa trong danh sách để thêm.":
        "⚠️ No workshops yet. Add them in ➕ Add Task → ⚙️ Add / Remove.",
    "💾 LƯU THAY ĐỔI ROLE": "💾 SAVE ROLE",
    "💾 LƯU THAY ĐỔI": "💾 SAVE CHANGES",
    "Đã cấp **WOS Manager** cho workshop **": "Made **WOS Manager** of workshop **",
    "Đã thu hồi quyền, tài khoản trở về trạng thái Pending.": "Role revoked, the account is back to Pending.",
    "Chọn Role Mới:": "New role:",
    "Thuộc team của (Team Leader / Foreman):": "Team of (Team Leader / Foreman):",
    "— Chưa xếp team —": "— No team —",
    "Workshop này chưa có Team Leader / Foreman nào.": "This workshop has no Team Leader / Foreman yet.",
    "Đã cập nhật: **": "Updated: **",
    "Họ và Tên": "Full name",
    "Vai Trò (Role)": "Role",
    "Team của": "Team of",

    # --- Thùng rác ---
    "🧰 Thùng Rác Công Việc": "🧰 Task Trash",
    "👤 Thùng Rác Tài Khoản": "👤 Account Trash",
    "Thùng rác công việc đang trống.": "The task trash is empty.",
    "Thùng rác tài khoản đang trống.": "The account trash is empty.",
    "Chọn công việc để xử lý:": "Choose a task:",
    "♻️ KHÔI PHỤC CÔNG VIỆC": "♻️ RESTORE TASK",
    "💥 XÓA VĨNH VIỄN CÔNG VIỆC": "💥 DELETE TASK FOREVER",
    "Đã khôi phục công việc!": "Task restored!",
    "Đã xóa vĩnh viễn công việc!": "Task permanently deleted!",
    "Chọn tài khoản để xử lý:": "Choose an account:",
    "♻️ MỞ KHÓA / KHÔI PHỤC TÀI KHOẢN": "♻️ UNLOCK / RESTORE ACCOUNT",
    "💥 XÓA VĨNH VIỄN TÀI KHOẢN": "💥 DELETE ACCOUNT FOREVER",
    "Đã khôi phục tài khoản người dùng!": "Account restored!",
    "Đã xóa vĩnh viễn tài khoản khỏi cơ sở dữ liệu!": "Account permanently deleted from the database!",

    # --- Báo cáo ---
    "Chưa có dữ liệu công việc để tạo báo cáo. Vui lòng thêm công việc trước!": "No task data for a report yet. Please add tasks first!",
    "Tổng Số Công Việc": "Total tasks",
    "Tổng Công Việc": "Total tasks",
    "Tiến Độ Trung Bình (%)": "Average progress (%)",
    "Tiến Độ Trung Bình": "Average progress",
    "Đã Hoàn Thành (100%)": "Completed (100%)",
    "Đang Thực Hiện": "In progress",
    "Số Lượng Công Việc": "Number of tasks",
    "Trạng Thái": "Status",
    "📥 Tải Báo Cáo Bảng Công Việc (File CSV/Excel)": "📥 Download Task Report (CSV/Excel)",

    # --- Đổi mật khẩu ---
    "Mật khẩu hiện tại:": "Current password:",
    "Mật khẩu mới (ít nhất 8 ký tự):": "New password (at least 8 characters):",
    "Nhập lại mật khẩu mới:": "Repeat new password:",
    "💾 ĐỔI MẬT KHẨU": "💾 CHANGE PASSWORD",
    "Mật khẩu hiện tại không đúng!": "Current password is wrong!",
    "Mật khẩu mới phải có ít nhất 8 ký tự!": "New password must be at least 8 characters!",
    "Không được dùng lại mật khẩu mặc định!": "You can't reuse the default password!",
    "✅ Đã đổi mật khẩu thành công! Các thiết bị khác đã bị đăng xuất.": "✅ Password changed! Other devices have been logged out.",
    "⚠️ Tài khoản này vẫn dùng mật khẩu mặc định 'admin123'. Vào mục 🔑 Đổi Mật Khẩu và đổi NGAY!":
        "⚠️ This account still uses the default password 'admin123'. Go to 🔑 Change Password and change it NOW!",
    # --- Biểu đồ ---
    "Chưa bắt đầu (0%)": "Not started (0%)",
    "Đang làm (1-50%)": "In progress (1-50%)",
    "Sắp xong (51-99%)": "Almost done (51-99%)",
    "Hoàn thành (100%)": "Done (100%)",
    "Số Lượng": "Count",
    "Mật khẩu mới:": "New password:",
    "💬 Tin Nhắn": "💬 Messages",
    "📢 Kênh chung": "📢 General chat",
    "#### 👥 Mọi người": "#### 👥 People",
    "🔎 Tìm tài khoản": "🔎 Find account",
    "Không tìm thấy tài khoản.": "No account found.",
    "Mọi người trong app đều thấy kênh này.": "Everyone in the app can see this channel.",
    "Tin nhắn riêng với ": "Private chat with ",
    "Nhập tin nhắn...": "Type a message...",
    "➕ Tạo nhóm chat": "➕ Create group chat",
    "Tên nhóm": "Group name",
    "Ví dụ: Tổ hàn Block 170": "e.g. Welding crew Block 170",
    "Thêm người vào nhóm": "Add people to the group",
    "✨ TẠO NHÓM": "✨ CREATE GROUP",
    "Vui lòng đặt tên nhóm!": "Please name the group!",
    "Hãy chọn ít nhất 1 người!": "Choose at least 1 person!",
    "Đã tạo nhóm **": "Created group **",
    "#### 💬 Nhóm chat": "#### 💬 Group chats",
    "⚙️ Thành viên nhóm": "⚙️ Group members",
    "➕ THÊM VÀO NHÓM": "➕ ADD TO GROUP",
    " người vào nhóm!": " people to the group!",
    "Xóa khỏi nhóm": "Remove from group",
    "➖ XÓA KHỎI NHÓM": "➖ REMOVE FROM GROUP",
    "Đã xóa khỏi nhóm.": "Removed from the group.",
    "🗑️ XÓA NHÓM": "🗑️ DELETE GROUP",
    "Chỉ người tạo nhóm mới thêm / xóa được thành viên.": "Only the group creator can add / remove members.",
    "🚪 RỜI NHÓM": "🚪 LEAVE GROUP",
    "🎥 Online Meeting": "🎥 Online Meetings",
    "### 🔑 Vào cuộc họp": "### 🔑 Join a meeting",
    "ID phòng họp": "Meeting room ID",
    "Mật khẩu phòng": "Room password",
    "🎥 VÀO PHÒNG": "🎥 JOIN",
    "Bạn đã nhập sai quá nhiều lần. Hãy tải lại trang và thử lại sau.": "Too many wrong attempts. Reload the page and try again later.",
    "Sai ID phòng hoặc mật khẩu, hoặc cuộc họp đã kết thúc!": "Wrong room ID or password, or the meeting has ended!",
    "✅ Đúng mã phòng: **": "✅ Room found: **",
    "Chủ phòng": "Host",
    "🎥 MỞ PHÒNG HỌP (camera + micro)": "🎥 OPEN MEETING (camera + mic)",
    "Phòng họp mở trong tab mới. Lần đầu, trình duyệt sẽ hỏi quyền dùng camera và micro — hãy bấm Cho phép.": "The meeting opens in a new tab. The first time, your browser will ask to use the camera and microphone — tap Allow.",
    "### ➕ Tạo cuộc họp mới": "### ➕ Create a new meeting",
    "Tên cuộc họp *": "Meeting name *",
    "Ví dụ: Họp giao ca sáng": "e.g. Morning shift handover",
    "Ngày họp": "Date",
    "Giờ họp": "Time",
    "Mời người (gửi ID + mật khẩu qua 💬 Tin Nhắn)": "Invite people (sends ID + password via 💬 Messages)",
    "📢 Thông báo vào Kênh chung": "📢 Announce in General chat",
    "🎥 TẠO CUỘC HỌP": "🎥 CREATE MEETING",
    "Vui lòng nhập tên cuộc họp!": "Please enter a meeting name!",
    "🎥 Mời họp": "🎥 Meeting invite",
    "🕒 Thời gian": "🕒 Time",
    "🔢 ID phòng": "🔢 Room ID",
    "🔑 Mật khẩu": "🔑 Password",
    "👉 Vào mục 🎥 Online Meeting để tham gia.": "👉 Go to 🎥 Online Meetings to join.",
    "🎉 Đã tạo cuộc họp **": "🎉 Created meeting **",
    "### 📋 Cuộc họp của tôi": "### 📋 My meetings",
    "Bạn chưa có cuộc họp nào.": "You have no meetings yet.",
    "▶️ BẮT ĐẦU": "▶️ START",
    "⛔ KẾT THÚC": "⛔ END",
    "Chỉ Foreman và Team Leader mới tạo được cuộc họp. Bạn có thể vào họp khi có ID phòng và mật khẩu.": "Only Foremen and Team Leaders can create meetings. You can join with a room ID and password.",
    "Chọn người...": "Choose people...",
    "👑 Quản trị nhóm": "👑 Group admin",
    "🗑️ Xóa tin nhắn (quản trị)": "🗑️ Delete messages (admin)",
    "Chưa có tin nhắn để xóa.": "No messages to delete.",
    "Chọn tin nhắn để xóa": "Message to delete",
    "🗑️ XÓA TIN NHẮN": "🗑️ DELETE MESSAGE",
    "Đã xóa tin nhắn.": "Message deleted.",
    "Tôi chắc chắn muốn xóa toàn bộ tin nhắn": "I'm sure I want to delete all messages",
    "🧹 XÓA TOÀN BỘ TIN NHẮN": "🧹 DELETE ALL MESSAGES",
    "Đã xóa toàn bộ tin nhắn.": "All messages deleted.",
    "Mời người (họ sẽ thấy lời mời trong 🎥 Online Meeting)": "Invite people (they'll see the invite in 🎥 Online Meetings)",
    "### 📨 Lời mời họp của bạn": "### 📨 Your meeting invites",
    "🎥 VÀO HỌP": "🎥 JOIN MEETING",
    "💬 Tin Nhắn & Meeting": "💬 Messages & Meetings",
    "Cài đặt": "Settings",
    "About us": "About us",
    "Đăng nhập / Đăng ký": "Sign in / Register",
    "Phần mềm quản lý công việc cho xưởng đóng tàu": "Work management software for shipyards",
    "ShipControl giúp đội ngũ xưởng tàu giao việc, theo dõi tiến độ và liên lạc với nhau ở cùng một nơi — trên máy tính lẫn điện thoại.": "ShipControl helps shipyard teams assign work, track progress and stay in touch in one place — on computers and phones.",
    "📋 Quản lý công việc": "📋 Task management",
    "Tạo công việc theo Block, Deck, Frame và theo dõi tiến độ từng phần trăm.": "Create tasks by Block, Deck and Frame and track progress percent by percent.",
    "👥 Team & phân quyền": "👥 Teams & roles",
    "Admin, WOS Manager, Foreman, Team Leader và Worker — mỗi người thấy đúng phần việc của mình.": "Admin, WOS Manager, Foreman, Team Leader and Worker — everyone sees exactly their own work.",
    "💬 Tin nhắn & họp online": "💬 Messages & online meetings",
    "Kênh chung, nhóm chat và phòng họp video có ID + mật khẩu.": "General chat, group chats and video meeting rooms with an ID + password.",
    "🌐 15 ngôn ngữ & giọng nói": "🌐 15 languages & voice",
    "Tiếng Việt, English, 中文, 日本語, 한국어, ไทย, Indonesia, Español, Français, Deutsch, Русский… — và nói thay vì gõ trên điện thoại.": "Vietnamese, English, Chinese, Japanese, Korean, Thai, Indonesian, Spanish, French, German, Russian and more — and speak instead of typing on phones.",
    "Chưa có tài khoản? Đăng ký và chờ quản lý cấp quyền.": "No account yet? Register and wait for a manager to give you a role.",
    "🎬 Từ xưởng đến biển": "🎬 From shipyard to sea",
    "Xem một con tàu được đóng từ những khối thép đầu tiên đến khi hạ thủy.": "Watch a ship being built from the first steel blocks to launch day.",
    "Không xem được video? ": "Can't play the video? ",
    "▶ Mở trên YouTube": "▶ Open on YouTube",
    "Tạo nhóm chat": "Create group chat",
    "Chưa có nhóm nào — bấm ➕ để tạo.": "No groups yet — tap ➕ to create one.",
    "Thêm": "Add",
    "Xóa": "Remove",
    "Thêm Worker": "Add workers",
    "Xóa Khỏi Team": "Remove from team",
    "Thêm Workshop": "Add workshop",
    "Xóa Workshop": "Remove workshop",
    "Xóa nhóm": "Delete group",
    "Rời nhóm": "Leave group",
    "🌊 Đại dương (Ocean)": "🌊 Ocean",
    "🌅 Hoàng hôn (Sunset)": "🌅 Sunset",
    "🌲 Rừng xanh (Forest)": "🌲 Forest",
    "⚪ Tối giản (Minimal)": "⚪ Minimal",
    "👑 Hoàng gia (Royal)": "👑 Royal",
    "📸 Báo cáo hình ảnh theo tuần": "📸 Weekly photo report",
    "📈 Thống kê": "📈 Statistics",
    "Bạn chưa có công việc nào để báo cáo.": "You have no tasks to report on yet.",
    "Bắt đầu": "Start",
    "Kết thúc dự kiến": "Planned finish",
    "Tiến độ": "Progress",
    "Đã hoàn thành": "Finished",
    "Tuần hiện tại": "Current week",
    "Tải ảnh lên": "Upload photos",
    "Tuần": "Week",
    "Chọn ảnh từ máy tính": "Choose photos from your computer",
    "Ghi chú cho ảnh (không bắt buộc)": "Photo note (optional)",
    "📤 TẢI LÊN": "📤 UPLOAD",
    "Hãy chọn ít nhất 1 ảnh!": "Choose at least 1 photo!",
    "Không đọc được ảnh (hoặc ảnh lớn hơn 15 MB): ": "Couldn't read these photos (or they are larger than 15 MB): ",
    "Chưa có ảnh nào cho tuần này.": "No photos for this week yet.",
    "📥 Tải báo cáo hình ảnh (HTML)": "📥 Download photo report (HTML)",
    "Mở file rồi bấm Ctrl+P để in hoặc lưu thành PDF.": "Open the file and press Ctrl+P to print or save it as PDF.",
    "💡 Giữ Ctrl (hoặc Shift) khi chọn để lấy nhiều ảnh cùng lúc. Bấm Upload thêm lần nữa để thêm ảnh khác.": "💡 Hold Ctrl (or Shift) while choosing to pick many photos at once. Press Upload again to add more.",
    "📊 Tải PowerPoint (.pptx)": "📊 Download PowerPoint (.pptx)",
    "◀ Trước": "◀ Previous",
    "Tiếp ▶": "Next ▶",
    "Tổng số tuần": "Total weeks",
    "Tổng số ảnh": "Total photos",
    "Đang thực hiện": "In progress",
    "ảnh": "photos",
    "Tổng kết": "Summary",
}

# 🌏 Bản dịch Trung / Nhật / Hàn: cùng thứ tự với VI_EN (mỗi dòng khớp một câu tiếng Việt)
_ZH_LIST = [
    "🚢 SHIPCONTROL - 船舶工作管理",
    "🔐 登录",
    "请登录或注册以继续。",
    "🧰 任务看板",
    "➕ 添加任务",
    "📋 分配任务",
    "👥 我的团队",
    "✏️ 编辑 / 删除",
    "👥 角色与权限",
    "🗑️ 回收站",
    "📊 报告",
    "🔑 修改密码",
    "🚪 退出登录",
    "⚙️ 设置",
    "📲 设备",
    "📲 您正在使用：",
    "💻 电脑",
    "📱 手机",
    "🎤 点击文本框中的麦克风，用说话代替打字。",
    "🎤 手机模式：点击用户名框中的麦克风，用说话代替打字。",
    "此浏览器暂不支持语音输入。请使用 Chrome（安卓）或 Safari（iPhone）。",
    "麦克风权限被阻止。请在浏览器设置中允许本网站使用麦克风。",
    "🏭 车间：查看 / 添加 / 删除",
    "还没有车间。请在下方添加第一个。",
    "#### ➕ 添加新车间",
    "车间代码（例如 WOS_07）：",
    "➕ 添加车间",
    "代码最多 20 个字符，名称最多 80 个字符！",
    "#### 🗑️ 删除车间",
    "要删除的车间：",
    "🗑️ 删除车间",
    "车间名称",
    "账号数",
    "任务数",
    "🎨 主题",
    "✨ 现代",
    "🚀 未来",
    "🎈 活泼",
    "📜 复古",
    "未来主题始终为深色。",
    ">账号<",
    "🔒 您的账号刚在另一个标签页或设备上打开，因此您已在此处退出。",
    "如果不是您本人，请重新登录并立即修改密码。",
    ">角色<",
    ">已登录<",
    "🌙 深色模式",
    "🌐 语言",
    "📋 任务进度看板",
    "➕ 添加新任务",
    "✏️ 编辑与删除",
    "👥 账号与角色（角色列表）",
    "📊 进度报告与统计",
    "🗑️ 回收站与恢复",
    "### ⚡ 快速更新进度",
    "### 🏭 按车间统计",
    "### 📈 任务进度分布",
    "### 📌 给工人分配任务",
    "### 📥 下载报告数据",
    "### 🔄 更改账号角色",
    "#### ➕ 添加工人",
    "#### ➖ 移出团队",
    "#### 👷 成员",
    "#### 🚪 退出团队",
    "🔑 登录",
    "📝 注册",
    "用户名：",
    "密码：",
    "🚀 登录",
    "✨ 立即注册",
    "姓名：",
    "新用户名：",
    "确认密码：",
    "请输入用户名和密码！",
    "用户名或密码错误，或账号已被锁定/删除！",
    "⏳ 您的账号正在等待管理员 / WOS Manager 分配角色。请稍后再来！",
    "欢迎 ",
    "）回来了！",
    "请填写所有信息！",
    "密码至少需要 8 个字符！",
    "用户名或姓名太长！",
    "两次输入的密码不一致！",
    "该用户名已存在，请换一个！",
    "🎉 注册成功！您的账号正在等待 WOS Manager 批准角色。",
    "👷 这些是分配给您的任务。",
    "只显示我团队的任务",
    "暂无可显示的任务。",
    "您没有需要更新的任务。",
    "选择要更新的任务：",
    "新进度 (%)",
    "施工备注：",
    "🚀 更新进度",
    "任务进度已更新！",
    "（当前：",
    "⚙️ 添加 / 删除分段和 WS 成本代码",
    "新分段（例如 170150）：",
    "➕ 添加分段",
    "请输入分段名称！",
    "已添加分段 **",
    "要删除的分段：",
    "🗑️ 删除分段",
    "已删除分段 **",
    "分段列表为空。",
    "代码（例如 WOS_07）：",
    "车间 / 部门名称：",
    "➕ 添加成本代码",
    "请输入代码和名称！",
    "要删除的成本代码：",
    "🗑️ 删除成本代码",
    "成本代码列表为空。",
    "无法删除：还有 **",
    "** 个账号属于此车间。请先在 👥 角色与权限 中将他们移到其他车间。",
    "** 已从列表中删除。现有任务保持不变。",
    "** 已在列表中！",
    "代码 **",
    "已添加 **",
    "已删除 **",
    "描述：",
    "创建人：",
    "创建日期：",
    "负责人：",
    "备注：",
    "— 无 —",
    "📌 此任务将属于您的团队。前往 📋 分配任务 把它交给工人。",
    "💾 保存新任务",
    "请填写必填项：Task ID 和 Task Name！",
    "计划完成日期不能早于计划开始日期！",
    "** 已存在（可能在回收站中）。请使用其他 ID！",
    "已保存任务 **",
    "⚠️ 还没有 WS 成本代码：Foreman / WOS Manager 需先在上方 ⚙️ 部分添加！",
    "👑 选择一个任务并分配给 **Team Leader** 或 **Foreman**。他们会再分配给团队中的工人。",
    "还没有任务。前往 ➕ 添加任务 创建。",
    "还没有 Team Leader / Foreman。请先在 👥 角色与权限 中分配角色。",
    "只显示未分配的任务",
    "所有任务都已分配！",
    "选择任务：",
    "分配给 Team Leader / Foreman：",
    "分配给工人：",
    "— 未分配 —",
    "未分配",
    "💾 保存分配",
    "任务已分配给 **",
    "您的团队还没有工人。前往 👥 我的团队 创建团队并添加工人。",
    "您还没有被分配任何任务。",
    "分配任务前，团队中需要有工人。",
    "（未命名）",
    "⚠️ 您的账号还没有车间。请先让 WOS Manager 分配车间再创建团队。",
    "您还没有团队。请起个名字来创建团队。",
    "团队名称 *",
    "例如：170 分段焊接队",
    "➕ 创建团队",
    "请输入团队名称！",
    "团队名称最多 60 个字符！",
    "已创建团队 **",
    "✏️ 重命名团队",
    "新团队名称：",
    "💾 保存名称",
    "团队名称无效（1–60 个字符）！",
    "团队已重命名！",
    "🗑️ 删除团队",
    "删除团队后：所有工人将离开团队，他们的任务变为未分配。任务仍归您所有，之后可以创建新团队。",
    "我确定要删除这个团队",
    "🗑️ 删除团队",
    "团队已删除。",
    "团队还没有工人。",
    "车间 ",
    " 中已没有未加入团队的工人。",
    "选择工人（同一车间，尚无团队）：",
    "➕ 加入团队",
    "请至少选择 1 名工人！",
    " 名工人加入团队！",
    "已将 ",
    "没有可移除的成员。",
    "选择工人：",
    "➖ 移出团队",
    "已将工人移出团队。其团队任务现为未分配。",
    "您还没有加入团队。您车间的 Team Leader / Foreman 可以把您加入团队。",
    "（未命名团队）",
    "**队长：** ",
    "离开后，分配给您的团队任务将变为未分配。",
    "我确定要离开这个团队",
    "🚪 离开团队",
    "您已离开团队。",
    "🧰 将任务移到回收站",
    "👤 锁定用户账号",
    "没有可编辑或删除的任务。",
    "要移到回收站的任务：",
    "🗑️ 将任务移到回收站",
    "任务已移到回收站！",
    "没有其他可删除的账号。",
    "要锁定的账号：",
    "🗑️ 锁定此账号",
    "账号已锁定并移到回收站！",
    "🛡️ 管理员只授予 **WOS Manager** 角色并选择其负责的 **车间**。Worker / Team Leader / Foreman 角色由 WOS Manager 分配。",
    "👑 您可以分配 **Worker / Team Leader / Foreman** 角色、选择 **车间**，并把工人加入 Team Leader / Foreman 的 **团队**。",
    "没有可分配角色的账号。",
    "要更改的账号：",
    "操作：",
    "👑 设为 WOS Manager",
    "⛔ 撤销角色（恢复为待审核）",
    "该经理负责的车间：",
    "⚠️ 还没有车间。请在 ➕ 添加任务 → ⚙️ 添加 / 删除 中添加。",
    "💾 保存角色",
    "💾 保存更改",
    "已设为 **WOS Manager**，车间 **",
    "角色已撤销，账号恢复为待审核。",
    "新角色：",
    "所属团队（Team Leader / Foreman）：",
    "— 无团队 —",
    "该车间还没有 Team Leader / Foreman。",
    "已更新：**",
    "姓名",
    "角色",
    "所属团队",
    "🧰 任务回收站",
    "👤 账号回收站",
    "任务回收站为空。",
    "账号回收站为空。",
    "选择任务：",
    "♻️ 恢复任务",
    "💥 永久删除任务",
    "任务已恢复！",
    "任务已永久删除！",
    "选择账号：",
    "♻️ 解锁 / 恢复账号",
    "💥 永久删除账号",
    "账号已恢复！",
    "账号已从数据库中永久删除！",
    "还没有可生成报告的任务数据。请先添加任务！",
    "任务总数",
    "任务总数",
    "平均进度 (%)",
    "平均进度",
    "已完成 (100%)",
    "进行中",
    "任务数量",
    "状态",
    "📥 下载任务报告（CSV/Excel）",
    "当前密码：",
    "新密码（至少 8 个字符）：",
    "再次输入新密码：",
    "💾 修改密码",
    "当前密码错误！",
    "新密码至少需要 8 个字符！",
    "不能再次使用默认密码！",
    "✅ 密码已修改！其他设备已退出登录。",
    "⚠️ 此账号仍在使用默认密码 'admin123'。请立即前往 🔑 修改密码 进行修改！",
    "未开始 (0%)",
    "进行中 (1-50%)",
    "即将完成 (51-99%)",
    "已完成 (100%)",
    "数量",
    "新密码：",
    "💬 消息",
    "📢 公共频道",
    "#### 👥 成员",
    "🔎 查找账号",
    "未找到账号。",
    "应用中的所有人都能看到此频道。",
    "私聊：",
    "输入消息...",
    "➕ 创建群聊",
    "群名称",
    "例如：170 分段焊接组",
    "添加成员",
    "✨ 创建群组",
    "请为群组命名！",
    "请至少选择 1 人！",
    "已创建群组 **",
    "#### 💬 群聊",
    "⚙️ 群成员",
    "➕ 加入群组",
    " 人加入群组！",
    "移出群组",
    "➖ 移出群组",
    "已移出群组。",
    "🗑️ 删除群组",
    "只有群主可以添加 / 移除成员。",
    "🚪 退出群组",
    "🎥 在线会议",
    "### 🔑 加入会议",
    "会议室 ID",
    "会议密码",
    "🎥 加入",
    "输入错误次数过多。请刷新页面后再试。",
    "会议室 ID 或密码错误，或会议已结束！",
    "✅ 已找到会议：**",
    "主持人",
    "🎥 打开会议（摄像头 + 麦克风）",
    "会议将在新标签页中打开。首次使用时，浏览器会请求摄像头和麦克风权限——请点击允许。",
    "### ➕ 创建新会议",
    "会议名称 *",
    "例如：早班交接会",
    "日期",
    "时间",
    "邀请成员（通过 💬 消息发送 ID 和密码）",
    "📢 在公共频道通知",
    "🎥 创建会议",
    "请输入会议名称！",
    "🎥 会议邀请",
    "🕒 时间",
    "🔢 会议室 ID",
    "🔑 密码",
    "👉 前往 🎥 在线会议 加入。",
    "🎉 已创建会议 **",
    "### 📋 我的会议",
    "您还没有会议。",
    "▶️ 开始",
    "⛔ 结束",
    "只有 Foreman 和 Team Leader 可以创建会议。有会议室 ID 和密码即可加入。",
    "选择成员...",
    "👑 群管理员",
    "🗑️ 删除消息（管理员）",
    "没有可删除的消息。",
    "选择要删除的消息",
    "🗑️ 删除消息",
    "消息已删除。",
    "我确定要删除所有消息",
    "🧹 删除所有消息",
    "已删除所有消息。",
    "邀请成员（他们会在 🎥 在线会议 中看到邀请）",
    "### 📨 您的会议邀请",
    "🎥 加入会议",
    "💬 消息与会议",
    "设置",
    "关于我们",
    "登录 / 注册",
    "船厂工作管理软件",
    "ShipControl 让船厂团队在同一个地方分配工作、跟踪进度并保持沟通——电脑和手机都能用。",
    "📋 任务管理",
    "按分段、甲板、肋位创建任务，并逐个百分比跟踪进度。",
    "👥 团队与权限",
    "Admin、WOS Manager、Foreman、Team Leader 和 Worker——每个人只看到自己的工作。",
    "💬 消息与在线会议",
    "公共频道、群聊，以及带 ID 和密码的视频会议室。",
    "🌐 15 种语言与语音输入",
    "越南语、英语、中文、日语、韩语、泰语、印尼语、西班牙语、法语、德语、俄语等——在手机上还可以用说话代替打字。",
    "还没有账号？请注册并等待管理员分配角色。",
    "🎬 从船厂到大海",
    "观看一艘船从第一块钢板到下水的建造过程。",
    "无法播放视频？",
    "▶ 在 YouTube 上打开",
    "创建群聊",
    "还没有群组——点击 ➕ 创建。",
    "添加",
    "删除",
    "添加工人",
    "移出团队",
    "添加车间",
    "删除车间",
    "删除群组",
    "退出群组",
    "🌊 海洋",
    "🌅 日落",
    "🌲 森林",
    "⚪ 极简",
    "👑 皇家",
    "📸 每周图片报告",
    "📈 统计",
    "你还没有可以报告的工作。",
    "开始",
    "预计完成",
    "进度",
    "已完成",
    "本周",
    "上传图片",
    "第",
    "从电脑选择图片",
    "图片备注（可选）",
    "📤 上传",
    "请至少选择 1 张图片！",
    "无法读取这些图片（或大于 15 MB）：",
    "本周还没有图片。",
    "📥 下载图片报告（HTML）",
    "打开文件后按 Ctrl+P 打印或另存为 PDF。",
    "💡 选择时按住 Ctrl（或 Shift）可一次选多张图片。再按一次 Upload 可继续添加。",
    "📊 下载 PowerPoint (.pptx)",
    "◀ 上一页",
    "下一页 ▶",
    "总周数",
    "图片总数",
    "进行中",
    "张图片",
    "总结",
]
_JA_LIST = [
    "🚢 SHIPCONTROL - 船舶作業管理",
    "🔐 ログイン",
    "続けるにはログインまたは登録してください。",
    "🧰 タスクボード",
    "➕ タスク追加",
    "📋 タスク割り当て",
    "👥 マイチーム",
    "✏️ 編集 / 削除",
    "👥 役割と権限",
    "🗑️ ゴミ箱",
    "📊 レポート",
    "🔑 パスワード変更",
    "🚪 ログアウト",
    "⚙️ 設定",
    "📲 デバイス",
    "📲 使用中のデバイス：",
    "💻 パソコン",
    "📱 スマホ",
    "🎤 入力欄のマイクをタップすると、入力の代わりに話せます。",
    "🎤 スマホモード：ユーザー名欄のマイクをタップすると、入力の代わりに話せます。",
    "このブラウザは音声入力に対応していません。Chrome（Android）または Safari（iPhone）をご利用ください。",
    "マイクへのアクセスがブロックされています。ブラウザの設定でこのサイトのマイクを許可してください。",
    "🏭 ワークショップ：一覧 / 追加 / 削除",
    "ワークショップがまだありません。下から最初の1つを追加してください。",
    "#### ➕ 新しいワークショップを追加",
    "ワークショップコード（例：WOS_07）：",
    "➕ ワークショップを追加",
    "コードは最大20文字、名前は最大80文字です！",
    "#### 🗑️ ワークショップを削除",
    "削除するワークショップ：",
    "🗑️ ワークショップを削除",
    "ワークショップ名",
    "アカウント数",
    "タスク数",
    "🎨 テーマ",
    "✨ モダン",
    "🚀 フューチャー",
    "🎈 ポップ",
    "📜 レトロ",
    "フューチャーテーマは常にダークです。",
    ">アカウント<",
    "🔒 別のタブまたは端末でこのアカウントが開かれたため、ここではログアウトしました。",
    "心当たりがない場合は、再度ログインしてすぐにパスワードを変更してください。",
    ">役割<",
    ">ログイン中<",
    "🌙 ダークモード",
    "🌐 言語",
    "📋 タスク進捗ボード",
    "➕ 新しいタスクを追加",
    "✏️ 編集と削除",
    "👥 アカウントと役割（役割一覧）",
    "📊 進捗レポートと統計",
    "🗑️ ゴミ箱と復元",
    "### ⚡ 進捗をすばやく更新",
    "### 🏭 ワークショップ別の統計",
    "### 📈 タスク進捗の分布",
    "### 📌 作業員にタスクを割り当て",
    "### 📥 レポートデータをダウンロード",
    "### 🔄 アカウントの役割を変更",
    "#### ➕ 作業員を追加",
    "#### ➖ チームから外す",
    "#### 👷 メンバー",
    "#### 🚪 チームを抜ける",
    "🔑 ログイン",
    "📝 新規登録",
    "ユーザー名：",
    "パスワード：",
    "🚀 ログイン",
    "✨ 今すぐ登録",
    "氏名：",
    "新しいユーザー名：",
    "パスワード（確認）：",
    "ユーザー名とパスワードを入力してください！",
    "ユーザー名かパスワードが違うか、アカウントがロック/削除されています！",
    "⏳ アカウントは Admin / WOS Manager による役割の割り当てを待っています。後でもう一度お試しください！",
    "ようこそ ",
    "）さん、おかえりなさい！",
    "すべての項目を入力してください！",
    "パスワードは8文字以上にしてください！",
    "ユーザー名または氏名が長すぎます！",
    "パスワードが一致しません！",
    "このユーザー名は既に使われています。別の名前を選んでください！",
    "🎉 登録完了！アカウントは WOS Manager による役割の承認を待っています。",
    "👷 あなたに割り当てられたタスクです。",
    "自分のチームのタスクのみ表示",
    "表示するタスクはまだありません。",
    "更新できるタスクがありません。",
    "更新するタスクを選択：",
    "新しい進捗 (%)",
    "作業メモ（Remark）：",
    "🚀 進捗を更新",
    "タスクの進捗を更新しました！",
    "（現在：",
    "⚙️ ブロックと WS コストコードの追加 / 削除",
    "新しいブロック（例：170150）：",
    "➕ ブロックを追加",
    "ブロック名を入力してください！",
    "ブロックを追加しました **",
    "削除するブロック：",
    "🗑️ ブロックを削除",
    "ブロックを削除しました **",
    "ブロック一覧は空です。",
    "コード（例：WOS_07）：",
    "ワークショップ / 部署名：",
    "➕ コストコードを追加",
    "コードと名前の両方を入力してください！",
    "削除するコストコード：",
    "🗑️ コストコードを削除",
    "コストコード一覧は空です。",
    "削除できません：まだ **",
    "** 件のアカウントがこのワークショップに所属しています。先に 👥 役割と権限 で別のワークショップへ移してください。",
    "** を一覧から削除しました。既存のタスクはそのままです。",
    "** は既に一覧にあります！",
    "コード **",
    "追加しました **",
    "削除しました **",
    "説明：",
    "作成者：",
    "作成日：",
    "担当者：",
    "備考：",
    "— なし —",
    "📌 このタスクはあなたのチームのものになります。📋 タスク割り当て で作業員に渡してください。",
    "💾 新しいタスクを保存",
    "必須項目を入力してください：Task ID と Task Name！",
    "計画完了日を計画開始日より前にはできません！",
    "** は既に存在します（ゴミ箱にある可能性があります）。別の ID を使ってください！",
    "タスクを保存しました **",
    "⚠️ WS コストコードがまだありません：Foreman / WOS Manager が上の ⚙️ で追加してください！",
    "👑 タスクを選んで **Team Leader** または **Foreman** に割り当てます。その人がチームの作業員に割り振ります。",
    "タスクがまだありません。➕ タスク追加 で作成してください。",
    "Team Leader / Foreman がまだいません。先に 👥 役割と権限 で役割を割り当ててください。",
    "未割り当てのタスクのみ表示",
    "すべてのタスクが割り当て済みです！",
    "タスクを選択：",
    "Team Leader / Foreman に割り当て：",
    "作業員に割り当て：",
    "— 未割り当て —",
    "未割り当て",
    "💾 割り当てを保存",
    "タスクを割り当てました **",
    "チームに作業員がいません。👥 マイチーム でチームを作り、作業員を追加してください。",
    "まだタスクが割り当てられていません。",
    "タスクを割り当てる前に、チームに作業員が必要です。",
    "（名前なし）",
    "⚠️ アカウントにワークショップがありません。チームを作る前に WOS Manager に割り当ててもらってください。",
    "まだチームがありません。名前を付けてチームを作成してください。",
    "チーム名 *",
    "例：ブロック170 溶接チーム",
    "➕ チームを作成",
    "チーム名を入力してください！",
    "チーム名は最大60文字です！",
    "チームを作成しました **",
    "✏️ チーム名を変更",
    "新しいチーム名：",
    "💾 名前を保存",
    "チーム名が無効です（1〜60文字）！",
    "チーム名を変更しました！",
    "🗑️ チームを削除",
    "チームを削除すると、作業員は全員チームを抜け、割り当て中のタスクは未割り当てになります。タスクはあなたに残り、後で新しいチームを作れます。",
    "このチームを削除します",
    "🗑️ チームを削除",
    "チームを削除しました。",
    "チームにはまだ作業員がいません。",
    "ワークショップ ",
    " にはチーム未所属の作業員がいません。",
    "作業員を選択（同じワークショップ、チーム未所属）：",
    "➕ チームに追加",
    "作業員を1人以上選んでください！",
    " 人の作業員をチームに追加しました！",
    "",
    "外せるメンバーがいません。",
    "作業員を選択：",
    "➖ チームから外す",
    "作業員をチームから外しました。その人のチームタスクは未割り当てになりました。",
    "まだチームに入っていません。同じワークショップの Team Leader / Foreman が追加できます。",
    "（名前のないチーム）",
    "**チームリーダー：** ",
    "チームを抜けると、あなたに割り当てられたチームタスクは未割り当てになります。",
    "このチームを抜けます",
    "🚪 チームを抜ける",
    "チームを抜けました。",
    "🧰 タスクをゴミ箱へ",
    "👤 ユーザーアカウントをロック",
    "編集・削除できるタスクがありません。",
    "ゴミ箱へ移すタスク：",
    "🗑️ タスクをゴミ箱へ移す",
    "タスクをゴミ箱へ移しました！",
    "削除できる他のアカウントがありません。",
    "ロックするアカウント：",
    "🗑️ このアカウントをロック",
    "アカウントをロックしてゴミ箱へ移しました！",
    "🛡️ Admin は **WOS Manager** の役割を付与し、担当する **ワークショップ** を選ぶだけです。Worker / Team Leader / Foreman の役割は WOS Manager が付与します。",
    "👑 **Worker / Team Leader / Foreman** の役割付与、**ワークショップ** の選択、Team Leader / Foreman の **チーム** への作業員配置ができます。",
    "役割を割り当てるアカウントがありません。",
    "変更するアカウント：",
    "操作：",
    "👑 WOS Manager にする",
    "⛔ 役割を取り消す（承認待ちに戻す）",
    "このマネージャーが担当するワークショップ：",
    "⚠️ ワークショップがまだありません。➕ タスク追加 → ⚙️ 追加 / 削除 で追加してください。",
    "💾 役割を保存",
    "💾 変更を保存",
    "**WOS Manager** に設定しました。ワークショップ **",
    "役割を取り消し、アカウントは承認待ちに戻りました。",
    "新しい役割：",
    "所属チーム（Team Leader / Foreman）：",
    "— チームなし —",
    "このワークショップには Team Leader / Foreman がまだいません。",
    "更新しました：**",
    "氏名",
    "役割",
    "所属チーム",
    "🧰 タスクのゴミ箱",
    "👤 アカウントのゴミ箱",
    "タスクのゴミ箱は空です。",
    "アカウントのゴミ箱は空です。",
    "タスクを選択：",
    "♻️ タスクを復元",
    "💥 タスクを完全に削除",
    "タスクを復元しました！",
    "タスクを完全に削除しました！",
    "アカウントを選択：",
    "♻️ アカウントのロック解除 / 復元",
    "💥 アカウントを完全に削除",
    "アカウントを復元しました！",
    "アカウントをデータベースから完全に削除しました！",
    "レポートを作るタスクデータがまだありません。先にタスクを追加してください！",
    "タスク総数",
    "タスク総数",
    "平均進捗 (%)",
    "平均進捗",
    "完了 (100%)",
    "進行中",
    "タスク数",
    "状態",
    "📥 タスクレポートをダウンロード（CSV/Excel）",
    "現在のパスワード：",
    "新しいパスワード（8文字以上）：",
    "新しいパスワード（再入力）：",
    "💾 パスワードを変更",
    "現在のパスワードが違います！",
    "新しいパスワードは8文字以上にしてください！",
    "初期パスワードは再利用できません！",
    "✅ パスワードを変更しました！他の端末はログアウトされました。",
    "⚠️ このアカウントはまだ初期パスワード 'admin123' を使っています。今すぐ 🔑 パスワード変更 で変更してください！",
    "未着手 (0%)",
    "作業中 (1-50%)",
    "もうすぐ完了 (51-99%)",
    "完了 (100%)",
    "件数",
    "新しいパスワード：",
    "💬 メッセージ",
    "📢 全体チャット",
    "#### 👥 メンバー",
    "🔎 アカウント検索",
    "アカウントが見つかりません。",
    "アプリの全員がこのチャンネルを見られます。",
    "個別チャット：",
    "メッセージを入力...",
    "➕ グループチャットを作成",
    "グループ名",
    "例：ブロック170 溶接班",
    "メンバーを追加",
    "✨ グループを作成",
    "グループ名を入力してください！",
    "1人以上選んでください！",
    "グループを作成しました **",
    "#### 💬 グループチャット",
    "⚙️ グループメンバー",
    "➕ グループに追加",
    " 人をグループに追加しました！",
    "グループから外す",
    "➖ グループから外す",
    "グループから外しました。",
    "🗑️ グループを削除",
    "メンバーの追加 / 削除はグループ作成者だけができます。",
    "🚪 グループを抜ける",
    "🎥 オンライン会議",
    "### 🔑 会議に参加",
    "会議室 ID",
    "会議パスワード",
    "🎥 参加",
    "入力ミスが多すぎます。ページを再読み込みして後でもう一度お試しください。",
    "会議室 ID またはパスワードが違うか、会議は終了しています！",
    "✅ 会議が見つかりました：**",
    "主催者",
    "🎥 会議を開く（カメラ + マイク）",
    "会議は新しいタブで開きます。初回はブラウザがカメラとマイクの使用許可を求めるので「許可」をタップしてください。",
    "### ➕ 新しい会議を作成",
    "会議名 *",
    "例：朝の引き継ぎ会議",
    "日付",
    "時刻",
    "参加者を招待（💬 メッセージで ID とパスワードを送信）",
    "📢 全体チャットでお知らせ",
    "🎥 会議を作成",
    "会議名を入力してください！",
    "🎥 会議の招待",
    "🕒 日時",
    "🔢 会議室 ID",
    "🔑 パスワード",
    "👉 🎥 オンライン会議 から参加してください。",
    "🎉 会議を作成しました **",
    "### 📋 自分の会議",
    "まだ会議はありません。",
    "▶️ 開始",
    "⛔ 終了",
    "会議を作成できるのは Foreman と Team Leader だけです。会議室 ID とパスワードがあれば参加できます。",
    "メンバーを選択...",
    "👑 グループ管理者",
    "🗑️ メッセージ削除（管理者）",
    "削除できるメッセージはありません。",
    "削除するメッセージ",
    "🗑️ メッセージを削除",
    "メッセージを削除しました。",
    "すべてのメッセージを削除します",
    "🧹 すべてのメッセージを削除",
    "すべてのメッセージを削除しました。",
    "参加者を招待（🎥 オンライン会議 に招待が表示されます）",
    "### 📨 あなたへの会議の招待",
    "🎥 会議に参加",
    "💬 メッセージと会議",
    "設定",
    "私たちについて",
    "ログイン / 登録",
    "造船所のための作業管理ソフト",
    "ShipControl は造船所のチームが作業の割り当て、進捗管理、連絡をひとつの場所で行えるようにします。パソコンでもスマホでも使えます。",
    "📋 タスク管理",
    "ブロック・デッキ・フレームごとにタスクを作成し、進捗を1%単位で管理。",
    "👥 チームと権限",
    "Admin、WOS Manager、Foreman、Team Leader、Worker ——それぞれが自分の仕事だけを見られます。",
    "💬 メッセージとオンライン会議",
    "全体チャット、グループチャット、ID とパスワード付きのビデオ会議室。",
    "🌐 15か国語と音声入力",
    "ベトナム語・英語・中国語・日本語・韓国語・タイ語・インドネシア語・スペイン語・フランス語・ドイツ語・ロシア語など。スマホでは入力の代わりに話せます。",
    "アカウントがない場合は、登録して管理者による役割の付与を待ってください。",
    "🎬 造船所から海へ",
    "最初の鋼材ブロックから進水まで、船が造られる様子をご覧ください。",
    "動画が再生できませんか？",
    "▶ YouTube で開く",
    "グループチャットを作成",
    "グループはまだありません — ➕ をタップして作成。",
    "追加",
    "削除",
    "作業員を追加",
    "チームから外す",
    "ワークショップを追加",
    "ワークショップを削除",
    "グループを削除",
    "グループを退出",
    "🌊 オーシャン",
    "🌅 サンセット",
    "🌲 フォレスト",
    "⚪ ミニマル",
    "👑 ロイヤル",
    "📸 週次写真レポート",
    "📈 統計",
    "報告できる作業はまだありません。",
    "開始",
    "完了予定",
    "進捗",
    "完了",
    "今週",
    "写真をアップロード",
    "第",
    "パソコンから写真を選択",
    "写真メモ（任意）",
    "📤 アップロード",
    "写真を1枚以上選んでください！",
    "読み込めなかった写真（または15 MB超）：",
    "この週の写真はまだありません。",
    "📥 写真レポートをダウンロード（HTML）",
    "ファイルを開いて Ctrl+P で印刷または PDF 保存できます。",
    "💡 Ctrl（または Shift）を押しながら選ぶと複数の写真を一度に選べます。Upload をもう一度押すと追加できます。",
    "📊 PowerPoint をダウンロード (.pptx)",
    "◀ 前へ",
    "次へ ▶",
    "合計週数",
    "写真の合計",
    "進行中",
    "枚",
    "まとめ",
]
_KO_LIST = [
    "🚢 SHIPCONTROL - 선박 작업 관리",
    "🔐 로그인",
    "계속하려면 로그인하거나 회원가입하세요.",
    "🧰 작업 보드",
    "➕ 작업 추가",
    "📋 작업 배정",
    "👥 내 팀",
    "✏️ 편집 / 삭제",
    "👥 역할 및 권한",
    "🗑️ 휴지통",
    "📊 보고서",
    "🔑 비밀번호 변경",
    "🚪 로그아웃",
    "⚙️ 설정",
    "📲 기기",
    "📲 사용 중인 기기:",
    "💻 컴퓨터",
    "📱 휴대폰",
    "🎤 입력창의 마이크를 누르면 타이핑 대신 말로 입력할 수 있습니다.",
    "🎤 휴대폰 모드: 사용자 이름 칸의 마이크를 누르면 타이핑 대신 말로 입력할 수 있습니다.",
    "이 브라우저는 아직 음성 입력을 지원하지 않습니다. Chrome(안드로이드) 또는 Safari(iPhone)를 사용하세요.",
    "마이크 접근이 차단되었습니다. 브라우저 설정에서 이 사이트의 마이크를 허용하세요.",
    "🏭 작업장: 보기 / 추가 / 삭제",
    "아직 작업장이 없습니다. 아래에서 첫 작업장을 추가하세요.",
    "#### ➕ 새 작업장 추가",
    "작업장 코드 (예: WOS_07):",
    "➕ 작업장 추가",
    "코드는 최대 20자, 이름은 최대 80자입니다!",
    "#### 🗑️ 작업장 삭제",
    "삭제할 작업장:",
    "🗑️ 작업장 삭제",
    "작업장 이름",
    "계정 수",
    "작업 수",
    "🎨 테마",
    "✨ 모던",
    "🚀 퓨처리스틱",
    "🎈 플레이풀",
    "📜 레트로",
    "퓨처리스틱 테마는 항상 어두운 모드입니다.",
    ">계정<",
    "🔒 다른 탭이나 기기에서 계정이 열려 이곳에서는 로그아웃되었습니다. ",
    "본인이 아니라면 다시 로그인하여 즉시 비밀번호를 변경하세요.",
    ">역할<",
    ">로그인됨<",
    "🌙 다크 모드",
    "🌐 언어",
    "📋 작업 진행 보드",
    "➕ 새 작업 추가",
    "✏️ 편집 및 삭제",
    "👥 계정 및 역할 (역할 목록)",
    "📊 진행 보고서 및 통계",
    "🗑️ 휴지통 및 복원",
    "### ⚡ 빠른 진행 상황 업데이트",
    "### 🏭 작업장별 통계",
    "### 📈 작업 진행 분포",
    "### 📌 작업자에게 작업 배정",
    "### 📥 보고서 데이터 다운로드",
    "### 🔄 계정 역할 변경",
    "#### ➕ 작업자 추가",
    "#### ➖ 팀에서 제외",
    "#### 👷 팀원",
    "#### 🚪 팀 나가기",
    "🔑 로그인",
    "📝 회원가입",
    "사용자 이름:",
    "비밀번호:",
    "🚀 로그인",
    "✨ 지금 가입",
    "이름:",
    "새 사용자 이름:",
    "비밀번호 확인:",
    "사용자 이름과 비밀번호를 모두 입력하세요!",
    "사용자 이름 또는 비밀번호가 틀렸거나 계정이 잠김/삭제되었습니다!",
    "⏳ 계정이 Admin / WOS Manager의 역할 배정을 기다리고 있습니다. 나중에 다시 오세요!",
    "환영합니다 ",
    ")님, 다시 오셨네요!",
    "모든 항목을 입력하세요!",
    "비밀번호는 8자 이상이어야 합니다!",
    "사용자 이름 또는 이름이 너무 깁니다!",
    "비밀번호가 일치하지 않습니다!",
    "이미 사용 중인 사용자 이름입니다. 다른 이름을 선택하세요!",
    "🎉 가입 완료! 계정이 WOS Manager의 역할 승인을 기다리고 있습니다.",
    "👷 나에게 배정된 작업입니다.",
    "내 팀의 작업만 보기",
    "표시할 작업이 없습니다.",
    "업데이트할 작업이 없습니다.",
    "업데이트할 작업 선택:",
    "새 진행률 (%)",
    "작업 메모 (Remark):",
    "🚀 진행률 업데이트",
    "작업 진행률이 업데이트되었습니다!",
    " (현재: ",
    "⚙️ 블록 및 WS 원가 코드 추가 / 삭제",
    "새 블록 (예: 170150):",
    "➕ 블록 추가",
    "블록 이름을 입력하세요!",
    "블록을 추가했습니다 **",
    "삭제할 블록:",
    "🗑️ 블록 삭제",
    "블록을 삭제했습니다 **",
    "블록 목록이 비어 있습니다.",
    "코드 (예: WOS_07):",
    "작업장 / 부서 이름:",
    "➕ 원가 코드 추가",
    "코드와 이름을 모두 입력하세요!",
    "삭제할 원가 코드:",
    "🗑️ 원가 코드 삭제",
    "원가 코드 목록이 비어 있습니다.",
    "삭제할 수 없습니다: 아직 **",
    "**개의 계정이 이 작업장에 속해 있습니다. 먼저 👥 역할 및 권한에서 다른 작업장으로 옮기세요.",
    "** 을(를) 목록에서 삭제했습니다. 기존 작업은 유지됩니다.",
    "** 은(는) 이미 목록에 있습니다!",
    "코드 **",
    "추가했습니다 **",
    "삭제했습니다 **",
    "설명:",
    "작성자:",
    "작성일:",
    "담당자:",
    "비고:",
    "— 없음 —",
    "📌 이 작업은 내 팀에 속합니다. 📋 작업 배정에서 작업자에게 넘기세요.",
    "💾 새 작업 저장",
    "필수 항목을 입력하세요: Task ID와 Task Name!",
    "계획 완료일은 계획 시작일보다 빠를 수 없습니다!",
    "** 은(는) 이미 존재합니다(휴지통에 있을 수 있음). 다른 ID를 사용하세요!",
    "작업을 저장했습니다 **",
    "⚠️ 아직 WS 원가 코드가 없습니다: Foreman / WOS Manager가 위의 ⚙️ 에서 먼저 추가해야 합니다!",
    "👑 작업을 골라 **Team Leader** 또는 **Foreman**에게 배정하세요. 그들이 팀의 작업자에게 나눠 줍니다.",
    "아직 작업이 없습니다. ➕ 작업 추가에서 만드세요.",
    "아직 Team Leader / Foreman이 없습니다. 먼저 👥 역할 및 권한에서 역할을 배정하세요.",
    "미배정 작업만 보기",
    "모든 작업이 배정되었습니다!",
    "작업 선택:",
    "Team Leader / Foreman에게 배정:",
    "작업자에게 배정:",
    "— 미배정 —",
    "미배정",
    "💾 배정 저장",
    "작업을 배정했습니다 **",
    "팀에 작업자가 없습니다. 👥 내 팀에서 팀을 만들고 작업자를 추가하세요.",
    "아직 배정된 작업이 없습니다.",
    "작업을 배정하려면 팀에 작업자가 있어야 합니다.",
    "(이름 없음)",
    "⚠️ 계정에 작업장이 없습니다. 팀을 만들기 전에 WOS Manager에게 작업장 배정을 요청하세요.",
    "아직 팀이 없습니다. 이름을 정해 팀을 만드세요.",
    "팀 이름 *",
    "예: 170블록 용접팀",
    "➕ 팀 만들기",
    "팀 이름을 입력하세요!",
    "팀 이름은 최대 60자입니다!",
    "팀을 만들었습니다 **",
    "✏️ 팀 이름 변경",
    "새 팀 이름:",
    "💾 이름 저장",
    "팀 이름이 올바르지 않습니다 (1–60자)!",
    "팀 이름을 변경했습니다!",
    "🗑️ 팀 삭제",
    "팀을 삭제하면 모든 작업자가 팀에서 나가고 그들의 작업은 미배정이 됩니다. 작업은 그대로 내게 남으며 나중에 새 팀을 만들 수 있습니다.",
    "이 팀을 삭제하겠습니다",
    "🗑️ 팀 삭제",
    "팀을 삭제했습니다.",
    "팀에 아직 작업자가 없습니다.",
    "작업장 ",
    "에는 팀이 없는 작업자가 더 이상 없습니다.",
    "작업자 선택 (같은 작업장, 팀 없음):",
    "➕ 팀에 추가",
    "작업자를 1명 이상 선택하세요!",
    "명의 작업자를 팀에 추가했습니다!",
    "",
    "제외할 팀원이 없습니다.",
    "작업자 선택:",
    "➖ 팀에서 제외",
    "작업자를 팀에서 제외했습니다. 그 작업자의 팀 작업은 미배정이 되었습니다.",
    "아직 팀에 속해 있지 않습니다. 같은 작업장의 Team Leader / Foreman이 추가할 수 있습니다.",
    "(이름 없는 팀)",
    "**팀장:** ",
    "팀을 나가면 나에게 배정된 팀 작업은 미배정이 됩니다.",
    "이 팀에서 나가겠습니다",
    "🚪 팀 나가기",
    "팀에서 나갔습니다.",
    "🧰 작업을 휴지통으로",
    "👤 사용자 계정 잠금",
    "편집하거나 삭제할 작업이 없습니다.",
    "휴지통으로 옮길 작업:",
    "🗑️ 작업을 휴지통으로 옮기기",
    "작업을 휴지통으로 옮겼습니다!",
    "삭제할 수 있는 다른 계정이 없습니다.",
    "잠글 계정:",
    "🗑️ 이 계정 잠그기",
    "계정을 잠그고 휴지통으로 옮겼습니다!",
    "🛡️ Admin은 **WOS Manager** 역할을 부여하고 담당 **작업장**만 선택합니다. Worker / Team Leader / Foreman 역할은 WOS Manager가 부여합니다.",
    "👑 **Worker / Team Leader / Foreman** 역할 부여, **작업장** 선택, Team Leader / Foreman의 **팀**에 작업자 배치를 할 수 있습니다.",
    "역할을 배정할 계정이 없습니다.",
    "변경할 계정:",
    "작업:",
    "👑 WOS Manager로 지정",
    "⛔ 역할 회수 (대기 상태로)",
    "이 매니저가 담당하는 작업장:",
    "⚠️ 아직 작업장이 없습니다. ➕ 작업 추가 → ⚙️ 추가 / 삭제에서 추가하세요.",
    "💾 역할 저장",
    "💾 변경 사항 저장",
    "**WOS Manager**로 지정했습니다. 작업장 **",
    "역할을 회수했고 계정이 대기 상태로 돌아갔습니다.",
    "새 역할:",
    "소속 팀 (Team Leader / Foreman):",
    "— 팀 없음 —",
    "이 작업장에는 아직 Team Leader / Foreman이 없습니다.",
    "업데이트됨: **",
    "이름",
    "역할",
    "소속 팀",
    "🧰 작업 휴지통",
    "👤 계정 휴지통",
    "작업 휴지통이 비어 있습니다.",
    "계정 휴지통이 비어 있습니다.",
    "작업 선택:",
    "♻️ 작업 복원",
    "💥 작업 영구 삭제",
    "작업을 복원했습니다!",
    "작업을 영구 삭제했습니다!",
    "계정 선택:",
    "♻️ 계정 잠금 해제 / 복원",
    "💥 계정 영구 삭제",
    "계정을 복원했습니다!",
    "계정을 데이터베이스에서 영구 삭제했습니다!",
    "보고서를 만들 작업 데이터가 없습니다. 먼저 작업을 추가하세요!",
    "전체 작업 수",
    "전체 작업 수",
    "평균 진행률 (%)",
    "평균 진행률",
    "완료 (100%)",
    "진행 중",
    "작업 수",
    "상태",
    "📥 작업 보고서 다운로드 (CSV/Excel)",
    "현재 비밀번호:",
    "새 비밀번호 (8자 이상):",
    "새 비밀번호 다시 입력:",
    "💾 비밀번호 변경",
    "현재 비밀번호가 틀렸습니다!",
    "새 비밀번호는 8자 이상이어야 합니다!",
    "기본 비밀번호는 다시 사용할 수 없습니다!",
    "✅ 비밀번호를 변경했습니다! 다른 기기는 로그아웃되었습니다.",
    "⚠️ 이 계정은 아직 기본 비밀번호 'admin123'을 사용 중입니다. 지금 바로 🔑 비밀번호 변경에서 변경하세요!",
    "시작 전 (0%)",
    "진행 중 (1-50%)",
    "거의 완료 (51-99%)",
    "완료 (100%)",
    "개수",
    "새 비밀번호:",
    "💬 메시지",
    "📢 전체 채팅",
    "#### 👥 사람들",
    "🔎 계정 찾기",
    "계정을 찾을 수 없습니다.",
    "앱의 모든 사람이 이 채널을 볼 수 있습니다.",
    "개인 채팅: ",
    "메시지를 입력하세요...",
    "➕ 그룹 채팅 만들기",
    "그룹 이름",
    "예: 170블록 용접조",
    "그룹에 사람 추가",
    "✨ 그룹 만들기",
    "그룹 이름을 입력하세요!",
    "1명 이상 선택하세요!",
    "그룹을 만들었습니다 **",
    "#### 💬 그룹 채팅",
    "⚙️ 그룹 멤버",
    "➕ 그룹에 추가",
    "명을 그룹에 추가했습니다!",
    "그룹에서 제외",
    "➖ 그룹에서 제외",
    "그룹에서 제외했습니다.",
    "🗑️ 그룹 삭제",
    "그룹을 만든 사람만 멤버를 추가 / 제외할 수 있습니다.",
    "🚪 그룹 나가기",
    "🎥 온라인 회의",
    "### 🔑 회의 참가",
    "회의실 ID",
    "회의실 비밀번호",
    "🎥 참가",
    "잘못 입력한 횟수가 너무 많습니다. 페이지를 새로고침하고 나중에 다시 시도하세요.",
    "회의실 ID 또는 비밀번호가 틀렸거나 회의가 종료되었습니다!",
    "✅ 회의를 찾았습니다: **",
    "호스트",
    "🎥 회의 열기 (카메라 + 마이크)",
    "회의는 새 탭에서 열립니다. 처음에는 브라우저가 카메라와 마이크 사용 권한을 요청하니 허용을 누르세요.",
    "### ➕ 새 회의 만들기",
    "회의 이름 *",
    "예: 아침 교대 회의",
    "날짜",
    "시간",
    "사람 초대 (💬 메시지로 ID와 비밀번호 전송)",
    "📢 전체 채팅에 알리기",
    "🎥 회의 만들기",
    "회의 이름을 입력하세요!",
    "🎥 회의 초대",
    "🕒 시간",
    "🔢 회의실 ID",
    "🔑 비밀번호",
    "👉 🎥 온라인 회의에서 참가하세요.",
    "🎉 회의를 만들었습니다 **",
    "### 📋 내 회의",
    "아직 회의가 없습니다.",
    "▶️ 시작",
    "⛔ 종료",
    "Foreman과 Team Leader만 회의를 만들 수 있습니다. 회의실 ID와 비밀번호가 있으면 참가할 수 있습니다.",
    "사람 선택...",
    "👑 그룹 관리자",
    "🗑️ 메시지 삭제 (관리자)",
    "삭제할 메시지가 없습니다.",
    "삭제할 메시지",
    "🗑️ 메시지 삭제",
    "메시지를 삭제했습니다.",
    "모든 메시지를 삭제하겠습니다",
    "🧹 모든 메시지 삭제",
    "모든 메시지를 삭제했습니다.",
    "사람 초대 (🎥 온라인 회의에서 초대를 보게 됩니다)",
    "### 📨 나의 회의 초대",
    "🎥 회의 참가",
    "💬 메시지 및 회의",
    "설정",
    "소개",
    "로그인 / 회원가입",
    "조선소를 위한 작업 관리 소프트웨어",
    "ShipControl은 조선소 팀이 한곳에서 작업을 배정하고 진행 상황을 추적하며 소통할 수 있게 해 줍니다. 컴퓨터와 휴대폰 모두에서 사용할 수 있습니다.",
    "📋 작업 관리",
    "블록·데크·프레임별로 작업을 만들고 진행률을 퍼센트 단위로 추적합니다.",
    "👥 팀 및 권한",
    "Admin, WOS Manager, Foreman, Team Leader, Worker — 각자 자기 일만 정확히 봅니다.",
    "💬 메시지 및 온라인 회의",
    "전체 채팅, 그룹 채팅, ID와 비밀번호가 있는 화상 회의실.",
    "🌐 15개 언어 및 음성 입력",
    "베트남어, 영어, 중국어, 일본어, 한국어, 태국어, 인도네시아어, 스페인어, 프랑스어, 독일어, 러시아어 등 — 휴대폰에서는 타이핑 대신 말로 입력할 수 있습니다.",
    "계정이 없나요? 가입한 뒤 관리자가 역할을 배정할 때까지 기다리세요.",
    "🎬 조선소에서 바다로",
    "첫 강철 블록부터 진수까지 배가 만들어지는 과정을 보세요.",
    "영상이 재생되지 않나요? ",
    "▶ YouTube에서 열기",
    "그룹 채팅 만들기",
    "아직 그룹이 없습니다 — ➕ 를 눌러 만드세요.",
    "추가",
    "삭제",
    "작업자 추가",
    "팀에서 제외",
    "작업장 추가",
    "작업장 삭제",
    "그룹 삭제",
    "그룹 나가기",
    "🌊 바다",
    "🌅 노을",
    "🌲 숲",
    "⚪ 미니멀",
    "👑 로열",
    "📸 주간 사진 보고서",
    "📈 통계",
    "보고할 작업이 아직 없습니다.",
    "시작",
    "완료 예정",
    "진행률",
    "완료됨",
    "이번 주",
    "사진 올리기",
    "주차",
    "컴퓨터에서 사진 선택",
    "사진 메모 (선택)",
    "📤 업로드",
    "사진을 1장 이상 선택하세요!",
    "읽을 수 없는 사진 (또는 15 MB 초과): ",
    "이번 주 사진이 아직 없습니다.",
    "📥 사진 보고서 다운로드 (HTML)",
    "파일을 열고 Ctrl+P를 눌러 인쇄하거나 PDF로 저장하세요.",
    "💡 Ctrl(또는 Shift)를 누른 채 선택하면 여러 사진을 한 번에 고를 수 있습니다. Upload를 다시 누르면 더 추가됩니다.",
    "📊 PowerPoint 다운로드 (.pptx)",
    "◀ 이전",
    "다음 ▶",
    "전체 주 수",
    "전체 사진 수",
    "진행 중",
    "장",
    "요약",
]

# 🌏 10 ngôn ngữ thêm (cùng thứ tự với VI_EN)
_TH_LIST = [
    "🚢 SHIPCONTROL - ระบบจัดการงานต่อเรือ",
    "🔐 เข้าสู่ระบบ",
    "กรุณาเข้าสู่ระบบหรือสมัครสมาชิกเพื่อดำเนินการต่อ",
    "🧰 กระดานงาน",
    "➕ เพิ่มงาน",
    "📋 มอบหมายงาน",
    "👥 ทีมของฉัน",
    "✏️ แก้ไข / ลบ",
    "👥 บทบาทและสิทธิ์",
    "🗑️ ถังขยะ",
    "📊 รายงาน",
    "🔑 เปลี่ยนรหัสผ่าน",
    "🚪 ออกจากระบบ",
    "⚙️ การตั้งค่า",
    "📲 อุปกรณ์",
    "📲 คุณกำลังใช้:",
    "💻 คอมพิวเตอร์",
    "📱 โทรศัพท์",
    "🎤 แตะไมโครโฟนในช่องข้อความเพื่อพูดแทนการพิมพ์",
    "🎤 โหมดโทรศัพท์: แตะไมโครโฟนในช่องชื่อผู้ใช้เพื่อพูดแทนการพิมพ์",
    "เบราว์เซอร์นี้ยังไม่รองรับการป้อนด้วยเสียง กรุณาใช้ Chrome (Android) หรือ Safari (iPhone)",
    "การเข้าถึงไมโครโฟนถูกบล็อก กรุณาอนุญาตไมโครโฟนสำหรับเว็บนี้ในการตั้งค่าเบราว์เซอร์",
    "🏭 โรงงาน: ดู / เพิ่ม / ลบ",
    "ยังไม่มีโรงงาน เพิ่มโรงงานแรกด้านล่าง",
    "#### ➕ เพิ่มโรงงานใหม่",
    "รหัสโรงงาน (เช่น WOS_07):",
    "➕ เพิ่มโรงงาน",
    "รหัสไม่เกิน 20 ตัวอักษร ชื่อไม่เกิน 80 ตัวอักษร!",
    "#### 🗑️ ลบโรงงาน",
    "โรงงานที่จะลบ:",
    "🗑️ ลบโรงงาน",
    "ชื่อโรงงาน",
    "บัญชี",
    "งาน",
    "🎨 ธีม",
    "✨ ทันสมัย",
    "🚀 ล้ำอนาคต",
    "🎈 สนุกสนาน",
    "📜 ย้อนยุค",
    "ธีมล้ำอนาคตเป็นโหมดมืดเสมอ",
    ">บัญชี<",
    "🔒 บัญชีของคุณเพิ่งถูกเปิดในแท็บอื่นหรืออุปกรณ์อื่น จึงออกจากระบบที่นี่ ",
    "ถ้าไม่ใช่คุณ ให้เข้าสู่ระบบอีกครั้งและเปลี่ยนรหัสผ่านทันที",
    ">บทบาท<",
    ">เข้าสู่ระบบแล้ว<",
    "🌙 โหมดมืด",
    "🌐 ภาษา",
    "📋 กระดานความคืบหน้างาน",
    "➕ เพิ่มงานใหม่",
    "✏️ แก้ไขและลบ",
    "👥 บัญชีและบทบาท (รายการบทบาท)",
    "📊 รายงานความคืบหน้าและสถิติ",
    "🗑️ ถังขยะและการกู้คืน",
    "### ⚡ อัปเดตความคืบหน้าด่วน",
    "### 🏭 สถิติตามโรงงาน",
    "### 📈 การกระจายความคืบหน้า",
    "### 📌 มอบหมายงานให้คนงาน",
    "### 📥 ดาวน์โหลดข้อมูลรายงาน",
    "### 🔄 เปลี่ยนบทบาทบัญชี",
    "#### ➕ เพิ่มคนงาน",
    "#### ➖ นำออกจากทีม",
    "#### 👷 สมาชิก",
    "#### 🚪 ออกจากทีม",
    "🔑 เข้าสู่ระบบ",
    "📝 สมัครสมาชิก",
    "ชื่อผู้ใช้:",
    "รหัสผ่าน:",
    "🚀 เข้าสู่ระบบ",
    "✨ สมัครเลย",
    "ชื่อ-นามสกุล:",
    "ชื่อผู้ใช้ใหม่:",
    "ยืนยันรหัสผ่าน:",
    "กรุณากรอกชื่อผู้ใช้และรหัสผ่าน!",
    "ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง หรือบัญชีถูกล็อก/ลบแล้ว!",
    "⏳ บัญชีของคุณกำลังรอ Admin / WOS Manager กำหนดบทบาท กรุณากลับมาใหม่ภายหลัง!",
    "ยินดีต้อนรับ ",
    ") กลับมาแล้ว!",
    "กรุณากรอกข้อมูลให้ครบ!",
    "รหัสผ่านต้องมีอย่างน้อย 8 ตัวอักษร!",
    "ชื่อผู้ใช้หรือชื่อยาวเกินไป!",
    "รหัสผ่านไม่ตรงกัน!",
    "ชื่อผู้ใช้นี้มีอยู่แล้ว กรุณาเลือกชื่ออื่น!",
    "🎉 สมัครสำเร็จ! บัญชีของคุณกำลังรอ WOS Manager อนุมัติบทบาท",
    "👷 นี่คืองานที่มอบหมายให้คุณ",
    "แสดงเฉพาะงานของทีมฉัน",
    "ยังไม่มีงานให้แสดง",
    "คุณไม่มีงานที่ต้องอัปเดต",
    "เลือกงานที่จะอัปเดต:",
    "ความคืบหน้าใหม่ (%)",
    "บันทึกการทำงาน (Remark):",
    "🚀 อัปเดตความคืบหน้า",
    "อัปเดตความคืบหน้าแล้ว!",
    " (ปัจจุบัน: ",
    "⚙️ เพิ่ม / ลบ Block และ WS Cost Code",
    "Block ใหม่ (เช่น 170150):",
    "➕ เพิ่ม BLOCK",
    "กรุณากรอกชื่อ Block!",
    "เพิ่ม Block แล้ว **",
    "Block ที่จะลบ:",
    "🗑️ ลบ BLOCK",
    "ลบ Block แล้ว **",
    "รายการ Block ว่างเปล่า",
    "รหัส (เช่น WOS_07):",
    "ชื่อโรงงาน / แผนก:",
    "➕ เพิ่ม COST CODE",
    "กรุณากรอกทั้งรหัสและชื่อ!",
    "Cost Code ที่จะลบ:",
    "🗑️ ลบ COST CODE",
    "รายการ Cost Code ว่างเปล่า",
    "ลบไม่ได้: **",
    "** บัญชียังอยู่ในโรงงานนี้ ย้ายไปโรงงานอื่นใน 👥 บทบาทและสิทธิ์ก่อน",
    "** ออกจากรายการแล้ว งานเดิมยังคงอยู่",
    "** มีอยู่ในรายการแล้ว!",
    "รหัส **",
    "เพิ่มแล้ว **",
    "ลบแล้ว **",
    "คำอธิบาย:",
    "สร้างโดย:",
    "วันที่สร้าง:",
    "ผู้รับผิดชอบ:",
    "หมายเหตุ:",
    "— ไม่เลือก —",
    "📌 งานนี้จะเป็นของทีมคุณ ไปที่ 📋 มอบหมายงาน เพื่อมอบให้คนงาน",
    "💾 บันทึกงานใหม่",
    "กรุณากรอกช่องที่จำเป็น: Task ID และ Task Name!",
    "วันสิ้นสุดตามแผนต้องไม่ก่อนวันเริ่ม!",
    "** มีอยู่แล้ว (อาจอยู่ในถังขยะ) กรุณาใช้ ID อื่น!",
    "บันทึกงานแล้ว **",
    "⚠️ ยังไม่มี WS COST CODE: Foreman / WOS Manager ต้องเพิ่มด้านบนก่อนจึงจะเพิ่มงานได้!",
    "👑 เลือกงานแล้วมอบหมายให้ **Team Leader** หรือ **Foreman** พวกเขาจะส่งต่อให้คนงานในทีม",
    "ยังไม่มีงาน ไปที่ ➕ เพิ่มงาน เพื่อสร้างงาน",
    "ยังไม่มี Team Leader / Foreman กำหนดบทบาทใน 👥 บทบาทและสิทธิ์ก่อน",
    "แสดงเฉพาะงานที่ยังไม่มอบหมาย",
    "มอบหมายงานครบทุกงานแล้ว!",
    "เลือกงาน:",
    "มอบหมายให้ Team Leader / Foreman:",
    "มอบหมายให้คนงาน:",
    "— ยังไม่มอบหมาย —",
    "ยังไม่มอบหมาย",
    "💾 บันทึกการมอบหมาย",
    "มอบหมายงานให้ **",
    "ทีมของคุณยังไม่มีคนงาน ไปที่ 👥 ทีมของฉัน เพื่อสร้างทีมและเพิ่มคนงาน",
    "คุณยังไม่ได้รับงานใดๆ",
    "ต้องมีคนงานในทีมก่อนจึงจะมอบหมายงานได้",
    "(ไม่มีชื่อ)",
    "⚠️ บัญชีของคุณยังไม่มีโรงงาน ขอให้ WOS Manager กำหนดให้ก่อนสร้างทีม",
    "คุณยังไม่มีทีม ตั้งชื่อเพื่อสร้างทีม",
    "ชื่อทีม *",
    "เช่น ทีมเชื่อม Block 170",
    "➕ สร้างทีม",
    "กรุณากรอกชื่อทีม!",
    "ชื่อทีมยาวได้ไม่เกิน 60 ตัวอักษร!",
    "สร้างทีมแล้ว **",
    "✏️ เปลี่ยนชื่อทีม",
    "ชื่อทีมใหม่:",
    "💾 บันทึกชื่อ",
    "ชื่อทีมไม่ถูกต้อง (1–60 ตัวอักษร)!",
    "เปลี่ยนชื่อทีมแล้ว!",
    "🗑️ ลบทีม",
    "เมื่อลบทีม: คนงานทุกคนจะออกจากทีมและงานของพวกเขาจะกลับเป็นยังไม่มอบหมาย งานยังเป็นของคุณ และคุณสร้างทีมใหม่ได้ภายหลัง",
    "ฉันแน่ใจว่าต้องการลบทีมนี้",
    "🗑️ ลบทีม",
    "ลบทีมแล้ว",
    "ทีมยังไม่มีคนงาน",
    "ไม่มีคนงานเหลือในโรงงาน ",
    " ที่ยังไม่มีทีม",
    "เลือกคนงาน (โรงงานเดียวกัน ยังไม่มีทีม):",
    "➕ เพิ่มเข้าทีม",
    "เลือกคนงานอย่างน้อย 1 คน!",
    " คนเข้าทีมแล้ว!",
    "เพิ่ม ",
    "ไม่มีสมาชิกให้นำออก",
    "เลือกคนงาน:",
    "➖ นำออกจากทีม",
    "นำคนงานออกจากทีมแล้ว งานของทีมที่มอบให้คนนี้กลับเป็นยังไม่มอบหมาย",
    "คุณยังไม่อยู่ในทีมใด Team Leader / Foreman ในโรงงานของคุณสามารถเพิ่มคุณได้",
    "(ทีมไม่มีชื่อ)",
    "**หัวหน้าทีม:** ",
    "เมื่อคุณออก งานของทีมที่มอบให้คุณจะกลับเป็นยังไม่มอบหมาย",
    "ฉันแน่ใจว่าต้องการออกจากทีมนี้",
    "🚪 ออกจากทีม",
    "คุณออกจากทีมแล้ว",
    "🧰 ย้ายงานไปถังขยะ",
    "👤 ล็อกบัญชีผู้ใช้",
    "ไม่มีงานให้แก้ไขหรือลบ",
    "งานที่จะย้ายไปถังขยะ:",
    "🗑️ ย้ายงานไปถังขยะ",
    "ย้ายงานไปถังขยะแล้ว!",
    "ไม่มีบัญชีอื่นให้ลบ",
    "บัญชีที่จะล็อก:",
    "🗑️ ล็อกบัญชีนี้",
    "ล็อกบัญชีและย้ายไปถังขยะแล้ว!",
    "🛡️ Admin ให้ได้เฉพาะบทบาท **WOS Manager** และเลือก **โรงงาน** ที่ผู้จัดการนั้นดูแล บทบาท Worker / Team Leader / Foreman ให้โดย WOS Manager",
    "👑 คุณให้บทบาท **Worker / Team Leader / Foreman** เลือก **โรงงาน** และจัดคนงานเข้า **ทีม** ของ Team Leader / Foreman ได้",
    "ไม่มีบัญชีให้กำหนดบทบาท",
    "บัญชีที่จะเปลี่ยน:",
    "การดำเนินการ:",
    "👑 ตั้งเป็น WOS Manager",
    "⛔ ถอนบทบาท (กลับเป็น Pending)",
    "โรงงานที่ผู้จัดการนี้ดูแล:",
    "⚠️ ยังไม่มีโรงงาน เพิ่มได้ใน ➕ เพิ่มงาน → ⚙️ เพิ่ม / ลบ",
    "💾 บันทึกบทบาท",
    "💾 บันทึกการเปลี่ยนแปลง",
    "ตั้งเป็น **WOS Manager** ของโรงงาน **",
    "ถอนบทบาทแล้ว บัญชีกลับเป็น Pending",
    "บทบาทใหม่:",
    "ทีมของ (Team Leader / Foreman):",
    "— ไม่มีทีม —",
    "โรงงานนี้ยังไม่มี Team Leader / Foreman",
    "อัปเดตแล้ว: **",
    "ชื่อ-นามสกุล",
    "บทบาท",
    "ทีมของ",
    "🧰 ถังขยะงาน",
    "👤 ถังขยะบัญชี",
    "ถังขยะงานว่างเปล่า",
    "ถังขยะบัญชีว่างเปล่า",
    "เลือกงาน:",
    "♻️ กู้คืนงาน",
    "💥 ลบงานถาวร",
    "กู้คืนงานแล้ว!",
    "ลบงานถาวรแล้ว!",
    "เลือกบัญชี:",
    "♻️ ปลดล็อก / กู้คืนบัญชี",
    "💥 ลบบัญชีถาวร",
    "กู้คืนบัญชีแล้ว!",
    "ลบบัญชีออกจากฐานข้อมูลถาวรแล้ว!",
    "ยังไม่มีข้อมูลงานสำหรับรายงาน กรุณาเพิ่มงานก่อน!",
    "งานทั้งหมด",
    "งานทั้งหมด",
    "ความคืบหน้าเฉลี่ย (%)",
    "ความคืบหน้าเฉลี่ย",
    "เสร็จแล้ว (100%)",
    "กำลังทำ",
    "จำนวนงาน",
    "สถานะ",
    "📥 ดาวน์โหลดรายงานงาน (CSV/Excel)",
    "รหัสผ่านปัจจุบัน:",
    "รหัสผ่านใหม่ (อย่างน้อย 8 ตัวอักษร):",
    "ยืนยันรหัสผ่านใหม่:",
    "💾 เปลี่ยนรหัสผ่าน",
    "รหัสผ่านปัจจุบันไม่ถูกต้อง!",
    "รหัสผ่านใหม่ต้องมีอย่างน้อย 8 ตัวอักษร!",
    "ห้ามใช้รหัสผ่านเริ่มต้นซ้ำ!",
    "✅ เปลี่ยนรหัสผ่านแล้ว! อุปกรณ์อื่นถูกออกจากระบบแล้ว",
    "⚠️ บัญชีนี้ยังใช้รหัสผ่านเริ่มต้น 'admin123' ไปที่ 🔑 เปลี่ยนรหัสผ่าน แล้วเปลี่ยนเดี๋ยวนี้!",
    "ยังไม่เริ่ม (0%)",
    "กำลังทำ (1-50%)",
    "ใกล้เสร็จ (51-99%)",
    "เสร็จแล้ว (100%)",
    "จำนวน",
    "รหัสผ่านใหม่:",
    "💬 ข้อความ",
    "📢 แชทรวม",
    "#### 👥 ผู้คน",
    "🔎 ค้นหาบัญชี",
    "ไม่พบบัญชี",
    "ทุกคนในแอปเห็นช่องนี้",
    "แชทส่วนตัวกับ ",
    "พิมพ์ข้อความ...",
    "➕ สร้างแชทกลุ่ม",
    "ชื่อกลุ่ม",
    "เช่น ทีมเชื่อม Block 170",
    "เพิ่มคนเข้ากลุ่ม",
    "✨ สร้างกลุ่ม",
    "กรุณาตั้งชื่อกลุ่ม!",
    "เลือกอย่างน้อย 1 คน!",
    "สร้างกลุ่มแล้ว **",
    "#### 💬 แชทกลุ่ม",
    "⚙️ สมาชิกกลุ่ม",
    "➕ เพิ่มเข้ากลุ่ม",
    " คนเข้ากลุ่มแล้ว!",
    "นำออกจากกลุ่ม",
    "➖ นำออกจากกลุ่ม",
    "นำออกจากกลุ่มแล้ว",
    "🗑️ ลบกลุ่ม",
    "เฉพาะผู้สร้างกลุ่มเท่านั้นที่เพิ่ม / นำสมาชิกออกได้",
    "🚪 ออกจากกลุ่ม",
    "🎥 ประชุมออนไลน์",
    "### 🔑 เข้าร่วมการประชุม",
    "รหัสห้องประชุม",
    "รหัสผ่านห้อง",
    "🎥 เข้าร่วม",
    "ลองผิดหลายครั้งเกินไป โหลดหน้าใหม่แล้วลองอีกครั้งภายหลัง",
    "รหัสห้องหรือรหัสผ่านไม่ถูกต้อง หรือการประชุมจบแล้ว!",
    "✅ พบห้อง: **",
    "ผู้จัด",
    "🎥 เปิดการประชุม (กล้อง + ไมค์)",
    "การประชุมจะเปิดในแท็บใหม่ ครั้งแรกเบราว์เซอร์จะขอใช้กล้องและไมโครโฟน — แตะอนุญาต",
    "### ➕ สร้างการประชุมใหม่",
    "ชื่อการประชุม *",
    "เช่น ส่งกะเช้า",
    "วันที่",
    "เวลา",
    "เชิญคน (ส่ง ID + รหัสผ่านทาง 💬 ข้อความ)",
    "📢 ประกาศในแชทรวม",
    "🎥 สร้างการประชุม",
    "กรุณากรอกชื่อการประชุม!",
    "🎥 คำเชิญประชุม",
    "🕒 เวลา",
    "🔢 รหัสห้อง",
    "🔑 รหัสผ่าน",
    "👉 ไปที่ 🎥 ประชุมออนไลน์ เพื่อเข้าร่วม",
    "🎉 สร้างการประชุมแล้ว **",
    "### 📋 การประชุมของฉัน",
    "คุณยังไม่มีการประชุม",
    "▶️ เริ่ม",
    "⛔ จบ",
    "เฉพาะ Foreman และ Team Leader เท่านั้นที่สร้างการประชุมได้ คุณเข้าร่วมได้ด้วยรหัสห้องและรหัสผ่าน",
    "เลือกคน...",
    "👑 ผู้ดูแลกลุ่ม",
    "🗑️ ลบข้อความ (ผู้ดูแล)",
    "ไม่มีข้อความให้ลบ",
    "ข้อความที่จะลบ",
    "🗑️ ลบข้อความ",
    "ลบข้อความแล้ว",
    "ฉันแน่ใจว่าต้องการลบข้อความทั้งหมด",
    "🧹 ลบข้อความทั้งหมด",
    "ลบข้อความทั้งหมดแล้ว",
    "เชิญคน (พวกเขาจะเห็นคำเชิญใน 🎥 ประชุมออนไลน์)",
    "### 📨 คำเชิญประชุมของคุณ",
    "🎥 เข้าร่วมประชุม",
    "💬 ข้อความและการประชุม",
    "การตั้งค่า",
    "เกี่ยวกับเรา",
    "เข้าสู่ระบบ / สมัครสมาชิก",
    "ซอฟต์แวร์จัดการงานสำหรับอู่ต่อเรือ",
    "ShipControl ช่วยทีมในอู่ต่อเรือมอบหมายงาน ติดตามความคืบหน้า และติดต่อกันได้ในที่เดียว — ทั้งบนคอมพิวเตอร์และโทรศัพท์",
    "📋 การจัดการงาน",
    "สร้างงานตาม Block, Deck และ Frame แล้วติดตามความคืบหน้าทีละเปอร์เซ็นต์",
    "👥 ทีมและบทบาท",
    "Admin, WOS Manager, Foreman, Team Leader และ Worker — ทุกคนเห็นเฉพาะงานของตัวเอง",
    "💬 ข้อความและประชุมออนไลน์",
    "แชทรวม แชทกลุ่ม และห้องประชุมวิดีโอด้วย ID + รหัสผ่าน",
    "🌐 15 ภาษาและเสียง",
    "เวียดนาม อังกฤษ จีน ญี่ปุ่น เกาหลี ไทย อินโดนีเซีย สเปน ฝรั่งเศส เยอรมัน รัสเซีย และอื่นๆ — และพูดแทนการพิมพ์บนโทรศัพท์",
    "ยังไม่มีบัญชี? สมัครแล้วรอผู้จัดการกำหนดบทบาทให้",
    "🎬 จากอู่ต่อเรือสู่ทะเล",
    "ชมการต่อเรือตั้งแต่บล็อกเหล็กชิ้นแรกจนถึงวันปล่อยเรือลงน้ำ",
    "เล่นวิดีโอไม่ได้? ",
    "▶ เปิดบน YouTube",
    "สร้างแชทกลุ่ม",
    "ยังไม่มีกลุ่ม — แตะ ➕ เพื่อสร้าง",
    "เพิ่ม",
    "ลบ",
    "เพิ่มคนงาน",
    "นำออกจากทีม",
    "เพิ่มโรงงาน",
    "ลบโรงงาน",
    "ลบกลุ่ม",
    "ออกจากกลุ่ม",
    "🌊 มหาสมุทร",
    "🌅 พระอาทิตย์ตก",
    "🌲 ป่าไม้",
    "⚪ มินิมอล",
    "👑 ราชวงศ์",
    "📸 รายงานภาพถ่ายรายสัปดาห์",
    "📈 สถิติ",
    "คุณยังไม่มีงานที่จะรายงาน",
    "เริ่ม",
    "คาดว่าจะเสร็จ",
    "ความคืบหน้า",
    "เสร็จแล้ว",
    "สัปดาห์ปัจจุบัน",
    "อัปโหลดรูป",
    "สัปดาห์",
    "เลือกรูปจากคอมพิวเตอร์",
    "หมายเหตุรูป (ไม่บังคับ)",
    "📤 อัปโหลด",
    "เลือกรูปอย่างน้อย 1 รูป!",
    "อ่านรูปเหล่านี้ไม่ได้ (หรือใหญ่เกิน 15 MB): ",
    "ยังไม่มีรูปสำหรับสัปดาห์นี้",
    "📥 ดาวน์โหลดรายงานภาพ (HTML)",
    "เปิดไฟล์แล้วกด Ctrl+P เพื่อพิมพ์หรือบันทึกเป็น PDF",
    "💡 กด Ctrl (หรือ Shift) ค้างไว้ขณะเลือกเพื่อเลือกหลายรูปพร้อมกัน กด Upload อีกครั้งเพื่อเพิ่มรูป",
    "📊 ดาวน์โหลด PowerPoint (.pptx)",
    "◀ ก่อนหน้า",
    "ถัดไป ▶",
    "จำนวนสัปดาห์ทั้งหมด",
    "จำนวนรูปทั้งหมด",
    "กำลังดำเนินการ",
    "รูป",
    "สรุป",
]
_ID_LIST = [
    "🚢 SHIPCONTROL - MANAJEMEN PEKERJAAN KAPAL",
    "🔐 Masuk",
    "Silakan masuk atau daftar untuk melanjutkan.",
    "🧰 Papan Tugas",
    "➕ Tambah Tugas",
    "📋 Bagikan Tugas",
    "👥 Tim Saya",
    "✏️ Ubah / Hapus",
    "👥 Peran & Izin",
    "🗑️ Sampah",
    "📊 Laporan",
    "🔑 Ganti Kata Sandi",
    "🚪 Keluar",
    "⚙️ Pengaturan",
    "📲 Perangkat",
    "📲 Anda menggunakan:",
    "💻 Komputer",
    "📱 Ponsel",
    "🎤 Ketuk mikrofon di kotak teks untuk berbicara alih-alih mengetik.",
    "🎤 Mode ponsel: ketuk mikrofon di kotak Nama pengguna untuk berbicara alih-alih mengetik.",
    "Browser ini belum mendukung input suara. Gunakan Chrome (Android) atau Safari (iPhone).",
    "Akses mikrofon diblokir. Izinkan mikrofon untuk situs ini di pengaturan browser.",
    "🏭 Bengkel: lihat / tambah / hapus",
    "Belum ada bengkel. Tambahkan yang pertama di bawah.",
    "#### ➕ Tambah Bengkel Baru",
    "Kode bengkel (mis. WOS_07):",
    "➕ TAMBAH BENGKEL",
    "Kode maks. 20 karakter, nama maks. 80 karakter!",
    "#### 🗑️ Hapus Bengkel",
    "Bengkel yang dihapus:",
    "🗑️ HAPUS BENGKEL",
    "Nama bengkel",
    "Akun",
    "Tugas",
    "🎨 Tema",
    "✨ Modern",
    "🚀 Futuristik",
    "🎈 Ceria",
    "📜 Klasik",
    "Tema Futuristik selalu gelap.",
    ">Akun<",
    "🔒 Akun Anda baru saja dibuka di tab atau perangkat lain, jadi Anda dikeluarkan di sini. ",
    "Jika itu bukan Anda, masuk lagi dan segera ganti kata sandi.",
    ">Peran<",
    ">Sedang masuk<",
    "🌙 Mode gelap",
    "🌐 Bahasa",
    "📋 Papan Progres Tugas",
    "➕ Tambah Tugas Baru",
    "✏️ Ubah & Hapus",
    "👥 Akun & Peran (Daftar Peran)",
    "📊 Laporan Progres & Statistik",
    "🗑️ Sampah & Pulihkan",
    "### ⚡ Perbarui Progres Cepat",
    "### 🏭 Statistik per Bengkel",
    "### 📈 Sebaran Progres Tugas",
    "### 📌 Bagikan Tugas ke Pekerja",
    "### 📥 Unduh Data Laporan",
    "### 🔄 Ubah Peran Akun",
    "#### ➕ Tambah Pekerja",
    "#### ➖ Keluarkan dari Tim",
    "#### 👷 Anggota",
    "#### 🚪 Keluar dari Tim",
    "🔑 Masuk",
    "📝 Daftar",
    "Nama pengguna:",
    "Kata sandi:",
    "🚀 MASUK",
    "✨ DAFTAR SEKARANG",
    "Nama lengkap:",
    "Nama pengguna baru:",
    "Konfirmasi kata sandi:",
    "Masukkan nama pengguna dan kata sandi!",
    "Nama pengguna atau kata sandi salah, atau akun terkunci/dihapus!",
    "⏳ Akun Anda menunggu Admin / WOS Manager memberi peran. Silakan kembali nanti!",
    "Selamat datang ",
    ") sudah kembali!",
    "Isi semua kolom!",
    "Kata sandi minimal 8 karakter!",
    "Nama pengguna atau nama lengkap terlalu panjang!",
    "Kata sandi tidak cocok!",
    "Nama pengguna ini sudah ada, pilih yang lain!",
    "🎉 Terdaftar! Akun Anda menunggu WOS Manager menyetujui peran.",
    "👷 Ini tugas yang diberikan kepada Anda.",
    "Tampilkan hanya tugas tim saya",
    "Belum ada tugas untuk ditampilkan.",
    "Anda tidak punya tugas untuk diperbarui.",
    "Pilih tugas untuk diperbarui:",
    "Progres baru (%)",
    "Catatan kerja (Remark):",
    "🚀 PERBARUI PROGRES",
    "Progres diperbarui!",
    " (Saat ini: ",
    "⚙️ Tambah / Hapus Block dan WS Cost Code",
    "Block baru (mis. 170150):",
    "➕ TAMBAH BLOCK",
    "Masukkan nama block!",
    "Block ditambahkan **",
    "Block yang dihapus:",
    "🗑️ HAPUS BLOCK",
    "Block dihapus **",
    "Daftar block kosong.",
    "Kode (mis. WOS_07):",
    "Nama bengkel / departemen:",
    "➕ TAMBAH COST CODE",
    "Masukkan kode dan nama!",
    "Cost code yang dihapus:",
    "🗑️ HAPUS COST CODE",
    "Daftar cost code kosong.",
    "Tidak bisa dihapus: **",
    "** akun masih ada di bengkel ini. Pindahkan dulu ke bengkel lain di 👥 Peran & Izin.",
    "** dari daftar. Tugas lama tetap disimpan.",
    "** sudah ada di daftar!",
    "Kode **",
    "Ditambahkan **",
    "Dihapus **",
    "Deskripsi:",
    "Dibuat oleh:",
    "Tanggal dibuat:",
    "Penanggung jawab:",
    "Catatan:",
    "— Tidak ada —",
    "📌 Tugas ini akan milik tim Anda. Buka 📋 Bagikan Tugas untuk memberikannya ke pekerja.",
    "💾 SIMPAN TUGAS BARU",
    "Isi kolom wajib: Task ID dan Task Name!",
    "Tanggal selesai rencana tidak boleh sebelum tanggal mulai!",
    "** sudah ada (mungkin di Sampah). Gunakan ID lain!",
    "Tugas disimpan **",
    "⚠️ BELUM ADA WS COST CODE: Foreman / WOS Manager harus menambahkannya di atas sebelum menambah tugas!",
    "👑 Pilih tugas dan berikan ke **Team Leader** atau **Foreman**. Mereka akan meneruskannya ke pekerja di timnya.",
    "Belum ada tugas. Buka ➕ Tambah Tugas untuk membuatnya.",
    "Belum ada Team Leader / Foreman. Berikan peran di 👥 Peran & Izin dulu.",
    "Tampilkan hanya tugas yang belum dibagikan",
    "Semua tugas sudah dibagikan!",
    "Pilih tugas:",
    "Berikan ke Team Leader / Foreman:",
    "Berikan ke pekerja:",
    "— Belum dibagikan —",
    "Belum dibagikan",
    "💾 SIMPAN PEMBAGIAN",
    "Tugas diberikan ke **",
    "Tim Anda belum punya pekerja. Buka 👥 Tim Saya untuk membuat tim dan menambah pekerja.",
    "Anda belum diberi tugas apa pun.",
    "Tim Anda harus punya pekerja sebelum membagikan tugas.",
    "(tanpa nama)",
    "⚠️ Akun Anda belum punya bengkel. Minta WOS Manager menetapkannya sebelum membuat tim.",
    "Anda belum punya tim. Beri nama untuk membuatnya.",
    "Nama tim *",
    "mis. Tim Las Block 170",
    "➕ BUAT TIM",
    "Masukkan nama tim!",
    "Nama tim maksimal 60 karakter!",
    "Tim dibuat **",
    "✏️ Ganti nama tim",
    "Nama tim baru:",
    "💾 SIMPAN NAMA",
    "Nama tim tidak valid (1–60 karakter)!",
    "Nama tim diganti!",
    "🗑️ Hapus tim",
    "Menghapus tim: semua pekerja keluar dari tim dan tugas mereka menjadi belum dibagikan. Tugas tetap milik Anda, dan Anda bisa membuat tim baru nanti.",
    "Saya yakin ingin menghapus tim ini",
    "🗑️ HAPUS TIM",
    "Tim dihapus.",
    "Tim belum punya pekerja.",
    "Tidak ada lagi pekerja di bengkel ",
    " yang belum punya tim.",
    "Pilih pekerja (bengkel sama, belum punya tim):",
    "➕ TAMBAH KE TIM",
    "Pilih minimal 1 pekerja!",
    " pekerja ke tim!",
    "Menambahkan ",
    "Tidak ada anggota untuk dikeluarkan.",
    "Pilih pekerja:",
    "➖ KELUARKAN DARI TIM",
    "Pekerja dikeluarkan dari tim. Tugas timnya kini belum dibagikan.",
    "Anda belum masuk tim. Team Leader / Foreman di bengkel Anda dapat menambahkan Anda.",
    "(Tim tanpa nama)",
    "**Ketua tim:** ",
    "Jika Anda keluar, tugas tim yang diberikan kepada Anda menjadi belum dibagikan.",
    "Saya yakin ingin keluar dari tim ini",
    "🚪 KELUAR DARI TIM",
    "Anda sudah keluar dari tim.",
    "🧰 Pindahkan Tugas ke Sampah",
    "👤 Kunci Akun Pengguna",
    "Tidak ada tugas untuk diubah atau dihapus.",
    "Tugas yang dipindah ke Sampah:",
    "🗑️ Pindahkan Tugas ke Sampah",
    "Tugas dipindahkan ke Sampah!",
    "Tidak ada akun lain untuk dihapus.",
    "Akun yang dikunci:",
    "🗑️ Kunci Akun Ini",
    "Akun dikunci dan dipindahkan ke Sampah!",
    "🛡️ Admin hanya memberi peran **WOS Manager** dan memilih **Bengkel** yang dipimpinnya. Peran Worker / Team Leader / Foreman diberikan oleh WOS Manager.",
    "👑 Anda bisa memberi peran **Worker / Team Leader / Foreman**, memilih **Bengkel**, dan memasukkan pekerja ke **tim** seorang Team Leader / Foreman.",
    "Tidak ada akun untuk diberi peran.",
    "Akun yang diubah:",
    "Tindakan:",
    "👑 Jadikan WOS Manager",
    "⛔ Cabut peran (kembali ke Pending)",
    "Bengkel yang dipimpin manajer ini:",
    "⚠️ Belum ada bengkel. Tambahkan di ➕ Tambah Tugas → ⚙️ Tambah / Hapus.",
    "💾 SIMPAN PERAN",
    "💾 SIMPAN PERUBAHAN",
    "Dijadikan **WOS Manager** bengkel **",
    "Peran dicabut, akun kembali ke Pending.",
    "Peran baru:",
    "Tim milik (Team Leader / Foreman):",
    "— Tanpa tim —",
    "Bengkel ini belum punya Team Leader / Foreman.",
    "Diperbarui: **",
    "Nama lengkap",
    "Peran",
    "Tim milik",
    "🧰 Sampah Tugas",
    "👤 Sampah Akun",
    "Sampah tugas kosong.",
    "Sampah akun kosong.",
    "Pilih tugas:",
    "♻️ PULIHKAN TUGAS",
    "💥 HAPUS TUGAS PERMANEN",
    "Tugas dipulihkan!",
    "Tugas dihapus permanen!",
    "Pilih akun:",
    "♻️ BUKA KUNCI / PULIHKAN AKUN",
    "💥 HAPUS AKUN PERMANEN",
    "Akun dipulihkan!",
    "Akun dihapus permanen dari database!",
    "Belum ada data tugas untuk laporan. Tambahkan tugas dulu!",
    "Total tugas",
    "Total tugas",
    "Rata-rata progres (%)",
    "Rata-rata progres",
    "Selesai (100%)",
    "Sedang berjalan",
    "Jumlah tugas",
    "Status",
    "📥 Unduh Laporan Tugas (CSV/Excel)",
    "Kata sandi saat ini:",
    "Kata sandi baru (minimal 8 karakter):",
    "Ulangi kata sandi baru:",
    "💾 GANTI KATA SANDI",
    "Kata sandi saat ini salah!",
    "Kata sandi baru minimal 8 karakter!",
    "Anda tidak boleh memakai lagi kata sandi bawaan!",
    "✅ Kata sandi diganti! Perangkat lain sudah dikeluarkan.",
    "⚠️ Akun ini masih memakai kata sandi bawaan 'admin123'. Buka 🔑 Ganti Kata Sandi dan ganti SEKARANG!",
    "Belum dimulai (0%)",
    "Sedang berjalan (1-50%)",
    "Hampir selesai (51-99%)",
    "Selesai (100%)",
    "Jumlah",
    "Kata sandi baru:",
    "💬 Pesan",
    "📢 Obrolan umum",
    "#### 👥 Orang",
    "🔎 Cari akun",
    "Akun tidak ditemukan.",
    "Semua orang di aplikasi bisa melihat kanal ini.",
    "Obrolan pribadi dengan ",
    "Ketik pesan...",
    "➕ Buat obrolan grup",
    "Nama grup",
    "mis. Kru las Block 170",
    "Tambahkan orang ke grup",
    "✨ BUAT GRUP",
    "Beri nama grup!",
    "Pilih minimal 1 orang!",
    "Grup dibuat **",
    "#### 💬 Obrolan grup",
    "⚙️ Anggota grup",
    "➕ TAMBAH KE GRUP",
    " orang ke grup!",
    "Keluarkan dari grup",
    "➖ KELUARKAN DARI GRUP",
    "Dikeluarkan dari grup.",
    "🗑️ HAPUS GRUP",
    "Hanya pembuat grup yang bisa menambah / mengeluarkan anggota.",
    "🚪 KELUAR DARI GRUP",
    "🎥 Rapat Online",
    "### 🔑 Gabung rapat",
    "ID ruang rapat",
    "Kata sandi ruang",
    "🎥 GABUNG",
    "Terlalu banyak percobaan salah. Muat ulang halaman dan coba lagi nanti.",
    "ID ruang atau kata sandi salah, atau rapat sudah selesai!",
    "✅ Ruang ditemukan: **",
    "Tuan rumah",
    "🎥 BUKA RAPAT (kamera + mikrofon)",
    "Rapat terbuka di tab baru. Pertama kali, browser akan meminta izin kamera dan mikrofon — ketuk Izinkan.",
    "### ➕ Buat rapat baru",
    "Nama rapat *",
    "mis. Serah terima shift pagi",
    "Tanggal",
    "Waktu",
    "Undang orang (mengirim ID + kata sandi lewat 💬 Pesan)",
    "📢 Umumkan di obrolan umum",
    "🎥 BUAT RAPAT",
    "Masukkan nama rapat!",
    "🎥 Undangan rapat",
    "🕒 Waktu",
    "🔢 ID ruang",
    "🔑 Kata sandi",
    "👉 Buka 🎥 Rapat Online untuk bergabung.",
    "🎉 Rapat dibuat **",
    "### 📋 Rapat saya",
    "Anda belum punya rapat.",
    "▶️ MULAI",
    "⛔ AKHIRI",
    "Hanya Foreman dan Team Leader yang bisa membuat rapat. Anda bisa bergabung dengan ID ruang dan kata sandi.",
    "Pilih orang...",
    "👑 Admin grup",
    "🗑️ Hapus pesan (admin)",
    "Tidak ada pesan untuk dihapus.",
    "Pesan yang dihapus",
    "🗑️ HAPUS PESAN",
    "Pesan dihapus.",
    "Saya yakin ingin menghapus semua pesan",
    "🧹 HAPUS SEMUA PESAN",
    "Semua pesan dihapus.",
    "Undang orang (mereka akan melihat undangan di 🎥 Rapat Online)",
    "### 📨 Undangan rapat Anda",
    "🎥 GABUNG RAPAT",
    "💬 Pesan & Rapat",
    "Pengaturan",
    "Tentang kami",
    "Masuk / Daftar",
    "Perangkat lunak manajemen pekerjaan untuk galangan kapal",
    "ShipControl membantu tim galangan kapal membagikan pekerjaan, memantau progres, dan tetap terhubung di satu tempat — di komputer dan ponsel.",
    "📋 Manajemen tugas",
    "Buat tugas berdasarkan Block, Deck, dan Frame lalu pantau progres persen demi persen.",
    "👥 Tim & peran",
    "Admin, WOS Manager, Foreman, Team Leader, dan Worker — setiap orang melihat pekerjaannya sendiri.",
    "💬 Pesan & rapat online",
    "Obrolan umum, obrolan grup, dan ruang rapat video dengan ID + kata sandi.",
    "🌐 15 bahasa & suara",
    "Vietnam, Inggris, Mandarin, Jepang, Korea, Thai, Indonesia, Spanyol, Prancis, Jerman, Rusia, dan lainnya — dan bicara alih-alih mengetik di ponsel.",
    "Belum punya akun? Daftar dan tunggu manajer memberi Anda peran.",
    "🎬 Dari galangan ke laut",
    "Saksikan kapal dibangun dari blok baja pertama hingga hari peluncuran.",
    "Video tidak bisa diputar? ",
    "▶ Buka di YouTube",
    "Buat obrolan grup",
    "Belum ada grup — ketuk ➕ untuk membuatnya.",
    "Tambah",
    "Hapus",
    "Tambah pekerja",
    "Keluarkan dari tim",
    "Tambah bengkel",
    "Hapus bengkel",
    "Hapus grup",
    "Keluar dari grup",
    "🌊 Samudra",
    "🌅 Senja",
    "🌲 Hutan",
    "⚪ Minimalis",
    "👑 Kerajaan",
    "📸 Laporan foto mingguan",
    "📈 Statistik",
    "Anda belum punya tugas untuk dilaporkan.",
    "Mulai",
    "Rencana selesai",
    "Progres",
    "Selesai",
    "Minggu ini",
    "Unggah foto",
    "Minggu",
    "Pilih foto dari komputer",
    "Catatan foto (opsional)",
    "📤 UNGGAH",
    "Pilih minimal 1 foto!",
    "Foto ini tidak bisa dibaca (atau lebih dari 15 MB): ",
    "Belum ada foto untuk minggu ini.",
    "📥 Unduh laporan foto (HTML)",
    "Buka file lalu tekan Ctrl+P untuk mencetak atau menyimpan sebagai PDF.",
    "💡 Tahan Ctrl (atau Shift) saat memilih untuk mengambil banyak foto sekaligus. Tekan Upload lagi untuk menambah.",
    "📊 Unduh PowerPoint (.pptx)",
    "◀ Sebelumnya",
    "Berikutnya ▶",
    "Total minggu",
    "Total foto",
    "Sedang berjalan",
    "foto",
    "Ringkasan",
]
_MS_LIST = [
    "🚢 SHIPCONTROL - PENGURUSAN KERJA KAPAL",
    "🔐 Log Masuk",
    "Sila log masuk atau daftar untuk teruskan.",
    "🧰 Papan Tugasan",
    "➕ Tambah Tugasan",
    "📋 Agih Tugasan",
    "👥 Pasukan Saya",
    "✏️ Sunting / Padam",
    "👥 Peranan & Kebenaran",
    "🗑️ Tong Sampah",
    "📊 Laporan",
    "🔑 Tukar Kata Laluan",
    "🚪 Log Keluar",
    "⚙️ Tetapan",
    "📲 Peranti",
    "📲 Anda sedang guna:",
    "💻 Komputer",
    "📱 Telefon",
    "🎤 Ketik mikrofon dalam kotak teks untuk bercakap dan bukannya menaip.",
    "🎤 Mod telefon: ketik mikrofon dalam kotak Nama pengguna untuk bercakap dan bukannya menaip.",
    "Pelayar ini belum menyokong input suara. Sila guna Chrome (Android) atau Safari (iPhone).",
    "Akses mikrofon disekat. Benarkan mikrofon untuk laman ini dalam tetapan pelayar.",
    "🏭 Bengkel: lihat / tambah / padam",
    "Belum ada bengkel. Tambah yang pertama di bawah.",
    "#### ➕ Tambah Bengkel Baharu",
    "Kod bengkel (cth. WOS_07):",
    "➕ TAMBAH BENGKEL",
    "Kod maks. 20 aksara, nama maks. 80 aksara!",
    "#### 🗑️ Padam Bengkel",
    "Bengkel untuk dipadam:",
    "🗑️ PADAM BENGKEL",
    "Nama bengkel",
    "Akaun",
    "Tugasan",
    "🎨 Tema",
    "✨ Moden",
    "🚀 Futuristik",
    "🎈 Ceria",
    "📜 Klasik",
    "Tema Futuristik sentiasa gelap.",
    ">Akaun<",
    "🔒 Akaun anda baru dibuka dalam tab atau peranti lain, jadi anda telah dilog keluar di sini. ",
    "Jika itu bukan anda, log masuk semula dan tukar kata laluan dengan segera.",
    ">Peranan<",
    ">Sedang log masuk<",
    "🌙 Mod gelap",
    "🌐 Bahasa",
    "📋 Papan Kemajuan Tugasan",
    "➕ Tambah Tugasan Baharu",
    "✏️ Sunting & Padam",
    "👥 Akaun & Peranan (Senarai Peranan)",
    "📊 Laporan Kemajuan & Statistik",
    "🗑️ Tong Sampah & Pulih",
    "### ⚡ Kemas Kini Kemajuan Pantas",
    "### 🏭 Statistik mengikut Bengkel",
    "### 📈 Taburan Kemajuan Tugasan",
    "### 📌 Agih Tugasan kepada Pekerja",
    "### 📥 Muat Turun Data Laporan",
    "### 🔄 Tukar Peranan Akaun",
    "#### ➕ Tambah Pekerja",
    "#### ➖ Keluarkan daripada Pasukan",
    "#### 👷 Ahli",
    "#### 🚪 Keluar Pasukan",
    "🔑 Log Masuk",
    "📝 Daftar",
    "Nama pengguna:",
    "Kata laluan:",
    "🚀 LOG MASUK",
    "✨ DAFTAR SEKARANG",
    "Nama penuh:",
    "Nama pengguna baharu:",
    "Sahkan kata laluan:",
    "Sila masukkan nama pengguna dan kata laluan!",
    "Nama pengguna atau kata laluan salah, atau akaun dikunci/dipadam!",
    "⏳ Akaun anda sedang menunggu Admin / WOS Manager memberi peranan. Sila datang semula nanti!",
    "Selamat kembali ",
    ")!",
    "Sila isi semua ruangan!",
    "Kata laluan mesti sekurang-kurangnya 8 aksara!",
    "Nama pengguna atau nama penuh terlalu panjang!",
    "Kata laluan tidak sepadan!",
    "Nama pengguna ini sudah wujud, sila pilih yang lain!",
    "🎉 Berjaya didaftar! Akaun anda sedang menunggu WOS Manager meluluskan peranan.",
    "👷 Ini tugasan yang diberikan kepada anda.",
    "Tunjuk tugasan pasukan saya sahaja",
    "Belum ada tugasan untuk ditunjukkan.",
    "Anda tiada tugasan untuk dikemas kini.",
    "Pilih tugasan untuk dikemas kini:",
    "Kemajuan baharu (%)",
    "Nota kerja (Remark):",
    "🚀 KEMAS KINI KEMAJUAN",
    "Kemajuan dikemas kini!",
    " (Semasa: ",
    "⚙️ Tambah / Buang Block dan WS Cost Code",
    "Block baharu (cth. 170150):",
    "➕ TAMBAH BLOCK",
    "Sila masukkan nama block!",
    "Block ditambah **",
    "Block untuk dipadam:",
    "🗑️ PADAM BLOCK",
    "Block dipadam **",
    "Senarai block kosong.",
    "Kod (cth. WOS_07):",
    "Nama bengkel / jabatan:",
    "➕ TAMBAH COST CODE",
    "Sila masukkan kod dan nama!",
    "Cost code untuk dipadam:",
    "🗑️ PADAM COST CODE",
    "Senarai cost code kosong.",
    "Tidak boleh dipadam: **",
    "** akaun masih dalam bengkel ini. Pindahkan mereka ke bengkel lain di 👥 Peranan & Kebenaran dahulu.",
    "** daripada senarai. Tugasan sedia ada dikekalkan.",
    "** sudah ada dalam senarai!",
    "Kod **",
    "Ditambah **",
    "Dipadam **",
    "Penerangan:",
    "Dicipta oleh:",
    "Tarikh dicipta:",
    "Penanggungjawab:",
    "Catatan:",
    "— Tiada —",
    "📌 Tugasan ini akan menjadi milik pasukan anda. Pergi ke 📋 Agih Tugasan untuk memberikannya kepada pekerja.",
    "💾 SIMPAN TUGASAN BAHARU",
    "Sila isi ruangan wajib: Task ID dan Task Name!",
    "Tarikh tamat rancangan tidak boleh sebelum tarikh mula!",
    "** sudah wujud (mungkin dalam Tong Sampah). Sila guna ID lain!",
    "Tugasan disimpan **",
    "⚠️ BELUM ADA WS COST CODE: Foreman / WOS Manager mesti menambahnya di atas sebelum menambah tugasan!",
    "👑 Pilih tugasan dan agihkan kepada **Team Leader** atau **Foreman**. Mereka akan menyerahkannya kepada pekerja dalam pasukan.",
    "Belum ada tugasan. Pergi ke ➕ Tambah Tugasan untuk mencipta satu.",
    "Belum ada Team Leader / Foreman. Beri peranan di 👥 Peranan & Kebenaran dahulu.",
    "Tunjuk tugasan yang belum diagih sahaja",
    "Semua tugasan sudah diagihkan!",
    "Pilih tugasan:",
    "Agih kepada Team Leader / Foreman:",
    "Agih kepada pekerja:",
    "— Belum diagih —",
    "Belum diagih",
    "💾 SIMPAN AGIHAN",
    "Tugasan diagihkan kepada **",
    "Pasukan anda belum ada pekerja. Pergi ke 👥 Pasukan Saya untuk mencipta pasukan dan menambah pekerja.",
    "Anda belum diberi sebarang tugasan.",
    "Pasukan anda perlu ada pekerja sebelum mengagih tugasan.",
    "(tiada nama)",
    "⚠️ Akaun anda belum ada bengkel. Minta WOS Manager menetapkannya sebelum mencipta pasukan.",
    "Anda belum ada pasukan. Beri nama untuk menciptanya.",
    "Nama pasukan *",
    "cth. Pasukan Kimpalan Block 170",
    "➕ CIPTA PASUKAN",
    "Sila masukkan nama pasukan!",
    "Nama pasukan maksimum 60 aksara!",
    "Pasukan dicipta **",
    "✏️ Tukar nama pasukan",
    "Nama pasukan baharu:",
    "💾 SIMPAN NAMA",
    "Nama pasukan tidak sah (1–60 aksara)!",
    "Nama pasukan ditukar!",
    "🗑️ Padam pasukan",
    "Memadam pasukan: semua pekerja keluar daripada pasukan dan tugasan mereka menjadi belum diagih. Tugasan kekal milik anda, dan anda boleh mencipta pasukan baharu kemudian.",
    "Saya pasti mahu memadam pasukan ini",
    "🗑️ PADAM PASUKAN",
    "Pasukan dipadam.",
    "Pasukan belum ada pekerja.",
    "Tiada lagi pekerja di bengkel ",
    " yang belum ada pasukan.",
    "Pilih pekerja (bengkel sama, belum ada pasukan):",
    "➕ TAMBAH KE PASUKAN",
    "Pilih sekurang-kurangnya 1 pekerja!",
    " pekerja ke pasukan!",
    "Menambah ",
    "Tiada ahli untuk dikeluarkan.",
    "Pilih pekerja:",
    "➖ KELUARKAN DARIPADA PASUKAN",
    "Pekerja dikeluarkan daripada pasukan. Tugasan pasukannya kini belum diagih.",
    "Anda belum dalam mana-mana pasukan. Team Leader / Foreman di bengkel anda boleh menambah anda.",
    "(Pasukan tiada nama)",
    "**Ketua pasukan:** ",
    "Jika anda keluar, tugasan pasukan yang diberikan kepada anda menjadi belum diagih.",
    "Saya pasti mahu keluar daripada pasukan ini",
    "🚪 KELUAR PASUKAN",
    "Anda telah keluar daripada pasukan.",
    "🧰 Pindah Tugasan ke Tong Sampah",
    "👤 Kunci Akaun Pengguna",
    "Tiada tugasan untuk disunting atau dipadam.",
    "Tugasan untuk dipindah ke Tong Sampah:",
    "🗑️ Pindah Tugasan ke Tong Sampah",
    "Tugasan dipindah ke Tong Sampah!",
    "Tiada akaun lain untuk dipadam.",
    "Akaun untuk dikunci:",
    "🗑️ Kunci Akaun Ini",
    "Akaun dikunci dan dipindah ke Tong Sampah!",
    "🛡️ Admin hanya memberi peranan **WOS Manager** dan memilih **Bengkel** yang diuruskannya. Peranan Worker / Team Leader / Foreman diberi oleh WOS Manager.",
    "👑 Anda boleh memberi peranan **Worker / Team Leader / Foreman**, memilih **Bengkel**, dan memasukkan pekerja ke dalam **pasukan** Team Leader / Foreman.",
    "Tiada akaun untuk diberi peranan.",
    "Akaun untuk ditukar:",
    "Tindakan:",
    "👑 Jadikan WOS Manager",
    "⛔ Tarik balik peranan (kembali ke Pending)",
    "Bengkel yang diuruskan pengurus ini:",
    "⚠️ Belum ada bengkel. Tambahnya di ➕ Tambah Tugasan → ⚙️ Tambah / Buang.",
    "💾 SIMPAN PERANAN",
    "💾 SIMPAN PERUBAHAN",
    "Dilantik **WOS Manager** bengkel **",
    "Peranan ditarik balik, akaun kembali ke Pending.",
    "Peranan baharu:",
    "Pasukan milik (Team Leader / Foreman):",
    "— Tiada pasukan —",
    "Bengkel ini belum ada Team Leader / Foreman.",
    "Dikemas kini: **",
    "Nama penuh",
    "Peranan",
    "Pasukan milik",
    "🧰 Tong Sampah Tugasan",
    "👤 Tong Sampah Akaun",
    "Tong sampah tugasan kosong.",
    "Tong sampah akaun kosong.",
    "Pilih tugasan:",
    "♻️ PULIHKAN TUGASAN",
    "💥 PADAM TUGASAN SELAMANYA",
    "Tugasan dipulihkan!",
    "Tugasan dipadam selamanya!",
    "Pilih akaun:",
    "♻️ BUKA KUNCI / PULIHKAN AKAUN",
    "💥 PADAM AKAUN SELAMANYA",
    "Akaun dipulihkan!",
    "Akaun dipadam selamanya daripada pangkalan data!",
    "Belum ada data tugasan untuk laporan. Sila tambah tugasan dahulu!",
    "Jumlah tugasan",
    "Jumlah tugasan",
    "Purata kemajuan (%)",
    "Purata kemajuan",
    "Selesai (100%)",
    "Sedang dijalankan",
    "Bilangan tugasan",
    "Status",
    "📥 Muat Turun Laporan Tugasan (CSV/Excel)",
    "Kata laluan semasa:",
    "Kata laluan baharu (sekurang-kurangnya 8 aksara):",
    "Ulang kata laluan baharu:",
    "💾 TUKAR KATA LALUAN",
    "Kata laluan semasa salah!",
    "Kata laluan baharu mesti sekurang-kurangnya 8 aksara!",
    "Anda tidak boleh guna semula kata laluan lalai!",
    "✅ Kata laluan ditukar! Peranti lain telah dilog keluar.",
    "⚠️ Akaun ini masih guna kata laluan lalai 'admin123'. Pergi ke 🔑 Tukar Kata Laluan dan tukar SEKARANG!",
    "Belum bermula (0%)",
    "Sedang dijalankan (1-50%)",
    "Hampir siap (51-99%)",
    "Siap (100%)",
    "Bilangan",
    "Kata laluan baharu:",
    "💬 Mesej",
    "📢 Sembang umum",
    "#### 👥 Orang",
    "🔎 Cari akaun",
    "Akaun tidak dijumpai.",
    "Semua orang dalam aplikasi boleh melihat saluran ini.",
    "Sembang peribadi dengan ",
    "Taip mesej...",
    "➕ Cipta sembang kumpulan",
    "Nama kumpulan",
    "cth. Krew kimpalan Block 170",
    "Tambah orang ke kumpulan",
    "✨ CIPTA KUMPULAN",
    "Sila namakan kumpulan!",
    "Pilih sekurang-kurangnya 1 orang!",
    "Kumpulan dicipta **",
    "#### 💬 Sembang kumpulan",
    "⚙️ Ahli kumpulan",
    "➕ TAMBAH KE KUMPULAN",
    " orang ke kumpulan!",
    "Keluarkan daripada kumpulan",
    "➖ KELUARKAN DARIPADA KUMPULAN",
    "Dikeluarkan daripada kumpulan.",
    "🗑️ PADAM KUMPULAN",
    "Hanya pencipta kumpulan boleh menambah / mengeluarkan ahli.",
    "🚪 KELUAR KUMPULAN",
    "🎥 Mesyuarat Dalam Talian",
    "### 🔑 Sertai mesyuarat",
    "ID bilik mesyuarat",
    "Kata laluan bilik",
    "🎥 SERTAI",
    "Terlalu banyak cubaan salah. Muat semula halaman dan cuba lagi nanti.",
    "ID bilik atau kata laluan salah, atau mesyuarat sudah tamat!",
    "✅ Bilik dijumpai: **",
    "Hos",
    "🎥 BUKA MESYUARAT (kamera + mikrofon)",
    "Mesyuarat dibuka dalam tab baharu. Kali pertama, pelayar akan meminta kamera dan mikrofon — ketik Benarkan.",
    "### ➕ Cipta mesyuarat baharu",
    "Nama mesyuarat *",
    "cth. Serahan syif pagi",
    "Tarikh",
    "Masa",
    "Jemput orang (hantar ID + kata laluan melalui 💬 Mesej)",
    "📢 Umumkan dalam sembang umum",
    "🎥 CIPTA MESYUARAT",
    "Sila masukkan nama mesyuarat!",
    "🎥 Jemputan mesyuarat",
    "🕒 Masa",
    "🔢 ID bilik",
    "🔑 Kata laluan",
    "👉 Pergi ke 🎥 Mesyuarat Dalam Talian untuk menyertai.",
    "🎉 Mesyuarat dicipta **",
    "### 📋 Mesyuarat saya",
    "Anda belum ada mesyuarat.",
    "▶️ MULA",
    "⛔ TAMAT",
    "Hanya Foreman dan Team Leader boleh mencipta mesyuarat. Anda boleh menyertai dengan ID bilik dan kata laluan.",
    "Pilih orang...",
    "👑 Admin kumpulan",
    "🗑️ Padam mesej (admin)",
    "Tiada mesej untuk dipadam.",
    "Mesej untuk dipadam",
    "🗑️ PADAM MESEJ",
    "Mesej dipadam.",
    "Saya pasti mahu memadam semua mesej",
    "🧹 PADAM SEMUA MESEJ",
    "Semua mesej dipadam.",
    "Jemput orang (mereka akan melihat jemputan di 🎥 Mesyuarat Dalam Talian)",
    "### 📨 Jemputan mesyuarat anda",
    "🎥 SERTAI MESYUARAT",
    "💬 Mesej & Mesyuarat",
    "Tetapan",
    "Tentang kami",
    "Log masuk / Daftar",
    "Perisian pengurusan kerja untuk limbungan kapal",
    "ShipControl membantu pasukan limbungan kapal mengagih kerja, menjejak kemajuan dan kekal berhubung di satu tempat — di komputer dan telefon.",
    "📋 Pengurusan tugasan",
    "Cipta tugasan mengikut Block, Deck dan Frame dan jejak kemajuan peratus demi peratus.",
    "👥 Pasukan & peranan",
    "Admin, WOS Manager, Foreman, Team Leader dan Worker — setiap orang melihat kerjanya sendiri.",
    "💬 Mesej & mesyuarat dalam talian",
    "Sembang umum, sembang kumpulan dan bilik mesyuarat video dengan ID + kata laluan.",
    "🌐 15 bahasa & suara",
    "Vietnam, Inggeris, Cina, Jepun, Korea, Thai, Indonesia, Sepanyol, Perancis, Jerman, Rusia dan lain-lain — dan bercakap dan bukannya menaip di telefon.",
    "Belum ada akaun? Daftar dan tunggu pengurus memberi anda peranan.",
    "🎬 Dari limbungan ke laut",
    "Tonton kapal dibina daripada blok keluli pertama hingga hari pelancaran.",
    "Video tidak boleh dimainkan? ",
    "▶ Buka di YouTube",
    "Cipta sembang kumpulan",
    "Belum ada kumpulan — ketik ➕ untuk mencipta.",
    "Tambah",
    "Buang",
    "Tambah pekerja",
    "Keluarkan daripada pasukan",
    "Tambah bengkel",
    "Buang bengkel",
    "Padam kumpulan",
    "Keluar kumpulan",
    "🌊 Lautan",
    "🌅 Senja",
    "🌲 Hutan",
    "⚪ Minimalis",
    "👑 Diraja",
    "📸 Laporan foto mingguan",
    "📈 Statistik",
    "Anda belum ada tugasan untuk dilaporkan.",
    "Mula",
    "Jangka siap",
    "Kemajuan",
    "Selesai",
    "Minggu semasa",
    "Muat naik foto",
    "Minggu",
    "Pilih foto dari komputer",
    "Nota foto (pilihan)",
    "📤 MUAT NAIK",
    "Pilih sekurang-kurangnya 1 foto!",
    "Foto ini tidak dapat dibaca (atau melebihi 15 MB): ",
    "Belum ada foto untuk minggu ini.",
    "📥 Muat turun laporan foto (HTML)",
    "Buka fail dan tekan Ctrl+P untuk mencetak atau simpan sebagai PDF.",
    "💡 Tahan Ctrl (atau Shift) semasa memilih untuk ambil banyak foto sekali gus. Tekan Upload sekali lagi untuk menambah.",
    "📊 Muat turun PowerPoint (.pptx)",
    "◀ Sebelum",
    "Seterusnya ▶",
    "Jumlah minggu",
    "Jumlah foto",
    "Sedang dijalankan",
    "foto",
    "Ringkasan",
]
_TL_LIST = [
    "🚢 SHIPCONTROL - PAMAMAHALA NG TRABAHO SA BARKO",
    "🔐 Mag-sign In",
    "Mag-sign in o magparehistro para magpatuloy.",
    "🧰 Board ng mga Gawain",
    "➕ Magdagdag ng Gawain",
    "📋 Magtalaga ng Gawain",
    "👥 Aking Team",
    "✏️ I-edit / Burahin",
    "👥 Mga Role at Pahintulot",
    "🗑️ Basurahan",
    "📊 Mga Ulat",
    "🔑 Palitan ang Password",
    "🚪 Mag-log Out",
    "⚙️ Mga Setting",
    "📲 Device",
    "📲 Ginagamit mo ang:",
    "💻 Computer",
    "📱 Telepono",
    "🎤 I-tap ang mikropono sa text box para magsalita sa halip na mag-type.",
    "🎤 Phone mode: i-tap ang mikropono sa Username box para magsalita sa halip na mag-type.",
    "Hindi pa sinusuportahan ng browser na ito ang voice input. Gamitin ang Chrome (Android) o Safari (iPhone).",
    "Naka-block ang access sa mikropono. Payagan ang mikropono para sa site na ito sa settings ng browser.",
    "🏭 Mga Workshop: tingnan / idagdag / burahin",
    "Wala pang workshop. Idagdag ang una sa ibaba.",
    "#### ➕ Magdagdag ng Bagong Workshop",
    "Code ng workshop (hal. WOS_07):",
    "➕ IDAGDAG ANG WORKSHOP",
    "Code hanggang 20 character, pangalan hanggang 80 character!",
    "#### 🗑️ Burahin ang Workshop",
    "Workshop na buburahin:",
    "🗑️ BURAHIN ANG WORKSHOP",
    "Pangalan ng workshop",
    "Mga account",
    "Mga gawain",
    "🎨 Tema",
    "✨ Moderno",
    "🚀 Futuristic",
    "🎈 Masaya",
    "📜 Makaluma",
    "Laging madilim ang Futuristic na tema.",
    ">Account<",
    "🔒 Kabubukas lang ng account mo sa ibang tab o device, kaya na-sign out ka rito. ",
    "Kung hindi ikaw iyon, mag-sign in ulit at palitan agad ang password mo.",
    ">Role<",
    ">Naka-sign in<",
    "🌙 Dark mode",
    "🌐 Wika",
    "📋 Board ng Progreso ng Gawain",
    "➕ Magdagdag ng Bagong Gawain",
    "✏️ I-edit at Burahin",
    "👥 Mga Account at Role (Listahan ng Role)",
    "📊 Mga Ulat ng Progreso at Istatistika",
    "🗑️ Basurahan at Pagbalik",
    "### ⚡ Mabilis na Update ng Progreso",
    "### 🏭 Istatistika ayon sa Workshop",
    "### 📈 Distribusyon ng Progreso",
    "### 📌 Magtalaga ng Gawain sa mga Manggagawa",
    "### 📥 I-download ang Data ng Ulat",
    "### 🔄 Palitan ang Role ng Account",
    "#### ➕ Magdagdag ng mga Manggagawa",
    "#### ➖ Alisin sa Team",
    "#### 👷 Mga Miyembro",
    "#### 🚪 Umalis sa Team",
    "🔑 Mag-sign In",
    "📝 Magparehistro",
    "Username:",
    "Password:",
    "🚀 MAG-SIGN IN",
    "✨ MAGPAREHISTRO NGAYON",
    "Buong pangalan:",
    "Bagong username:",
    "Kumpirmahin ang password:",
    "Ilagay ang username at password!",
    "Mali ang username o password, o naka-lock/nabura ang account!",
    "⏳ Naghihintay ang account mo na bigyan ng role ng Admin / WOS Manager. Bumalik mamaya!",
    "Maligayang pagbabalik, ",
    ")!",
    "Punan ang lahat ng field!",
    "Dapat hindi bababa sa 8 character ang password!",
    "Masyadong mahaba ang username o pangalan!",
    "Hindi magkatugma ang mga password!",
    "May gumagamit na ng username na ito, pumili ng iba!",
    "🎉 Nakarehistro na! Naghihintay ang account mo na aprubahan ng WOS Manager ang role.",
    "👷 Ito ang mga gawaing itinalaga sa iyo.",
    "Ipakita lang ang mga gawain ng team ko",
    "Wala pang gawaing maipapakita.",
    "Wala kang gawaing ia-update.",
    "Pumili ng gawaing ia-update:",
    "Bagong progreso (%)",
    "Tala sa trabaho (Remark):",
    "🚀 I-UPDATE ANG PROGRESO",
    "Na-update ang progreso!",
    " (Kasalukuyan: ",
    "⚙️ Magdagdag / Mag-alis ng Block at WS Cost Code",
    "Bagong block (hal. 170150):",
    "➕ IDAGDAG ANG BLOCK",
    "Ilagay ang pangalan ng block!",
    "Naidagdag ang block **",
    "Block na buburahin:",
    "🗑️ BURAHIN ANG BLOCK",
    "Nabura ang block **",
    "Walang laman ang listahan ng block.",
    "Code (hal. WOS_07):",
    "Pangalan ng workshop / departamento:",
    "➕ IDAGDAG ANG COST CODE",
    "Ilagay ang code at pangalan!",
    "Cost code na buburahin:",
    "🗑️ BURAHIN ANG COST CODE",
    "Walang laman ang listahan ng cost code.",
    "Hindi mabura: **",
    "** account pa ang nasa workshop na ito. Ilipat muna sila sa ibang workshop sa 👥 Mga Role at Pahintulot.",
    "** sa listahan. Nananatili ang mga dating gawain.",
    "** ay nasa listahan na!",
    "Code **",
    "Naidagdag **",
    "Nabura **",
    "Paglalarawan:",
    "Ginawa ni:",
    "Petsa ng paggawa:",
    "Namamahala:",
    "Puna:",
    "— Wala —",
    "📌 Mapupunta ang gawaing ito sa team mo. Pumunta sa 📋 Magtalaga ng Gawain para ibigay ito sa manggagawa.",
    "💾 I-SAVE ANG BAGONG GAWAIN",
    "Punan ang mga kailangang field: Task ID at Task Name!",
    "Hindi puwedeng mauna ang Plan Finish Date sa Plan Start Date!",
    "** ay mayroon na (baka nasa Basurahan). Gumamit ng ibang ID!",
    "Na-save ang gawain **",
    "⚠️ WALA PANG WS COST CODE: kailangang magdagdag muna ang Foreman / WOS Manager sa itaas bago magdagdag ng gawain!",
    "👑 Pumili ng gawain at italaga sa isang **Team Leader** o **Foreman**. Ipapasa nila ito sa mga manggagawa sa team nila.",
    "Wala pang gawain. Pumunta sa ➕ Magdagdag ng Gawain para gumawa.",
    "Wala pang Team Leader / Foreman. Magbigay muna ng role sa 👥 Mga Role at Pahintulot.",
    "Ipakita lang ang mga hindi pa naitatalagang gawain",
    "Naitalaga na ang lahat ng gawain!",
    "Pumili ng gawain:",
    "Italaga sa Team Leader / Foreman:",
    "Italaga sa manggagawa:",
    "— Hindi pa naitatalaga —",
    "Hindi pa naitatalaga",
    "💾 I-SAVE ANG PAGTATALAGA",
    "Naitalaga ang gawain kay **",
    "Wala pang manggagawa ang team mo. Pumunta sa 👥 Aking Team para gumawa ng team at magdagdag ng manggagawa.",
    "Wala ka pang natatanggap na gawain.",
    "Kailangan mo ng manggagawa sa team bago magtalaga ng gawain.",
    "(walang pangalan)",
    "⚠️ Wala pang workshop ang account mo. Hilingin sa WOS Manager na magtakda bago gumawa ng team.",
    "Wala ka pang team. Bigyan ito ng pangalan para gawin.",
    "Pangalan ng team *",
    "hal. Welding Team Block 170",
    "➕ GUMAWA NG TEAM",
    "Ilagay ang pangalan ng team!",
    "Hanggang 60 character lang ang pangalan ng team!",
    "Nagawa ang team **",
    "✏️ Palitan ang pangalan ng team",
    "Bagong pangalan ng team:",
    "💾 I-SAVE ANG PANGALAN",
    "Hindi wastong pangalan ng team (1–60 character)!",
    "Napalitan ang pangalan ng team!",
    "🗑️ Burahin ang team",
    "Kapag binura ang team: aalis ang lahat ng manggagawa at magiging hindi naitalaga ang mga gawain nila. Mananatili sa iyo ang mga gawain at puwede kang gumawa ng bagong team mamaya.",
    "Sigurado akong buburahin ko ang team na ito",
    "🗑️ BURAHIN ANG TEAM",
    "Nabura ang team.",
    "Wala pang manggagawa ang team.",
    "Wala nang manggagawa sa workshop ",
    " na walang team.",
    "Pumili ng manggagawa (parehong workshop, wala pang team):",
    "➕ IDAGDAG SA TEAM",
    "Pumili ng kahit 1 manggagawa!",
    " manggagawa sa team!",
    "Naidagdag ang ",
    "Walang miyembrong maaalis.",
    "Pumili ng manggagawa:",
    "➖ ALISIN SA TEAM",
    "Inalis ang manggagawa sa team. Hindi na naitalaga ang mga gawain niya sa team.",
    "Wala ka pa sa team. Puwede kang idagdag ng Team Leader / Foreman sa workshop mo.",
    "(Team na walang pangalan)",
    "**Lider ng team:** ",
    "Kapag umalis ka, magiging hindi naitalaga ang mga gawain ng team na nakatalaga sa iyo.",
    "Sigurado akong aalis ako sa team na ito",
    "🚪 UMALIS SA TEAM",
    "Umalis ka na sa team.",
    "🧰 Ilipat ang Gawain sa Basurahan",
    "👤 I-lock ang mga Account",
    "Walang gawaing ie-edit o buburahin.",
    "Gawaing ililipat sa Basurahan:",
    "🗑️ Ilipat ang Gawain sa Basurahan",
    "Nailipat ang gawain sa Basurahan!",
    "Walang ibang account na mabubura.",
    "Account na ila-lock:",
    "🗑️ I-lock ang Account na Ito",
    "Na-lock ang account at nailipat sa Basurahan!",
    "🛡️ Ang Admin ay nagbibigay lang ng role na **WOS Manager** at pumipili ng **Workshop** na hawak nito. Ang mga role na Worker / Team Leader / Foreman ay ibinibigay ng mga WOS Manager.",
    "👑 Puwede kang magbigay ng role na **Worker / Team Leader / Foreman**, pumili ng **Workshop**, at ilagay ang mga manggagawa sa **team** ng isang Team Leader / Foreman.",
    "Walang account na mabibigyan ng role.",
    "Account na papalitan:",
    "Aksyon:",
    "👑 Gawing WOS Manager",
    "⛔ Bawiin ang role (balik sa Pending)",
    "Workshop na hawak ng manager na ito:",
    "⚠️ Wala pang workshop. Idagdag sa ➕ Magdagdag ng Gawain → ⚙️ Magdagdag / Mag-alis.",
    "💾 I-SAVE ANG ROLE",
    "💾 I-SAVE ANG MGA PAGBABAGO",
    "Ginawang **WOS Manager** ng workshop **",
    "Binawi ang role, bumalik sa Pending ang account.",
    "Bagong role:",
    "Team ni (Team Leader / Foreman):",
    "— Walang team —",
    "Wala pang Team Leader / Foreman ang workshop na ito.",
    "Na-update: **",
    "Buong pangalan",
    "Role",
    "Team ni",
    "🧰 Basurahan ng Gawain",
    "👤 Basurahan ng Account",
    "Walang laman ang basurahan ng gawain.",
    "Walang laman ang basurahan ng account.",
    "Pumili ng gawain:",
    "♻️ IBALIK ANG GAWAIN",
    "💥 BURAHIN ANG GAWAIN NANG TULUYAN",
    "Naibalik ang gawain!",
    "Tuluyang nabura ang gawain!",
    "Pumili ng account:",
    "♻️ I-UNLOCK / IBALIK ANG ACCOUNT",
    "💥 BURAHIN ANG ACCOUNT NANG TULUYAN",
    "Naibalik ang account!",
    "Tuluyang nabura ang account sa database!",
    "Wala pang data ng gawain para sa ulat. Magdagdag muna ng gawain!",
    "Kabuuang gawain",
    "Kabuuang gawain",
    "Average na progreso (%)",
    "Average na progreso",
    "Tapos na (100%)",
    "Ginagawa",
    "Bilang ng gawain",
    "Status",
    "📥 I-download ang Ulat ng Gawain (CSV/Excel)",
    "Kasalukuyang password:",
    "Bagong password (hindi bababa sa 8 character):",
    "Ulitin ang bagong password:",
    "💾 PALITAN ANG PASSWORD",
    "Mali ang kasalukuyang password!",
    "Dapat hindi bababa sa 8 character ang bagong password!",
    "Hindi mo puwedeng gamitin ulit ang default na password!",
    "✅ Napalitan ang password! Na-log out na ang ibang device.",
    "⚠️ Gumagamit pa ang account na ito ng default na password na 'admin123'. Pumunta sa 🔑 Palitan ang Password at palitan NGAYON!",
    "Hindi pa nasisimulan (0%)",
    "Ginagawa (1-50%)",
    "Halos tapos na (51-99%)",
    "Tapos na (100%)",
    "Bilang",
    "Bagong password:",
    "💬 Mga Mensahe",
    "📢 Pangkalahatang chat",
    "#### 👥 Mga Tao",
    "🔎 Hanapin ang account",
    "Walang nahanap na account.",
    "Nakikita ng lahat sa app ang channel na ito.",
    "Pribadong chat kay ",
    "Mag-type ng mensahe...",
    "➕ Gumawa ng group chat",
    "Pangalan ng grupo",
    "hal. Welding crew Block 170",
    "Magdagdag ng tao sa grupo",
    "✨ GUMAWA NG GRUPO",
    "Pangalanan ang grupo!",
    "Pumili ng kahit 1 tao!",
    "Nagawa ang grupo **",
    "#### 💬 Mga group chat",
    "⚙️ Mga miyembro ng grupo",
    "➕ IDAGDAG SA GRUPO",
    " tao sa grupo!",
    "Alisin sa grupo",
    "➖ ALISIN SA GRUPO",
    "Inalis sa grupo.",
    "🗑️ BURAHIN ANG GRUPO",
    "Ang gumawa lang ng grupo ang puwedeng magdagdag / mag-alis ng miyembro.",
    "🚪 UMALIS SA GRUPO",
    "🎥 Online na Pulong",
    "### 🔑 Sumali sa pulong",
    "Room ID ng pulong",
    "Password ng room",
    "🎥 SUMALI",
    "Masyadong maraming maling subok. I-reload ang page at subukan ulit mamaya.",
    "Mali ang room ID o password, o tapos na ang pulong!",
    "✅ Nahanap ang room: **",
    "Host",
    "🎥 BUKSAN ANG PULONG (camera + mic)",
    "Bubukas ang pulong sa bagong tab. Sa unang beses, hihingi ang browser ng camera at mikropono — i-tap ang Allow.",
    "### ➕ Gumawa ng bagong pulong",
    "Pangalan ng pulong *",
    "hal. Turnover ng shift sa umaga",
    "Petsa",
    "Oras",
    "Mag-imbita ng tao (ipapadala ang ID + password sa 💬 Mga Mensahe)",
    "📢 Ianunsyo sa pangkalahatang chat",
    "🎥 GUMAWA NG PULONG",
    "Ilagay ang pangalan ng pulong!",
    "🎥 Imbitasyon sa pulong",
    "🕒 Oras",
    "🔢 Room ID",
    "🔑 Password",
    "👉 Pumunta sa 🎥 Online na Pulong para sumali.",
    "🎉 Nagawa ang pulong **",
    "### 📋 Aking mga pulong",
    "Wala ka pang pulong.",
    "▶️ SIMULAN",
    "⛔ TAPUSIN",
    "Ang mga Foreman at Team Leader lang ang puwedeng gumawa ng pulong. Puwede kang sumali gamit ang room ID at password.",
    "Pumili ng tao...",
    "👑 Admin ng grupo",
    "🗑️ Burahin ang mga mensahe (admin)",
    "Walang mensaheng mabubura.",
    "Mensaheng buburahin",
    "🗑️ BURAHIN ANG MENSAHE",
    "Nabura ang mensahe.",
    "Sigurado akong buburahin ko ang lahat ng mensahe",
    "🧹 BURAHIN ANG LAHAT NG MENSAHE",
    "Nabura ang lahat ng mensahe.",
    "Mag-imbita ng tao (makikita nila ang imbitasyon sa 🎥 Online na Pulong)",
    "### 📨 Mga imbitasyon mo sa pulong",
    "🎥 SUMALI SA PULONG",
    "💬 Mga Mensahe at Pulong",
    "Mga Setting",
    "Tungkol sa amin",
    "Mag-sign in / Magparehistro",
    "Software sa pamamahala ng trabaho para sa shipyard",
    "Tinutulungan ng ShipControl ang mga team sa shipyard na magtalaga ng trabaho, subaybayan ang progreso at manatiling konektado sa iisang lugar — sa computer at telepono.",
    "📋 Pamamahala ng gawain",
    "Gumawa ng gawain ayon sa Block, Deck at Frame at subaybayan ang progreso bawat porsyento.",
    "👥 Mga team at role",
    "Admin, WOS Manager, Foreman, Team Leader at Worker — nakikita ng bawat isa ang sarili niyang trabaho.",
    "💬 Mga mensahe at online na pulong",
    "Pangkalahatang chat, group chat at video meeting room na may ID + password.",
    "🌐 15 wika at boses",
    "Vietnamese, English, Chinese, Japanese, Korean, Thai, Indonesian, Spanish, French, German, Russian at iba pa — at magsalita sa halip na mag-type sa telepono.",
    "Wala ka pang account? Magparehistro at hintaying bigyan ka ng manager ng role.",
    "🎬 Mula shipyard hanggang dagat",
    "Panoorin ang paggawa ng barko mula sa unang bakal na block hanggang sa araw ng paglulunsad.",
    "Hindi ma-play ang video? ",
    "▶ Buksan sa YouTube",
    "Gumawa ng group chat",
    "Wala pang grupo — i-tap ang ➕ para gumawa.",
    "Idagdag",
    "Alisin",
    "Magdagdag ng manggagawa",
    "Alisin sa team",
    "Magdagdag ng workshop",
    "Alisin ang workshop",
    "Burahin ang grupo",
    "Umalis sa grupo",
    "🌊 Karagatan",
    "🌅 Paglubog ng araw",
    "🌲 Gubat",
    "⚪ Minimal",
    "👑 Maharlika",
    "📸 Lingguhang ulat ng larawan",
    "📈 Istatistika",
    "Wala ka pang gawaing maiuulat.",
    "Simula",
    "Planong tapos",
    "Progreso",
    "Tapos na",
    "Kasalukuyang linggo",
    "Mag-upload ng larawan",
    "Linggo",
    "Pumili ng larawan mula sa computer",
    "Tala sa larawan (opsyonal)",
    "📤 I-UPLOAD",
    "Pumili ng kahit 1 larawan!",
    "Hindi mabasa ang mga larawang ito (o lampas 15 MB): ",
    "Wala pang larawan para sa linggong ito.",
    "📥 I-download ang ulat ng larawan (HTML)",
    "Buksan ang file at pindutin ang Ctrl+P para i-print o i-save bilang PDF.",
    "💡 Pindutin nang matagal ang Ctrl (o Shift) habang pumipili para kumuha ng maraming larawan nang sabay. Pindutin ulit ang Upload para magdagdag.",
    "📊 I-download ang PowerPoint (.pptx)",
    "◀ Nakaraan",
    "Susunod ▶",
    "Kabuuang linggo",
    "Kabuuang larawan",
    "Ginagawa pa",
    "larawan",
    "Buod",
]
_HI_LIST = [
    "🚢 SHIPCONTROL - जहाज़ कार्य प्रबंधन",
    "🔐 साइन इन",
    "जारी रखने के लिए साइन इन या रजिस्टर करें।",
    "🧰 कार्य बोर्ड",
    "➕ कार्य जोड़ें",
    "📋 कार्य सौंपें",
    "👥 मेरी टीम",
    "✏️ संपादित करें / हटाएँ",
    "👥 भूमिकाएँ और अनुमतियाँ",
    "🗑️ रीसायकल बिन",
    "📊 रिपोर्ट",
    "🔑 पासवर्ड बदलें",
    "🚪 लॉग आउट",
    "⚙️ सेटिंग्स",
    "📲 डिवाइस",
    "📲 आप उपयोग कर रहे हैं:",
    "💻 कंप्यूटर",
    "📱 फ़ोन",
    "🎤 टाइप करने के बजाय बोलने के लिए टेक्स्ट बॉक्स में माइक्रोफ़ोन पर टैप करें।",
    "🎤 फ़ोन मोड: टाइप करने के बजाय बोलने के लिए यूज़रनेम बॉक्स में माइक्रोफ़ोन पर टैप करें।",
    "यह ब्राउज़र अभी वॉइस इनपुट सपोर्ट नहीं करता। कृपया Chrome (Android) या Safari (iPhone) इस्तेमाल करें।",
    "माइक्रोफ़ोन की अनुमति बंद है। ब्राउज़र सेटिंग्स में इस साइट के लिए माइक्रोफ़ोन की अनुमति दें।",
    "🏭 वर्कशॉप: देखें / जोड़ें / हटाएँ",
    "अभी कोई वर्कशॉप नहीं है। नीचे पहली वर्कशॉप जोड़ें।",
    "#### ➕ नई वर्कशॉप जोड़ें",
    "वर्कशॉप कोड (जैसे WOS_07):",
    "➕ वर्कशॉप जोड़ें",
    "कोड अधिकतम 20 अक्षर, नाम अधिकतम 80 अक्षर!",
    "#### 🗑️ वर्कशॉप हटाएँ",
    "हटाने के लिए वर्कशॉप:",
    "🗑️ वर्कशॉप हटाएँ",
    "वर्कशॉप का नाम",
    "खाते",
    "कार्य",
    "🎨 थीम",
    "✨ आधुनिक",
    "🚀 भविष्यवादी",
    "🎈 मज़ेदार",
    "📜 पुराना ज़माना",
    "भविष्यवादी थीम हमेशा डार्क रहती है।",
    ">खाता<",
    "🔒 आपका खाता अभी किसी दूसरे टैब या डिवाइस पर खोला गया, इसलिए यहाँ से आपको साइन आउट कर दिया गया। ",
    "अगर वह आप नहीं थे, तो फिर से साइन इन करें और तुरंत पासवर्ड बदलें।",
    ">भूमिका<",
    ">साइन इन है<",
    "🌙 डार्क मोड",
    "🌐 भाषा",
    "📋 कार्य प्रगति बोर्ड",
    "➕ नया कार्य जोड़ें",
    "✏️ संपादित करें और हटाएँ",
    "👥 खाते और भूमिकाएँ (भूमिका सूची)",
    "📊 प्रगति रिपोर्ट और आँकड़े",
    "🗑️ रीसायकल बिन और पुनर्स्थापना",
    "### ⚡ त्वरित प्रगति अपडेट",
    "### 🏭 वर्कशॉप के अनुसार आँकड़े",
    "### 📈 कार्य प्रगति का वितरण",
    "### 📌 मज़दूरों को कार्य सौंपें",
    "### 📥 रिपोर्ट डेटा डाउनलोड करें",
    "### 🔄 खाते की भूमिका बदलें",
    "#### ➕ मज़दूर जोड़ें",
    "#### ➖ टीम से हटाएँ",
    "#### 👷 सदस्य",
    "#### 🚪 टीम छोड़ें",
    "🔑 साइन इन",
    "📝 रजिस्टर",
    "यूज़रनेम:",
    "पासवर्ड:",
    "🚀 साइन इन करें",
    "✨ अभी रजिस्टर करें",
    "पूरा नाम:",
    "नया यूज़रनेम:",
    "पासवर्ड की पुष्टि करें:",
    "कृपया यूज़रनेम और पासवर्ड दोनों डालें!",
    "गलत यूज़रनेम या पासवर्ड, या खाता लॉक/हटाया गया है!",
    "⏳ आपका खाता Admin / WOS Manager द्वारा भूमिका मिलने का इंतज़ार कर रहा है। कृपया बाद में आएँ!",
    "स्वागत है ",
    ") फिर से आ गए!",
    "कृपया सभी फ़ील्ड भरें!",
    "पासवर्ड कम से कम 8 अक्षर का होना चाहिए!",
    "यूज़रनेम या पूरा नाम बहुत लंबा है!",
    "पासवर्ड मेल नहीं खाते!",
    "यह यूज़रनेम पहले से है, दूसरा चुनें!",
    "🎉 रजिस्टर हो गया! आपका खाता WOS Manager की भूमिका मंज़ूरी का इंतज़ार कर रहा है।",
    "👷 ये आपको सौंपे गए कार्य हैं।",
    "सिर्फ़ मेरी टीम के कार्य दिखाएँ",
    "अभी दिखाने के लिए कोई कार्य नहीं है।",
    "आपके पास अपडेट करने के लिए कोई कार्य नहीं है।",
    "अपडेट करने के लिए कार्य चुनें:",
    "नई प्रगति (%)",
    "कार्य नोट (Remark):",
    "🚀 प्रगति अपडेट करें",
    "प्रगति अपडेट हो गई!",
    " (वर्तमान: ",
    "⚙️ Block और WS Cost Code जोड़ें / हटाएँ",
    "नया block (जैसे 170150):",
    "➕ BLOCK जोड़ें",
    "कृपया block का नाम डालें!",
    "Block जोड़ा गया **",
    "हटाने के लिए block:",
    "🗑️ BLOCK हटाएँ",
    "Block हटाया गया **",
    "Block सूची खाली है।",
    "कोड (जैसे WOS_07):",
    "वर्कशॉप / विभाग का नाम:",
    "➕ COST CODE जोड़ें",
    "कृपया कोड और नाम दोनों डालें!",
    "हटाने के लिए cost code:",
    "🗑️ COST CODE हटाएँ",
    "Cost code सूची खाली है।",
    "हटा नहीं सकते: **",
    "** खाते अभी भी इस वर्कशॉप में हैं। पहले उन्हें 👥 भूमिकाएँ और अनुमतियाँ में दूसरी वर्कशॉप में भेजें।",
    "** सूची से। पुराने कार्य बने रहेंगे।",
    "** पहले से सूची में है!",
    "कोड **",
    "जोड़ा गया **",
    "हटाया गया **",
    "विवरण:",
    "बनाने वाले:",
    "बनाने की तारीख:",
    "ज़िम्मेदार:",
    "टिप्पणी:",
    "— कोई नहीं —",
    "📌 यह कार्य आपकी टीम का होगा। इसे किसी मज़दूर को देने के लिए 📋 कार्य सौंपें पर जाएँ।",
    "💾 नया कार्य सहेजें",
    "कृपया ज़रूरी फ़ील्ड भरें: Task ID और Task Name!",
    "योजना की समाप्ति तारीख शुरुआत की तारीख से पहले नहीं हो सकती!",
    "** पहले से मौजूद है (शायद रीसायकल बिन में)। कृपया दूसरी ID इस्तेमाल करें!",
    "कार्य सहेजा गया **",
    "⚠️ अभी कोई WS COST CODE नहीं: कार्य जोड़ने से पहले Foreman / WOS Manager को ऊपर एक जोड़ना होगा!",
    "👑 एक कार्य चुनें और उसे **Team Leader** या **Foreman** को सौंपें। वे इसे अपनी टीम के मज़दूरों को देंगे।",
    "अभी कोई कार्य नहीं है। बनाने के लिए ➕ कार्य जोड़ें पर जाएँ।",
    "अभी कोई Team Leader / Foreman नहीं है। पहले 👥 भूमिकाएँ और अनुमतियाँ में भूमिकाएँ दें।",
    "सिर्फ़ बिना सौंपे कार्य दिखाएँ",
    "सभी कार्य सौंप दिए गए हैं!",
    "कार्य चुनें:",
    "Team Leader / Foreman को सौंपें:",
    "मज़दूर को सौंपें:",
    "— नहीं सौंपा गया —",
    "नहीं सौंपा गया",
    "💾 सौंपना सहेजें",
    "कार्य सौंपा गया **",
    "आपकी टीम में अभी कोई मज़दूर नहीं है। टीम बनाने और मज़दूर जोड़ने के लिए 👥 मेरी टीम पर जाएँ।",
    "आपको अभी तक कोई कार्य नहीं दिया गया है।",
    "कार्य सौंपने से पहले आपकी टीम में मज़दूर होने चाहिए।",
    "(बिना नाम)",
    "⚠️ आपके खाते में अभी कोई वर्कशॉप नहीं है। टीम बनाने से पहले WOS Manager से वर्कशॉप दिलवाएँ।",
    "आपकी अभी कोई टीम नहीं है। बनाने के लिए उसे नाम दें।",
    "टीम का नाम *",
    "जैसे वेल्डिंग टीम Block 170",
    "➕ टीम बनाएँ",
    "कृपया टीम का नाम डालें!",
    "टीम का नाम अधिकतम 60 अक्षर हो सकता है!",
    "टीम बनाई गई **",
    "✏️ टीम का नाम बदलें",
    "टीम का नया नाम:",
    "💾 नाम सहेजें",
    "अमान्य टीम नाम (1–60 अक्षर)!",
    "टीम का नाम बदल दिया गया!",
    "🗑️ टीम हटाएँ",
    "टीम हटाने पर: सभी मज़दूर टीम से निकल जाएँगे और उनके कार्य बिना सौंपे हो जाएँगे। कार्य आपके पास रहेंगे, और आप बाद में नई टीम बना सकते हैं।",
    "मुझे यकीन है कि मैं यह टीम हटाना चाहता हूँ",
    "🗑️ टीम हटाएँ",
    "टीम हटा दी गई।",
    "टीम में अभी कोई मज़दूर नहीं है।",
    "वर्कशॉप में कोई मज़दूर नहीं बचा ",
    " जिसकी टीम न हो।",
    "मज़दूर चुनें (एक ही वर्कशॉप, अभी बिना टीम):",
    "➕ टीम में जोड़ें",
    "कम से कम 1 मज़दूर चुनें!",
    " मज़दूर टीम में जोड़े गए!",
    "",
    "हटाने के लिए कोई सदस्य नहीं।",
    "मज़दूर चुनें:",
    "➖ टीम से हटाएँ",
    "मज़दूर को टीम से हटा दिया गया। उसके टीम कार्य अब बिना सौंपे हैं।",
    "आप अभी किसी टीम में नहीं हैं। आपकी वर्कशॉप का कोई Team Leader / Foreman आपको जोड़ सकता है।",
    "(बिना नाम की टीम)",
    "**टीम लीडर:** ",
    "टीम छोड़ने पर, आपको सौंपे गए टीम कार्य बिना सौंपे हो जाएँगे।",
    "मुझे यकीन है कि मैं यह टीम छोड़ना चाहता हूँ",
    "🚪 टीम छोड़ें",
    "आपने टीम छोड़ दी।",
    "🧰 कार्य रीसायकल बिन में डालें",
    "👤 उपयोगकर्ता खाते लॉक करें",
    "संपादित करने या हटाने के लिए कोई कार्य नहीं।",
    "रीसायकल बिन में डालने के लिए कार्य:",
    "🗑️ कार्य रीसायकल बिन में डालें",
    "कार्य रीसायकल बिन में डाल दिया गया!",
    "हटाने के लिए कोई दूसरा खाता नहीं।",
    "लॉक करने के लिए खाता:",
    "🗑️ यह खाता लॉक करें",
    "खाता लॉक करके रीसायकल बिन में डाल दिया गया!",
    "🛡️ Admin सिर्फ़ **WOS Manager** की भूमिका देता है और वह **वर्कशॉप** चुनता है जिसे वह मैनेजर चलाएगा। Worker / Team Leader / Foreman की भूमिकाएँ WOS Manager देते हैं।",
    "👑 आप **Worker / Team Leader / Foreman** की भूमिकाएँ दे सकते हैं, **वर्कशॉप** चुन सकते हैं, और मज़दूरों को किसी Team Leader / Foreman की **टीम** में डाल सकते हैं।",
    "भूमिका देने के लिए कोई खाता नहीं।",
    "बदलने के लिए खाता:",
    "कार्रवाई:",
    "👑 WOS Manager बनाएँ",
    "⛔ भूमिका वापस लें (Pending पर वापस)",
    "यह मैनेजर जिस वर्कशॉप को चलाएगा:",
    "⚠️ अभी कोई वर्कशॉप नहीं है। ➕ कार्य जोड़ें → ⚙️ जोड़ें / हटाएँ में जोड़ें।",
    "💾 भूमिका सहेजें",
    "💾 बदलाव सहेजें",
    "वर्कशॉप का **WOS Manager** बनाया गया **",
    "भूमिका वापस ली गई, खाता फिर से Pending है।",
    "नई भूमिका:",
    "किसकी टीम (Team Leader / Foreman):",
    "— कोई टीम नहीं —",
    "इस वर्कशॉप में अभी कोई Team Leader / Foreman नहीं है।",
    "अपडेट किया गया: **",
    "पूरा नाम",
    "भूमिका",
    "किसकी टीम",
    "🧰 कार्य रीसायकल बिन",
    "👤 खाता रीसायकल बिन",
    "कार्य रीसायकल बिन खाली है।",
    "खाता रीसायकल बिन खाली है।",
    "कार्य चुनें:",
    "♻️ कार्य पुनर्स्थापित करें",
    "💥 कार्य हमेशा के लिए हटाएँ",
    "कार्य पुनर्स्थापित हो गया!",
    "कार्य हमेशा के लिए हटा दिया गया!",
    "खाता चुनें:",
    "♻️ खाता अनलॉक / पुनर्स्थापित करें",
    "💥 खाता हमेशा के लिए हटाएँ",
    "खाता पुनर्स्थापित हो गया!",
    "खाता डेटाबेस से हमेशा के लिए हटा दिया गया!",
    "रिपोर्ट के लिए अभी कोई कार्य डेटा नहीं है। पहले कार्य जोड़ें!",
    "कुल कार्य",
    "कुल कार्य",
    "औसत प्रगति (%)",
    "औसत प्रगति",
    "पूरे (100%)",
    "चल रहे",
    "कार्यों की संख्या",
    "स्थिति",
    "📥 कार्य रिपोर्ट डाउनलोड करें (CSV/Excel)",
    "वर्तमान पासवर्ड:",
    "नया पासवर्ड (कम से कम 8 अक्षर):",
    "नया पासवर्ड दोबारा डालें:",
    "💾 पासवर्ड बदलें",
    "वर्तमान पासवर्ड गलत है!",
    "नया पासवर्ड कम से कम 8 अक्षर का होना चाहिए!",
    "आप डिफ़ॉल्ट पासवर्ड दोबारा इस्तेमाल नहीं कर सकते!",
    "✅ पासवर्ड बदल गया! दूसरे डिवाइस लॉग आउट कर दिए गए।",
    "⚠️ यह खाता अभी भी डिफ़ॉल्ट पासवर्ड 'admin123' इस्तेमाल कर रहा है। 🔑 पासवर्ड बदलें पर जाएँ और अभी बदलें!",
    "शुरू नहीं हुआ (0%)",
    "चल रहा है (1-50%)",
    "लगभग पूरा (51-99%)",
    "पूरा (100%)",
    "संख्या",
    "नया पासवर्ड:",
    "💬 संदेश",
    "📢 सामान्य चैट",
    "#### 👥 लोग",
    "🔎 खाता खोजें",
    "कोई खाता नहीं मिला।",
    "ऐप में सभी लोग यह चैनल देख सकते हैं।",
    "निजी चैट: ",
    "संदेश लिखें...",
    "➕ ग्रुप चैट बनाएँ",
    "ग्रुप का नाम",
    "जैसे वेल्डिंग दल Block 170",
    "ग्रुप में लोग जोड़ें",
    "✨ ग्रुप बनाएँ",
    "कृपया ग्रुप को नाम दें!",
    "कम से कम 1 व्यक्ति चुनें!",
    "ग्रुप बनाया गया **",
    "#### 💬 ग्रुप चैट",
    "⚙️ ग्रुप के सदस्य",
    "➕ ग्रुप में जोड़ें",
    " लोग ग्रुप में जोड़े गए!",
    "ग्रुप से हटाएँ",
    "➖ ग्रुप से हटाएँ",
    "ग्रुप से हटा दिया गया।",
    "🗑️ ग्रुप हटाएँ",
    "सिर्फ़ ग्रुप बनाने वाला ही सदस्य जोड़ / हटा सकता है।",
    "🚪 ग्रुप छोड़ें",
    "🎥 ऑनलाइन मीटिंग",
    "### 🔑 मीटिंग में शामिल हों",
    "मीटिंग रूम ID",
    "रूम पासवर्ड",
    "🎥 शामिल हों",
    "बहुत ज़्यादा गलत प्रयास। पेज रीलोड करें और बाद में फिर कोशिश करें।",
    "गलत रूम ID या पासवर्ड, या मीटिंग खत्म हो चुकी है!",
    "✅ रूम मिल गया: **",
    "होस्ट",
    "🎥 मीटिंग खोलें (कैमरा + माइक)",
    "मीटिंग नए टैब में खुलेगी। पहली बार ब्राउज़र कैमरा और माइक्रोफ़ोन की अनुमति माँगेगा — Allow पर टैप करें।",
    "### ➕ नई मीटिंग बनाएँ",
    "मीटिंग का नाम *",
    "जैसे सुबह की शिफ़्ट हैंडओवर",
    "तारीख",
    "समय",
    "लोगों को आमंत्रित करें (💬 संदेश से ID + पासवर्ड भेजता है)",
    "📢 सामान्य चैट में घोषणा करें",
    "🎥 मीटिंग बनाएँ",
    "कृपया मीटिंग का नाम डालें!",
    "🎥 मीटिंग निमंत्रण",
    "🕒 समय",
    "🔢 रूम ID",
    "🔑 पासवर्ड",
    "👉 शामिल होने के लिए 🎥 ऑनलाइन मीटिंग पर जाएँ।",
    "🎉 मीटिंग बनाई गई **",
    "### 📋 मेरी मीटिंग",
    "आपकी अभी कोई मीटिंग नहीं है।",
    "▶️ शुरू करें",
    "⛔ समाप्त करें",
    "सिर्फ़ Foreman और Team Leader मीटिंग बना सकते हैं। आप रूम ID और पासवर्ड से शामिल हो सकते हैं।",
    "लोग चुनें...",
    "👑 ग्रुप एडमिन",
    "🗑️ संदेश हटाएँ (एडमिन)",
    "हटाने के लिए कोई संदेश नहीं।",
    "हटाने के लिए संदेश",
    "🗑️ संदेश हटाएँ",
    "संदेश हटा दिया गया।",
    "मुझे यकीन है कि मैं सभी संदेश हटाना चाहता हूँ",
    "🧹 सभी संदेश हटाएँ",
    "सभी संदेश हटा दिए गए।",
    "लोगों को आमंत्रित करें (वे निमंत्रण 🎥 ऑनलाइन मीटिंग में देखेंगे)",
    "### 📨 आपके मीटिंग निमंत्रण",
    "🎥 मीटिंग में शामिल हों",
    "💬 संदेश और मीटिंग",
    "सेटिंग्स",
    "हमारे बारे में",
    "साइन इन / रजिस्टर",
    "शिपयार्ड के लिए कार्य प्रबंधन सॉफ़्टवेयर",
    "ShipControl शिपयार्ड की टीमों को काम सौंपने, प्रगति पर नज़र रखने और एक ही जगह जुड़े रहने में मदद करता है — कंप्यूटर और फ़ोन दोनों पर।",
    "📋 कार्य प्रबंधन",
    "Block, Deck और Frame के अनुसार कार्य बनाएँ और प्रतिशत-दर-प्रतिशत प्रगति देखें।",
    "👥 टीमें और भूमिकाएँ",
    "Admin, WOS Manager, Foreman, Team Leader और Worker — हर कोई सिर्फ़ अपना काम देखता है।",
    "💬 संदेश और ऑनलाइन मीटिंग",
    "सामान्य चैट, ग्रुप चैट और ID + पासवर्ड वाले वीडियो मीटिंग रूम।",
    "🌐 15 भाषाएँ और आवाज़",
    "वियतनामी, अंग्रेज़ी, चीनी, जापानी, कोरियाई, थाई, इंडोनेशियाई, स्पेनिश, फ़्रेंच, जर्मन, रूसी और भी — और फ़ोन पर टाइप करने के बजाय बोलें।",
    "अभी खाता नहीं है? रजिस्टर करें और मैनेजर से भूमिका मिलने का इंतज़ार करें।",
    "🎬 शिपयार्ड से समुद्र तक",
    "देखें कैसे एक जहाज़ पहले स्टील ब्लॉक से लेकर लॉन्च के दिन तक बनता है।",
    "वीडियो नहीं चल रहा? ",
    "▶ YouTube पर खोलें",
    "ग्रुप चैट बनाएँ",
    "अभी कोई ग्रुप नहीं — बनाने के लिए ➕ पर टैप करें।",
    "जोड़ें",
    "हटाएँ",
    "मज़दूर जोड़ें",
    "टीम से हटाएँ",
    "वर्कशॉप जोड़ें",
    "वर्कशॉप हटाएँ",
    "ग्रुप हटाएँ",
    "ग्रुप छोड़ें",
    "🌊 महासागर",
    "🌅 सूर्यास्त",
    "🌲 जंगल",
    "⚪ मिनिमल",
    "👑 शाही",
    "📸 साप्ताहिक फ़ोटो रिपोर्ट",
    "📈 आँकड़े",
    "आपके पास रिपोर्ट करने के लिए अभी कोई कार्य नहीं है।",
    "शुरुआत",
    "अनुमानित समाप्ति",
    "प्रगति",
    "पूरा हुआ",
    "वर्तमान सप्ताह",
    "फ़ोटो अपलोड करें",
    "सप्ताह",
    "कंप्यूटर से फ़ोटो चुनें",
    "फ़ोटो नोट (वैकल्पिक)",
    "📤 अपलोड करें",
    "कम से कम 1 फ़ोटो चुनें!",
    "ये फ़ोटो पढ़ी नहीं जा सकीं (या 15 MB से बड़ी हैं): ",
    "इस सप्ताह की अभी कोई फ़ोटो नहीं है।",
    "📥 फ़ोटो रिपोर्ट डाउनलोड करें (HTML)",
    "फ़ाइल खोलें और प्रिंट करने या PDF के रूप में सहेजने के लिए Ctrl+P दबाएँ।",
    "💡 एक साथ कई फ़ोटो चुनने के लिए चुनते समय Ctrl (या Shift) दबाए रखें। और जोड़ने के लिए फिर से Upload दबाएँ।",
    "📊 PowerPoint डाउनलोड करें (.pptx)",
    "◀ पिछला",
    "अगला ▶",
    "कुल सप्ताह",
    "कुल फ़ोटो",
    "चल रहा है",
    "फ़ोटो",
    "सारांश",
]
_ES_LIST = [
    "🚢 SHIPCONTROL - GESTIÓN DE TRABAJOS NAVALES",
    "🔐 Iniciar sesión",
    "Inicia sesión o regístrate para continuar.",
    "🧰 Tablero de tareas",
    "➕ Añadir tarea",
    "📋 Asignar tareas",
    "👥 Mi equipo",
    "✏️ Editar / Eliminar",
    "👥 Roles y permisos",
    "🗑️ Papelera",
    "📊 Informes",
    "🔑 Cambiar contraseña",
    "🚪 Cerrar sesión",
    "⚙️ Ajustes",
    "📲 Dispositivo",
    "📲 Estás usando:",
    "💻 Ordenador",
    "📱 Teléfono",
    "🎤 Toca el micrófono en un cuadro de texto para hablar en lugar de escribir.",
    "🎤 Modo teléfono: toca el micrófono en el cuadro de usuario para hablar en lugar de escribir.",
    "Este navegador aún no admite entrada por voz. Usa Chrome (Android) o Safari (iPhone).",
    "El acceso al micrófono está bloqueado. Permite el micrófono para este sitio en los ajustes del navegador.",
    "🏭 Talleres: ver / añadir / eliminar",
    "Aún no hay talleres. Añade el primero abajo.",
    "#### ➕ Añadir nuevo taller",
    "Código del taller (ej. WOS_07):",
    "➕ AÑADIR TALLER",
    "¡Código máx. 20 caracteres, nombre máx. 80 caracteres!",
    "#### 🗑️ Eliminar taller",
    "Taller a eliminar:",
    "🗑️ ELIMINAR TALLER",
    "Nombre del taller",
    "Cuentas",
    "Tareas",
    "🎨 Tema",
    "✨ Moderno",
    "🚀 Futurista",
    "🎈 Divertido",
    "📜 Antiguo",
    "El tema Futurista siempre es oscuro.",
    ">Cuenta<",
    "🔒 Tu cuenta se acaba de abrir en otra pestaña u otro dispositivo, así que se cerró la sesión aquí. ",
    "Si no fuiste tú, inicia sesión de nuevo y cambia tu contraseña enseguida.",
    ">Rol<",
    ">Sesión iniciada<",
    "🌙 Modo oscuro",
    "🌐 Idioma",
    "📋 Tablero de progreso de tareas",
    "➕ Añadir nueva tarea",
    "✏️ Editar y eliminar",
    "👥 Cuentas y roles (lista de roles)",
    "📊 Informes de progreso y estadísticas",
    "🗑️ Papelera y restaurar",
    "### ⚡ Actualización rápida de progreso",
    "### 🏭 Estadísticas por taller",
    "### 📈 Distribución del progreso",
    "### 📌 Asignar tareas a trabajadores",
    "### 📥 Descargar datos del informe",
    "### 🔄 Cambiar rol de la cuenta",
    "#### ➕ Añadir trabajadores",
    "#### ➖ Quitar del equipo",
    "#### 👷 Miembros",
    "#### 🚪 Salir del equipo",
    "🔑 Iniciar sesión",
    "📝 Registrarse",
    "Usuario:",
    "Contraseña:",
    "🚀 ENTRAR",
    "✨ REGISTRARSE AHORA",
    "Nombre completo:",
    "Nuevo usuario:",
    "Confirmar contraseña:",
    "¡Introduce usuario y contraseña!",
    "¡Usuario o contraseña incorrectos, o la cuenta está bloqueada/eliminada!",
    "⏳ Tu cuenta espera a que un Admin / WOS Manager le asigne un rol. ¡Vuelve más tarde!",
    "Bienvenido/a ",
    ") ha vuelto!",
    "¡Rellena todos los campos!",
    "¡La contraseña debe tener al menos 8 caracteres!",
    "¡El usuario o el nombre es demasiado largo!",
    "¡Las contraseñas no coinciden!",
    "¡Este usuario ya existe, elige otro!",
    "🎉 ¡Registrado! Tu cuenta espera a que un WOS Manager apruebe un rol.",
    "👷 Estas son las tareas asignadas a ti.",
    "Mostrar solo las tareas de mi equipo",
    "Aún no hay tareas que mostrar.",
    "No tienes tareas que actualizar.",
    "Elige una tarea para actualizar:",
    "Nuevo progreso (%)",
    "Nota de trabajo (Remark):",
    "🚀 ACTUALIZAR PROGRESO",
    "¡Progreso actualizado!",
    " (Actual: ",
    "⚙️ Añadir / quitar Blocks y WS Cost Codes",
    "Nuevo block (ej. 170150):",
    "➕ AÑADIR BLOCK",
    "¡Introduce un nombre de block!",
    "Block añadido **",
    "Block a eliminar:",
    "🗑️ ELIMINAR BLOCK",
    "Block eliminado **",
    "La lista de blocks está vacía.",
    "Código (ej. WOS_07):",
    "Nombre del taller / departamento:",
    "➕ AÑADIR COST CODE",
    "¡Introduce código y nombre!",
    "Cost code a eliminar:",
    "🗑️ ELIMINAR COST CODE",
    "La lista de cost codes está vacía.",
    "No se puede eliminar: **",
    "** cuenta(s) siguen en este taller. Muévelas a otro taller en 👥 Roles y permisos primero.",
    "** de la lista. Las tareas existentes se conservan.",
    "** ya está en la lista!",
    "Código **",
    "Añadido **",
    "Eliminado **",
    "Descripción:",
    "Iniciado por:",
    "Fecha de inicio:",
    "Responsable:",
    "Observación:",
    "— Ninguno —",
    "📌 Esta tarea pertenecerá a tu equipo. Ve a 📋 Asignar tareas para dársela a un trabajador.",
    "💾 GUARDAR NUEVA TAREA",
    "¡Rellena los campos obligatorios: Task ID y Task Name!",
    "¡La fecha de fin planificada no puede ser anterior a la de inicio!",
    "** ya existe (puede estar en la Papelera). ¡Usa otro ID!",
    "Tarea guardada **",
    "⚠️ AÚN NO HAY WS COST CODE: ¡un Foreman / WOS Manager debe añadir uno arriba antes de añadir tareas!",
    "👑 Elige una tarea y asígnala a un **Team Leader** o **Foreman**. Ellos la pasarán a los trabajadores de su equipo.",
    "Aún no hay tareas. Ve a ➕ Añadir tarea para crear una.",
    "Aún no hay Team Leaders / Foremen. Asigna roles en 👥 Roles y permisos primero.",
    "Mostrar solo tareas sin asignar",
    "¡Todas las tareas están asignadas!",
    "Elige una tarea:",
    "Asignar a Team Leader / Foreman:",
    "Asignar a trabajador:",
    "— Sin asignar —",
    "Sin asignar",
    "💾 GUARDAR ASIGNACIÓN",
    "Tarea asignada a **",
    "Tu equipo aún no tiene trabajadores. Ve a 👥 Mi equipo para crear un equipo y añadir trabajadores.",
    "Aún no te han dado ninguna tarea.",
    "Necesitas trabajadores en tu equipo antes de asignar tareas.",
    "(sin nombre)",
    "⚠️ Tu cuenta aún no tiene taller. Pide a un WOS Manager que te asigne uno antes de crear un equipo.",
    "Aún no tienes equipo. Dale un nombre para crearlo.",
    "Nombre del equipo *",
    "ej. Equipo de soldadura Block 170",
    "➕ CREAR EQUIPO",
    "¡Introduce un nombre de equipo!",
    "¡El nombre del equipo puede tener como máximo 60 caracteres!",
    "Equipo creado **",
    "✏️ Renombrar equipo",
    "Nuevo nombre del equipo:",
    "💾 GUARDAR NOMBRE",
    "¡Nombre de equipo no válido (1–60 caracteres)!",
    "¡Equipo renombrado!",
    "🗑️ Eliminar equipo",
    "Al eliminar el equipo: todos los trabajadores salen del equipo y sus tareas quedan sin asignar. Las tareas siguen siendo tuyas y puedes crear un equipo nuevo después.",
    "Estoy seguro de que quiero eliminar este equipo",
    "🗑️ ELIMINAR EQUIPO",
    "Equipo eliminado.",
    "El equipo aún no tiene trabajadores.",
    "No quedan trabajadores en el taller ",
    " sin equipo.",
    "Elige trabajadores (mismo taller, sin equipo):",
    "➕ AÑADIR AL EQUIPO",
    "¡Elige al menos 1 trabajador!",
    " trabajador(es) al equipo!",
    "Añadidos ",
    "No hay miembros para quitar.",
    "Elige un trabajador:",
    "➖ QUITAR DEL EQUIPO",
    "Trabajador quitado del equipo. Sus tareas del equipo quedan sin asignar.",
    "Aún no estás en un equipo. Un Team Leader / Foreman de tu taller puede añadirte.",
    "(Equipo sin nombre)",
    "**Jefe de equipo:** ",
    "Al salir, las tareas del equipo asignadas a ti quedarán sin asignar.",
    "Estoy seguro de que quiero salir de este equipo",
    "🚪 SALIR DEL EQUIPO",
    "Has salido del equipo.",
    "🧰 Mover tareas a la papelera",
    "👤 Bloquear cuentas",
    "No hay tareas para editar o eliminar.",
    "Tarea a mover a la papelera:",
    "🗑️ Mover tarea a la papelera",
    "¡Tarea movida a la papelera!",
    "No hay otras cuentas para eliminar.",
    "Cuenta a bloquear:",
    "🗑️ Bloquear esta cuenta",
    "¡Cuenta bloqueada y movida a la papelera!",
    "🛡️ El Admin solo concede el rol **WOS Manager** y elige el **Taller** que dirige. Los roles Worker / Team Leader / Foreman los dan los WOS Managers.",
    "👑 Puedes dar los roles **Worker / Team Leader / Foreman**, elegir el **Taller** y poner trabajadores en el **equipo** de un Team Leader / Foreman.",
    "No hay cuentas a las que asignar roles.",
    "Cuenta a cambiar:",
    "Acción:",
    "👑 Hacer WOS Manager",
    "⛔ Quitar rol (volver a Pending)",
    "Taller que dirige este manager:",
    "⚠️ Aún no hay talleres. Añádelos en ➕ Añadir tarea → ⚙️ Añadir / quitar.",
    "💾 GUARDAR ROL",
    "💾 GUARDAR CAMBIOS",
    "Nombrado **WOS Manager** del taller **",
    "Rol retirado, la cuenta vuelve a Pending.",
    "Nuevo rol:",
    "Equipo de (Team Leader / Foreman):",
    "— Sin equipo —",
    "Este taller aún no tiene Team Leader / Foreman.",
    "Actualizado: **",
    "Nombre completo",
    "Rol",
    "Equipo de",
    "🧰 Papelera de tareas",
    "👤 Papelera de cuentas",
    "La papelera de tareas está vacía.",
    "La papelera de cuentas está vacía.",
    "Elige una tarea:",
    "♻️ RESTAURAR TAREA",
    "💥 ELIMINAR TAREA PARA SIEMPRE",
    "¡Tarea restaurada!",
    "¡Tarea eliminada definitivamente!",
    "Elige una cuenta:",
    "♻️ DESBLOQUEAR / RESTAURAR CUENTA",
    "💥 ELIMINAR CUENTA PARA SIEMPRE",
    "¡Cuenta restaurada!",
    "¡Cuenta eliminada definitivamente de la base de datos!",
    "Aún no hay datos de tareas para un informe. ¡Añade tareas primero!",
    "Total de tareas",
    "Total de tareas",
    "Progreso medio (%)",
    "Progreso medio",
    "Completadas (100%)",
    "En curso",
    "Número de tareas",
    "Estado",
    "📥 Descargar informe de tareas (CSV/Excel)",
    "Contraseña actual:",
    "Nueva contraseña (mínimo 8 caracteres):",
    "Repite la nueva contraseña:",
    "💾 CAMBIAR CONTRASEÑA",
    "¡La contraseña actual es incorrecta!",
    "¡La nueva contraseña debe tener al menos 8 caracteres!",
    "¡No puedes reutilizar la contraseña predeterminada!",
    "✅ ¡Contraseña cambiada! Se cerró la sesión en los otros dispositivos.",
    "⚠️ Esta cuenta todavía usa la contraseña predeterminada 'admin123'. ¡Ve a 🔑 Cambiar contraseña y cámbiala YA!",
    "Sin empezar (0%)",
    "En curso (1-50%)",
    "Casi terminado (51-99%)",
    "Terminado (100%)",
    "Cantidad",
    "Nueva contraseña:",
    "💬 Mensajes",
    "📢 Chat general",
    "#### 👥 Personas",
    "🔎 Buscar cuenta",
    "No se encontró ninguna cuenta.",
    "Todos en la app pueden ver este canal.",
    "Chat privado con ",
    "Escribe un mensaje...",
    "➕ Crear chat de grupo",
    "Nombre del grupo",
    "ej. Cuadrilla de soldadura Block 170",
    "Añadir personas al grupo",
    "✨ CREAR GRUPO",
    "¡Ponle nombre al grupo!",
    "¡Elige al menos 1 persona!",
    "Grupo creado **",
    "#### 💬 Chats de grupo",
    "⚙️ Miembros del grupo",
    "➕ AÑADIR AL GRUPO",
    " personas al grupo!",
    "Quitar del grupo",
    "➖ QUITAR DEL GRUPO",
    "Quitado del grupo.",
    "🗑️ ELIMINAR GRUPO",
    "Solo el creador del grupo puede añadir / quitar miembros.",
    "🚪 SALIR DEL GRUPO",
    "🎥 Reuniones en línea",
    "### 🔑 Unirse a una reunión",
    "ID de la sala",
    "Contraseña de la sala",
    "🎥 UNIRSE",
    "Demasiados intentos fallidos. Recarga la página e inténtalo más tarde.",
    "¡ID o contraseña incorrectos, o la reunión ha terminado!",
    "✅ Sala encontrada: **",
    "Anfitrión",
    "🎥 ABRIR REUNIÓN (cámara + micrófono)",
    "La reunión se abre en una pestaña nueva. La primera vez, el navegador pedirá usar la cámara y el micrófono — toca Permitir.",
    "### ➕ Crear una nueva reunión",
    "Nombre de la reunión *",
    "ej. Relevo del turno de mañana",
    "Fecha",
    "Hora",
    "Invitar personas (envía ID + contraseña por 💬 Mensajes)",
    "📢 Anunciar en el chat general",
    "🎥 CREAR REUNIÓN",
    "¡Introduce un nombre para la reunión!",
    "🎥 Invitación a reunión",
    "🕒 Hora",
    "🔢 ID de la sala",
    "🔑 Contraseña",
    "👉 Ve a 🎥 Reuniones en línea para unirte.",
    "🎉 Reunión creada **",
    "### 📋 Mis reuniones",
    "Aún no tienes reuniones.",
    "▶️ INICIAR",
    "⛔ TERMINAR",
    "Solo los Foremen y Team Leaders pueden crear reuniones. Puedes unirte con un ID y una contraseña.",
    "Elige personas...",
    "👑 Admin del grupo",
    "🗑️ Eliminar mensajes (admin)",
    "No hay mensajes para eliminar.",
    "Mensaje a eliminar",
    "🗑️ ELIMINAR MENSAJE",
    "Mensaje eliminado.",
    "Estoy seguro de que quiero eliminar todos los mensajes",
    "🧹 ELIMINAR TODOS LOS MENSAJES",
    "Todos los mensajes eliminados.",
    "Invitar personas (verán la invitación en 🎥 Reuniones en línea)",
    "### 📨 Tus invitaciones a reuniones",
    "🎥 UNIRSE A LA REUNIÓN",
    "💬 Mensajes y reuniones",
    "Ajustes",
    "Sobre nosotros",
    "Entrar / Registrarse",
    "Software de gestión de trabajos para astilleros",
    "ShipControl ayuda a los equipos del astillero a asignar trabajo, seguir el progreso y mantenerse en contacto en un solo lugar — en ordenadores y teléfonos.",
    "📋 Gestión de tareas",
    "Crea tareas por Block, Deck y Frame y sigue el progreso porcentaje a porcentaje.",
    "👥 Equipos y roles",
    "Admin, WOS Manager, Foreman, Team Leader y Worker — cada uno ve exactamente su propio trabajo.",
    "💬 Mensajes y reuniones en línea",
    "Chat general, chats de grupo y salas de videoreunión con ID + contraseña.",
    "🌐 15 idiomas y voz",
    "Vietnamita, inglés, chino, japonés, coreano, tailandés, indonesio, español, francés, alemán, ruso y más — y habla en lugar de escribir en el teléfono.",
    "¿Aún no tienes cuenta? Regístrate y espera a que un responsable te dé un rol.",
    "🎬 Del astillero al mar",
    "Mira cómo se construye un barco desde los primeros bloques de acero hasta la botadura.",
    "¿No se reproduce el video? ",
    "▶ Abrir en YouTube",
    "Crear chat de grupo",
    "Aún no hay grupos — toca ➕ para crear uno.",
    "Añadir",
    "Quitar",
    "Añadir trabajadores",
    "Quitar del equipo",
    "Añadir taller",
    "Quitar taller",
    "Eliminar grupo",
    "Salir del grupo",
    "🌊 Océano",
    "🌅 Atardecer",
    "🌲 Bosque",
    "⚪ Minimalista",
    "👑 Real",
    "📸 Informe fotográfico semanal",
    "📈 Estadísticas",
    "Aún no tienes tareas sobre las que informar.",
    "Inicio",
    "Fin previsto",
    "Progreso",
    "Terminada",
    "Semana actual",
    "Subir fotos",
    "Semana",
    "Elige fotos del ordenador",
    "Nota de la foto (opcional)",
    "📤 SUBIR",
    "¡Elige al menos 1 foto!",
    "No se pudieron leer estas fotos (o pesan más de 15 MB): ",
    "Aún no hay fotos de esta semana.",
    "📥 Descargar informe fotográfico (HTML)",
    "Abre el archivo y pulsa Ctrl+P para imprimirlo o guardarlo como PDF.",
    "💡 Mantén Ctrl (o Shift) al elegir para seleccionar varias fotos a la vez. Pulsa Upload otra vez para añadir más.",
    "📊 Descargar PowerPoint (.pptx)",
    "◀ Anterior",
    "Siguiente ▶",
    "Total de semanas",
    "Total de fotos",
    "En curso",
    "fotos",
    "Resumen",
]
_FR_LIST = [
    "🚢 SHIPCONTROL - GESTION DES TRAVAUX NAVALS",
    "🔐 Connexion",
    "Connectez-vous ou inscrivez-vous pour continuer.",
    "🧰 Tableau des tâches",
    "➕ Ajouter une tâche",
    "📋 Attribuer des tâches",
    "👥 Mon équipe",
    "✏️ Modifier / Supprimer",
    "👥 Rôles et permissions",
    "🗑️ Corbeille",
    "📊 Rapports",
    "🔑 Changer le mot de passe",
    "🚪 Se déconnecter",
    "⚙️ Paramètres",
    "📲 Appareil",
    "📲 Vous utilisez :",
    "💻 Ordinateur",
    "📱 Téléphone",
    "🎤 Touchez le micro dans une zone de texte pour parler au lieu de taper.",
    "🎤 Mode téléphone : touchez le micro dans la zone Identifiant pour parler au lieu de taper.",
    "Ce navigateur ne prend pas encore en charge la saisie vocale. Utilisez Chrome (Android) ou Safari (iPhone).",
    "L'accès au micro est bloqué. Autorisez le micro pour ce site dans les paramètres du navigateur.",
    "🏭 Ateliers : voir / ajouter / supprimer",
    "Aucun atelier pour l'instant. Ajoutez le premier ci-dessous.",
    "#### ➕ Ajouter un atelier",
    "Code de l'atelier (ex. WOS_07) :",
    "➕ AJOUTER L'ATELIER",
    "Code : 20 caractères max, nom : 80 caractères max !",
    "#### 🗑️ Supprimer un atelier",
    "Atelier à supprimer :",
    "🗑️ SUPPRIMER L'ATELIER",
    "Nom de l'atelier",
    "Comptes",
    "Tâches",
    "🎨 Thème",
    "✨ Moderne",
    "🚀 Futuriste",
    "🎈 Ludique",
    "📜 Rétro",
    "Le thème Futuriste est toujours sombre.",
    ">Compte<",
    "🔒 Votre compte vient d'être ouvert dans un autre onglet ou sur un autre appareil, vous avez donc été déconnecté ici. ",
    "Si ce n'était pas vous, reconnectez-vous et changez votre mot de passe tout de suite.",
    ">Rôle<",
    ">Connecté<",
    "🌙 Mode sombre",
    "🌐 Langue",
    "📋 Tableau d'avancement des tâches",
    "➕ Ajouter une nouvelle tâche",
    "✏️ Modifier et supprimer",
    "👥 Comptes et rôles (liste des rôles)",
    "📊 Rapports d'avancement et statistiques",
    "🗑️ Corbeille et restauration",
    "### ⚡ Mise à jour rapide de l'avancement",
    "### 🏭 Statistiques par atelier",
    "### 📈 Répartition de l'avancement",
    "### 📌 Attribuer des tâches aux ouvriers",
    "### 📥 Télécharger les données du rapport",
    "### 🔄 Changer le rôle du compte",
    "#### ➕ Ajouter des ouvriers",
    "#### ➖ Retirer de l'équipe",
    "#### 👷 Membres",
    "#### 🚪 Quitter l'équipe",
    "🔑 Connexion",
    "📝 Inscription",
    "Identifiant :",
    "Mot de passe :",
    "🚀 SE CONNECTER",
    "✨ S'INSCRIRE",
    "Nom complet :",
    "Nouvel identifiant :",
    "Confirmer le mot de passe :",
    "Saisissez l'identifiant et le mot de passe !",
    "Identifiant ou mot de passe incorrect, ou compte verrouillé/supprimé !",
    "⏳ Votre compte attend qu'un Admin / WOS Manager lui attribue un rôle. Revenez plus tard !",
    "Bienvenue ",
    ") est de retour !",
    "Remplissez tous les champs !",
    "Le mot de passe doit contenir au moins 8 caractères !",
    "L'identifiant ou le nom est trop long !",
    "Les mots de passe ne correspondent pas !",
    "Cet identifiant existe déjà, choisissez-en un autre !",
    "🎉 Inscrit ! Votre compte attend qu'un WOS Manager approuve un rôle.",
    "👷 Voici les tâches qui vous sont attribuées.",
    "Afficher seulement les tâches de mon équipe",
    "Aucune tâche à afficher pour l'instant.",
    "Vous n'avez aucune tâche à mettre à jour.",
    "Choisissez une tâche à mettre à jour :",
    "Nouvel avancement (%)",
    "Note de travail (Remark) :",
    "🚀 METTRE À JOUR",
    "Avancement mis à jour !",
    " (Actuel : ",
    "⚙️ Ajouter / retirer des Blocks et WS Cost Codes",
    "Nouveau block (ex. 170150) :",
    "➕ AJOUTER LE BLOCK",
    "Saisissez un nom de block !",
    "Block ajouté **",
    "Block à supprimer :",
    "🗑️ SUPPRIMER LE BLOCK",
    "Block supprimé **",
    "La liste des blocks est vide.",
    "Code (ex. WOS_07) :",
    "Nom de l'atelier / du service :",
    "➕ AJOUTER LE COST CODE",
    "Saisissez le code et le nom !",
    "Cost code à supprimer :",
    "🗑️ SUPPRIMER LE COST CODE",
    "La liste des cost codes est vide.",
    "Suppression impossible : **",
    "** compte(s) appartiennent encore à cet atelier. Déplacez-les d'abord vers un autre atelier dans 👥 Rôles et permissions.",
    "** de la liste. Les tâches existantes sont conservées.",
    "** est déjà dans la liste !",
    "Code **",
    "Ajouté **",
    "Supprimé **",
    "Description :",
    "Créé par :",
    "Date de création :",
    "Responsable :",
    "Remarque :",
    "— Aucun —",
    "📌 Cette tâche appartiendra à votre équipe. Allez dans 📋 Attribuer des tâches pour la donner à un ouvrier.",
    "💾 ENREGISTRER LA TÂCHE",
    "Remplissez les champs obligatoires : Task ID et Task Name !",
    "La date de fin prévue ne peut pas être avant la date de début !",
    "** existe déjà (peut-être dans la Corbeille). Utilisez un autre ID !",
    "Tâche enregistrée **",
    "⚠️ PAS ENCORE DE WS COST CODE : un Foreman / WOS Manager doit en ajouter un ci-dessus avant d'ajouter des tâches !",
    "👑 Choisissez une tâche et attribuez-la à un **Team Leader** ou **Foreman**. Ils la transmettront aux ouvriers de leur équipe.",
    "Aucune tâche pour l'instant. Allez dans ➕ Ajouter une tâche pour en créer une.",
    "Aucun Team Leader / Foreman pour l'instant. Attribuez d'abord des rôles dans 👥 Rôles et permissions.",
    "Afficher seulement les tâches non attribuées",
    "Toutes les tâches sont attribuées !",
    "Choisissez une tâche :",
    "Attribuer à un Team Leader / Foreman :",
    "Attribuer à un ouvrier :",
    "— Non attribuée —",
    "Non attribuée",
    "💾 ENREGISTRER L'ATTRIBUTION",
    "Tâche attribuée à **",
    "Votre équipe n'a pas encore d'ouvriers. Allez dans 👥 Mon équipe pour créer une équipe et ajouter des ouvriers.",
    "On ne vous a encore donné aucune tâche.",
    "Il faut des ouvriers dans votre équipe avant d'attribuer des tâches.",
    "(sans nom)",
    "⚠️ Votre compte n'a pas encore d'atelier. Demandez à un WOS Manager d'en attribuer un avant de créer une équipe.",
    "Vous n'avez pas encore d'équipe. Donnez-lui un nom pour la créer.",
    "Nom de l'équipe *",
    "ex. Équipe soudure Block 170",
    "➕ CRÉER L'ÉQUIPE",
    "Saisissez un nom d'équipe !",
    "Le nom de l'équipe peut avoir 60 caractères au maximum !",
    "Équipe créée **",
    "✏️ Renommer l'équipe",
    "Nouveau nom de l'équipe :",
    "💾 ENREGISTRER LE NOM",
    "Nom d'équipe invalide (1–60 caractères) !",
    "Équipe renommée !",
    "🗑️ Supprimer l'équipe",
    "Supprimer l'équipe : tous les ouvriers quittent l'équipe et leurs tâches ne sont plus attribuées. Les tâches restent à vous et vous pourrez créer une nouvelle équipe plus tard.",
    "Je suis sûr de vouloir supprimer cette équipe",
    "🗑️ SUPPRIMER L'ÉQUIPE",
    "Équipe supprimée.",
    "L'équipe n'a pas encore d'ouvriers.",
    "Plus aucun ouvrier dans l'atelier ",
    " sans équipe.",
    "Choisissez des ouvriers (même atelier, sans équipe) :",
    "➕ AJOUTER À L'ÉQUIPE",
    "Choisissez au moins 1 ouvrier !",
    " ouvrier(s) à l'équipe !",
    "Ajouté ",
    "Aucun membre à retirer.",
    "Choisissez un ouvrier :",
    "➖ RETIRER DE L'ÉQUIPE",
    "Ouvrier retiré de l'équipe. Ses tâches d'équipe ne sont plus attribuées.",
    "Vous n'êtes pas encore dans une équipe. Un Team Leader / Foreman de votre atelier peut vous ajouter.",
    "(Équipe sans nom)",
    "**Chef d'équipe :** ",
    "En partant, les tâches d'équipe qui vous sont attribuées ne le seront plus.",
    "Je suis sûr de vouloir quitter cette équipe",
    "🚪 QUITTER L'ÉQUIPE",
    "Vous avez quitté l'équipe.",
    "🧰 Mettre des tâches à la corbeille",
    "👤 Verrouiller des comptes",
    "Aucune tâche à modifier ou supprimer.",
    "Tâche à mettre à la corbeille :",
    "🗑️ Mettre la tâche à la corbeille",
    "Tâche mise à la corbeille !",
    "Aucun autre compte à supprimer.",
    "Compte à verrouiller :",
    "🗑️ Verrouiller ce compte",
    "Compte verrouillé et mis à la corbeille !",
    "🛡️ L'Admin attribue seulement le rôle **WOS Manager** et choisit l'**Atelier** qu'il dirige. Les rôles Worker / Team Leader / Foreman sont donnés par les WOS Managers.",
    "👑 Vous pouvez donner les rôles **Worker / Team Leader / Foreman**, choisir l'**Atelier** et placer des ouvriers dans l'**équipe** d'un Team Leader / Foreman.",
    "Aucun compte à qui attribuer un rôle.",
    "Compte à modifier :",
    "Action :",
    "👑 Nommer WOS Manager",
    "⛔ Retirer le rôle (retour à Pending)",
    "Atelier dirigé par ce manager :",
    "⚠️ Aucun atelier pour l'instant. Ajoutez-en dans ➕ Ajouter une tâche → ⚙️ Ajouter / retirer.",
    "💾 ENREGISTRER LE RÔLE",
    "💾 ENREGISTRER",
    "Nommé **WOS Manager** de l'atelier **",
    "Rôle retiré, le compte revient à Pending.",
    "Nouveau rôle :",
    "Équipe de (Team Leader / Foreman) :",
    "— Sans équipe —",
    "Cet atelier n'a pas encore de Team Leader / Foreman.",
    "Mis à jour : **",
    "Nom complet",
    "Rôle",
    "Équipe de",
    "🧰 Corbeille des tâches",
    "👤 Corbeille des comptes",
    "La corbeille des tâches est vide.",
    "La corbeille des comptes est vide.",
    "Choisissez une tâche :",
    "♻️ RESTAURER LA TÂCHE",
    "💥 SUPPRIMER LA TÂCHE DÉFINITIVEMENT",
    "Tâche restaurée !",
    "Tâche supprimée définitivement !",
    "Choisissez un compte :",
    "♻️ DÉVERROUILLER / RESTAURER LE COMPTE",
    "💥 SUPPRIMER LE COMPTE DÉFINITIVEMENT",
    "Compte restauré !",
    "Compte supprimé définitivement de la base de données !",
    "Pas encore de données de tâches pour un rapport. Ajoutez d'abord des tâches !",
    "Total des tâches",
    "Total des tâches",
    "Avancement moyen (%)",
    "Avancement moyen",
    "Terminées (100%)",
    "En cours",
    "Nombre de tâches",
    "Statut",
    "📥 Télécharger le rapport des tâches (CSV/Excel)",
    "Mot de passe actuel :",
    "Nouveau mot de passe (au moins 8 caractères) :",
    "Répétez le nouveau mot de passe :",
    "💾 CHANGER LE MOT DE PASSE",
    "Le mot de passe actuel est incorrect !",
    "Le nouveau mot de passe doit contenir au moins 8 caractères !",
    "Vous ne pouvez pas réutiliser le mot de passe par défaut !",
    "✅ Mot de passe changé ! Les autres appareils ont été déconnectés.",
    "⚠️ Ce compte utilise encore le mot de passe par défaut 'admin123'. Allez dans 🔑 Changer le mot de passe et changez-le MAINTENANT !",
    "Pas commencée (0%)",
    "En cours (1-50%)",
    "Presque finie (51-99%)",
    "Terminée (100%)",
    "Nombre",
    "Nouveau mot de passe :",
    "💬 Messages",
    "📢 Discussion générale",
    "#### 👥 Personnes",
    "🔎 Chercher un compte",
    "Aucun compte trouvé.",
    "Tout le monde dans l'app voit ce canal.",
    "Discussion privée avec ",
    "Écrivez un message...",
    "➕ Créer un groupe",
    "Nom du groupe",
    "ex. Équipe soudure Block 170",
    "Ajouter des personnes au groupe",
    "✨ CRÉER LE GROUPE",
    "Donnez un nom au groupe !",
    "Choisissez au moins 1 personne !",
    "Groupe créé **",
    "#### 💬 Groupes",
    "⚙️ Membres du groupe",
    "➕ AJOUTER AU GROUPE",
    " personnes au groupe !",
    "Retirer du groupe",
    "➖ RETIRER DU GROUPE",
    "Retiré du groupe.",
    "🗑️ SUPPRIMER LE GROUPE",
    "Seul le créateur du groupe peut ajouter / retirer des membres.",
    "🚪 QUITTER LE GROUPE",
    "🎥 Réunions en ligne",
    "### 🔑 Rejoindre une réunion",
    "ID de la salle",
    "Mot de passe de la salle",
    "🎥 REJOINDRE",
    "Trop d'essais erronés. Rechargez la page et réessayez plus tard.",
    "ID ou mot de passe incorrect, ou la réunion est terminée !",
    "✅ Salle trouvée : **",
    "Organisateur",
    "🎥 OUVRIR LA RÉUNION (caméra + micro)",
    "La réunion s'ouvre dans un nouvel onglet. La première fois, le navigateur demandera l'accès à la caméra et au micro — touchez Autoriser.",
    "### ➕ Créer une nouvelle réunion",
    "Nom de la réunion *",
    "ex. Relève de l'équipe du matin",
    "Date",
    "Heure",
    "Inviter des personnes (envoie ID + mot de passe via 💬 Messages)",
    "📢 Annoncer dans la discussion générale",
    "🎥 CRÉER LA RÉUNION",
    "Saisissez un nom de réunion !",
    "🎥 Invitation à une réunion",
    "🕒 Heure",
    "🔢 ID de la salle",
    "🔑 Mot de passe",
    "👉 Allez dans 🎥 Réunions en ligne pour rejoindre.",
    "🎉 Réunion créée **",
    "### 📋 Mes réunions",
    "Vous n'avez pas encore de réunions.",
    "▶️ DÉMARRER",
    "⛔ TERMINER",
    "Seuls les Foremen et Team Leaders peuvent créer des réunions. Vous pouvez rejoindre avec un ID et un mot de passe.",
    "Choisir des personnes...",
    "👑 Admin du groupe",
    "🗑️ Supprimer des messages (admin)",
    "Aucun message à supprimer.",
    "Message à supprimer",
    "🗑️ SUPPRIMER LE MESSAGE",
    "Message supprimé.",
    "Je suis sûr de vouloir supprimer tous les messages",
    "🧹 SUPPRIMER TOUS LES MESSAGES",
    "Tous les messages ont été supprimés.",
    "Inviter des personnes (elles verront l'invitation dans 🎥 Réunions en ligne)",
    "### 📨 Vos invitations",
    "🎥 REJOINDRE LA RÉUNION",
    "💬 Messages et réunions",
    "Paramètres",
    "À propos",
    "Connexion / Inscription",
    "Logiciel de gestion des travaux pour chantiers navals",
    "ShipControl aide les équipes du chantier naval à attribuer le travail, suivre l'avancement et rester en contact au même endroit — sur ordinateur et téléphone.",
    "📋 Gestion des tâches",
    "Créez des tâches par Block, Deck et Frame et suivez l'avancement pour cent par pour cent.",
    "👥 Équipes et rôles",
    "Admin, WOS Manager, Foreman, Team Leader et Worker — chacun voit exactement son propre travail.",
    "💬 Messages et réunions en ligne",
    "Discussion générale, groupes et salles de visioconférence avec ID + mot de passe.",
    "🌐 15 langues et voix",
    "Vietnamien, anglais, chinois, japonais, coréen, thaï, indonésien, espagnol, français, allemand, russe et plus — et parlez au lieu de taper sur téléphone.",
    "Pas encore de compte ? Inscrivez-vous et attendez qu'un responsable vous donne un rôle.",
    "🎬 Du chantier à la mer",
    "Regardez un navire se construire, des premiers blocs d'acier jusqu'au lancement.",
    "La vidéo ne se lit pas ? ",
    "▶ Ouvrir sur YouTube",
    "Créer un groupe",
    "Pas encore de groupe — touchez ➕ pour en créer un.",
    "Ajouter",
    "Retirer",
    "Ajouter des ouvriers",
    "Retirer de l'équipe",
    "Ajouter un atelier",
    "Retirer un atelier",
    "Supprimer le groupe",
    "Quitter le groupe",
    "🌊 Océan",
    "🌅 Coucher de soleil",
    "🌲 Forêt",
    "⚪ Minimaliste",
    "👑 Royal",
    "📸 Rapport photo hebdomadaire",
    "📈 Statistiques",
    "Vous n'avez encore aucune tâche à rapporter.",
    "Début",
    "Fin prévue",
    "Avancement",
    "Terminée",
    "Semaine en cours",
    "Ajouter des photos",
    "Semaine",
    "Choisir des photos sur l'ordinateur",
    "Note sur la photo (facultatif)",
    "📤 ENVOYER",
    "Choisissez au moins 1 photo !",
    "Impossible de lire ces photos (ou plus de 15 Mo) : ",
    "Pas encore de photos pour cette semaine.",
    "📥 Télécharger le rapport photo (HTML)",
    "Ouvrez le fichier et appuyez sur Ctrl+P pour l'imprimer ou l'enregistrer en PDF.",
    "💡 Maintenez Ctrl (ou Maj) pendant la sélection pour choisir plusieurs photos à la fois. Appuyez de nouveau sur Upload pour en ajouter.",
    "📊 Télécharger le PowerPoint (.pptx)",
    "◀ Précédent",
    "Suivant ▶",
    "Nombre de semaines",
    "Nombre de photos",
    "En cours",
    "photos",
    "Bilan",
]
_DE_LIST = [
    "🚢 SHIPCONTROL - SCHIFFBAU-ARBEITSVERWALTUNG",
    "🔐 Anmelden",
    "Bitte melde dich an oder registriere dich, um fortzufahren.",
    "🧰 Aufgabenübersicht",
    "➕ Aufgabe hinzufügen",
    "📋 Aufgaben zuweisen",
    "👥 Mein Team",
    "✏️ Bearbeiten / Löschen",
    "👥 Rollen & Rechte",
    "🗑️ Papierkorb",
    "📊 Berichte",
    "🔑 Passwort ändern",
    "🚪 Abmelden",
    "⚙️ Einstellungen",
    "📲 Gerät",
    "📲 Du verwendest:",
    "💻 Computer",
    "📱 Handy",
    "🎤 Tippe auf das Mikrofon in einem Textfeld, um zu sprechen statt zu tippen.",
    "🎤 Handymodus: Tippe auf das Mikrofon im Benutzernamen-Feld, um zu sprechen statt zu tippen.",
    "Dieser Browser unterstützt noch keine Spracheingabe. Bitte nutze Chrome (Android) oder Safari (iPhone).",
    "Der Mikrofonzugriff ist blockiert. Erlaube das Mikrofon für diese Seite in den Browser-Einstellungen.",
    "🏭 Werkstätten: anzeigen / hinzufügen / löschen",
    "Noch keine Werkstätten. Füge unten die erste hinzu.",
    "#### ➕ Neue Werkstatt hinzufügen",
    "Werkstatt-Code (z. B. WOS_07):",
    "➕ WERKSTATT HINZUFÜGEN",
    "Code max. 20 Zeichen, Name max. 80 Zeichen!",
    "#### 🗑️ Werkstatt löschen",
    "Zu löschende Werkstatt:",
    "🗑️ WERKSTATT LÖSCHEN",
    "Werkstattname",
    "Konten",
    "Aufgaben",
    "🎨 Design",
    "✨ Modern",
    "🚀 Futuristisch",
    "🎈 Verspielt",
    "📜 Nostalgisch",
    "Das Futuristisch-Design ist immer dunkel.",
    ">Konto<",
    "🔒 Dein Konto wurde gerade in einem anderen Tab oder auf einem anderen Gerät geöffnet, daher wurdest du hier abgemeldet. ",
    "Wenn du das nicht warst, melde dich erneut an und ändere sofort dein Passwort.",
    ">Rolle<",
    ">Angemeldet<",
    "🌙 Dunkelmodus",
    "🌐 Sprache",
    "📋 Fortschrittsübersicht der Aufgaben",
    "➕ Neue Aufgabe hinzufügen",
    "✏️ Bearbeiten & Löschen",
    "👥 Konten & Rollen (Rollenliste)",
    "📊 Fortschrittsberichte & Statistiken",
    "🗑️ Papierkorb & Wiederherstellen",
    "### ⚡ Schnelle Fortschrittsaktualisierung",
    "### 🏭 Statistik nach Werkstatt",
    "### 📈 Verteilung des Fortschritts",
    "### 📌 Aufgaben an Arbeiter zuweisen",
    "### 📥 Berichtsdaten herunterladen",
    "### 🔄 Kontorolle ändern",
    "#### ➕ Arbeiter hinzufügen",
    "#### ➖ Aus dem Team entfernen",
    "#### 👷 Mitglieder",
    "#### 🚪 Team verlassen",
    "🔑 Anmelden",
    "📝 Registrieren",
    "Benutzername:",
    "Passwort:",
    "🚀 ANMELDEN",
    "✨ JETZT REGISTRIEREN",
    "Vollständiger Name:",
    "Neuer Benutzername:",
    "Passwort bestätigen:",
    "Bitte Benutzername und Passwort eingeben!",
    "Benutzername oder Passwort falsch, oder das Konto ist gesperrt/gelöscht!",
    "⏳ Dein Konto wartet darauf, dass ein Admin / WOS Manager eine Rolle zuweist. Bitte komm später wieder!",
    "Willkommen ",
    ") ist zurück!",
    "Bitte alle Felder ausfüllen!",
    "Das Passwort muss mindestens 8 Zeichen haben!",
    "Benutzername oder Name ist zu lang!",
    "Die Passwörter stimmen nicht überein!",
    "Dieser Benutzername existiert bereits, bitte wähle einen anderen!",
    "🎉 Registriert! Dein Konto wartet auf die Rollenfreigabe durch einen WOS Manager.",
    "👷 Das sind die dir zugewiesenen Aufgaben.",
    "Nur Aufgaben meines Teams anzeigen",
    "Noch keine Aufgaben anzuzeigen.",
    "Du hast keine Aufgaben zum Aktualisieren.",
    "Aufgabe zum Aktualisieren wählen:",
    "Neuer Fortschritt (%)",
    "Arbeitsnotiz (Remark):",
    "🚀 FORTSCHRITT AKTUALISIEREN",
    "Fortschritt aktualisiert!",
    " (Aktuell: ",
    "⚙️ Blocks und WS Cost Codes hinzufügen / entfernen",
    "Neuer Block (z. B. 170150):",
    "➕ BLOCK HINZUFÜGEN",
    "Bitte einen Block-Namen eingeben!",
    "Block hinzugefügt **",
    "Zu löschender Block:",
    "🗑️ BLOCK LÖSCHEN",
    "Block gelöscht **",
    "Die Block-Liste ist leer.",
    "Code (z. B. WOS_07):",
    "Name der Werkstatt / Abteilung:",
    "➕ COST CODE HINZUFÜGEN",
    "Bitte Code und Namen eingeben!",
    "Zu löschender Cost Code:",
    "🗑️ COST CODE LÖSCHEN",
    "Die Cost-Code-Liste ist leer.",
    "Löschen nicht möglich: **",
    "** Konto/Konten gehören noch zu dieser Werkstatt. Verschiebe sie zuerst in 👥 Rollen & Rechte in eine andere Werkstatt.",
    "** aus der Liste. Bestehende Aufgaben bleiben erhalten.",
    "** ist bereits in der Liste!",
    "Code **",
    "Hinzugefügt **",
    "Gelöscht **",
    "Beschreibung:",
    "Erstellt von:",
    "Erstellungsdatum:",
    "Verantwortlich:",
    "Bemerkung:",
    "— Keiner —",
    "📌 Diese Aufgabe gehört zu deinem Team. Gehe zu 📋 Aufgaben zuweisen, um sie einem Arbeiter zu geben.",
    "💾 NEUE AUFGABE SPEICHERN",
    "Bitte Pflichtfelder ausfüllen: Task ID und Task Name!",
    "Das geplante Enddatum darf nicht vor dem Startdatum liegen!",
    "** existiert bereits (evtl. im Papierkorb). Bitte eine andere ID verwenden!",
    "Aufgabe gespeichert **",
    "⚠️ NOCH KEIN WS COST CODE: Ein Foreman / WOS Manager muss oben erst einen hinzufügen, bevor Aufgaben angelegt werden können!",
    "👑 Wähle eine Aufgabe und weise sie einem **Team Leader** oder **Foreman** zu. Diese geben sie an die Arbeiter ihres Teams weiter.",
    "Noch keine Aufgaben. Gehe zu ➕ Aufgabe hinzufügen, um eine anzulegen.",
    "Noch keine Team Leader / Foremen. Weise zuerst Rollen in 👥 Rollen & Rechte zu.",
    "Nur nicht zugewiesene Aufgaben anzeigen",
    "Alle Aufgaben sind zugewiesen!",
    "Aufgabe wählen:",
    "Zuweisen an Team Leader / Foreman:",
    "Zuweisen an Arbeiter:",
    "— Nicht zugewiesen —",
    "Nicht zugewiesen",
    "💾 ZUWEISUNG SPEICHERN",
    "Aufgabe zugewiesen an **",
    "Dein Team hat noch keine Arbeiter. Gehe zu 👥 Mein Team, um ein Team zu erstellen und Arbeiter hinzuzufügen.",
    "Dir wurden noch keine Aufgaben gegeben.",
    "Du brauchst Arbeiter im Team, bevor du Aufgaben zuweisen kannst.",
    "(ohne Namen)",
    "⚠️ Dein Konto hat noch keine Werkstatt. Bitte einen WOS Manager, eine zuzuweisen, bevor du ein Team erstellst.",
    "Du hast noch kein Team. Gib ihm einen Namen, um es zu erstellen.",
    "Teamname *",
    "z. B. Schweißteam Block 170",
    "➕ TEAM ERSTELLEN",
    "Bitte einen Teamnamen eingeben!",
    "Der Teamname darf höchstens 60 Zeichen haben!",
    "Team erstellt **",
    "✏️ Team umbenennen",
    "Neuer Teamname:",
    "💾 NAMEN SPEICHERN",
    "Ungültiger Teamname (1–60 Zeichen)!",
    "Team umbenannt!",
    "🗑️ Team löschen",
    "Beim Löschen des Teams verlassen alle Arbeiter das Team und ihre Aufgaben werden nicht mehr zugewiesen. Die Aufgaben bleiben bei dir und du kannst später ein neues Team erstellen.",
    "Ich bin sicher, dass ich dieses Team löschen will",
    "🗑️ TEAM LÖSCHEN",
    "Team gelöscht.",
    "Das Team hat noch keine Arbeiter.",
    "Keine Arbeiter mehr in Werkstatt ",
    " ohne Team.",
    "Arbeiter wählen (gleiche Werkstatt, noch ohne Team):",
    "➕ ZUM TEAM HINZUFÜGEN",
    "Wähle mindestens 1 Arbeiter!",
    " Arbeiter zum Team hinzugefügt!",
    "Es wurden ",
    "Keine Mitglieder zum Entfernen.",
    "Arbeiter wählen:",
    "➖ AUS DEM TEAM ENTFERNEN",
    "Arbeiter aus dem Team entfernt. Seine Teamaufgaben sind jetzt nicht mehr zugewiesen.",
    "Du bist noch in keinem Team. Ein Team Leader / Foreman deiner Werkstatt kann dich hinzufügen.",
    "(Team ohne Namen)",
    "**Teamleiter:** ",
    "Wenn du gehst, werden dir zugewiesene Teamaufgaben nicht mehr zugewiesen.",
    "Ich bin sicher, dass ich dieses Team verlassen will",
    "🚪 TEAM VERLASSEN",
    "Du hast das Team verlassen.",
    "🧰 Aufgaben in den Papierkorb",
    "👤 Benutzerkonten sperren",
    "Keine Aufgaben zum Bearbeiten oder Löschen.",
    "Aufgabe in den Papierkorb verschieben:",
    "🗑️ Aufgabe in den Papierkorb",
    "Aufgabe in den Papierkorb verschoben!",
    "Keine anderen Konten zum Löschen.",
    "Zu sperrendes Konto:",
    "🗑️ Dieses Konto sperren",
    "Konto gesperrt und in den Papierkorb verschoben!",
    "🛡️ Der Admin vergibt nur die Rolle **WOS Manager** und wählt die **Werkstatt**, die dieser leitet. Die Rollen Worker / Team Leader / Foreman vergeben die WOS Manager.",
    "👑 Du kannst die Rollen **Worker / Team Leader / Foreman** vergeben, die **Werkstatt** wählen und Arbeiter in das **Team** eines Team Leaders / Foremans setzen.",
    "Keine Konten für eine Rollenvergabe.",
    "Zu änderndes Konto:",
    "Aktion:",
    "👑 Zum WOS Manager machen",
    "⛔ Rolle entziehen (zurück zu Pending)",
    "Werkstatt, die dieser Manager leitet:",
    "⚠️ Noch keine Werkstätten. Füge sie in ➕ Aufgabe hinzufügen → ⚙️ Hinzufügen / Entfernen hinzu.",
    "💾 ROLLE SPEICHERN",
    "💾 ÄNDERUNGEN SPEICHERN",
    "Zum **WOS Manager** der Werkstatt ernannt **",
    "Rolle entzogen, das Konto ist wieder Pending.",
    "Neue Rolle:",
    "Team von (Team Leader / Foreman):",
    "— Kein Team —",
    "Diese Werkstatt hat noch keinen Team Leader / Foreman.",
    "Aktualisiert: **",
    "Vollständiger Name",
    "Rolle",
    "Team von",
    "🧰 Aufgaben-Papierkorb",
    "👤 Konten-Papierkorb",
    "Der Aufgaben-Papierkorb ist leer.",
    "Der Konten-Papierkorb ist leer.",
    "Aufgabe wählen:",
    "♻️ AUFGABE WIEDERHERSTELLEN",
    "💥 AUFGABE ENDGÜLTIG LÖSCHEN",
    "Aufgabe wiederhergestellt!",
    "Aufgabe endgültig gelöscht!",
    "Konto wählen:",
    "♻️ KONTO ENTSPERREN / WIEDERHERSTELLEN",
    "💥 KONTO ENDGÜLTIG LÖSCHEN",
    "Konto wiederhergestellt!",
    "Konto endgültig aus der Datenbank gelöscht!",
    "Noch keine Aufgabendaten für einen Bericht. Bitte zuerst Aufgaben hinzufügen!",
    "Aufgaben gesamt",
    "Aufgaben gesamt",
    "Durchschnittlicher Fortschritt (%)",
    "Durchschnittlicher Fortschritt",
    "Erledigt (100%)",
    "In Arbeit",
    "Anzahl der Aufgaben",
    "Status",
    "📥 Aufgabenbericht herunterladen (CSV/Excel)",
    "Aktuelles Passwort:",
    "Neues Passwort (mindestens 8 Zeichen):",
    "Neues Passwort wiederholen:",
    "💾 PASSWORT ÄNDERN",
    "Das aktuelle Passwort ist falsch!",
    "Das neue Passwort muss mindestens 8 Zeichen haben!",
    "Das Standardpasswort darf nicht wiederverwendet werden!",
    "✅ Passwort geändert! Andere Geräte wurden abgemeldet.",
    "⚠️ Dieses Konto nutzt noch das Standardpasswort 'admin123'. Gehe zu 🔑 Passwort ändern und ändere es SOFORT!",
    "Nicht begonnen (0%)",
    "In Arbeit (1-50%)",
    "Fast fertig (51-99%)",
    "Fertig (100%)",
    "Anzahl",
    "Neues Passwort:",
    "💬 Nachrichten",
    "📢 Allgemeiner Chat",
    "#### 👥 Personen",
    "🔎 Konto suchen",
    "Kein Konto gefunden.",
    "Alle in der App sehen diesen Kanal.",
    "Privatchat mit ",
    "Nachricht eingeben...",
    "➕ Gruppenchat erstellen",
    "Gruppenname",
    "z. B. Schweißtrupp Block 170",
    "Personen zur Gruppe hinzufügen",
    "✨ GRUPPE ERSTELLEN",
    "Bitte gib der Gruppe einen Namen!",
    "Wähle mindestens 1 Person!",
    "Gruppe erstellt **",
    "#### 💬 Gruppenchats",
    "⚙️ Gruppenmitglieder",
    "➕ ZUR GRUPPE HINZUFÜGEN",
    " Personen zur Gruppe hinzugefügt!",
    "Aus der Gruppe entfernen",
    "➖ AUS DER GRUPPE ENTFERNEN",
    "Aus der Gruppe entfernt.",
    "🗑️ GRUPPE LÖSCHEN",
    "Nur der Ersteller der Gruppe kann Mitglieder hinzufügen / entfernen.",
    "🚪 GRUPPE VERLASSEN",
    "🎥 Online-Meetings",
    "### 🔑 Einem Meeting beitreten",
    "Raum-ID",
    "Raum-Passwort",
    "🎥 BEITRETEN",
    "Zu viele falsche Versuche. Lade die Seite neu und versuche es später erneut.",
    "Raum-ID oder Passwort falsch, oder das Meeting ist beendet!",
    "✅ Raum gefunden: **",
    "Gastgeber",
    "🎥 MEETING ÖFFNEN (Kamera + Mikro)",
    "Das Meeting öffnet sich in einem neuen Tab. Beim ersten Mal fragt der Browser nach Kamera und Mikrofon — tippe auf Zulassen.",
    "### ➕ Neues Meeting erstellen",
    "Name des Meetings *",
    "z. B. Schichtübergabe Frühschicht",
    "Datum",
    "Uhrzeit",
    "Personen einladen (sendet ID + Passwort per 💬 Nachrichten)",
    "📢 Im allgemeinen Chat ankündigen",
    "🎥 MEETING ERSTELLEN",
    "Bitte einen Namen für das Meeting eingeben!",
    "🎥 Meeting-Einladung",
    "🕒 Uhrzeit",
    "🔢 Raum-ID",
    "🔑 Passwort",
    "👉 Gehe zu 🎥 Online-Meetings, um beizutreten.",
    "🎉 Meeting erstellt **",
    "### 📋 Meine Meetings",
    "Du hast noch keine Meetings.",
    "▶️ STARTEN",
    "⛔ BEENDEN",
    "Nur Foremen und Team Leader können Meetings erstellen. Du kannst mit Raum-ID und Passwort beitreten.",
    "Personen wählen...",
    "👑 Gruppen-Admin",
    "🗑️ Nachrichten löschen (Admin)",
    "Keine Nachrichten zum Löschen.",
    "Zu löschende Nachricht",
    "🗑️ NACHRICHT LÖSCHEN",
    "Nachricht gelöscht.",
    "Ich bin sicher, dass ich alle Nachrichten löschen will",
    "🧹 ALLE NACHRICHTEN LÖSCHEN",
    "Alle Nachrichten gelöscht.",
    "Personen einladen (sie sehen die Einladung in 🎥 Online-Meetings)",
    "### 📨 Deine Meeting-Einladungen",
    "🎥 MEETING BEITRETEN",
    "💬 Nachrichten & Meetings",
    "Einstellungen",
    "Über uns",
    "Anmelden / Registrieren",
    "Arbeitsverwaltungssoftware für Werften",
    "ShipControl hilft Werft-Teams, Arbeit zu verteilen, den Fortschritt zu verfolgen und an einem Ort in Kontakt zu bleiben — am Computer und am Handy.",
    "📋 Aufgabenverwaltung",
    "Erstelle Aufgaben nach Block, Deck und Frame und verfolge den Fortschritt Prozent für Prozent.",
    "👥 Teams & Rollen",
    "Admin, WOS Manager, Foreman, Team Leader und Worker — jeder sieht genau seine eigene Arbeit.",
    "💬 Nachrichten & Online-Meetings",
    "Allgemeiner Chat, Gruppenchats und Videomeeting-Räume mit ID + Passwort.",
    "🌐 15 Sprachen & Sprache",
    "Vietnamesisch, Englisch, Chinesisch, Japanisch, Koreanisch, Thai, Indonesisch, Spanisch, Französisch, Deutsch, Russisch und mehr — und am Handy sprechen statt tippen.",
    "Noch kein Konto? Registriere dich und warte, bis dir ein Manager eine Rolle gibt.",
    "🎬 Von der Werft aufs Meer",
    "Sieh zu, wie ein Schiff von den ersten Stahlblöcken bis zum Stapellauf gebaut wird.",
    "Video lässt sich nicht abspielen? ",
    "▶ Auf YouTube öffnen",
    "Gruppenchat erstellen",
    "Noch keine Gruppen — tippe auf ➕, um eine zu erstellen.",
    "Hinzufügen",
    "Entfernen",
    "Arbeiter hinzufügen",
    "Aus dem Team entfernen",
    "Werkstatt hinzufügen",
    "Werkstatt entfernen",
    "Gruppe löschen",
    "Gruppe verlassen",
    "🌊 Ozean",
    "🌅 Sonnenuntergang",
    "🌲 Wald",
    "⚪ Minimal",
    "👑 Königlich",
    "📸 Wöchentlicher Fotobericht",
    "📈 Statistik",
    "Du hast noch keine Aufgaben für einen Bericht.",
    "Start",
    "Geplantes Ende",
    "Fortschritt",
    "Fertig",
    "Aktuelle Woche",
    "Fotos hochladen",
    "Woche",
    "Fotos vom Computer auswählen",
    "Notiz zum Foto (optional)",
    "📤 HOCHLADEN",
    "Wähle mindestens 1 Foto!",
    "Diese Fotos konnten nicht gelesen werden (oder sind größer als 15 MB): ",
    "Für diese Woche gibt es noch keine Fotos.",
    "📥 Fotobericht herunterladen (HTML)",
    "Datei öffnen und Strg+P drücken, um zu drucken oder als PDF zu speichern.",
    "💡 Halte beim Auswählen Strg (oder Umschalt) gedrückt, um mehrere Fotos auf einmal zu wählen. Drücke erneut auf Upload, um weitere hinzuzufügen.",
    "📊 PowerPoint herunterladen (.pptx)",
    "◀ Zurück",
    "Weiter ▶",
    "Wochen gesamt",
    "Fotos gesamt",
    "In Arbeit",
    "Fotos",
    "Zusammenfassung",
]
_RU_LIST = [
    "🚢 SHIPCONTROL - УПРАВЛЕНИЕ СУДОСТРОИТЕЛЬНЫМИ РАБОТАМИ",
    "🔐 Вход",
    "Войдите или зарегистрируйтесь, чтобы продолжить.",
    "🧰 Доска задач",
    "➕ Добавить задачу",
    "📋 Назначить задачи",
    "👥 Моя бригада",
    "✏️ Изменить / Удалить",
    "👥 Роли и права",
    "🗑️ Корзина",
    "📊 Отчёты",
    "🔑 Сменить пароль",
    "🚪 Выйти",
    "⚙️ Настройки",
    "📲 Устройство",
    "📲 Вы используете:",
    "💻 Компьютер",
    "📱 Телефон",
    "🎤 Нажмите на микрофон в текстовом поле, чтобы говорить вместо набора текста.",
    "🎤 Режим телефона: нажмите на микрофон в поле логина, чтобы говорить вместо набора текста.",
    "Этот браузер пока не поддерживает голосовой ввод. Используйте Chrome (Android) или Safari (iPhone).",
    "Доступ к микрофону заблокирован. Разрешите микрофон для этого сайта в настройках браузера.",
    "🏭 Цеха: просмотр / добавление / удаление",
    "Цехов пока нет. Добавьте первый ниже.",
    "#### ➕ Добавить новый цех",
    "Код цеха (напр. WOS_07):",
    "➕ ДОБАВИТЬ ЦЕХ",
    "Код — не более 20 символов, название — не более 80 символов!",
    "#### 🗑️ Удалить цех",
    "Цех для удаления:",
    "🗑️ УДАЛИТЬ ЦЕХ",
    "Название цеха",
    "Аккаунты",
    "Задачи",
    "🎨 Тема",
    "✨ Современная",
    "🚀 Футуристическая",
    "🎈 Игривая",
    "📜 Ретро",
    "Футуристическая тема всегда тёмная.",
    ">Аккаунт<",
    "🔒 Ваш аккаунт только что открыли в другой вкладке или на другом устройстве, поэтому здесь вы вышли из системы. ",
    "Если это были не вы, войдите снова и сразу смените пароль.",
    ">Роль<",
    ">В системе<",
    "🌙 Тёмный режим",
    "🌐 Язык",
    "📋 Доска прогресса задач",
    "➕ Добавить новую задачу",
    "✏️ Изменение и удаление",
    "👥 Аккаунты и роли (список ролей)",
    "📊 Отчёты о прогрессе и статистика",
    "🗑️ Корзина и восстановление",
    "### ⚡ Быстрое обновление прогресса",
    "### 🏭 Статистика по цехам",
    "### 📈 Распределение прогресса",
    "### 📌 Назначить задачи рабочим",
    "### 📥 Скачать данные отчёта",
    "### 🔄 Изменить роль аккаунта",
    "#### ➕ Добавить рабочих",
    "#### ➖ Убрать из бригады",
    "#### 👷 Участники",
    "#### 🚪 Покинуть бригаду",
    "🔑 Вход",
    "📝 Регистрация",
    "Логин:",
    "Пароль:",
    "🚀 ВОЙТИ",
    "✨ ЗАРЕГИСТРИРОВАТЬСЯ",
    "Полное имя:",
    "Новый логин:",
    "Подтвердите пароль:",
    "Введите логин и пароль!",
    "Неверный логин или пароль, либо аккаунт заблокирован/удалён!",
    "⏳ Ваш аккаунт ждёт, пока Admin / WOS Manager назначит роль. Зайдите позже!",
    "С возвращением, ",
    ")!",
    "Заполните все поля!",
    "Пароль должен быть не короче 8 символов!",
    "Логин или имя слишком длинные!",
    "Пароли не совпадают!",
    "Такой логин уже есть, выберите другой!",
    "🎉 Готово! Ваш аккаунт ждёт, пока WOS Manager одобрит роль.",
    "👷 Это задачи, назначенные вам.",
    "Показывать только задачи моей бригады",
    "Пока нет задач для показа.",
    "У вас нет задач для обновления.",
    "Выберите задачу для обновления:",
    "Новый прогресс (%)",
    "Заметка о работе (Remark):",
    "🚀 ОБНОВИТЬ ПРОГРЕСС",
    "Прогресс обновлён!",
    " (Сейчас: ",
    "⚙️ Добавить / убрать Blocks и WS Cost Codes",
    "Новый block (напр. 170150):",
    "➕ ДОБАВИТЬ BLOCK",
    "Введите название block!",
    "Добавлен block **",
    "Block для удаления:",
    "🗑️ УДАЛИТЬ BLOCK",
    "Удалён block **",
    "Список blocks пуст.",
    "Код (напр. WOS_07):",
    "Название цеха / отдела:",
    "➕ ДОБАВИТЬ COST CODE",
    "Введите код и название!",
    "Cost code для удаления:",
    "🗑️ УДАЛИТЬ COST CODE",
    "Список cost codes пуст.",
    "Нельзя удалить: **",
    "** аккаунт(ов) всё ещё относятся к этому цеху. Сначала переведите их в другой цех в 👥 Роли и права.",
    "** из списка. Существующие задачи сохранены.",
    "** уже есть в списке!",
    "Код **",
    "Добавлено **",
    "Удалено **",
    "Описание:",
    "Создал:",
    "Дата создания:",
    "Ответственный:",
    "Примечание:",
    "— Нет —",
    "📌 Эта задача будет принадлежать вашей бригаде. Перейдите в 📋 Назначить задачи, чтобы дать её рабочему.",
    "💾 СОХРАНИТЬ ЗАДАЧУ",
    "Заполните обязательные поля: Task ID и Task Name!",
    "Плановая дата окончания не может быть раньше даты начала!",
    "** уже существует (возможно, в Корзине). Используйте другой ID!",
    "Задача сохранена **",
    "⚠️ ПОКА НЕТ WS COST CODE: Foreman / WOS Manager должен добавить его выше, прежде чем добавлять задачи!",
    "👑 Выберите задачу и назначьте её **Team Leader** или **Foreman**. Они передадут её рабочим своей бригады.",
    "Задач пока нет. Перейдите в ➕ Добавить задачу, чтобы создать.",
    "Пока нет Team Leader / Foreman. Сначала назначьте роли в 👥 Роли и права.",
    "Показывать только неназначенные задачи",
    "Все задачи назначены!",
    "Выберите задачу:",
    "Назначить Team Leader / Foreman:",
    "Назначить рабочему:",
    "— Не назначено —",
    "Не назначено",
    "💾 СОХРАНИТЬ НАЗНАЧЕНИЕ",
    "Задача назначена **",
    "В вашей бригаде пока нет рабочих. Перейдите в 👥 Моя бригада, чтобы создать бригаду и добавить рабочих.",
    "Вам пока не дали ни одной задачи.",
    "Чтобы назначать задачи, в бригаде должны быть рабочие.",
    "(без имени)",
    "⚠️ У вашего аккаунта пока нет цеха. Попросите WOS Manager назначить его перед созданием бригады.",
    "У вас пока нет бригады. Дайте ей название, чтобы создать.",
    "Название бригады *",
    "напр. Сварочная бригада Block 170",
    "➕ СОЗДАТЬ БРИГАДУ",
    "Введите название бригады!",
    "Название бригады — не более 60 символов!",
    "Бригада создана **",
    "✏️ Переименовать бригаду",
    "Новое название бригады:",
    "💾 СОХРАНИТЬ НАЗВАНИЕ",
    "Неверное название бригады (1–60 символов)!",
    "Бригада переименована!",
    "🗑️ Удалить бригаду",
    "При удалении бригады все рабочие покидают её, а их задачи становятся неназначенными. Задачи остаются за вами, позже можно создать новую бригаду.",
    "Я уверен, что хочу удалить эту бригаду",
    "🗑️ УДАЛИТЬ БРИГАДУ",
    "Бригада удалена.",
    "В бригаде пока нет рабочих.",
    "Не осталось рабочих в цехе ",
    " без бригады.",
    "Выберите рабочих (тот же цех, без бригады):",
    "➕ ДОБАВИТЬ В БРИГАДУ",
    "Выберите хотя бы 1 рабочего!",
    " рабочих в бригаду!",
    "Добавлено ",
    "Некого убирать.",
    "Выберите рабочего:",
    "➖ УБРАТЬ ИЗ БРИГАДЫ",
    "Рабочий убран из бригады. Его задачи бригады теперь не назначены.",
    "Вы пока не в бригаде. Team Leader / Foreman вашего цеха может вас добавить.",
    "(Бригада без названия)",
    "**Бригадир:** ",
    "Если вы уйдёте, назначенные вам задачи бригады станут неназначенными.",
    "Я уверен, что хочу покинуть эту бригаду",
    "🚪 ПОКИНУТЬ БРИГАДУ",
    "Вы покинули бригаду.",
    "🧰 Переместить задачи в корзину",
    "👤 Заблокировать аккаунты",
    "Нет задач для изменения или удаления.",
    "Задача для переноса в корзину:",
    "🗑️ Переместить задачу в корзину",
    "Задача перемещена в корзину!",
    "Нет других аккаунтов для удаления.",
    "Аккаунт для блокировки:",
    "🗑️ Заблокировать этот аккаунт",
    "Аккаунт заблокирован и перемещён в корзину!",
    "🛡️ Admin назначает только роль **WOS Manager** и выбирает **Цех**, которым тот руководит. Роли Worker / Team Leader / Foreman назначают WOS Manager.",
    "👑 Вы можете давать роли **Worker / Team Leader / Foreman**, выбирать **Цех** и включать рабочих в **бригаду** Team Leader / Foreman.",
    "Нет аккаунтов для назначения ролей.",
    "Аккаунт для изменения:",
    "Действие:",
    "👑 Сделать WOS Manager",
    "⛔ Снять роль (вернуть в Pending)",
    "Цех, которым руководит этот менеджер:",
    "⚠️ Цехов пока нет. Добавьте их в ➕ Добавить задачу → ⚙️ Добавить / убрать.",
    "💾 СОХРАНИТЬ РОЛЬ",
    "💾 СОХРАНИТЬ ИЗМЕНЕНИЯ",
    "Назначен **WOS Manager** цеха **",
    "Роль снята, аккаунт снова в Pending.",
    "Новая роль:",
    "Бригада (Team Leader / Foreman):",
    "— Без бригады —",
    "В этом цехе пока нет Team Leader / Foreman.",
    "Обновлено: **",
    "Полное имя",
    "Роль",
    "Бригада",
    "🧰 Корзина задач",
    "👤 Корзина аккаунтов",
    "Корзина задач пуста.",
    "Корзина аккаунтов пуста.",
    "Выберите задачу:",
    "♻️ ВОССТАНОВИТЬ ЗАДАЧУ",
    "💥 УДАЛИТЬ ЗАДАЧУ НАВСЕГДА",
    "Задача восстановлена!",
    "Задача удалена навсегда!",
    "Выберите аккаунт:",
    "♻️ РАЗБЛОКИРОВАТЬ / ВОССТАНОВИТЬ АККАУНТ",
    "💥 УДАЛИТЬ АККАУНТ НАВСЕГДА",
    "Аккаунт восстановлен!",
    "Аккаунт удалён из базы данных навсегда!",
    "Пока нет данных о задачах для отчёта. Сначала добавьте задачи!",
    "Всего задач",
    "Всего задач",
    "Средний прогресс (%)",
    "Средний прогресс",
    "Завершено (100%)",
    "В работе",
    "Количество задач",
    "Статус",
    "📥 Скачать отчёт по задачам (CSV/Excel)",
    "Текущий пароль:",
    "Новый пароль (не менее 8 символов):",
    "Повторите новый пароль:",
    "💾 СМЕНИТЬ ПАРОЛЬ",
    "Текущий пароль неверен!",
    "Новый пароль должен быть не короче 8 символов!",
    "Нельзя снова использовать пароль по умолчанию!",
    "✅ Пароль изменён! На других устройствах выполнен выход.",
    "⚠️ Этот аккаунт всё ещё использует пароль по умолчанию 'admin123'. Перейдите в 🔑 Сменить пароль и смените его СЕЙЧАС!",
    "Не начато (0%)",
    "В работе (1-50%)",
    "Почти готово (51-99%)",
    "Готово (100%)",
    "Количество",
    "Новый пароль:",
    "💬 Сообщения",
    "📢 Общий чат",
    "#### 👥 Люди",
    "🔎 Найти аккаунт",
    "Аккаунт не найден.",
    "Этот канал видят все в приложении.",
    "Личный чат с ",
    "Введите сообщение...",
    "➕ Создать групповой чат",
    "Название группы",
    "напр. Сварщики Block 170",
    "Добавить людей в группу",
    "✨ СОЗДАТЬ ГРУППУ",
    "Дайте группе название!",
    "Выберите хотя бы 1 человека!",
    "Группа создана **",
    "#### 💬 Групповые чаты",
    "⚙️ Участники группы",
    "➕ ДОБАВИТЬ В ГРУППУ",
    " чел. в группу!",
    "Убрать из группы",
    "➖ УБРАТЬ ИЗ ГРУППЫ",
    "Убран из группы.",
    "🗑️ УДАЛИТЬ ГРУППУ",
    "Только создатель группы может добавлять / убирать участников.",
    "🚪 ПОКИНУТЬ ГРУППУ",
    "🎥 Онлайн-совещания",
    "### 🔑 Войти в совещание",
    "ID комнаты",
    "Пароль комнаты",
    "🎥 ВОЙТИ",
    "Слишком много неверных попыток. Обновите страницу и попробуйте позже.",
    "Неверный ID или пароль, либо совещание закончилось!",
    "✅ Комната найдена: **",
    "Организатор",
    "🎥 ОТКРЫТЬ СОВЕЩАНИЕ (камера + микрофон)",
    "Совещание откроется в новой вкладке. В первый раз браузер попросит доступ к камере и микрофону — нажмите «Разрешить».",
    "### ➕ Создать новое совещание",
    "Название совещания *",
    "напр. Передача утренней смены",
    "Дата",
    "Время",
    "Пригласить людей (отправит ID + пароль через 💬 Сообщения)",
    "📢 Объявить в общем чате",
    "🎥 СОЗДАТЬ СОВЕЩАНИЕ",
    "Введите название совещания!",
    "🎥 Приглашение на совещание",
    "🕒 Время",
    "🔢 ID комнаты",
    "🔑 Пароль",
    "👉 Перейдите в 🎥 Онлайн-совещания, чтобы войти.",
    "🎉 Совещание создано **",
    "### 📋 Мои совещания",
    "У вас пока нет совещаний.",
    "▶️ НАЧАТЬ",
    "⛔ ЗАВЕРШИТЬ",
    "Создавать совещания могут только Foreman и Team Leader. Вы можете войти по ID и паролю.",
    "Выберите людей...",
    "👑 Админ группы",
    "🗑️ Удалить сообщения (админ)",
    "Нет сообщений для удаления.",
    "Сообщение для удаления",
    "🗑️ УДАЛИТЬ СООБЩЕНИЕ",
    "Сообщение удалено.",
    "Я уверен, что хочу удалить все сообщения",
    "🧹 УДАЛИТЬ ВСЕ СООБЩЕНИЯ",
    "Все сообщения удалены.",
    "Пригласить людей (они увидят приглашение в 🎥 Онлайн-совещания)",
    "### 📨 Ваши приглашения на совещания",
    "🎥 ВОЙТИ В СОВЕЩАНИЕ",
    "💬 Сообщения и совещания",
    "Настройки",
    "О нас",
    "Вход / Регистрация",
    "Программа управления работами для верфей",
    "ShipControl помогает командам верфи распределять работу, следить за прогрессом и оставаться на связи в одном месте — на компьютере и телефоне.",
    "📋 Управление задачами",
    "Создавайте задачи по Block, Deck и Frame и отслеживайте прогресс процент за процентом.",
    "👥 Бригады и роли",
    "Admin, WOS Manager, Foreman, Team Leader и Worker — каждый видит именно свою работу.",
    "💬 Сообщения и онлайн-совещания",
    "Общий чат, групповые чаты и видеокомнаты с ID + паролем.",
    "🌐 15 языков и голос",
    "Вьетнамский, английский, китайский, японский, корейский, тайский, индонезийский, испанский, французский, немецкий, русский и другие — и говорите вместо набора текста на телефоне.",
    "Ещё нет аккаунта? Зарегистрируйтесь и дождитесь, пока руководитель даст вам роль.",
    "🎬 С верфи в море",
    "Посмотрите, как строится судно — от первых стальных блоков до спуска на воду.",
    "Видео не воспроизводится? ",
    "▶ Открыть на YouTube",
    "Создать групповой чат",
    "Групп пока нет — нажмите ➕, чтобы создать.",
    "Добавить",
    "Убрать",
    "Добавить рабочих",
    "Убрать из бригады",
    "Добавить цех",
    "Убрать цех",
    "Удалить группу",
    "Покинуть группу",
    "🌊 Океан",
    "🌅 Закат",
    "🌲 Лес",
    "⚪ Минимализм",
    "👑 Королевская",
    "📸 Еженедельный фотоотчёт",
    "📈 Статистика",
    "У вас пока нет задач для отчёта.",
    "Начало",
    "Плановое окончание",
    "Прогресс",
    "Завершено",
    "Текущая неделя",
    "Загрузить фото",
    "Неделя",
    "Выберите фото с компьютера",
    "Подпись к фото (необязательно)",
    "📤 ЗАГРУЗИТЬ",
    "Выберите хотя бы 1 фото!",
    "Не удалось прочитать эти фото (или они больше 15 МБ): ",
    "За эту неделю фото пока нет.",
    "📥 Скачать фотоотчёт (HTML)",
    "Откройте файл и нажмите Ctrl+P, чтобы распечатать или сохранить в PDF.",
    "💡 Удерживайте Ctrl (или Shift) при выборе, чтобы выбрать сразу несколько фото. Нажмите Upload ещё раз, чтобы добавить.",
    "📊 Скачать PowerPoint (.pptx)",
    "◀ Назад",
    "Далее ▶",
    "Всего недель",
    "Всего фото",
    "В работе",
    "фото",
    "Итоги",
]
_PT_LIST = [
    "🚢 SHIPCONTROL - GESTÃO DE TRABALHOS NAVAIS",
    "🔐 Entrar",
    "Entre ou cadastre-se para continuar.",
    "🧰 Quadro de tarefas",
    "➕ Adicionar tarefa",
    "📋 Atribuir tarefas",
    "👥 Minha equipe",
    "✏️ Editar / Excluir",
    "👥 Funções e permissões",
    "🗑️ Lixeira",
    "📊 Relatórios",
    "🔑 Alterar senha",
    "🚪 Sair",
    "⚙️ Configurações",
    "📲 Dispositivo",
    "📲 Você está usando:",
    "💻 Computador",
    "📱 Celular",
    "🎤 Toque no microfone de uma caixa de texto para falar em vez de digitar.",
    "🎤 Modo celular: toque no microfone na caixa de usuário para falar em vez de digitar.",
    "Este navegador ainda não aceita entrada por voz. Use o Chrome (Android) ou o Safari (iPhone).",
    "O acesso ao microfone está bloqueado. Permita o microfone para este site nas configurações do navegador.",
    "🏭 Oficinas: ver / adicionar / excluir",
    "Ainda não há oficinas. Adicione a primeira abaixo.",
    "#### ➕ Adicionar nova oficina",
    "Código da oficina (ex.: WOS_07):",
    "➕ ADICIONAR OFICINA",
    "Código com no máx. 20 caracteres, nome com no máx. 80 caracteres!",
    "#### 🗑️ Excluir oficina",
    "Oficina a excluir:",
    "🗑️ EXCLUIR OFICINA",
    "Nome da oficina",
    "Contas",
    "Tarefas",
    "🎨 Tema",
    "✨ Moderno",
    "🚀 Futurista",
    "🎈 Divertido",
    "📜 Antigo",
    "O tema Futurista é sempre escuro.",
    ">Conta<",
    "🔒 Sua conta acabou de ser aberta em outra aba ou outro dispositivo, então você foi desconectado aqui. ",
    "Se não foi você, entre de novo e troque sua senha imediatamente.",
    ">Função<",
    ">Conectado<",
    "🌙 Modo escuro",
    "🌐 Idioma",
    "📋 Quadro de progresso das tarefas",
    "➕ Adicionar nova tarefa",
    "✏️ Editar e excluir",
    "👥 Contas e funções (lista de funções)",
    "📊 Relatórios de progresso e estatísticas",
    "🗑️ Lixeira e restauração",
    "### ⚡ Atualização rápida do progresso",
    "### 🏭 Estatísticas por oficina",
    "### 📈 Distribuição do progresso",
    "### 📌 Atribuir tarefas aos trabalhadores",
    "### 📥 Baixar dados do relatório",
    "### 🔄 Alterar função da conta",
    "#### ➕ Adicionar trabalhadores",
    "#### ➖ Remover da equipe",
    "#### 👷 Membros",
    "#### 🚪 Sair da equipe",
    "🔑 Entrar",
    "📝 Cadastrar",
    "Usuário:",
    "Senha:",
    "🚀 ENTRAR",
    "✨ CADASTRAR AGORA",
    "Nome completo:",
    "Novo usuário:",
    "Confirmar senha:",
    "Digite o usuário e a senha!",
    "Usuário ou senha incorretos, ou a conta está bloqueada/excluída!",
    "⏳ Sua conta está aguardando um Admin / WOS Manager atribuir uma função. Volte mais tarde!",
    "Bem-vindo(a) ",
    ") voltou!",
    "Preencha todos os campos!",
    "A senha deve ter pelo menos 8 caracteres!",
    "O usuário ou o nome é longo demais!",
    "As senhas não coincidem!",
    "Este usuário já existe, escolha outro!",
    "🎉 Cadastrado! Sua conta está aguardando um WOS Manager aprovar uma função.",
    "👷 Estas são as tarefas atribuídas a você.",
    "Mostrar só as tarefas da minha equipe",
    "Ainda não há tarefas para mostrar.",
    "Você não tem tarefas para atualizar.",
    "Escolha uma tarefa para atualizar:",
    "Novo progresso (%)",
    "Nota de trabalho (Remark):",
    "🚀 ATUALIZAR PROGRESSO",
    "Progresso atualizado!",
    " (Atual: ",
    "⚙️ Adicionar / remover Blocks e WS Cost Codes",
    "Novo block (ex.: 170150):",
    "➕ ADICIONAR BLOCK",
    "Digite o nome do block!",
    "Block adicionado **",
    "Block a excluir:",
    "🗑️ EXCLUIR BLOCK",
    "Block excluído **",
    "A lista de blocks está vazia.",
    "Código (ex.: WOS_07):",
    "Nome da oficina / departamento:",
    "➕ ADICIONAR COST CODE",
    "Digite o código e o nome!",
    "Cost code a excluir:",
    "🗑️ EXCLUIR COST CODE",
    "A lista de cost codes está vazia.",
    "Não é possível excluir: **",
    "** conta(s) ainda pertencem a esta oficina. Mova-as para outra oficina em 👥 Funções e permissões primeiro.",
    "** da lista. As tarefas existentes são mantidas.",
    "** já está na lista!",
    "Código **",
    "Adicionado **",
    "Excluído **",
    "Descrição:",
    "Criado por:",
    "Data de criação:",
    "Responsável:",
    "Observação:",
    "— Nenhum —",
    "📌 Esta tarefa pertencerá à sua equipe. Vá em 📋 Atribuir tarefas para entregá-la a um trabalhador.",
    "💾 SALVAR NOVA TAREFA",
    "Preencha os campos obrigatórios: Task ID e Task Name!",
    "A data final planejada não pode ser antes da data inicial!",
    "** já existe (pode estar na Lixeira). Use outro ID!",
    "Tarefa salva **",
    "⚠️ AINDA NÃO HÁ WS COST CODE: um Foreman / WOS Manager precisa adicionar um acima antes de criar tarefas!",
    "👑 Escolha uma tarefa e atribua a um **Team Leader** ou **Foreman**. Eles a repassam aos trabalhadores da equipe.",
    "Ainda não há tarefas. Vá em ➕ Adicionar tarefa para criar uma.",
    "Ainda não há Team Leaders / Foremen. Atribua funções em 👥 Funções e permissões primeiro.",
    "Mostrar só tarefas não atribuídas",
    "Todas as tarefas foram atribuídas!",
    "Escolha uma tarefa:",
    "Atribuir a Team Leader / Foreman:",
    "Atribuir a trabalhador:",
    "— Não atribuída —",
    "Não atribuída",
    "💾 SALVAR ATRIBUIÇÃO",
    "Tarefa atribuída a **",
    "Sua equipe ainda não tem trabalhadores. Vá em 👥 Minha equipe para criar uma equipe e adicionar trabalhadores.",
    "Você ainda não recebeu nenhuma tarefa.",
    "Você precisa de trabalhadores na equipe antes de atribuir tarefas.",
    "(sem nome)",
    "⚠️ Sua conta ainda não tem oficina. Peça a um WOS Manager para atribuir uma antes de criar uma equipe.",
    "Você ainda não tem equipe. Dê um nome para criá-la.",
    "Nome da equipe *",
    "ex.: Equipe de solda Block 170",
    "➕ CRIAR EQUIPE",
    "Digite o nome da equipe!",
    "O nome da equipe pode ter no máximo 60 caracteres!",
    "Equipe criada **",
    "✏️ Renomear equipe",
    "Novo nome da equipe:",
    "💾 SALVAR NOME",
    "Nome de equipe inválido (1–60 caracteres)!",
    "Equipe renomeada!",
    "🗑️ Excluir equipe",
    "Ao excluir a equipe: todos os trabalhadores saem da equipe e as tarefas deles ficam sem atribuição. As tarefas continuam com você e você pode criar uma nova equipe depois.",
    "Tenho certeza de que quero excluir esta equipe",
    "🗑️ EXCLUIR EQUIPE",
    "Equipe excluída.",
    "A equipe ainda não tem trabalhadores.",
    "Não há mais trabalhadores na oficina ",
    " sem equipe.",
    "Escolha trabalhadores (mesma oficina, sem equipe):",
    "➕ ADICIONAR À EQUIPE",
    "Escolha pelo menos 1 trabalhador!",
    " trabalhador(es) à equipe!",
    "Adicionado(s) ",
    "Não há membros para remover.",
    "Escolha um trabalhador:",
    "➖ REMOVER DA EQUIPE",
    "Trabalhador removido da equipe. As tarefas da equipe dele ficaram sem atribuição.",
    "Você ainda não está em uma equipe. Um Team Leader / Foreman da sua oficina pode adicionar você.",
    "(Equipe sem nome)",
    "**Líder da equipe:** ",
    "Ao sair, as tarefas da equipe atribuídas a você ficam sem atribuição.",
    "Tenho certeza de que quero sair desta equipe",
    "🚪 SAIR DA EQUIPE",
    "Você saiu da equipe.",
    "🧰 Mover tarefas para a lixeira",
    "👤 Bloquear contas",
    "Não há tarefas para editar ou excluir.",
    "Tarefa a mover para a lixeira:",
    "🗑️ Mover tarefa para a lixeira",
    "Tarefa movida para a lixeira!",
    "Não há outras contas para excluir.",
    "Conta a bloquear:",
    "🗑️ Bloquear esta conta",
    "Conta bloqueada e movida para a lixeira!",
    "🛡️ O Admin só concede a função **WOS Manager** e escolhe a **Oficina** que ele gerencia. As funções Worker / Team Leader / Foreman são dadas pelos WOS Managers.",
    "👑 Você pode dar as funções **Worker / Team Leader / Foreman**, escolher a **Oficina** e colocar trabalhadores na **equipe** de um Team Leader / Foreman.",
    "Não há contas para atribuir funções.",
    "Conta a alterar:",
    "Ação:",
    "👑 Tornar WOS Manager",
    "⛔ Revogar função (voltar a Pending)",
    "Oficina que este gerente comanda:",
    "⚠️ Ainda não há oficinas. Adicione em ➕ Adicionar tarefa → ⚙️ Adicionar / remover.",
    "💾 SALVAR FUNÇÃO",
    "💾 SALVAR ALTERAÇÕES",
    "Tornado **WOS Manager** da oficina **",
    "Função revogada, a conta voltou a Pending.",
    "Nova função:",
    "Equipe de (Team Leader / Foreman):",
    "— Sem equipe —",
    "Esta oficina ainda não tem Team Leader / Foreman.",
    "Atualizado: **",
    "Nome completo",
    "Função",
    "Equipe de",
    "🧰 Lixeira de tarefas",
    "👤 Lixeira de contas",
    "A lixeira de tarefas está vazia.",
    "A lixeira de contas está vazia.",
    "Escolha uma tarefa:",
    "♻️ RESTAURAR TAREFA",
    "💥 EXCLUIR TAREFA PARA SEMPRE",
    "Tarefa restaurada!",
    "Tarefa excluída definitivamente!",
    "Escolha uma conta:",
    "♻️ DESBLOQUEAR / RESTAURAR CONTA",
    "💥 EXCLUIR CONTA PARA SEMPRE",
    "Conta restaurada!",
    "Conta excluída definitivamente do banco de dados!",
    "Ainda não há dados de tarefas para um relatório. Adicione tarefas primeiro!",
    "Total de tarefas",
    "Total de tarefas",
    "Progresso médio (%)",
    "Progresso médio",
    "Concluídas (100%)",
    "Em andamento",
    "Número de tarefas",
    "Situação",
    "📥 Baixar relatório de tarefas (CSV/Excel)",
    "Senha atual:",
    "Nova senha (pelo menos 8 caracteres):",
    "Repita a nova senha:",
    "💾 ALTERAR SENHA",
    "A senha atual está errada!",
    "A nova senha deve ter pelo menos 8 caracteres!",
    "Você não pode reutilizar a senha padrão!",
    "✅ Senha alterada! Os outros dispositivos foram desconectados.",
    "⚠️ Esta conta ainda usa a senha padrão 'admin123'. Vá em 🔑 Alterar senha e troque AGORA!",
    "Não iniciada (0%)",
    "Em andamento (1-50%)",
    "Quase pronta (51-99%)",
    "Concluída (100%)",
    "Quantidade",
    "Nova senha:",
    "💬 Mensagens",
    "📢 Chat geral",
    "#### 👥 Pessoas",
    "🔎 Buscar conta",
    "Nenhuma conta encontrada.",
    "Todos no app veem este canal.",
    "Chat privado com ",
    "Digite uma mensagem...",
    "➕ Criar chat em grupo",
    "Nome do grupo",
    "ex.: Turma de solda Block 170",
    "Adicionar pessoas ao grupo",
    "✨ CRIAR GRUPO",
    "Dê um nome ao grupo!",
    "Escolha pelo menos 1 pessoa!",
    "Grupo criado **",
    "#### 💬 Chats em grupo",
    "⚙️ Membros do grupo",
    "➕ ADICIONAR AO GRUPO",
    " pessoas ao grupo!",
    "Remover do grupo",
    "➖ REMOVER DO GRUPO",
    "Removido do grupo.",
    "🗑️ EXCLUIR GRUPO",
    "Só o criador do grupo pode adicionar / remover membros.",
    "🚪 SAIR DO GRUPO",
    "🎥 Reuniões online",
    "### 🔑 Entrar em uma reunião",
    "ID da sala",
    "Senha da sala",
    "🎥 ENTRAR",
    "Muitas tentativas erradas. Recarregue a página e tente mais tarde.",
    "ID ou senha da sala incorretos, ou a reunião terminou!",
    "✅ Sala encontrada: **",
    "Anfitrião",
    "🎥 ABRIR REUNIÃO (câmera + microfone)",
    "A reunião abre em uma nova aba. Na primeira vez, o navegador vai pedir a câmera e o microfone — toque em Permitir.",
    "### ➕ Criar nova reunião",
    "Nome da reunião *",
    "ex.: Passagem do turno da manhã",
    "Data",
    "Hora",
    "Convidar pessoas (envia ID + senha por 💬 Mensagens)",
    "📢 Anunciar no chat geral",
    "🎥 CRIAR REUNIÃO",
    "Digite o nome da reunião!",
    "🎥 Convite para reunião",
    "🕒 Hora",
    "🔢 ID da sala",
    "🔑 Senha",
    "👉 Vá em 🎥 Reuniões online para entrar.",
    "🎉 Reunião criada **",
    "### 📋 Minhas reuniões",
    "Você ainda não tem reuniões.",
    "▶️ INICIAR",
    "⛔ ENCERRAR",
    "Só Foremen e Team Leaders podem criar reuniões. Você pode entrar com ID e senha.",
    "Escolha pessoas...",
    "👑 Admin do grupo",
    "🗑️ Excluir mensagens (admin)",
    "Não há mensagens para excluir.",
    "Mensagem a excluir",
    "🗑️ EXCLUIR MENSAGEM",
    "Mensagem excluída.",
    "Tenho certeza de que quero excluir todas as mensagens",
    "🧹 EXCLUIR TODAS AS MENSAGENS",
    "Todas as mensagens foram excluídas.",
    "Convidar pessoas (elas verão o convite em 🎥 Reuniões online)",
    "### 📨 Seus convites de reunião",
    "🎥 ENTRAR NA REUNIÃO",
    "💬 Mensagens e reuniões",
    "Configurações",
    "Sobre nós",
    "Entrar / Cadastrar",
    "Software de gestão de trabalhos para estaleiros",
    "O ShipControl ajuda as equipes do estaleiro a distribuir o trabalho, acompanhar o progresso e manter contato em um só lugar — no computador e no celular.",
    "📋 Gestão de tarefas",
    "Crie tarefas por Block, Deck e Frame e acompanhe o progresso por cento a por cento.",
    "👥 Equipes e funções",
    "Admin, WOS Manager, Foreman, Team Leader e Worker — cada um vê exatamente o seu próprio trabalho.",
    "💬 Mensagens e reuniões online",
    "Chat geral, chats em grupo e salas de videochamada com ID + senha.",
    "🌐 15 idiomas e voz",
    "Vietnamita, inglês, chinês, japonês, coreano, tailandês, indonésio, espanhol, francês, alemão, russo e mais — e fale em vez de digitar no celular.",
    "Ainda não tem conta? Cadastre-se e aguarde um gerente lhe dar uma função.",
    "🎬 Do estaleiro ao mar",
    "Veja um navio sendo construído dos primeiros blocos de aço até o lançamento.",
    "O vídeo não abre? ",
    "▶ Abrir no YouTube",
    "Criar chat em grupo",
    "Ainda não há grupos — toque em ➕ para criar um.",
    "Adicionar",
    "Remover",
    "Adicionar trabalhadores",
    "Remover da equipe",
    "Adicionar oficina",
    "Remover oficina",
    "Excluir grupo",
    "Sair do grupo",
    "🌊 Oceano",
    "🌅 Pôr do sol",
    "🌲 Floresta",
    "⚪ Minimalista",
    "👑 Real",
    "📸 Relatório fotográfico semanal",
    "📈 Estatísticas",
    "Você ainda não tem tarefas para relatar.",
    "Início",
    "Término previsto",
    "Progresso",
    "Concluída",
    "Semana atual",
    "Enviar fotos",
    "Semana",
    "Escolha fotos do computador",
    "Nota da foto (opcional)",
    "📤 ENVIAR",
    "Escolha pelo menos 1 foto!",
    "Não foi possível ler estas fotos (ou têm mais de 15 MB): ",
    "Ainda não há fotos desta semana.",
    "📥 Baixar relatório fotográfico (HTML)",
    "Abra o arquivo e pressione Ctrl+P para imprimir ou salvar em PDF.",
    "💡 Segure Ctrl (ou Shift) ao escolher para pegar várias fotos de uma vez. Toque em Upload de novo para adicionar mais.",
    "📊 Baixar PowerPoint (.pptx)",
    "◀ Anterior",
    "Próximo ▶",
    "Total de semanas",
    "Total de fotos",
    "Em andamento",
    "fotos",
    "Resumo",
]

_VI_KEYS = list(VI_EN.keys())
TRANSLATIONS = {
    "en": VI_EN,
    "zh": dict(zip(_VI_KEYS, _ZH_LIST)),
    "ja": dict(zip(_VI_KEYS, _JA_LIST)),
    "ko": dict(zip(_VI_KEYS, _KO_LIST)),
    "th": dict(zip(_VI_KEYS, _TH_LIST)),
    "id": dict(zip(_VI_KEYS, _ID_LIST)),
    "ms": dict(zip(_VI_KEYS, _MS_LIST)),
    "tl": dict(zip(_VI_KEYS, _TL_LIST)),
    "hi": dict(zip(_VI_KEYS, _HI_LIST)),
    "es": dict(zip(_VI_KEYS, _ES_LIST)),
    "fr": dict(zip(_VI_KEYS, _FR_LIST)),
    "de": dict(zip(_VI_KEYS, _DE_LIST)),
    "ru": dict(zip(_VI_KEYS, _RU_LIST)),
    "pt": dict(zip(_VI_KEYS, _PT_LIST)),
}
LANG_OPTIONS = ["vi", "en", "zh", "ja", "ko", "th", "id", "ms", "tl", "hi", "es", "fr", "de", "ru", "pt"]
LANG_LABELS = {"vi": "🇻🇳 Tiếng Việt", "en": "🇬🇧 English", "zh": "🇨🇳 中文", "ja": "🇯🇵 日本語", "ko": "🇰🇷 한국어",
               "th": "🇹🇭 ไทย", "id": "🇮🇩 Bahasa Indonesia", "ms": "🇲🇾 Bahasa Melayu", "tl": "🇵🇭 Filipino",
               "hi": "🇮🇳 हिन्दी", "es": "🇪🇸 Español", "fr": "🇫🇷 Français", "de": "🇩🇪 Deutsch",
               "ru": "🇷🇺 Русский", "pt": "🇧🇷 Português"}
SPEECH_LANG = {"vi": "vi-VN", "en": "en-US", "zh": "zh-CN", "ja": "ja-JP", "ko": "ko-KR",
               "th": "th-TH", "id": "id-ID", "ms": "ms-MY", "tl": "fil-PH", "hi": "hi-IN",
               "es": "es-ES", "fr": "fr-FR", "de": "de-DE", "ru": "ru-RU", "pt": "pt-BR"}


import re as _re
import functools as _functools
from streamlit.delta_generator import DeltaGenerator as _DG

_VI_EN_RX = _re.compile("|".join(_re.escape(k) for k in sorted(VI_EN, key=len, reverse=True)))
LANG_COOKIE = "shipcontrol_lang"


def ui_lang():
    return st.session_state.get("ui_lang", "vi")


def tr(text):
    """Dịch một đoạn chữ sang ngôn ngữ đang chọn (giữ nguyên nếu là Tiếng Việt)."""
    lang = ui_lang()
    if not isinstance(text, str) or not text or lang == "vi" or lang not in TRANSLATIONS:
        return text
    if text.lstrip().startswith("<style"):
        return text
    table = TRANSLATIONS[lang]
    return _VI_EN_RX.sub(lambda m: table.get(m.group(0), VI_EN[m.group(0)]), text)


def _tr_data(data):
    if ui_lang() != "vi" and isinstance(data, pd.DataFrame):
        return data.rename(columns=lambda c: tr(c) if isinstance(c, str) else c)
    return data


def _install_translator():
    """Tự động dịch chữ hiển thị của các thành phần Streamlit.
    Chỉ đổi phần HIỂN THỊ; giá trị trả về và logic của app giữ nguyên tiếng Việt.

    QUAN TRỌNG: Streamlit giữ nguyên chương trình Python giữa các lần chạy (kể cả khi bạn
    upload app.py mới). Vì vậy các hàm bọc KHÔNG được giữ bản 'tr' cũ — chúng luôn gọi
    bản mới nhất qua _DG._sc_live_* (gán lại ở mỗi lần chạy, xem bên dưới)."""
    # Giữ bản gốc của markdown (không dịch) để hiện nội dung người dùng viết, ví dụ tin nhắn
    if not hasattr(_DG, "_sc_orig_markdown"):
        _orig = _DG.markdown
        while getattr(_orig, "_sc_wrapped", False) and hasattr(_orig, "__wrapped__"):
            _orig = _orig.__wrapped__
        _DG._sc_orig_markdown = _orig

    def _L():
        return _DG._sc_live_lang()

    def _T(x):
        return _DG._sc_live_tr(x)

    def text_first(fn):
        @_functools.wraps(fn)
        def wrapper(self, *args, **kwargs):
            if _L() != "vi":
                if args and isinstance(args[0], str):
                    args = (_T(args[0]),) + args[1:]
                for k in ("label", "body", "placeholder", "help"):
                    if isinstance(kwargs.get(k), str):
                        kwargs[k] = _T(kwargs[k])
            return fn(self, *args, **kwargs)
        return wrapper

    def with_options(fn):
        @_functools.wraps(fn)
        def wrapper(self, *args, **kwargs):
            if _L() != "vi":
                if args and isinstance(args[0], str):
                    args = (_T(args[0]),) + args[1:]
                for k in ("label", "placeholder", "help"):
                    if isinstance(kwargs.get(k), str):
                        kwargs[k] = _T(kwargs[k])
                _ff = kwargs.get("format_func") or str
                kwargs["format_func"] = lambda o, _ff=_ff: _T(_ff(o))
            return fn(self, *args, **kwargs)
        return wrapper

    def data_first(fn):
        @_functools.wraps(fn)
        def wrapper(self, *args, **kwargs):
            if args:
                args = (_DG._sc_live_trdata(args[0]),) + args[1:]
            elif "data" in kwargs:
                kwargs["data"] = _DG._sc_live_trdata(kwargs["data"])
            return fn(self, *args, **kwargs)
        return wrapper

    def tabs_first(fn):
        @_functools.wraps(fn)
        def wrapper(self, tabs, *args, **kwargs):
            return fn(self, [_T(t) for t in tabs], *args, **kwargs)
        return wrapper

    groups = {
        text_first: ["markdown", "caption", "info", "success", "warning", "error", "title", "header",
                     "subheader", "text", "button", "form_submit_button", "download_button", "text_input",
                     "text_area", "number_input", "date_input", "time_input", "checkbox", "toggle",
                     "expander", "popover", "metric", "slider", "chat_input", "link_button"],
        with_options: ["selectbox", "radio", "multiselect", "select_slider"],
        data_first: ["dataframe", "table"],
        tabs_first: ["tabs"],
    }
    _VER = 2
    for wrap, names in groups.items():
        for name in names:
            fn = getattr(_DG, name, None)
            if fn is None or getattr(fn, "_sc_ver", 0) == _VER:
                continue                      # bản bọc mới nhất đã có → bỏ qua
            # Gỡ các lớp bọc cũ (từ phiên bản app trước) để không bị dịch bằng từ điển cũ
            while getattr(fn, "_sc_wrapped", False) and hasattr(fn, "__wrapped__"):
                fn = fn.__wrapped__
            wrapped = wrap(fn)
            wrapped._sc_wrapped = True
            wrapped._sc_ver = _VER
            setattr(_DG, name, wrapped)


# Bộ dịch MỚI NHẤT cho lần chạy này (các hàm bọc luôn gọi qua đây)
_DG._sc_live_tr = tr
_DG._sc_live_lang = ui_lang
_DG._sc_live_trdata = _tr_data


_install_translator()
# st.button, st.markdown... được gắn sẵn vào khung chính lúc import → gắn lại để dùng bản đã dịch
for _name in ["markdown", "caption", "info", "success", "warning", "error", "title", "header", "subheader",
              "text", "button", "form_submit_button", "download_button", "text_input", "text_area",
              "number_input", "date_input", "time_input", "checkbox", "toggle", "expander", "popover",
              "metric", "slider", "selectbox", "radio", "multiselect", "select_slider", "dataframe",
              "table", "tabs", "chat_input", "link_button"]:
    if hasattr(st, _name) and hasattr(st._main, _name):
        setattr(st, _name, getattr(st._main, _name))


# ==========================================
# 🎨 CHỦ ĐỀ GIAO DIỆN (Futuristic / Playful / Modern / Old times)
#    Mỗi chủ đề là một lớp CSS phủ lên giao diện Modern.
#    Dùng "html body ..." để luôn ưu tiên hơn các kiểu gốc.
# ==========================================
THEME_OPTIONS = ["modern", "futuristic", "playful", "oldtimes", "ocean", "sunset", "forest", "minimal", "royal"]
THEME_LABELS = {
    "modern": "✨ Hiện đại (Modern)",
    "futuristic": "🚀 Tương lai (Futuristic)",
    "playful": "🎈 Vui nhộn (Playful)",
    "oldtimes": "📜 Cổ điển (Old times)",
    "ocean": "🌊 Đại dương (Ocean)",
    "sunset": "🌅 Hoàng hôn (Sunset)",
    "forest": "🌲 Rừng xanh (Forest)",
    "minimal": "⚪ Tối giản (Minimal)",
    "royal": "👑 Hoàng gia (Royal)",
}

_THEME_FUTURISTIC = r'''
@import url('https://fonts.googleapis.com/css2?family=Exo+2:wght@400;500;600;700;800&display=swap');
html body .stApp, html body .stApp p, html body .stApp label, html body .stApp input, html body .stApp textarea,
html body .stApp button, html body .stApp h1, html body .stApp h2, html body .stApp h3, html body .stApp li,
html body .stApp span:not([data-testid="stIconMaterial"]) {
    font-family: 'Twemoji Country Flags', 'Exo 2', system-ui, sans-serif !important;
}
html body .stApp { background: #03060f !important; }
html body .stApp::before {
    background:
        linear-gradient(rgba(34,211,238,0.07) 1px, transparent 1px) 0 0 / 44px 44px,
        linear-gradient(90deg, rgba(34,211,238,0.07) 1px, transparent 1px) 0 0 / 44px 44px,
        radial-gradient(900px 500px at 85% 0%, rgba(232,121,249,0.18), transparent 60%),
        linear-gradient(180deg, rgba(3,6,15,0.70) 0%, rgba(3,6,15,0.90) 50%, rgba(3,6,15,0.96) 100%),
        url('__BG_URL__') center / cover no-repeat !important;
}
html body section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #050816 0%, #0a1030 100%) !important;
    border-right: 1px solid rgba(34,211,238,0.45) !important;
    box-shadow: 0 0 30px rgba(34,211,238,0.18) !important;
}
html body .sidebar-header, html body .main-title, html body .big-table-title {
    text-transform: uppercase !important;
    letter-spacing: 0.08em !important;
    color: #67e8f9 !important;
    text-shadow: 0 0 12px rgba(34,211,238,0.65) !important;
}
html body .main-title::after { background: linear-gradient(90deg, #22d3ee, #e879f9) !important; box-shadow: 0 0 14px #22d3ee; }
html body section[data-testid="stSidebar"] div.stButton > button[kind="primary"] {
    background: rgba(34,211,238,0.10) !important;
    border: 1px solid #22d3ee !important;
    border-radius: 6px !important;
    box-shadow: 0 0 16px rgba(34,211,238,0.55), inset 0 0 12px rgba(34,211,238,0.25) !important;
}
html body section[data-testid="stSidebar"] div.stButton > button[kind="primary"] * { color: #a5f3fc !important; -webkit-text-fill-color: #a5f3fc !important; }
html body section[data-testid="stSidebar"] div.stButton > button[kind="secondary"] { border-radius: 6px !important; }
html body section[data-testid="stSidebar"] div.stButton > button[kind="secondary"]:hover:not(:disabled) {
    background: rgba(232,121,249,0.10) !important; border-color: rgba(232,121,249,0.6) !important;
}
html body .stMain div.stButton > button[kind="primary"],
html body div[data-testid="stFormSubmitButton"] > button,
html body div.stDownloadButton > button {
    background: linear-gradient(90deg, #06b6d4, #a855f7) !important;
    border: 1px solid #67e8f9 !important;
    border-radius: 6px !important;
    box-shadow: 0 0 18px rgba(34,211,238,0.45) !important;
    text-transform: uppercase;
}
html body .stMain div.stButton > button[kind="primary"]:hover:not(:disabled),
html body div[data-testid="stFormSubmitButton"] > button:hover:not(:disabled) {
    box-shadow: 0 0 28px rgba(232,121,249,0.7) !important;
}
html body .stMain div.stButton > button[kind="secondary"] {
    background: rgba(8,16,40,0.85) !important;
    border: 1px solid rgba(34,211,238,0.6) !important;
    border-radius: 6px !important;
}
html body .stMain div.stButton > button[kind="secondary"] * { color: #a5f3fc !important; -webkit-text-fill-color: #a5f3fc !important; }
html body div[data-testid="stPopoverBody"] { background: #070c1f !important; }
html body div[data-testid="stForm"], html body div.st-key-sc_add_task_box, html body div[data-testid="stExpander"] details,
html body div[data-testid="stDataFrame"] {
    background: rgba(8,14,34,0.88) !important;
    border: 1px solid rgba(34,211,238,0.45) !important;
    border-radius: 8px !important;
    box-shadow: 0 0 22px rgba(34,211,238,0.15), inset 0 0 20px rgba(34,211,238,0.05) !important;
}
html body .profile-card { border-color: rgba(34,211,238,0.5) !important; border-radius: 8px !important; box-shadow: 0 0 16px rgba(34,211,238,0.2); }
html body .pc-avatar { background: linear-gradient(135deg, #22d3ee, #a855f7) !important; color: #fff !important; }
html body .made-by-minh { background: linear-gradient(90deg, #22d3ee, #e879f9) !important; border-radius: 6px !important; }
html body .st-key-sc_topbar { background: rgba(5,8,22,0.95) !important; border-bottom: 1px solid #22d3ee !important; box-shadow: 0 0 22px rgba(34,211,238,0.35) !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] { box-shadow: inset 0 -3px 0 #22d3ee !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] * { color: #a5f3fc !important; -webkit-text-fill-color: #a5f3fc !important; }
'''

_THEME_PLAYFUL = r'''
@import url('https://fonts.googleapis.com/css2?family=Baloo+2:wght@500;600;700;800&display=swap');
html body .stApp, html body .stApp p, html body .stApp label, html body .stApp input, html body .stApp textarea,
html body .stApp button, html body .stApp h1, html body .stApp h2, html body .stApp h3, html body .stApp li,
html body .stApp span:not([data-testid="stIconMaterial"]) {
    font-family: 'Twemoji Country Flags', 'Baloo 2', 'Comic Sans MS', system-ui, sans-serif !important;
}
html body .stApp::before {
    background:
        radial-gradient(circle at 12% 18%, rgba(251,191,36,__DOT__) 0 70px, transparent 71px),
        radial-gradient(circle at 88% 12%, rgba(244,114,182,__DOT__) 0 90px, transparent 91px),
        radial-gradient(circle at 80% 85%, rgba(52,211,153,__DOT__) 0 80px, transparent 81px),
        __PLAY_OVERLAY__,
        url('__BG_URL__') center / cover no-repeat !important;
}
html body section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #ec4899 0%, #a855f7 55%, #6366f1 100%) !important;
}
html body .sidebar-header { color: #fef08a !important; transform: rotate(-2deg); }
html body .main-title { color: __PLAY_TITLE__ !important; }
html body .main-title::after {
    height: 8px !important; width: 180px !important;
    background: repeating-linear-gradient(90deg, #f472b6 0 20px, #facc15 20px 40px, #34d399 40px 60px, #60a5fa 60px 80px) !important;
}
html body .big-table-title { color: #db2777 !important; }
html body section[data-testid="stSidebar"] div.stButton > button { border-radius: 999px !important; }
html body section[data-testid="stSidebar"] div.stButton > button[kind="primary"] {
    background: #ffffff !important;
    box-shadow: 4px 4px 0 #facc15 !important;
    transform: rotate(-1deg);
}
html body section[data-testid="stSidebar"] div.stButton > button[kind="primary"] * { color: #a21caf !important; -webkit-text-fill-color: #a21caf !important; }
html body section[data-testid="stSidebar"] div.stButton > button[kind="secondary"]:hover:not(:disabled) {
    transform: scale(1.04) rotate(-1deg) !important;
}
html body div.stButton > button, html body div[data-testid="stFormSubmitButton"] > button, html body div.stDownloadButton > button {
    border-radius: 999px !important;
    transition: transform 0.18s cubic-bezier(0.34, 1.56, 0.64, 1), box-shadow 0.18s ease !important;
}
html body .stMain div.stButton > button[kind="primary"],
html body div[data-testid="stFormSubmitButton"] > button,
html body div.stDownloadButton > button {
    background: linear-gradient(135deg, #fb923c, #ec4899) !important;
    border: 3px solid #1e1b4b !important;
    box-shadow: 5px 5px 0 #1e1b4b !important;
}
html body .stMain div.stButton > button[kind="primary"]:hover:not(:disabled),
html body div[data-testid="stFormSubmitButton"] > button:hover:not(:disabled) {
    transform: translate(-2px, -2px) rotate(-1deg) scale(1.03) !important;
    box-shadow: 7px 7px 0 #1e1b4b !important;
}
html body .stMain div.stButton > button[kind="secondary"] {
    border: 3px solid #1e1b4b !important;
    box-shadow: 4px 4px 0 #a855f7 !important;
}
html body div.stButton > button:active:not(:disabled),
html body div[data-testid="stFormSubmitButton"] > button:active:not(:disabled) {
    transform: translate(3px, 3px) !important; box-shadow: 1px 1px 0 #1e1b4b !important;
}
html body div[data-testid="stForm"], html body div.st-key-sc_add_task_box, html body div[data-testid="stExpander"] details,
html body div[data-testid="stDataFrame"], html body div[data-testid="stPopoverBody"] {
    border: 3px solid #1e1b4b !important;
    border-radius: 24px !important;
    box-shadow: 8px 8px 0 #f472b6 !important;
}
html body div[data-testid="stAlert"] { border: 3px solid #1e1b4b !important; border-radius: 20px !important; }
html body .profile-card { border-radius: 24px !important; border: 3px solid rgba(255,255,255,0.6) !important; }
html body .pc-avatar { background: #fef08a !important; color: #a21caf !important; transform: rotate(-6deg); }
html body .made-by-minh { transform: rotate(-2deg); box-shadow: 4px 4px 0 #1e1b4b !important; }
html body .st-key-sc_topbar { background: linear-gradient(90deg, #ec4899, #a855f7, #6366f1) !important; border-radius: 0 0 24px 24px !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button { border-radius: 999px !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] { background: #ffffff !important; box-shadow: 3px 3px 0 #facc15 !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] * { color: #a21caf !important; -webkit-text-fill-color: #a21caf !important; }
/* Phông Baloo 2 trông nhỏ hơn phông thường ở cùng cỡ chữ → tăng cỡ để menu to bằng các chủ đề khác */
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button * { font-size: 1.05rem !important; line-height: 1 !important; padding-top: 3px !important; }
html body .sc-brand { font-size: 1.2rem !important; }
'''

_THEME_OLDTIMES = r'''
@import url('https://fonts.googleapis.com/css2?family=Lora:wght@400;500;600;700&family=Playfair+Display:wght@700;800;900&display=swap');
html body .stApp, html body .stApp p, html body .stApp label, html body .stApp input, html body .stApp textarea,
html body .stApp button, html body .stApp li, html body .stApp span:not([data-testid="stIconMaterial"]) {
    font-family: 'Twemoji Country Flags', 'Lora', Georgia, 'Times New Roman', serif !important;
}
html body .stApp h1, html body .stApp h2, html body .stApp h3, html body .main-title,
html body .big-table-title, html body .sidebar-header {
    font-family: 'Twemoji Country Flags', 'Playfair Display', Georgia, serif !important;
    letter-spacing: 0.01em !important;
}
html body .stApp { background: __OLD_PAPER__ !important; }
html body .stApp::before {
    filter: sepia(0.85) contrast(0.95) brightness(__OLD_BRIGHT__);
    background:
        repeating-linear-gradient(0deg, rgba(120,90,50,0.035) 0 2px, transparent 2px 4px),
        __OLD_OVERLAY__,
        url('__BG_URL__') center / cover no-repeat !important;
}
html body section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #3b2716 0%, #4e3420 60%, #2e1e10 100%) !important;
    border-right: 4px double #c9a227 !important;
}
html body .sidebar-header { color: #e6c35c !important; }
html body .main-title, html body .big-table-title, html body .stApp h3 { color: __OLD_INK__ !important; }
html body .main-title { border-bottom: 3px double __OLD_INK__ !important; }
html body .main-title::after { display: none !important; }
html body .stMain [data-testid="stWidgetLabel"] p, html body .stMain p { color: __OLD_INK__ !important; -webkit-text-fill-color: __OLD_INK__ !important; }
html body section[data-testid="stSidebar"] div.stButton > button { border-radius: 3px !important; }
html body section[data-testid="stSidebar"] div.stButton > button[kind="primary"] {
    background: linear-gradient(180deg, #e6c35c, #c9a227) !important;
    border: 1px solid #8a6d12 !important;
    box-shadow: inset 0 0 0 2px rgba(255,255,255,0.35), 0 2px 0 #6b520c !important;
}
html body section[data-testid="stSidebar"] div.stButton > button[kind="primary"] * { color: #2e1e10 !important; -webkit-text-fill-color: #2e1e10 !important; }
html body section[data-testid="stSidebar"] div.stButton > button[kind="secondary"] * { color: #f3e6c4 !important; -webkit-text-fill-color: #f3e6c4 !important; }
html body div.stButton > button, html body div[data-testid="stFormSubmitButton"] > button, html body div.stDownloadButton > button {
    border-radius: 3px !important;
}
html body .stMain div.stButton > button[kind="primary"],
html body div[data-testid="stFormSubmitButton"] > button,
html body div.stDownloadButton > button {
    background: linear-gradient(180deg, #6b4a2b, #4e3420) !important;
    border: 3px double #c9a227 !important;
    box-shadow: 0 3px 0 #2e1e10 !important;
}
html body .stMain div.stButton > button[kind="primary"] *,
html body div[data-testid="stFormSubmitButton"] > button * { color: #f3e6c4 !important; -webkit-text-fill-color: #f3e6c4 !important; }
html body .stMain div.stButton > button[kind="secondary"] {
    background: __OLD_CARD__ !important;
    border: 2px solid #8b6b3d !important;
    box-shadow: 0 2px 0 #8b6b3d !important;
}
html body .stMain div.stButton > button[kind="secondary"] * { color: __OLD_INK__ !important; -webkit-text-fill-color: __OLD_INK__ !important; }
html body div[data-testid="stForm"], html body div.st-key-sc_add_task_box, html body div[data-testid="stExpander"] details,
html body div[data-testid="stDataFrame"], html body div[data-testid="stPopoverBody"] {
    background: __OLD_CARD__ !important;
    border: 3px double #8b6b3d !important;
    border-radius: 4px !important;
    box-shadow: 0 6px 18px rgba(62,44,28,0.25) !important;
}
html body div[data-testid="stTextInputRootElement"], html body div[data-testid="stTextAreaRootElement"],
html body div[data-testid="stNumberInputContainer"], html body div[data-testid="stDateInputField"],
html body div[data-testid="stSelectbox"] div:has(> input) {
    background-color: __OLD_INPUT__ !important;
    border: 1px solid #8b6b3d !important;
    border-radius: 3px !important;
}
html body .stApp input, html body .stApp textarea, html body div[data-testid="stSelectbox"] div:has(> input) * {
    color: __OLD_INK__ !important; -webkit-text-fill-color: __OLD_INK__ !important;
}
html body [data-testid="stChatInput"] > div {
    background-color: __OLD_INPUT__ !important;
    border: 1px solid #8b6b3d !important;
    border-radius: 4px !important;
}
html body .profile-card { border: 1px solid #c9a227 !important; border-radius: 4px !important; }
html body .pc-avatar { background: #c9a227 !important; color: #2e1e10 !important; border: 2px solid #f3e6c4; }
html body .made-by-minh { border-radius: 3px !important; background: linear-gradient(180deg, #e6c35c, #c9a227) !important; }
html body .st-key-sc_topbar { background: linear-gradient(180deg, #4e3420, #3b2716) !important; border-bottom: 4px double #c9a227 !important; border-radius: 0 !important; }
html body .sc-brand { color: #e6c35c !important; font-family: 'Playfair Display', Georgia, serif !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button * { color: #f3e6c4 !important; -webkit-text-fill-color: #f3e6c4 !important; font-family: 'Lora', Georgia, serif !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] { box-shadow: inset 0 -3px 0 #e6c35c !important; }
'''


# 🎨 5 chủ đề mới dùng chung một khuôn CSS, chỉ khác màu / phông / bo góc
_THEME_ACCENT = r'''
@import url('__FONT_URL__');
html body .stApp, html body .stApp p, html body .stApp label, html body .stApp input, html body .stApp textarea,
html body .stApp button, html body .stApp li, html body .stApp span:not([data-testid="stIconMaterial"]) {
    font-family: 'Twemoji Country Flags', __FONT__ !important;
}
html body .stApp h1, html body .stApp h2, html body .stApp h3, html body .stApp h4, html body .main-title,
html body .big-table-title, html body .sc-brand {
    font-family: 'Twemoji Country Flags', __HFONT__ !important;
}
html body .stApp::before {
    background: __OVERLAY__, url('__BG_URL__') center / cover no-repeat !important;
}
html body .main-title, html body .big-table-title { color: __TITLE__ !important; }
html body .main-title::after { background: __GRAD__ !important; }
html body .st-key-sc_topbar {
    background: __TOPBAR__ !important;
    border-radius: __TOP_R__ !important;
    box-shadow: __TOP_SHADOW__ !important;
}
html body .sc-brand { color: __BRAND__ !important; border-right-color: __NAV_LINE__ !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button * { color: __NAV_TX__ !important; -webkit-text-fill-color: __NAV_TX__ !important; }
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] {
    background: __NAV_ACTIVE_BG__ !important;
    box-shadow: inset 0 -3px 0 __ACCENT2__ !important;
}
html body div.st-key-sc_nav.st-key-sc_nav div.stButton > button[kind="primary"] * { color: __NAV_ACTIVE_TX__ !important; -webkit-text-fill-color: __NAV_ACTIVE_TX__ !important; }
html body section[data-testid="stSidebar"] { background: __TOPBAR__ !important; }
html body div.stButton > button, html body div[data-testid="stFormSubmitButton"] > button, html body div.stDownloadButton > button {
    border-radius: __BTN_R__ !important;
}
html body .stMain div.stButton > button[kind="primary"],
html body div[data-testid="stFormSubmitButton"] > button,
html body div.stDownloadButton > button {
    background: __GRAD__ !important;
    border: __BTN_BORDER__ !important;
    box-shadow: __BTN_SHADOW__ !important;
}
html body .stMain div.stButton > button[kind="primary"] *,
html body div[data-testid="stFormSubmitButton"] > button *,
html body div.stDownloadButton > button * { color: __BTN_TX__ !important; -webkit-text-fill-color: __BTN_TX__ !important; }
html body .stMain div.stButton > button[kind="secondary"] {
    border: 1.5px solid __ACCENT__ !important;
    box-shadow: none !important;
}
html body .stMain div.stButton > button[kind="secondary"]:hover:not(:disabled) { border-color: __ACCENT2__ !important; }
html body div[data-testid="stForm"], html body div.st-key-sc_add_task_box, html body div[data-testid="stExpander"] details,
html body div[data-testid="stDataFrame"], html body div[data-testid="stPopoverBody"] {
    border: __CARD_BORDER__ !important;
    border-radius: __CARD_R__ !important;
    box-shadow: __CARD_SHADOW__ !important;
}
html body div[data-testid="stTextInputRootElement"]:focus-within, html body div[data-testid="stTextAreaRootElement"]:focus-within,
html body div[data-testid="stSelectbox"] div:has(> input):focus-within {
    border-color: __ACCENT__ !important;
    box-shadow: 0 0 0 3px __FOCUS__ !important;
}
html body div[data-testid="stAlert"] { border-radius: __CARD_R__ !important; }
html body .pc-avatar { background: __ACCENT2__ !important; color: __AVATAR_TX__ !important; }
html body .profile-card { border-radius: __CARD_R__ !important; }
html body [data-testid="stSlider"] [role="slider"] { background: __ACCENT__ !important; }
html body [data-testid="stProgress"] > div > div > div > div { background: __GRAD__ !important; }
__EXTRA__
'''

_ACCENT_THEMES = {
    "ocean": {
        "__FONT_URL__": "https://fonts.googleapis.com/css2?family=Poppins:wght@400;500;600;700;800&display=swap",
        "__FONT__": "'Poppins', system-ui, sans-serif", "__HFONT__": "'Poppins', system-ui, sans-serif",
        "light": {"__OVERLAY__": "linear-gradient(180deg, rgba(236,254,255,0.55) 0%, rgba(224,242,254,0.85) 50%, rgba(240,249,255,0.93) 100%)",
                  "__TITLE__": "#075985"},
        "dark": {"__OVERLAY__": "linear-gradient(180deg, rgba(8,47,73,0.72) 0%, rgba(8,47,73,0.88) 50%, rgba(3,20,35,0.95) 100%)",
                 "__TITLE__": "#67e8f9"},
        "__GRAD__": "linear-gradient(135deg, #0ea5e9, #06b6d4 55%, #14b8a6)",
        "__TOPBAR__": "linear-gradient(90deg, #0c4a6e, #075985 45%, #0e7490)",
        "__TOP_R__": "0 0 18px 18px", "__TOP_SHADOW__": "0 8px 24px rgba(8,47,73,0.35)",
        "__BRAND__": "#a5f3fc", "__NAV_LINE__": "rgba(165,243,252,0.35)", "__NAV_TX__": "#e0f2fe",
        "__NAV_ACTIVE_BG__": "rgba(255,255,255,0.16)", "__NAV_ACTIVE_TX__": "#ffffff",
        "__ACCENT__": "#0891b2", "__ACCENT2__": "#22d3ee", "__FOCUS__": "rgba(34,211,238,0.25)",
        "__BTN_R__": "14px", "__BTN_BORDER__": "none", "__BTN_SHADOW__": "0 6px 16px rgba(8,145,178,0.35)", "__BTN_TX__": "#ffffff",
        "__CARD_BORDER__": "1px solid rgba(8,145,178,0.25)", "__CARD_R__": "18px", "__CARD_SHADOW__": "0 10px 28px rgba(8,47,73,0.12)",
        "__AVATAR_TX__": "#083344",
        "__EXTRA__": "html body .main-title::after { height: 6px !important; border-radius: 6px !important; }",
    },
    "sunset": {
        "__FONT_URL__": "https://fonts.googleapis.com/css2?family=Nunito:wght@400;600;700;800;900&display=swap",
        "__FONT__": "'Nunito', system-ui, sans-serif", "__HFONT__": "'Nunito', system-ui, sans-serif",
        "light": {"__OVERLAY__": "linear-gradient(180deg, rgba(255,237,213,0.55) 0%, rgba(254,226,226,0.82) 50%, rgba(255,247,237,0.93) 100%)",
                  "__TITLE__": "#c2410c"},
        "dark": {"__OVERLAY__": "linear-gradient(180deg, rgba(67,20,7,0.72) 0%, rgba(76,5,25,0.88) 50%, rgba(30,10,30,0.95) 100%)",
                 "__TITLE__": "#fdba74"},
        "__GRAD__": "linear-gradient(135deg, #f97316, #ef4444 50%, #db2777)",
        "__TOPBAR__": "linear-gradient(90deg, #9a3412, #be123c 50%, #7e22ce)",
        "__TOP_R__": "0 0 22px 22px", "__TOP_SHADOW__": "0 8px 24px rgba(190,18,60,0.30)",
        "__BRAND__": "#fed7aa", "__NAV_LINE__": "rgba(254,215,170,0.35)", "__NAV_TX__": "#fff1e6",
        "__NAV_ACTIVE_BG__": "rgba(255,255,255,0.18)", "__NAV_ACTIVE_TX__": "#ffffff",
        "__ACCENT__": "#ea580c", "__ACCENT2__": "#fbbf24", "__FOCUS__": "rgba(251,146,60,0.28)",
        "__BTN_R__": "999px", "__BTN_BORDER__": "none", "__BTN_SHADOW__": "0 6px 18px rgba(239,68,68,0.35)", "__BTN_TX__": "#ffffff",
        "__CARD_BORDER__": "1px solid rgba(234,88,12,0.22)", "__CARD_R__": "22px", "__CARD_SHADOW__": "0 12px 30px rgba(190,18,60,0.12)",
        "__AVATAR_TX__": "#7c2d12",
        "__EXTRA__": "",
    },
    "forest": {
        "__FONT_URL__": "https://fonts.googleapis.com/css2?family=Rubik:wght@400;500;600;700;800&display=swap",
        "__FONT__": "'Rubik', system-ui, sans-serif", "__HFONT__": "'Rubik', system-ui, sans-serif",
        "light": {"__OVERLAY__": "linear-gradient(180deg, rgba(236,253,245,0.55) 0%, rgba(220,252,231,0.84) 50%, rgba(240,253,244,0.93) 100%)",
                  "__TITLE__": "#166534"},
        "dark": {"__OVERLAY__": "linear-gradient(180deg, rgba(5,46,22,0.74) 0%, rgba(5,46,22,0.88) 50%, rgba(2,24,12,0.95) 100%)",
                 "__TITLE__": "#86efac"},
        "__GRAD__": "linear-gradient(135deg, #16a34a, #15803d 55%, #4d7c0f)",
        "__TOPBAR__": "linear-gradient(90deg, #14532d, #166534 50%, #3f6212)",
        "__TOP_R__": "0 0 12px 12px", "__TOP_SHADOW__": "0 8px 22px rgba(20,83,45,0.32)",
        "__BRAND__": "#bbf7d0", "__NAV_LINE__": "rgba(187,247,208,0.35)", "__NAV_TX__": "#ecfdf5",
        "__NAV_ACTIVE_BG__": "rgba(255,255,255,0.15)", "__NAV_ACTIVE_TX__": "#ffffff",
        "__ACCENT__": "#16a34a", "__ACCENT2__": "#a3e635", "__FOCUS__": "rgba(74,222,128,0.25)",
        "__BTN_R__": "10px", "__BTN_BORDER__": "none", "__BTN_SHADOW__": "0 5px 14px rgba(21,128,61,0.32)", "__BTN_TX__": "#ffffff",
        "__CARD_BORDER__": "1px solid rgba(22,163,74,0.25)", "__CARD_R__": "12px", "__CARD_SHADOW__": "0 10px 26px rgba(20,83,45,0.12)",
        "__AVATAR_TX__": "#14532d",
        "__EXTRA__": "",
    },
    "minimal": {
        "__FONT_URL__": "https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap",
        "__FONT__": "'Inter', system-ui, sans-serif", "__HFONT__": "'Inter', system-ui, sans-serif",
        "light": {"__OVERLAY__": "linear-gradient(180deg, rgba(255,255,255,0.90) 0%, rgba(255,255,255,0.96) 40%, #ffffff 100%)",
                  "__TITLE__": "#111827", "__GRAD__": "linear-gradient(135deg, #111827, #111827)", "__TOPBAR__": "#111827",
                  "__BTN_TX__": "#ffffff", "__BTN_BORDER__": "1px solid #111827", "__CARD_BORDER__": "1px solid #e5e7eb"},
        "dark": {"__OVERLAY__": "linear-gradient(180deg, rgba(10,10,10,0.90) 0%, rgba(10,10,10,0.96) 40%, #0a0a0a 100%)",
                 "__TITLE__": "#f5f5f5", "__GRAD__": "linear-gradient(135deg, #f5f5f5, #f5f5f5)", "__TOPBAR__": "#171717",
                 "__BTN_TX__": "#0a0a0a", "__BTN_BORDER__": "1px solid #f5f5f5", "__CARD_BORDER__": "1px solid #262626"},
        "__TOP_R__": "0", "__TOP_SHADOW__": "none",
        "__BRAND__": "#ffffff", "__NAV_LINE__": "rgba(255,255,255,0.2)", "__NAV_TX__": "#d1d5db",
        "__NAV_ACTIVE_BG__": "transparent", "__NAV_ACTIVE_TX__": "#ffffff",
        "__ACCENT__": "#6b7280", "__ACCENT2__": "#ffffff", "__FOCUS__": "rgba(107,114,128,0.20)",
        "__BTN_R__": "6px", "__BTN_SHADOW__": "none",
        "__CARD_R__": "8px", "__CARD_SHADOW__": "none",
        "__AVATAR_TX__": "#111827",
        "__EXTRA__": "html body .main-title::after { height: 2px !important; } html body .stApp div.stButton > button:hover { transform: none !important; }",
    },
    "royal": {
        "__FONT_URL__": "https://fonts.googleapis.com/css2?family=Cinzel:wght@600;700;800&family=Montserrat:wght@400;500;600;700&display=swap",
        "__FONT__": "'Montserrat', system-ui, sans-serif", "__HFONT__": "'Cinzel', Georgia, serif",
        "light": {"__OVERLAY__": "linear-gradient(180deg, rgba(245,243,255,0.55) 0%, rgba(237,233,254,0.85) 50%, rgba(250,245,255,0.93) 100%)",
                  "__TITLE__": "#4c1d95"},
        "dark": {"__OVERLAY__": "linear-gradient(180deg, rgba(30,10,60,0.74) 0%, rgba(30,10,60,0.88) 50%, rgba(15,5,30,0.95) 100%)",
                 "__TITLE__": "#fcd34d"},
        "__GRAD__": "linear-gradient(135deg, #6d28d9, #4c1d95)",
        "__TOPBAR__": "linear-gradient(90deg, #2e1065, #4c1d95 55%, #581c87)",
        "__TOP_R__": "0 0 16px 16px", "__TOP_SHADOW__": "0 8px 24px rgba(46,16,101,0.40)",
        "__BRAND__": "#fcd34d", "__NAV_LINE__": "rgba(252,211,77,0.40)", "__NAV_TX__": "#ede9fe",
        "__NAV_ACTIVE_BG__": "rgba(252,211,77,0.14)", "__NAV_ACTIVE_TX__": "#fde68a",
        "__ACCENT__": "#7c3aed", "__ACCENT2__": "#fbbf24", "__FOCUS__": "rgba(167,139,250,0.30)",
        "__BTN_R__": "8px", "__BTN_BORDER__": "1px solid #fbbf24", "__BTN_SHADOW__": "0 6px 16px rgba(76,29,149,0.35)", "__BTN_TX__": "#fef3c7",
        "__CARD_BORDER__": "1px solid rgba(251,191,36,0.45)", "__CARD_R__": "10px", "__CARD_SHADOW__": "0 10px 28px rgba(46,16,101,0.16)",
        "__AVATAR_TX__": "#2e1065",
        "__EXTRA__": "html body .main-title::after { background: linear-gradient(90deg, #fbbf24, #f59e0b) !important; }",
    },
}



def build_theme_css(theme, dark, bg_url):
    if theme == "futuristic":
        css = _THEME_FUTURISTIC
        tokens = {}
    elif theme == "playful":
        css = _THEME_PLAYFUL
        tokens = {
            "__DOT__": "0.22" if dark else "0.35",
            "__PLAY_OVERLAY__": ("linear-gradient(180deg, rgba(30,27,75,0.72) 0%, rgba(30,27,75,0.88) 50%, rgba(30,27,75,0.94) 100%)"
                                 if dark else
                                 "linear-gradient(180deg, rgba(255,247,237,0.45) 0%, rgba(255,247,237,0.78) 45%, rgba(253,242,248,0.9) 100%)"),
            "__PLAY_TITLE__": "#fde68a" if dark else "#7c3aed",
        }
    elif theme == "oldtimes":
        css = _THEME_OLDTIMES
        tokens = {
            "__OLD_PAPER__": "#1f160d" if dark else "#efe4c8",
            "__OLD_OVERLAY__": ("linear-gradient(180deg, rgba(31,22,13,0.70) 0%, rgba(31,22,13,0.88) 50%, rgba(31,22,13,0.94) 100%)"
                                if dark else
                                "linear-gradient(180deg, rgba(244,236,216,0.45) 0%, rgba(244,236,216,0.78) 45%, rgba(244,236,216,0.9) 100%)"),
            "__OLD_BRIGHT__": "0.9" if dark else "1",
            "__OLD_INK__": "#f3e6c4" if dark else "#3e2c1c",
            "__OLD_CARD__": "#2b1f13" if dark else "#fbf5e6",
            "__OLD_INPUT__": "#1f160d" if dark else "#fffaf0",
        }
    elif theme in _ACCENT_THEMES:
        cfg = _ACCENT_THEMES[theme]
        css = _THEME_ACCENT
        tokens = {k: v for k, v in cfg.items() if k.startswith("__")}
        tokens.update(cfg["dark" if dark else "light"])
    else:
        return ""
    tokens["__BG_URL__"] = bg_url
    for k, v in tokens.items():
        css = css.replace(k, v)
    return css


# ==========================================
# 🎤 NHẬP BẰNG GIỌNG NÓI (chế độ Điện thoại)
#    Gắn nút 🎤 vào mọi ô nhập chữ. Bấm → nói → chữ tự điền vào ô.
#    Dùng Web Speech API của trình duyệt (miễn phí, không cần API key).
# ==========================================
import streamlit.components.v1 as _components

_VOICE_JS = r"""
<script>
(function () {
  const win = window.parent, doc = win.document;
  // Tắt tính năng tự dịch của Chrome / Edge (app đã có nút đổi ngôn ngữ riêng; tự dịch làm chữ bị nửa Việt nửa Anh)
  try {
    doc.documentElement.setAttribute("translate", "no");
    doc.documentElement.classList.add("notranslate");
    if (!doc.querySelector('meta[name="google"][content="notranslate"]')) {
      const m = doc.createElement("meta"); m.name = "google"; m.content = "notranslate"; doc.head.appendChild(m);
    }
  } catch (e) {}
  const LANG = "__LANG__";
  const MSG_NOSUPPORT = "__MSG_NOSUPPORT__";
  const MSG_DENIED = "__MSG_DENIED__";
  const SR = win.SpeechRecognition || win.webkitSpeechRecognition;

  // Tắt bản cũ (nếu có) trước khi gắn bản mới
  if (win.__scVoice && win.__scVoice.cleanup) { try { win.__scVoice.cleanup(); } catch (e) {} }
  // Gỡ các nút cũ còn sót (của lần tải trước) để gắn lại nút mới hoạt động được
  doc.querySelectorAll(".sc-mic").forEach((b) => b.remove());
  doc.querySelectorAll("[data-sc-mic]").forEach((el) => el.removeAttribute("data-sc-mic"));

  if (!doc.getElementById("sc-mic-style")) {
    const st = doc.createElement("style");
    st.id = "sc-mic-style";
    st.textContent = `
      html body div .sc-mic { position:absolute !important; right:6px !important; top:50% !important;
        transform:translateY(-50%) !important; z-index:5 !important;
        width:38px !important; height:38px !important; min-height:0 !important; border-radius:50% !important;
        border:none !important; cursor:pointer; padding:0 !important; margin:0 !important;
        background:linear-gradient(135deg,#0ea5e9,#0284c7) !important; background-color:#0284c7 !important;
        color:#fff !important; font-size:18px !important; line-height:1 !important;
        box-shadow:0 3px 10px rgba(2,132,199,.4) !important; display:flex !important;
        align-items:center !important; justify-content:center !important; }
      html body div .sc-mic.sc-area { top:auto !important; bottom:8px !important; transform:none !important; }
      html body div .sc-mic.sc-chatmic { right:52px !important; width:34px !important; height:34px !important; }
      [data-testid="stChatInput"].sc-has-mic textarea { padding-right:96px !important; }
      html body div .sc-mic.sc-rec { background:linear-gradient(135deg,#ef4444,#dc2626) !important;
        background-color:#dc2626 !important; animation:scPulse 1s infinite; }
      @keyframes scPulse { 0%{box-shadow:0 0 0 0 rgba(239,68,68,.6)} 100%{box-shadow:0 0 0 14px rgba(239,68,68,0)} }
      .sc-has-mic input, .sc-has-mic textarea { padding-right:50px !important; }
    `;
    doc.head.appendChild(st);
  }

  let active = null;

  function setValue(el, value) {
    const proto = el.tagName === "TEXTAREA" ? win.HTMLTextAreaElement.prototype : win.HTMLInputElement.prototype;
    Object.getOwnPropertyDescriptor(proto, "value").set.call(el, value);
    el.dispatchEvent(new win.Event("input", { bubbles: true }));
  }
  function commit(el) {
    // Báo cho Streamlit là đã nhập xong (giống như bấm ra ngoài ô)
    el.dispatchEvent(new win.FocusEvent("focusout", { bubbles: true }));
    el.dispatchEvent(new win.FocusEvent("blur"));
  }

  function listen(el, btn) {
    if (!SR) { win.alert(MSG_NOSUPPORT); return; }
    if (active) { active.stop(); return; }
    const rec = new SR();
    rec.lang = LANG; rec.interimResults = true; rec.continuous = false; rec.maxAlternatives = 1;
    const base = el.value ? el.value.replace(/\s+$/, "") + " " : "";
    rec.onresult = (ev) => {
      let text = "";
      for (let i = 0; i < ev.results.length; i++) text += ev.results[i][0].transcript;
      setValue(el, base + text.trim());
    };
    rec.onerror = (ev) => {
      if (ev.error === "not-allowed" || ev.error === "service-not-allowed") win.alert(MSG_DENIED);
    };
    rec.onend = () => { btn.classList.remove("sc-rec"); btn.textContent = "🎤"; active = null; commit(el); };
    active = rec;
    btn.classList.add("sc-rec"); btn.textContent = "⏹";
    try { rec.start(); } catch (e) { active = null; btn.classList.remove("sc-rec"); btn.textContent = "🎤"; }
  }

  function attach(el) {
    if (el.dataset.scMic) return;
    const wrap = el.closest('[data-testid="stTextInputRootElement"], [data-testid="stTextAreaRootElement"], [data-testid="stChatInput"]');
    if (!wrap) return;
    if (el.type === "password") return;              // không gắn vào ô mật khẩu
    if (el.closest(".st-key-chat_action")) return;    // ô ẩn của menu tin nhắn
    el.dataset.scMic = "1";
    wrap.style.position = "relative";
    wrap.classList.add("sc-has-mic");
    const btn = doc.createElement("button");
    btn.type = "button";
    const isChat = !!el.closest('[data-testid="stChatInput"]');
    btn.className = "sc-mic" + (isChat ? " sc-chatmic" : (el.tagName === "TEXTAREA" ? " sc-area" : ""));
    btn.textContent = "🎤";
    btn.setAttribute("aria-label", "Voice input");
    btn.addEventListener("mousedown", (e) => e.preventDefault());   // giữ nguyên con trỏ trong ô
    btn.addEventListener("click", (e) => { e.preventDefault(); e.stopPropagation(); listen(el, btn); });
    wrap.appendChild(btn);
  }

  function scan() {
    doc.querySelectorAll('[data-testid="stTextInputRootElement"] input, [data-testid="stTextAreaRootElement"] textarea, [data-testid="stChatInput"] textarea')
       .forEach(attach);
  }

  scan();
  const obs = new win.MutationObserver(() => scan());
  obs.observe(doc.body, { childList: true, subtree: true });

  win.__scVoice = {
    cleanup() {
      obs.disconnect();
      if (active) { try { active.abort(); } catch (e) {} }
      doc.querySelectorAll(".sc-mic").forEach((b) => b.remove());
      doc.querySelectorAll("[data-sc-mic]").forEach((el) => { delete el.dataset.scMic; });
      doc.querySelectorAll(".sc-has-mic").forEach((w) => w.classList.remove("sc-has-mic"));
    }
  };
})();
</script>
"""

_VOICE_OFF_JS = r"""
<script>
(function () {
  const win = window.parent, doc = win.document;
  // Tắt tính năng tự dịch của Chrome / Edge (app đã có nút đổi ngôn ngữ riêng; tự dịch làm chữ bị nửa Việt nửa Anh)
  try {
    doc.documentElement.setAttribute("translate", "no");
    doc.documentElement.classList.add("notranslate");
    if (!doc.querySelector('meta[name="google"][content="notranslate"]')) {
      const m = doc.createElement("meta"); m.name = "google"; m.content = "notranslate"; doc.head.appendChild(m);
    }
  } catch (e) {}
  if (win.__scVoice && win.__scVoice.cleanup) { try { win.__scVoice.cleanup(); } catch (e) {} }
  win.__scVoice = null;
  // Tự dọn luôn (phòng khi bản cũ đã bị tắt trước): gỡ nút micro và các đánh dấu
  doc.querySelectorAll(".sc-mic").forEach((b) => b.remove());
  doc.querySelectorAll("[data-sc-mic]").forEach((el) => el.removeAttribute("data-sc-mic"));
  doc.querySelectorAll(".sc-has-mic").forEach((w) => w.classList.remove("sc-has-mic"));
})();
</script>
"""


def render_voice_input(enabled):
    """Bật (điện thoại) hoặc tắt (máy tính) nút micro trong các ô nhập chữ."""
    if enabled:
        js = (_VOICE_JS
              .replace("__LANG__", SPEECH_LANG.get(ui_lang(), "vi-VN"))
              .replace("__MSG_NOSUPPORT__", tr("Trình duyệt này chưa hỗ trợ nhập bằng giọng nói. Hãy dùng Chrome (Android) hoặc Safari (iPhone).").replace('"', "'"))
              .replace("__MSG_DENIED__", tr("Chưa được phép dùng micro. Hãy cho phép micro cho trang web này trong cài đặt trình duyệt.").replace('"', "'")))
    else:
        js = _VOICE_OFF_JS
    with _js_slot:
        _components.html(js, height=0)


DEVICE_OPTIONS = ["desktop", "mobile"]
DEVICE_LABELS = {"desktop": "💻 Máy tính", "mobile": "📱 Điện thoại"}
DEVICE_COOKIE = "shipcontrol_device"


def _guess_device():
    """Đoán thiết bị từ trình duyệt (người dùng vẫn đổi lại được)."""
    try:
        ua = st.context.headers.get("User-Agent", "") or ""
    except Exception:
        ua = ""
    _MOBILE_HINTS = ("Mobi", "Android", "iPhone", "iPad", "iPod", "Windows Phone", "IEMobile", "Opera Mini",
                     "BlackBerry", "webOS", "Silk/", "Kindle")
    return "mobile" if any(k in ua for k in _MOBILE_HINTS) else "desktop"


# ==========================================
# 💬 TIN NHẮN (kênh chung + nhắn riêng)
# ==========================================
CHAT_MENU = "💬 Tin Nhắn"
CONNECT_MENU = "💬 Tin Nhắn & Meeting"      # 1 nút menu gộp Tin nhắn + Họp online
CHAT_ROLE_ICONS = {"Admin": "🛡️", "WOS Manager": "👑", "Foreman": "👔", "Team Leader": "🧢", "Worker": "👷"}
_YOU_LABEL = {"vi": "Bạn", "en": "You", "zh": "我", "ja": "自分", "ko": "나",
              "th": "คุณ", "id": "Anda", "ms": "Anda", "tl": "Ikaw", "hi": "आप",
              "es": "Tú", "fr": "Vous", "de": "Du", "ru": "Вы", "pt": "Você"}
_EMPTY_CHAT = {
    "vi": "Chưa có tin nhắn nào. Hãy gửi lời chào đầu tiên! 👋",
    "en": "No messages yet. Say hi! 👋",
    "zh": "还没有消息。来打个招呼吧！👋",
    "ja": "まだメッセージはありません。最初のあいさつを送りましょう！👋",
    "ko": "아직 메시지가 없습니다. 먼저 인사해 보세요! 👋",
    "th": "ยังไม่มีข้อความ ทักทายกันก่อนเลย! 👋",
    "id": "Belum ada pesan. Ayo sapa duluan! 👋",
    "ms": "Belum ada mesej. Jom bertegur sapa! 👋",
    "tl": "Wala pang mensahe. Mag-hi ka na! 👋",
    "hi": "अभी कोई संदेश नहीं। पहला नमस्ते भेजें! 👋",
    "es": "Aún no hay mensajes. ¡Saluda! 👋",
    "fr": "Pas encore de messages. Dites bonjour ! 👋",
    "de": "Noch keine Nachrichten. Sag Hallo! 👋",
    "ru": "Сообщений пока нет. Поздоровайтесь! 👋",
    "pt": "Ainda não há mensagens. Diga oi! 👋",
}


def dm_conv_id(a, b):
    a, b = int(a), int(b)
    return f"dm:{min(a, b)}:{max(a, b)}"


def dm_other_id(conv, me):
    _, a, b = conv.split(":")
    return int(b) if int(a) == int(me) else int(a)


def chat_unread_by_conv(me):
    """Số tin chưa đọc theo từng cuộc trò chuyện (chỉ những cuộc mình được xem)."""
    rows = cursor.execute("""
        SELECT m.conv, COUNT(*) FROM chat_messages m
        LEFT JOIN chat_reads r ON r.user_id = ? AND r.conv = m.conv
        WHERE m.sender_id != ? AND m.id > COALESCE(r.last_read_id, 0)
          AND (m.conv = 'general'
               OR m.conv IN (SELECT 'group:' || group_id FROM chat_group_members WHERE user_id = ?))
        GROUP BY m.conv
    """, (me, me, me)).fetchall()
    return {c: n for c, n in rows}


def mark_chat_read(me, conv, last_id):
    try:
        cursor.execute("""
            INSERT INTO chat_reads (user_id, conv, last_read_id) VALUES (?, ?, ?)
            ON CONFLICT(user_id, conv) DO UPDATE SET last_read_id = MAX(last_read_id, excluded.last_read_id)
        """, (me, conv, last_id))
    except sqlite3.OperationalError:
        pass


def _chat_time(iso):
    try:
        t = datetime.fromisoformat(iso)
    except Exception:
        return ""
    return t.strftime("%H:%M") if t.date() == datetime.now().date() else t.strftime("%d/%m %H:%M")


def render_chat_messages(msgs, me, is_admin=False, pinned=None):
    """Vẽ khung tin nhắn (tin của mình bên phải màu xanh, của người khác bên trái).
    Nội dung tin nhắn KHÔNG đi qua bộ dịch, để giữ đúng chữ người gửi viết.
    Mỗi tin có data-... để menu chuột phải (giống Zalo) biết tin nào, của ai, đã ghim / đánh dấu chưa."""
    lang = ui_lang()
    pin_html = ""
    if pinned:
        rows = "".join(
            f"<div class='sc-pin-row'><b>{html.escape(pn)}</b>: {html.escape((pb or '').replace(chr(10), ' ')[:90])}</div>"
            for _, pb, pn in pinned)
        pin_html = f"<div class='sc-pinned'><div class='sc-pin-title'>{html.escape(_PINNED_TITLE.get(lang, _PINNED_TITLE['vi']))}</div>{rows}</div>"
    if not msgs:
        inner = f"<div class='sc-chat-empty'>{html.escape(_EMPTY_CHAT.get(lang, _EMPTY_CHAT['vi']))}</div>"
    else:
        parts = []
        for mid, sid, body, created, fname, uname, role, recalled, starred, is_pinned in msgs:   # mới nhất trước
            mine = (sid == me)
            sender = fname or uname or "?"
            name = _YOU_LABEL.get(lang, "Bạn") if mine else sender
            icon = "" if mine else CHAT_ROLE_ICONS.get(role, "👤") + " "
            marks = (" 📌" if is_pinned else "") + (" ⭐" if starred else "")
            try:
                full_time = datetime.fromisoformat(created).strftime("%d/%m/%Y %H:%M")
            except Exception:
                full_time = created or ""
            if recalled:
                bubble = f"<div class='sc-bubble sc-recalled'>{html.escape(_RECALLED_TEXT.get(lang, _RECALLED_TEXT['vi']))}</div>"
            else:
                bubble = f"<div class='sc-bubble'>{html.escape(body or '').replace(chr(10), '<br>')}</div>"
            attrs = (f" data-id='{mid}' data-mine='{1 if mine else 0}' data-admin='{1 if is_admin else 0}'"
                     f" data-recalled='{1 if recalled else 0}' data-starred='{1 if starred else 0}'"
                     f" data-pinned='{1 if is_pinned else 0}' data-sender='{html.escape(sender, quote=True)}'"
                     f" data-time='{html.escape(full_time, quote=True)}' data-text='{html.escape(body or '', quote=True)}'")
            parts.append(
                f"<div class='sc-msg {'sc-mine' if mine else 'sc-theirs'}'{attrs}>"
                f"<div class='sc-meta'>{html.escape(icon + name)} · {_chat_time(created)}{marks}</div>"
                f"{bubble}</div>")
        inner = "".join(parts)
    _DG._sc_orig_markdown(st._main, f"{pin_html}<div class='sc-chat'>{inner}</div>", unsafe_allow_html=True)


# ==========================================
# 🖱️ MENU CHUỘT PHẢI TRÊN TIN NHẮN (giống Zalo)
#    Chuột phải (máy tính) hoặc nhấn giữ (điện thoại) vào tin nhắn → menu:
#    Trả lời · Chia sẻ · Copy · Ghim · Đánh dấu · Xem chi tiết · Thu hồi · Xóa chỉ ở phía tôi
# ==========================================
CHAT_CTX_LABELS = {
    "vi": {"reply": "Trả lời", "share": "Chia sẻ", "copy": "Copy tin nhắn", "pin": "Ghim tin nhắn", "unpin": "Bỏ ghim",
           "star": "Đánh dấu tin nhắn", "unstar": "Bỏ đánh dấu", "details": "Xem chi tiết", "recall": "Thu hồi",
           "hide": "Xóa chỉ ở phía tôi", "admin_delete": "Xóa tin nhắn (quản trị)", "copied": "Đã sao chép",
           "sender": "Người gửi", "time": "Thời gian", "recall_confirm": "Thu hồi tin nhắn này với mọi người?",
           "admin_confirm": "Xóa tin nhắn này cho mọi người?"},
    "en": {"reply": "Reply", "share": "Share", "copy": "Copy message", "pin": "Pin message", "unpin": "Unpin",
           "star": "Star message", "unstar": "Remove star", "details": "View details", "recall": "Recall",
           "hide": "Delete for me only", "admin_delete": "Delete message (admin)", "copied": "Copied",
           "sender": "Sender", "time": "Time", "recall_confirm": "Recall this message for everyone?",
           "admin_confirm": "Delete this message for everyone?"},
    "zh": {"reply": "回复", "share": "分享", "copy": "复制消息", "pin": "置顶消息", "unpin": "取消置顶",
           "star": "标记消息", "unstar": "取消标记", "details": "查看详情", "recall": "撤回",
           "hide": "仅在我这边删除", "admin_delete": "删除消息（管理员）", "copied": "已复制",
           "sender": "发送者", "time": "时间", "recall_confirm": "要为所有人撤回这条消息吗？",
           "admin_confirm": "要为所有人删除这条消息吗？"},
    "ja": {"reply": "返信", "share": "共有", "copy": "メッセージをコピー", "pin": "メッセージを固定", "unpin": "固定を解除",
           "star": "メッセージにスター", "unstar": "スターを外す", "details": "詳細を表示", "recall": "送信取消",
           "hide": "自分だけ削除", "admin_delete": "メッセージを削除（管理者）", "copied": "コピーしました",
           "sender": "送信者", "time": "日時", "recall_confirm": "このメッセージを全員から取り消しますか？",
           "admin_confirm": "このメッセージを全員から削除しますか？"},
    "ko": {"reply": "답장", "share": "공유", "copy": "메시지 복사", "pin": "메시지 고정", "unpin": "고정 해제",
           "star": "메시지 표시", "unstar": "표시 해제", "details": "자세히 보기", "recall": "회수",
           "hide": "나에게서만 삭제", "admin_delete": "메시지 삭제 (관리자)", "copied": "복사됨",
           "sender": "보낸 사람", "time": "시간", "recall_confirm": "모든 사람에게서 이 메시지를 회수할까요?",
           "admin_confirm": "모든 사람에게서 이 메시지를 삭제할까요?"},
    "th": {"reply": "ตอบกลับ", "share": "แชร์", "copy": "คัดลอกข้อความ", "pin": "ปักหมุดข้อความ", "unpin": "เลิกปักหมุด", "star": "ติดดาวข้อความ", "unstar": "เอาดาวออก", "details": "ดูรายละเอียด", "recall": "เรียกคืน", "hide": "ลบเฉพาะฝั่งฉัน", "admin_delete": "ลบข้อความ (ผู้ดูแล)", "copied": "คัดลอกแล้ว", "sender": "ผู้ส่ง", "time": "เวลา", "recall_confirm": "เรียกคืนข้อความนี้สำหรับทุกคน?", "admin_confirm": "ลบข้อความนี้สำหรับทุกคน?"},
    "id": {"reply": "Balas", "share": "Bagikan", "copy": "Salin pesan", "pin": "Sematkan pesan", "unpin": "Lepas sematan", "star": "Beri bintang", "unstar": "Hapus bintang", "details": "Lihat detail", "recall": "Tarik", "hide": "Hapus untuk saya saja", "admin_delete": "Hapus pesan (admin)", "copied": "Disalin", "sender": "Pengirim", "time": "Waktu", "recall_confirm": "Tarik pesan ini untuk semua orang?", "admin_confirm": "Hapus pesan ini untuk semua orang?"},
    "ms": {"reply": "Balas", "share": "Kongsi", "copy": "Salin mesej", "pin": "Semat mesej", "unpin": "Nyahsemat", "star": "Bintangkan mesej", "unstar": "Buang bintang", "details": "Lihat butiran", "recall": "Tarik balik", "hide": "Padam untuk saya sahaja", "admin_delete": "Padam mesej (admin)", "copied": "Disalin", "sender": "Penghantar", "time": "Masa", "recall_confirm": "Tarik balik mesej ini untuk semua?", "admin_confirm": "Padam mesej ini untuk semua?"},
    "tl": {"reply": "Sumagot", "share": "Ibahagi", "copy": "Kopyahin ang mensahe", "pin": "I-pin ang mensahe", "unpin": "I-unpin", "star": "Lagyan ng star", "unstar": "Alisin ang star", "details": "Tingnan ang detalye", "recall": "Bawiin", "hide": "Burahin para sa akin lang", "admin_delete": "Burahin ang mensahe (admin)", "copied": "Nakopya", "sender": "Nagpadala", "time": "Oras", "recall_confirm": "Bawiin ang mensaheng ito para sa lahat?", "admin_confirm": "Burahin ang mensaheng ito para sa lahat?"},
    "hi": {"reply": "जवाब दें", "share": "शेयर करें", "copy": "संदेश कॉपी करें", "pin": "संदेश पिन करें", "unpin": "अनपिन करें", "star": "संदेश को स्टार करें", "unstar": "स्टार हटाएँ", "details": "विवरण देखें", "recall": "वापस लें", "hide": "सिर्फ़ मेरे लिए हटाएँ", "admin_delete": "संदेश हटाएँ (एडमिन)", "copied": "कॉपी हो गया", "sender": "भेजने वाला", "time": "समय", "recall_confirm": "यह संदेश सबके लिए वापस लें?", "admin_confirm": "यह संदेश सबके लिए हटाएँ?"},
    "es": {"reply": "Responder", "share": "Compartir", "copy": "Copiar mensaje", "pin": "Fijar mensaje", "unpin": "Desfijar", "star": "Destacar mensaje", "unstar": "Quitar destacado", "details": "Ver detalles", "recall": "Anular envío", "hide": "Eliminar solo para mí", "admin_delete": "Eliminar mensaje (admin)", "copied": "Copiado", "sender": "Remitente", "time": "Hora", "recall_confirm": "¿Anular este mensaje para todos?", "admin_confirm": "¿Eliminar este mensaje para todos?"},
    "fr": {"reply": "Répondre", "share": "Partager", "copy": "Copier le message", "pin": "Épingler le message", "unpin": "Désépingler", "star": "Marquer d'une étoile", "unstar": "Retirer l'étoile", "details": "Voir les détails", "recall": "Retirer", "hide": "Supprimer pour moi seulement", "admin_delete": "Supprimer le message (admin)", "copied": "Copié", "sender": "Expéditeur", "time": "Heure", "recall_confirm": "Retirer ce message pour tout le monde ?", "admin_confirm": "Supprimer ce message pour tout le monde ?"},
    "de": {"reply": "Antworten", "share": "Teilen", "copy": "Nachricht kopieren", "pin": "Nachricht anheften", "unpin": "Lösen", "star": "Mit Stern markieren", "unstar": "Stern entfernen", "details": "Details anzeigen", "recall": "Zurückziehen", "hide": "Nur für mich löschen", "admin_delete": "Nachricht löschen (Admin)", "copied": "Kopiert", "sender": "Absender", "time": "Zeit", "recall_confirm": "Diese Nachricht für alle zurückziehen?", "admin_confirm": "Diese Nachricht für alle löschen?"},
    "ru": {"reply": "Ответить", "share": "Поделиться", "copy": "Копировать сообщение", "pin": "Закрепить сообщение", "unpin": "Открепить", "star": "Отметить звёздочкой", "unstar": "Снять звёздочку", "details": "Подробнее", "recall": "Отозвать", "hide": "Удалить только у меня", "admin_delete": "Удалить сообщение (админ)", "copied": "Скопировано", "sender": "Отправитель", "time": "Время", "recall_confirm": "Отозвать это сообщение у всех?", "admin_confirm": "Удалить это сообщение у всех?"},
    "pt": {"reply": "Responder", "share": "Compartilhar", "copy": "Copiar mensagem", "pin": "Fixar mensagem", "unpin": "Desafixar", "star": "Favoritar mensagem", "unstar": "Remover favorito", "details": "Ver detalhes", "recall": "Apagar para todos", "hide": "Apagar só para mim", "admin_delete": "Excluir mensagem (admin)", "copied": "Copiado", "sender": "Remetente", "time": "Hora", "recall_confirm": "Apagar esta mensagem para todos?", "admin_confirm": "Excluir esta mensagem para todos?"},
}
_RECALLED_TEXT = {"vi": "Tin nhắn đã được thu hồi", "en": "Message recalled", "zh": "消息已撤回",
                  "ja": "メッセージは取り消されました", "ko": "회수된 메시지입니다",
                  "th": "ข้อความถูกเรียกคืนแล้ว", "id": "Pesan ditarik", "ms": "Mesej telah ditarik balik",
                  "tl": "Binawi ang mensahe", "hi": "संदेश वापस ले लिया गया", "es": "Mensaje anulado",
                  "fr": "Message retiré", "de": "Nachricht zurückgezogen", "ru": "Сообщение отозвано",
                  "pt": "Mensagem apagada para todos"}
_PINNED_TITLE = {"vi": "📌 Tin nhắn đã ghim", "en": "📌 Pinned messages", "zh": "📌 置顶消息",
                 "ja": "📌 固定されたメッセージ", "ko": "📌 고정된 메시지",
                 "th": "📌 ข้อความที่ปักหมุด", "id": "📌 Pesan yang disematkan", "ms": "📌 Mesej yang disemat",
                 "tl": "📌 Mga naka-pin na mensahe", "hi": "📌 पिन किए गए संदेश", "es": "📌 Mensajes fijados",
                 "fr": "📌 Messages épinglés", "de": "📌 Angeheftete Nachrichten", "ru": "📌 Закреплённые сообщения",
                 "pt": "📌 Mensagens fixadas"}

_CHAT_CTX_JS = r"""
<script>
(function () {
  const win = window.parent, doc = win.document;
  const L = __LABELS__;
  if (win.__scCtx && win.__scCtx.cleanup) { try { win.__scCtx.cleanup(); } catch (e) {} }

  if (!doc.getElementById("sc-ctx-style")) {
    const st = doc.createElement("style");
    st.id = "sc-ctx-style";
    st.textContent = `
      .sc-ctx { position:fixed; z-index:100000; min-width:230px; padding:6px 0; border-radius:14px;
        background:#ffffff; color:#1f2937; box-shadow:0 12px 36px rgba(15,23,42,.22), 0 0 0 1px rgba(15,23,42,.06);
        font-family:inherit; animation:scCtxIn .12s ease-out; user-select:none; }
      @keyframes scCtxIn { from { opacity:0; transform:scale(.96); } to { opacity:1; transform:none; } }
      .sc-ctx-item { display:flex; align-items:center; gap:12px; padding:9px 16px; font-size:15px; cursor:pointer; }
      .sc-ctx-item:hover { background:#f1f5f9; }
      .sc-ctx-item .ic { width:20px; text-align:center; font-size:16px; opacity:.85; }
      .sc-ctx-sep { height:1px; background:#e5e7eb; margin:5px 0; }
      .sc-ctx-danger { color:#dc2626; }
      .sc-ctx-toast { position:fixed; left:50%; bottom:90px; transform:translateX(-50%); z-index:100001;
        background:#111827; color:#fff; padding:8px 16px; border-radius:999px; font-size:14px; animation:scCtxIn .15s ease-out; }
      .sc-msg.sc-ctx-active .sc-bubble { outline:2px solid #0ea5e9; outline-offset:2px; }
    `;
    doc.head.appendChild(st);
  }

  let menu = null, activeMsg = null, pressTimer = null;

  function toast(text, ms) {
    doc.querySelectorAll(".sc-ctx-toast").forEach((x) => x.remove());
    const t = doc.createElement("div"); t.className = "sc-ctx-toast"; t.textContent = text;
    doc.body.appendChild(t); setTimeout(() => t.remove(), ms || 1400);
  }
  function close() {
    if (menu) { menu.remove(); menu = null; }
    if (activeMsg) { activeMsg.classList.remove("sc-ctx-active"); activeMsg = null; }
  }
  function sendAction(action, id) {
    // Gửi lệnh cho Python qua một ô nhập ẩn
    const inp = doc.querySelector(".st-key-chat_action input");
    if (!inp) return;
    const val = JSON.stringify({ a: action, id: id, n: Date.now() + Math.random() });
    Object.getOwnPropertyDescriptor(win.HTMLInputElement.prototype, "value").set.call(inp, val);
    inp.dispatchEvent(new win.Event("input", { bubbles: true }));
    inp.dispatchEvent(new win.KeyboardEvent("keydown", { key: "Enter", code: "Enter", keyCode: 13, which: 13, bubbles: true }));
  }
  function copyText(text) {
    try { win.navigator.clipboard.writeText(text).then(() => toast(L.copied)); }
    catch (e) {
      const ta = doc.createElement("textarea"); ta.value = text; doc.body.appendChild(ta); ta.select();
      try { doc.execCommand("copy"); toast(L.copied); } catch (e2) {} ta.remove();
    }
  }
  function reply(d) {
    const ta = doc.querySelector('[data-testid="stChatInput"] textarea');
    if (!ta) return;
    const snippet = d.text.length > 60 ? d.text.slice(0, 60) + "…" : d.text;
    const val = "↩️ " + d.sender + ": \u201C" + snippet.replace(/\s+/g, " ") + "\u201D\n";
    Object.getOwnPropertyDescriptor(win.HTMLTextAreaElement.prototype, "value").set.call(ta, val);
    ta.dispatchEvent(new win.Event("input", { bubbles: true }));
    ta.focus();
  }

  function item(ic, label, fn, danger) {
    const it = doc.createElement("div");
    it.className = "sc-ctx-item" + (danger ? " sc-ctx-danger" : "");
    it.innerHTML = '<span class="ic"></span><span></span>';
    it.firstChild.textContent = ic; it.lastChild.textContent = label;
    it.addEventListener("click", (e) => { e.stopPropagation(); close(); fn(); });
    return it;
  }
  function sep() { const s = doc.createElement("div"); s.className = "sc-ctx-sep"; return s; }

  function open(msgEl, x, y) {
    close();
    const d = msgEl.dataset;
    if (!d.id || d.recalled === "1") return;
    const data = { id: d.id, text: d.text || "", sender: d.sender || "", time: d.time || "" };
    activeMsg = msgEl; msgEl.classList.add("sc-ctx-active");
    menu = doc.createElement("div"); menu.className = "sc-ctx";
    menu.appendChild(item("❝", L.reply, () => reply(data)));
    menu.appendChild(item("↗", L.share, () => {
      if (win.navigator.share) { win.navigator.share({ text: data.text }).catch(() => {}); } else { copyText(data.text); }
    }));
    menu.appendChild(sep());
    menu.appendChild(item("⧉", L.copy, () => copyText(data.text)));
    menu.appendChild(sep());
    menu.appendChild(item("📌", d.pinned === "1" ? L.unpin : L.pin, () => sendAction(d.pinned === "1" ? "unpin" : "pin", data.id)));
    menu.appendChild(item("☆", d.starred === "1" ? L.unstar : L.star, () => sendAction(d.starred === "1" ? "unstar" : "star", data.id)));
    menu.appendChild(item("ⓘ", L.details, () => toast(L.sender + ": " + data.sender + "  ·  " + L.time + ": " + data.time, 3500)));
    menu.appendChild(sep());
    if (d.mine === "1") {
      menu.appendChild(item("↺", L.recall, () => sendAction("recall", data.id), true));
    } else if (d.admin === "1") {
      menu.appendChild(item("🗑", L.admin_delete, () => { if (win.confirm(L.admin_confirm)) sendAction("admin_delete", data.id); }, true));
    }
    menu.appendChild(item("🗑", L.hide, () => sendAction("hide", data.id), true));
    doc.body.appendChild(menu);
    // Menu mở LÊN TRÊN chỗ bấm (đáy menu ở ngay con trỏ), để không bị che ở cuối màn hình
    // (ví dụ nút "Manage app"). Chỉ khi phía trên không đủ chỗ mới mở xuống dưới.
    const r = { width: menu.offsetWidth, height: menu.offsetHeight };   // kích thước thật (không bị ảnh hưởng bởi hiệu ứng phóng to)
    const vw = win.innerWidth, vh = win.innerHeight;
    const BOTTOM_SAFE = 72;                                  // chừa chỗ cho thanh/nút ở đáy màn hình
    let top = y - r.height - 6;
    if (top < 8) top = Math.min(y + 6, vh - r.height - BOTTOM_SAFE);
    top = Math.max(8, top);
    menu.style.left = Math.max(8, Math.min(x, vw - r.width - 8)) + "px";
    menu.style.top = top + "px";
    menu.style.transformOrigin = (top < y ? "bottom" : "top") + " left";
  }

  function onContext(e) {
    const m = e.target.closest && e.target.closest(".sc-msg");
    if (!m) return;
    e.preventDefault();
    open(m, e.clientX, e.clientY);
  }
  // Điện thoại (iPhone không có chuột phải): nhấn giữ 0,5 giây
  function onTouchStart(e) {
    const m = e.target.closest && e.target.closest(".sc-msg");
    if (!m) return;
    const t = e.touches[0];
    pressTimer = setTimeout(() => { open(m, t.clientX, t.clientY); }, 500);
  }
  function cancelPress() { if (pressTimer) { clearTimeout(pressTimer); pressTimer = null; } }
  function onDocClick(e) { if (menu && !menu.contains(e.target)) close(); }
  function onKey(e) { if (e.key === "Escape") close(); }

  doc.addEventListener("contextmenu", onContext);
  doc.addEventListener("touchstart", onTouchStart, { passive: true });
  doc.addEventListener("touchend", cancelPress);
  doc.addEventListener("touchmove", cancelPress);
  doc.addEventListener("click", onDocClick, true);
  doc.addEventListener("keydown", onKey);
  win.addEventListener("scroll", close, true);

  win.__scCtx = {
    cleanup() {
      close();
      doc.removeEventListener("contextmenu", onContext);
      doc.removeEventListener("touchstart", onTouchStart);
      doc.removeEventListener("touchend", cancelPress);
      doc.removeEventListener("touchmove", cancelPress);
      doc.removeEventListener("click", onDocClick, true);
      doc.removeEventListener("keydown", onKey);
      win.removeEventListener("scroll", close, true);
    }
  };
})();
</script>
"""


def render_chat_context_menu():
    labels = CHAT_CTX_LABELS.get(ui_lang(), CHAT_CTX_LABELS["vi"])
    _components.html(_CHAT_CTX_JS.replace("__LABELS__", json.dumps(labels, ensure_ascii=False)), height=0)


def handle_chat_action(raw, conv, me, role):
    """Thực hiện lệnh từ menu chuột phải (kiểm tra quyền ở phía máy chủ)."""
    try:
        act = json.loads(raw)
        action, mid = act.get("a"), int(act.get("id"))
    except Exception:
        return
    row = cursor.execute("SELECT sender_id, conv FROM chat_messages WHERE id = ?", (mid,)).fetchone()
    if not row or row[1] != conv:
        return
    now = datetime.now().isoformat(timespec="seconds")
    if action == "pin":
        cursor.execute("INSERT OR IGNORE INTO chat_pins (conv, message_id, pinned_by, pinned_at) VALUES (?, ?, ?, ?)",
                       (conv, mid, me, now))
    elif action == "unpin":
        cursor.execute("DELETE FROM chat_pins WHERE conv = ? AND message_id = ?", (conv, mid))
    elif action == "star":
        cursor.execute("INSERT OR IGNORE INTO chat_stars (user_id, message_id) VALUES (?, ?)", (me, mid))
    elif action == "unstar":
        cursor.execute("DELETE FROM chat_stars WHERE user_id = ? AND message_id = ?", (me, mid))
    elif action == "hide":
        cursor.execute("INSERT OR IGNORE INTO chat_hidden (user_id, message_id) VALUES (?, ?)", (me, mid))
    elif action == "recall" and row[0] == me:
        cursor.execute("UPDATE chat_messages SET body = '', recalled = 1 WHERE id = ?", (mid,))
        cursor.execute("DELETE FROM chat_pins WHERE message_id = ?", (mid,))
    elif action == "admin_delete" and is_chat_admin(conv, me, role):
        cursor.execute("DELETE FROM chat_messages WHERE id = ?", (mid,))
        cursor.execute("DELETE FROM chat_pins WHERE message_id = ?", (mid,))
        cursor.execute("DELETE FROM chat_stars WHERE message_id = ?", (mid,))
        cursor.execute("DELETE FROM chat_hidden WHERE message_id = ?", (mid,))


def get_secret(key, default=None):
    """Đọc cấu hình bí mật từ .streamlit/secrets.toml (hoặc Secrets trên Streamlit Cloud)."""
    try:
        return st.secrets[key]
    except Exception:
        return default


# 📍 1. ĐƯỜNG LINK TRANG WEB (đặt APP_URL trong Secrets)
APP_URL = get_secret("APP_URL", "https://your-app.streamlit.app")

# ==========================================
# 💬 NHÓM CHAT + 🎥 HỌP ONLINE
# ==========================================
MEETING_MENU = "🎥 Online Meeting"
MEETING_HOST_ROLES = ["Foreman", "Team Leader"]
# Máy chủ họp video (Jitsi Meet). Có thể đổi trong Secrets: MEETING_SERVER = "https://..."
MEETING_SERVER = str(get_secret("MEETING_SERVER", "https://meet.jit.si")).rstrip("/")


def is_chat_admin(conv, uid, role):
    """Quản trị của cuộc trò chuyện: nhóm → người tạo nhóm; kênh chung → Admin / WOS Manager."""
    if conv.startswith("group:"):
        row = cursor.execute("SELECT created_by FROM chat_groups WHERE id = ?", (int(conv.split(":")[1]),)).fetchone()
        return bool(row and row[0] == uid)
    if conv == "general":
        return role in ("Admin", "WOS Manager")
    return False


def post_chat_message(conv, sender_id, body):
    cursor.execute("INSERT INTO chat_messages (conv, sender_id, body, created_at) VALUES (?, ?, ?, ?)",
                   (conv, sender_id, str(body).strip()[:2000], datetime.now().isoformat(timespec="seconds")))


def create_chat_group(name, owner_id, member_ids):
    cursor.execute("INSERT INTO chat_groups (name, created_by, created_at) VALUES (?, ?, ?)",
                   (name, owner_id, datetime.now().isoformat(timespec="seconds")))
    gid = cursor.lastrowid
    for uid in set([owner_id] + list(member_ids)):
        cursor.execute("INSERT OR IGNORE INTO chat_group_members (group_id, user_id) VALUES (?, ?)", (gid, uid))
    return gid


def user_chat_groups(me):
    return cursor.execute("""
        SELECT g.id, g.name, g.created_by,
               (SELECT COUNT(*) FROM chat_group_members x WHERE x.group_id = g.id) AS n
        FROM chat_groups g JOIN chat_group_members m ON m.group_id = g.id
        WHERE m.user_id = ? ORDER BY g.name COLLATE NOCASE
    """, (me,)).fetchall()


def is_group_member(gid, uid):
    return cursor.execute("SELECT 1 FROM chat_group_members WHERE group_id = ? AND user_id = ?", (gid, uid)).fetchone() is not None


def group_members(gid):
    return cursor.execute("""
        SELECT u.id, u.username, u.fullname, u.role FROM chat_group_members m
        JOIN users u ON u.id = m.user_id WHERE m.group_id = ? AND u.is_deleted = 0
        ORDER BY u.fullname COLLATE NOCASE
    """, (gid,)).fetchall()


def delete_chat_group(gid):
    cursor.execute("DELETE FROM chat_group_members WHERE group_id = ?", (gid,))
    cursor.execute("DELETE FROM chat_messages WHERE conv = ?", (f"group:{gid}",))
    cursor.execute("DELETE FROM chat_reads WHERE conv = ?", (f"group:{gid}",))
    cursor.execute("DELETE FROM chat_groups WHERE id = ?", (gid,))


def leave_chat_group(gid, uid):
    cursor.execute("DELETE FROM chat_group_members WHERE group_id = ? AND user_id = ?", (gid, uid))
    left = cursor.execute("SELECT user_id FROM chat_group_members WHERE group_id = ? ORDER BY user_id LIMIT 1", (gid,)).fetchone()
    if not left:
        delete_chat_group(gid)                      # không còn ai → xóa nhóm
    else:
        owner = cursor.execute("SELECT created_by FROM chat_groups WHERE id = ?", (gid,)).fetchone()
        if owner and owner[0] == uid:               # chủ nhóm rời → chuyển quyền cho người khác
            cursor.execute("UPDATE chat_groups SET created_by = ? WHERE id = ?", (left[0], gid))


def format_room_id(room_id):
    d = "".join(ch for ch in str(room_id) if ch.isdigit())
    return " ".join(d[i:i + 3] for i in range(0, len(d), 3))


def create_meeting(title, host_id, start_at):
    while True:
        room_id = str(secrets.randbelow(900_000_000) + 100_000_000)          # 9 chữ số, giống Zoom
        if not cursor.execute("SELECT 1 FROM meetings WHERE room_id = ?", (room_id,)).fetchone():
            break
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"                           # bỏ các ký tự dễ nhầm (0/O, 1/I)
    password = "".join(secrets.choice(alphabet) for _ in range(6))
    video_room = "ShipControl" + secrets.token_hex(10)                      # tên phòng video bí mật, không đoán được
    cursor.execute("""INSERT INTO meetings (room_id, password, title, host_id, video_room, start_at, created_at, is_active)
                      VALUES (?, ?, ?, ?, ?, ?, ?, 1)""",
                   (room_id, password, title, host_id, video_room, start_at, datetime.now().isoformat(timespec="seconds")))
    return room_id, password


def find_meeting(room_input, password_input):
    rid = "".join(ch for ch in str(room_input or "") if ch.isdigit())
    pw = str(password_input or "").strip().upper()
    if not rid or not pw:
        return None
    row = cursor.execute("SELECT id, password FROM meetings WHERE room_id = ? AND is_active = 1", (rid,)).fetchone()
    if row and hmac.compare_digest(row[1].upper(), pw):
        return row
    return None


def meeting_url(video_room, display_name):
    from urllib.parse import quote
    return f"{MEETING_SERVER}/{video_room}#userInfo.displayName=%22{quote(str(display_name))}%22"


def meeting_code_card(room_id, password, compact=False):
    lab_id, lab_pw = html.escape(tr("🔢 ID phòng")), html.escape(tr("🔑 Mật khẩu"))
    return (f"<div class='sc-meet-card{' sc-compact' if compact else ''}'>"
            f"<div><span>{lab_id}</span><b>{html.escape(format_room_id(room_id))}</b></div>"
            f"<div><span>{lab_pw}</span><b>{html.escape(password)}</b></div></div>")


# CẤU HÌNH GIAO DIỆN
st.set_page_config(
    page_title="ShipControl - Quản Lý Công Việc Tàu",
    page_icon="🚢",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# KHỞI TẠO COOKIE MANAGER
cookie_manager = stx.CookieManager(key="shipcontrol_cookie_mgr")
SESSION_COOKIE = "shipcontrol_session"
SESSION_DAYS = 30

# 2. KẾT NỐI CƠ SỞ DỮ LIỆU SQLITE & AUTO-MIGRATION
# 🗄️ KẾT NỐI DATABASE
#   - isolation_level=None (tự lưu từng lệnh): không còn "giao dịch treo" khi một lượt chạy bị ngắt giữa chừng
#     (đây là nguyên nhân lỗi "database is locked" khi nhiều người / nhiều tab dùng cùng lúc)
#   - timeout 30s + WAL: nhiều người đọc/ghi cùng lúc thì chờ nhau thay vì báo lỗi
conn = sqlite3.connect("ship_control.db", check_same_thread=False, timeout=30, isolation_level=None)
cursor = conn.cursor()
try:
    cursor.execute("PRAGMA journal_mode=WAL")
    cursor.execute("PRAGMA busy_timeout=30000")
    cursor.execute("PRAGMA synchronous=NORMAL")
except sqlite3.OperationalError:
    pass

cursor.execute('''
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE,
        password TEXT,
        fullname TEXT,
        role TEXT DEFAULT 'Pending',
        is_deleted INTEGER DEFAULT 0
    )
''')

cursor.execute('''
    CREATE TABLE IF NOT EXISTS custom_cost_codes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT UNIQUE,
        name TEXT,
        description TEXT,
        is_deleted INTEGER DEFAULT 0
    )
''')

cursor.execute('''
    CREATE TABLE IF NOT EXISTS tasks (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        task_id TEXT UNIQUE,
        task_name TEXT,
        task_cost_code TEXT,
        block TEXT,
        description TEXT,
        initial_by TEXT,
        initial_date TEXT,
        area TEXT,
        deck TEXT,
        frame TEXT,
        in_charge_by TEXT,
        plan_start_date TEXT,
        plan_finish_date TEXT,
        progress INTEGER DEFAULT 0,
        remark TEXT,
        image_path TEXT,
        is_deleted INTEGER DEFAULT 0
    )
''')

# AUTO-MIGRATION
try:
    cursor.execute("ALTER TABLE users ADD COLUMN role TEXT DEFAULT 'Pending'")
except sqlite3.OperationalError:
    pass

try:
    cursor.execute("ALTER TABLE users ADD COLUMN is_deleted INTEGER DEFAULT 0")
except sqlite3.OperationalError:
    pass

tasks_schema_updates = {
    "task_cost_code": "TEXT", "block": "TEXT", "description": "TEXT",
    "initial_by": "TEXT", "initial_date": "TEXT", "area": "TEXT",
    "deck": "TEXT", "frame": "TEXT", "in_charge_by": "TEXT",
    "plan_start_date": "TEXT", "plan_finish_date": "TEXT",
    "progress": "INTEGER DEFAULT 0", "remark": "TEXT",
    "image_path": "TEXT", "is_deleted": "INTEGER DEFAULT 0"
}
cursor.execute("PRAGMA table_info(tasks)")
existing_cols = [col[1] for col in cursor.fetchall()]
for col_name, col_type in tasks_schema_updates.items():
    if col_name not in existing_cols:
        try:
            cursor.execute(f"ALTER TABLE tasks ADD COLUMN {col_name} {col_type}")
        except Exception:
            pass


# 🏭 CỘT WORKSHOP CHO NGƯỜI DÙNG (mỗi tài khoản thuộc 1 workshop / phòng ban)
try:
    cursor.execute("ALTER TABLE users ADD COLUMN workshop TEXT")
except sqlite3.OperationalError:
    pass

# 👥 TEAM & GIAO VIỆC
#   users.leader_id          : Worker thuộc team của Team Leader / Foreman nào
#   tasks.assigned_leader_id : WOS Manager giao việc cho Team Leader / Foreman nào
#   tasks.assigned_worker_id : Team Leader / Foreman giao việc cho Worker nào
for _sql in ["ALTER TABLE users ADD COLUMN leader_id INTEGER",
             "ALTER TABLE users ADD COLUMN team_name TEXT",
             "ALTER TABLE tasks ADD COLUMN assigned_leader_id INTEGER",
             "ALTER TABLE tasks ADD COLUMN assigned_worker_id INTEGER"]:
    try:
        cursor.execute(_sql)
    except sqlite3.OperationalError:
        pass

LEADER_ROLES = ["Team Leader", "Foreman"]

# ==========================================
# 📸 BÁO CÁO HÌNH ẢNH THEO TUẦN
# ==========================================
PHOTO_MAX_MB = 15
PHOTO_MAX_SIDE = 1600
PHOTO_MAX_FILES = 50          # tối đa 50 ảnh mỗi lần tải lên


def _prog(v):
    try:
        f = float(v)
        return 0 if f != f else int(f)
    except Exception:
        return 0


def _parse_day(v):
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
    except Exception:
        return None


def compress_photo(raw):
    """Thu nhỏ ảnh (cạnh dài tối đa 1600px, JPEG) để cơ sở dữ liệu không bị phình to."""
    from PIL import Image, ImageOps
    import io
    im = Image.open(io.BytesIO(raw))
    im = ImageOps.exif_transpose(im)
    if im.mode not in ("RGB", "L"):
        bg = Image.new("RGB", im.size, (255, 255, 255))
        try:
            bg.paste(im, mask=im.convert("RGBA").split()[-1])
        except Exception:
            bg.paste(im.convert("RGB"))
        im = bg
    im.thumbnail((PHOTO_MAX_SIDE, PHOTO_MAX_SIDE))
    out = io.BytesIO()
    im.convert("RGB").save(out, format="JPEG", quality=82, optimize=True)
    return out.getvalue()


def task_week_info(task_row):
    """task_row: (id, task_id, task_name, plan_start, plan_finish, initial_date, progress)
    → (ngày bắt đầu, ngày kết thúc dự kiến, tuần hiện tại, đã xong?)"""
    start = _parse_day(task_row[3]) or _parse_day(task_row[5]) or date.today()
    finish = _parse_day(task_row[4])
    cur_week = max(1, (date.today() - start).days // 7 + 1)
    done = _prog(task_row[6]) >= 100
    return start, finish, cur_week, done


def build_photo_report_html(task_row, weeks, photos_by_week, names):
    import base64
    start, finish, cur_week, done = task_week_info(task_row)
    parts = [f"""<!doctype html><html><head><meta charset='utf-8'>
<title>{html.escape(str(task_row[1]))} - {html.escape(tr('📸 Báo cáo hình ảnh theo tuần'))}</title>
<style>
body{{font-family:Arial,Helvetica,sans-serif;margin:28px;color:#0f172a}}
h1{{font-size:22px;margin:0 0 4px}} .meta{{color:#475569;font-size:13px;margin-bottom:18px}}
h2{{font-size:17px;border-bottom:2px solid #0ea5e9;padding-bottom:4px;margin-top:26px;page-break-after:avoid}}
.grid{{display:grid;grid-template-columns:repeat(2,1fr);gap:12px}}
figure{{margin:0;border:1px solid #e2e8f0;border-radius:8px;padding:6px;page-break-inside:avoid}}
img{{width:100%;border-radius:6px}} figcaption{{font-size:12px;color:#334155;margin-top:4px}}
.empty{{color:#94a3b8;font-style:italic}}
</style></head><body>
<h1>🚢 {html.escape(str(task_row[1]))} — {html.escape(str(task_row[2]))}</h1>
<div class='meta'>{html.escape(tr('Bắt đầu'))}: {start:%d/%m/%Y}"""
             + (f" · {html.escape(tr('Kết thúc dự kiến'))}: {finish:%d/%m/%Y}" if finish else "")
             + f" · {html.escape(tr('Tiến độ'))}: {_prog(task_row[6])}%"
             + (f" · ✅ {html.escape(tr('Đã hoàn thành'))}" if done else "")
             + f" · {datetime.now():%d/%m/%Y %H:%M}</div>"]
    for w in weeks:
        d1 = start + timedelta(days=7 * (w - 1)); d2 = d1 + timedelta(days=6)
        parts.append(f"<h2>📅 {html.escape(tr('Tuần'))} {w} ({d1:%d/%m} – {d2:%d/%m})</h2>")
        items = photos_by_week.get(w, [])
        if not items:
            parts.append(f"<div class='empty'>{html.escape(tr('Chưa có ảnh nào cho tuần này.'))}</div>")
            continue
        parts.append("<div class='grid'>")
        for pid, img, mime, note, by, at in items:
            b64 = base64.b64encode(img).decode()
            cap = " · ".join(x for x in [html.escape(note or ""), html.escape(names.get(by, "")), html.escape(str(at or "")[:16])] if x)
            parts.append(f"<figure><img src='data:{mime or 'image/jpeg'};base64,{b64}'><figcaption>{cap}</figcaption></figure>")
        parts.append("</div>")
    parts.append("</body></html>")
    return "".join(parts).encode("utf-8")


def build_photo_report_pptx(task_row, weeks, photos_by_week, names):
    """Xuất báo cáo hình ảnh thành file PowerPoint (.pptx): trang bìa, mỗi tuần 1 slide (≤6 ảnh/slide), trang tổng kết."""
    import io
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from PIL import Image

    NAVY, NAVY2, AMBER = RGBColor(0x0B, 0x25, 0x45), RGBColor(0x13, 0x40, 0x74), RGBColor(0xF2, 0xA5, 0x41)
    STEEL, INK, MUTED = RGBColor(0x8D, 0xA9, 0xC4), RGBColor(0x0F, 0x1E, 0x33), RGBColor(0x64, 0x74, 0x8B)
    WHITE, TINT = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xEE, 0xF3, 0xF8)
    HEAD, BODY = "Cambria", "Calibri"

    def clean(t):   # bỏ emoji ở đầu chữ (PowerPoint hiển thị emoji không đẹp)
        t = str(t or "")
        while t and not t[0].isalnum():
            t = t[1:]
        return t.strip()

    start, finish, cur_week, done = task_week_info(task_row)
    prog = _prog(task_row[6])
    total_photos = sum(len(v) for v in photos_by_week.values())
    title_txt = f"{task_row[1]} — {task_row[2]}"

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def bg(slide, color):
        f = slide.background.fill
        f.solid(); f.fore_color.rgb = color

    def text(slide, x, y, w, h, t, size, color, bold=False, font=BODY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=None, name=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        para = tf.paragraphs[0]
        para.alignment = align
        r = para.add_run()
        r.text = t
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = color
        if spacing:
            r.font._element.set("spc", str(spacing))
        return tb

    def pill(slide, x, y, w, h, t, fill, color, size=14, name=None):
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = 0.5
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.fill.background()
        sh.shadow.inherit = False
        if name:
            sh.name = name
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
        r = para.add_run(); r.text = t
        r.font.size, r.font.bold, r.font.name = Pt(size), True, BODY
        r.font.color.rgb = color
        return sh

    def footer(slide, n, total):
        text(slide, 0.6, 6.98, 9.5, 0.3, f"ShipControl  ·  {title_txt}", 10, MUTED, name="Footer")
        text(slide, 11.2, 6.98, 1.53, 0.3, f"{n} / {total}", 10, MUTED, align=PP_ALIGN.RIGHT, name="Slide number")

    # Chia mỗi tuần thành các trang ≤ 6 ảnh
    pages = []
    for w in weeks:
        items = photos_by_week.get(w, [])
        chunks = [items[i:i + 6] for i in range(0, len(items), 6)] or [[]]
        for ci, ch in enumerate(chunks):
            pages.append((w, ch, ci + 1, len(chunks)))
    total_slides = len(pages) + 2

    # ---------- 1. Trang bìa ----------
    s1 = prs.slides.add_slide(blank); bg(s1, NAVY)
    text(s1, 0.8, 0.8, 7.5, 0.4, "SHIPCONTROL", 14, AMBER, bold=True, spacing=300, name="Kicker")
    text(s1, 0.8, 1.5, 7.6, 2.4, title_txt, 40, WHITE, bold=True, font=HEAD, anchor=MSO_ANCHOR.BOTTOM, name="Title")
    text(s1, 0.8, 4.05, 7.6, 0.6, clean(tr("📸 Báo cáo hình ảnh theo tuần")), 20, STEEL, name="Subtitle")
    meta = f"{clean(tr('Bắt đầu'))}: {start:%d/%m/%Y}"
    if finish:
        meta += f"     {clean(tr('Kết thúc dự kiến'))}: {finish:%d/%m/%Y}"
    text(s1, 0.8, 5.9, 7.6, 0.4, meta, 14, WHITE, name="Dates")
    text(s1, 0.8, 6.35, 7.6, 0.4, f"{datetime.now():%d/%m/%Y %H:%M}", 12, STEEL, name="Generated")
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.0), Inches(1.6), Inches(3.5), Inches(4.2))
    card.adjustments[0] = 0.08
    card.fill.solid(); card.fill.fore_color.rgb = NAVY2; card.line.fill.background(); card.name = "Progress card"; card.shadow.inherit = False
    text(s1, 9.0, 2.2, 3.5, 1.6, f"{prog}%", 72, AMBER, bold=True, font=HEAD, align=PP_ALIGN.CENTER, name="Progress value")
    text(s1, 9.0, 3.9, 3.5, 0.5, clean(tr("Tiến độ")), 16, WHITE, align=PP_ALIGN.CENTER, name="Progress label")
    pill(s1, 9.55, 4.75, 2.4, 0.5, clean(tr("Đã hoàn thành")) if done else clean(tr("Đang thực hiện")),
         AMBER if done else STEEL, NAVY, size=13, name="Status")

    # ---------- 2. Mỗi tuần 1 slide ----------
    for idx, (w, items, part, parts) in enumerate(pages, start=2):
        sl = prs.slides.add_slide(blank); bg(sl, WHITE)
        d1 = start + timedelta(days=7 * (w - 1)); d2 = d1 + timedelta(days=6)
        head = f"{clean(tr('Tuần'))} {w}" + (f"  ({part}/{parts})" if parts > 1 else "")
        text(sl, 0.6, 0.4, 8.5, 0.75, head, 36, NAVY, bold=True, font=HEAD, name="Title")
        text(sl, 0.6, 1.15, 8.5, 0.4, f"{d1:%d/%m/%Y} – {d2:%d/%m/%Y}", 16, MUTED, name="Dates")
        n_week = len(photos_by_week.get(w, []))
        pill(sl, 10.53, 0.55, 2.2, 0.5, f"{n_week} {clean(tr('ảnh'))}", NAVY, WHITE, name="Photo count")
        if done and w == weeks[-1] and part == parts:
            pill(sl, 8.13, 0.55, 2.2, 0.5, clean(tr("Đã hoàn thành")), AMBER, NAVY, size=13, name="Finished")

        if not items:
            box = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.85), Inches(12.13), Inches(4.85))
            box.adjustments[0] = 0.04
            box.fill.solid(); box.fill.fore_color.rgb = TINT; box.line.fill.background(); box.name = "Empty week"; box.shadow.inherit = False
            text(sl, 0.6, 1.85, 12.13, 4.85, clean(tr("Chưa có ảnh nào cho tuần này.")), 20, MUTED,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="Empty message")
        else:
            n = len(items)
            cols = 1 if n == 1 else 2 if n in (2, 4) else 3
            rows = (n + cols - 1) // cols
            gx, gy, cap_h = 0.3, 0.25, 0.4
            area_x, area_y, area_w, area_h = 0.6, 1.85, 12.13, 4.9
            cell_w = (area_w - gx * (cols - 1)) / cols
            cell_h = (area_h - gy * (rows - 1)) / rows
            img_h_max = cell_h - cap_h
            for i, (pid, img, mime, note_, by, at) in enumerate(items):
                r_, c_ = divmod(i, cols)
                cx = area_x + c_ * (cell_w + gx); cy = area_y + r_ * (cell_h + gy)
                try:
                    iw, ih = Image.open(io.BytesIO(img)).size
                except Exception:
                    iw, ih = 4, 3
                scale = min(cell_w / iw, img_h_max / ih)
                w_in, h_in = iw * scale, ih * scale
                pic = sl.shapes.add_picture(io.BytesIO(img), Inches(cx + (cell_w - w_in) / 2), Inches(cy),
                                            Inches(w_in), Inches(h_in))
                pic.name = f"Photo {i + 1}"
                cap = "  ·  ".join(x for x in [note_ or "", names.get(by, ""), str(at or "")[:16]] if x)
                text(sl, cx, cy + h_in + 0.06, cell_w, cap_h - 0.06, cap[:120], 10, MUTED,
                     align=PP_ALIGN.CENTER, name=f"Caption {i + 1}")
        footer(sl, idx, total_slides)

    # ---------- 3. Trang tổng kết ----------
    se = prs.slides.add_slide(blank); bg(se, NAVY)
    text(se, 0.8, 0.7, 11.7, 0.9, clean(tr("Tổng kết")), 40, WHITE, bold=True, font=HEAD, name="Title")
    text(se, 0.8, 1.55, 11.7, 0.5, title_txt, 18, STEEL, name="Task")
    stats = [(str(len(weeks)), clean(tr("Tổng số tuần"))), (str(total_photos), clean(tr("Tổng số ảnh"))),
             (f"{prog}%", clean(tr("Tiến độ")))]
    for i, (val, lab) in enumerate(stats):
        x = 0.8 + i * 4.0
        c = se.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.6), Inches(3.6), Inches(2.8))
        c.adjustments[0] = 0.08
        c.fill.solid(); c.fill.fore_color.rgb = NAVY2; c.line.fill.background(); c.name = f"Stat card {i + 1}"; c.shadow.inherit = False
        text(se, x, 2.9, 3.6, 1.4, val, 60, AMBER, bold=True, font=HEAD, align=PP_ALIGN.CENTER, name=f"Stat value {i + 1}")
        text(se, x, 4.35, 3.6, 0.6, lab, 16, WHITE, align=PP_ALIGN.CENTER, name=f"Stat label {i + 1}")
    pill(se, 0.8, 5.95, 3.0, 0.55, clean(tr("Đã hoàn thành")) if done else clean(tr("Đang thực hiện")),
         AMBER if done else STEEL, NAVY, size=14, name="Status")
    text(se, 9.5, 6.98, 3.23, 0.3, f"{total_slides} / {total_slides}", 10, STEEL, align=PP_ALIGN.RIGHT, name="Slide number")

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def render_weekly_photo_report(user_data, current_role):
    my_id = user_data["id"]
    base_sql = """SELECT id, task_id, task_name, plan_start_date, plan_finish_date, initial_date, progress
                  FROM tasks WHERE is_deleted = 0"""
    if current_role == "Worker":
        rows = cursor.execute(base_sql + " AND assigned_worker_id = ? ORDER BY id DESC", (my_id,)).fetchall()
    elif current_role in LEADER_ROLES:
        rows = cursor.execute(base_sql + " AND assigned_leader_id = ? ORDER BY id DESC", (my_id,)).fetchall()
    else:
        rows = cursor.execute(base_sql + " ORDER BY id DESC").fetchall()
    if not rows:
        st.info("Bạn chưa có công việc nào để báo cáo.")
        return

    by_id = {r[0]: r for r in rows}
    sel = st.selectbox("Chọn công việc:", list(by_id.keys()), key="wr_task",
                       format_func=lambda k: f"{by_id[k][1]} - {by_id[k][2]} ({_prog(by_id[k][6])}%)")
    task = by_id[sel]
    start, finish, cur_week, done = task_week_info(task)

    photos = cursor.execute("""SELECT id, week, image, mime, note, uploaded_by, uploaded_at
                               FROM task_photos WHERE task_id = ? ORDER BY week, id""", (sel,)).fetchall()
    photos_by_week = {}
    for pid, w, img, mime, note, by, at in photos:
        photos_by_week.setdefault(w, []).append((pid, img, mime, note, by, at))
    max_photo_week = max(photos_by_week) if photos_by_week else 0
    # Đang làm: Tuần 1 → tuần hiện tại. Đã xong (100%): dừng ở tuần cuối cùng có ảnh.
    last_week = (max_photo_week or cur_week) if done else max(cur_week, max_photo_week)
    weeks = list(range(1, last_week + 1))

    info = f"🗓️ {tr('Bắt đầu')}: {start:%d/%m/%Y}"
    if finish:
        info += f" · 🎯 {tr('Kết thúc dự kiến')}: {finish:%d/%m/%Y}"
    info += f" · 📈 {tr('Tiến độ')}: {_prog(task[6])}%"
    info += f" · ✅ {tr('Đã hoàn thành')}" if done else f" · 📍 {tr('Tuần hiện tại')}: {cur_week}"
    st.caption(info)

    names = {u[0]: (u[2] or u[1]) for u in cursor.execute("SELECT id, username, fullname FROM users").fetchall()}
    can_manage = current_role in ["Admin", "WOS Manager"] + LEADER_ROLES

    # ➕ Tải ảnh lên (nút mini)
    m = mini_bar(f"wr{sel}", "📸 Báo cáo hình ảnh theo tuần", [("add", "➕", "Tải ảnh lên")])
    if m == "add":
        with mini_panel(f"wr{sel}"):
            wk_opts = weeks if done else list(range(1, max(cur_week, max_photo_week) + 1))
            up_week = st.selectbox("Tuần", wk_opts, index=len(wk_opts) - 1, key=f"wr_week_{sel}",
                                   format_func=lambda w: f"{tr('Tuần')} {w} ({(start + timedelta(days=7 * (w - 1))):%d/%m} – {(start + timedelta(days=7 * (w - 1) + 6)):%d/%m})")
            files = st.file_uploader("Chọn ảnh từ máy tính", type=["jpg", "jpeg", "png", "webp", "bmp"],
                                     accept_multiple_files=True, key=f"wr_files_{sel}_{st.session_state.get('_wr_nonce', 0)}")
            st.caption("💡 Giữ Ctrl (hoặc Shift) khi chọn để lấy nhiều ảnh cùng lúc. Bấm Upload thêm lần nữa để thêm ảnh khác.")
            if files:
                # Xem trước các ảnh đã chọn (nhỏ) trước khi tải lên
                st.markdown(f"**📷 × {len(files)}**")
                _pv = st.columns(6)
                for _i, _f in enumerate(files[:PHOTO_MAX_FILES]):
                    with _pv[_i % 6]:
                        try:
                            st.image(_f.getvalue(), use_container_width=True)
                        except Exception:
                            st.caption(f"⚠️ {_f.name}")
            note = st.text_input("Ghi chú cho ảnh (không bắt buộc)", key=f"wr_note_{sel}", max_chars=200)
            if st.button("📤 TẢI LÊN", key=f"wr_upload_{sel}", type="primary", use_container_width=True):
                if not files:
                    st.error("Hãy chọn ít nhất 1 ảnh!")
                else:
                    saved, bad = 0, []
                    for f in files[:PHOTO_MAX_FILES]:
                        raw = f.getvalue()
                        if len(raw) > PHOTO_MAX_MB * 1024 * 1024:
                            bad.append(f.name); continue
                        try:
                            img = compress_photo(raw)
                        except Exception:
                            bad.append(f.name); continue
                        cursor.execute("""INSERT INTO task_photos (task_id, week, image, mime, note, uploaded_by, uploaded_at)
                                          VALUES (?, ?, ?, 'image/jpeg', ?, ?, ?)""",
                                       (sel, int(up_week), img, note.strip()[:200], my_id, datetime.now().strftime("%Y-%m-%d %H:%M")))
                        saved += 1
                    conn.commit()
                    if bad:
                        st.session_state["_wr_bad"] = bad
                    st.session_state["_wr_nonce"] = st.session_state.get("_wr_nonce", 0) + 1
                    st.session_state.pop(f"wr_note_{sel}", None)
                    st.toast(f"📸 +{saved}")
                    mini_close(f"wr{sel}")
                    st.rerun()
    if st.session_state.get("_wr_bad"):
        st.warning(tr("Không đọc được ảnh (hoặc ảnh lớn hơn 15 MB): ") + ", ".join(st.session_state.pop("_wr_bad")))

    # 🖥️ Xem kiểu trình chiếu PowerPoint: mỗi tuần là 1 slide, bấm ◀ ▶ hoặc chọn tuần
    pk = f"wr_pill_{sel}"
    if st.session_state.get(pk) not in weeks:
        st.session_state[pk] = weeks[-1]

    def _go(delta, _pk=pk, _weeks=weeks):
        cur = st.session_state.get(_pk) or _weeks[-1]
        i = min(max(_weeks.index(cur) + delta, 0), len(_weeks) - 1)
        st.session_state[_pk] = _weeks[i]

    w = st.session_state[pk]
    n_idx = weeks.index(w)
    nav1, nav2, nav3 = st.columns([1, 3, 1], vertical_alignment="center")
    with nav1:
        st.button("◀ Trước", key=f"wr_prev_{sel}", on_click=_go, args=(-1,), disabled=(n_idx == 0), use_container_width=True)
    with nav2:
        st.pills("Tuần", weeks, key=pk, label_visibility="collapsed",
                 format_func=lambda x: f"{tr('Tuần')} {x}")
    with nav3:
        st.button("Tiếp ▶", key=f"wr_next_{sel}", on_click=_go, args=(1,), disabled=(n_idx == len(weeks) - 1), use_container_width=True)
    w = st.session_state.get(pk) or weeks[-1]

    d1 = start + timedelta(days=7 * (w - 1)); d2 = d1 + timedelta(days=6)
    items = photos_by_week.get(w, [])
    badge = ""
    if done and w == last_week:
        badge = f"<span class='sc-slide-badge sc-done'>✅ {html.escape(tr('Đã hoàn thành'))}</span>"
    elif not done and w == cur_week:
        badge = f"<span class='sc-slide-badge'>📍 {html.escape(tr('Tuần hiện tại'))}</span>"
    _ncol = 0 if not items else (1 if len(items) == 1 else 2 if len(items) in (2, 4) else 3)
    with st.container(key=f"sc_slide_n{_ncol}_{sel}"):
        _DG._sc_orig_markdown(st._main, f"""
<div class='sc-slide-head'>
  <div><div class='sc-slide-kicker'>SHIPCONTROL · {html.escape(str(task[1]))}</div>
       <div class='sc-slide-title'>{html.escape(tr('Tuần'))} {w}</div>
       <div class='sc-slide-dates'>{d1:%d/%m/%Y} – {d2:%d/%m/%Y}</div></div>
  <div class='sc-slide-right'>{badge}<span class='sc-slide-badge sc-count'>📷 {len(items)}</span></div>
</div>""", unsafe_allow_html=True)
        if not items:
            _DG._sc_orig_markdown(st._main, f"<div class='sc-slide-empty'>{html.escape(tr('Chưa có ảnh nào cho tuần này.'))}</div>",
                                  unsafe_allow_html=True)
        else:
            n = len(items)
            ncol = 1 if n == 1 else 2 if n in (2, 4) else 3
            cols = st.columns(ncol, gap="medium")
            for i, (pid, img, mime, note_, by, at) in enumerate(items):
                with cols[i % ncol]:
                    st.image(img, use_container_width=True)
                    cap = " · ".join(x for x in [note_ or "", f"👤 {names.get(by, '?')}", f"🕒 {str(at or '')[:16]}"] if x)
                    _DG._sc_orig_markdown(st._main, f"<div class='sc-slide-cap'>{html.escape(cap)}</div>", unsafe_allow_html=True)
                    if by == my_id or can_manage:
                        if st.button("🗑️", key=f"wr_del_{pid}", help=tr("Xóa")):
                            cursor.execute("DELETE FROM task_photos WHERE id = ?", (pid,))
                            conn.commit()
                            st.toast("🗑️")
                            st.rerun()
        _DG._sc_orig_markdown(st._main, f"<div class='sc-slide-foot'><span>ShipControl · {html.escape(str(task[1]))} — {html.escape(str(task[2]))}</span>"
                                        f"<span>{n_idx + 1} / {len(weeks)}</span></div>", unsafe_allow_html=True)

    # 📊 Xuất PowerPoint (.pptx): bìa + mỗi tuần 1 slide + tổng kết
    try:
        _pptx = build_photo_report_pptx(task, weeks, photos_by_week, names)
        st.download_button("📊 Tải PowerPoint (.pptx)", data=_pptx,
                           file_name=f"Report_{task[1]}_{datetime.now():%Y%m%d}.pptx",
                           mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                           key=f"wr_pptx_{sel}", type="primary")
    except ImportError:
        st.caption("⚠️ python-pptx")

def remove_worker_from_team(worker_id, leader_id):
    """Đưa Worker ra khỏi team và bỏ giao các việc của team đó cho Worker này."""
    cursor.execute("UPDATE users SET leader_id = NULL WHERE id = ? AND leader_id = ?", (worker_id, leader_id))
    cursor.execute("UPDATE tasks SET assigned_worker_id = NULL WHERE assigned_worker_id = ? AND assigned_leader_id = ?",
                   (worker_id, leader_id))
    conn.commit()

# 🏭 DANH SÁCH WORKSHOP / PHÒNG BAN MẶC ĐỊNH (theo bảng Excel)
DEFAULT_WORKSHOPS = [
    ("WOS_01", "Fabrication WS"),
    ("WOS_02", "Hull WS 02"),
    ("WOS_03", "Hull WS 03"),
    ("WOS_04", "Outfitting WS"),
    ("WOS_05", "Piping WS"),
    ("WOS_06", "Painting WS"),
    ("DEP_01", "Technical Dept"),
    ("DEP_02", "QA QC Department"),
    ("DEP_03", "Safety Department"),
    ("DEP_04", "Planing Department"),
    ("DEP_05", "Project Department"),
    ("DEP_06", "Finance Department"),
    ("DEP_07", "HR Department"),
    ("DEP_08", "Security Department"),
]
# Danh sách trên chỉ được tạo sẵn MỘT LẦN khi database còn mới.
# Sau đó thêm / xóa WS Cost Code ngay trong trang ➕ Thêm Công Việc.
cursor.execute("CREATE TABLE IF NOT EXISTS app_flags (name TEXT PRIMARY KEY)")
if not cursor.execute("SELECT 1 FROM app_flags WHERE name = 'workshops_seeded'").fetchone():
    for ws_code, ws_name in DEFAULT_WORKSHOPS:
        cursor.execute("INSERT OR IGNORE INTO custom_cost_codes (code, name, description, is_deleted) VALUES (?, ?, '', 0)",
                       (ws_code, ws_name))
    cursor.execute("INSERT INTO app_flags (name) VALUES ('workshops_seeded')")

# 🧱 DANH SÁCH BLOCK
cursor.execute('''
    CREATE TABLE IF NOT EXISTS blocks (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE,
        is_deleted INTEGER DEFAULT 0
    )
''')
if not cursor.execute("SELECT 1 FROM app_flags WHERE name = 'blocks_seeded'").fetchone():
    # Lấy các Block đã dùng trong công việc cũ làm danh sách ban đầu
    for (b,) in cursor.execute("SELECT DISTINCT TRIM(block) FROM tasks WHERE block IS NOT NULL AND TRIM(block) != ''").fetchall():
        cursor.execute("INSERT OR IGNORE INTO blocks (name, is_deleted) VALUES (?, 0)", (b,))
    cursor.execute("INSERT INTO app_flags (name) VALUES ('blocks_seeded')")

# 💬 BẢNG TIN NHẮN
cursor.execute('''
    CREATE TABLE IF NOT EXISTS chat_messages (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        conv TEXT NOT NULL,
        sender_id INTEGER NOT NULL,
        body TEXT NOT NULL,
        created_at TEXT NOT NULL
    )
''')
cursor.execute("CREATE INDEX IF NOT EXISTS idx_chat_conv ON chat_messages (conv, id)")
try:
    cursor.execute("ALTER TABLE chat_messages ADD COLUMN recalled INTEGER DEFAULT 0")
except sqlite3.OperationalError:
    pass
cursor.execute("CREATE TABLE IF NOT EXISTS chat_pins (conv TEXT, message_id INTEGER, pinned_by INTEGER, pinned_at TEXT, PRIMARY KEY (conv, message_id))")
cursor.execute("CREATE TABLE IF NOT EXISTS chat_stars (user_id INTEGER, message_id INTEGER, PRIMARY KEY (user_id, message_id))")
cursor.execute("CREATE TABLE IF NOT EXISTS chat_hidden (user_id INTEGER, message_id INTEGER, PRIMARY KEY (user_id, message_id))")
cursor.execute('''
    CREATE TABLE IF NOT EXISTS chat_reads (
        user_id INTEGER NOT NULL,
        conv TEXT NOT NULL,
        last_read_id INTEGER NOT NULL DEFAULT 0,
        PRIMARY KEY (user_id, conv)
    )
''')

# 💬 NHÓM CHAT + 🎥 CUỘC HỌP
cursor.execute('''
    CREATE TABLE IF NOT EXISTS chat_groups (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        created_by INTEGER NOT NULL,
        created_at TEXT NOT NULL
    )
''')
cursor.execute('''
    CREATE TABLE IF NOT EXISTS chat_group_members (
        group_id INTEGER NOT NULL,
        user_id INTEGER NOT NULL,
        PRIMARY KEY (group_id, user_id)
    )
''')
cursor.execute('''
    CREATE TABLE IF NOT EXISTS meetings (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        room_id TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL,
        title TEXT NOT NULL,
        host_id INTEGER NOT NULL,
        video_room TEXT NOT NULL,
        start_at TEXT,
        created_at TEXT NOT NULL,
        is_active INTEGER DEFAULT 1
    )
''')

cursor.execute('''
    CREATE TABLE IF NOT EXISTS meeting_invites (
        meeting_id INTEGER NOT NULL,
        user_id INTEGER NOT NULL,
        PRIMARY KEY (meeting_id, user_id)
    )
''')

# 📸 ẢNH BÁO CÁO THEO TUẦN (mỗi ảnh gắn với 1 công việc + số tuần)
cursor.execute('''
    CREATE TABLE IF NOT EXISTS task_photos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        task_id INTEGER NOT NULL,
        week INTEGER NOT NULL,
        image BLOB NOT NULL,
        mime TEXT DEFAULT 'image/jpeg',
        note TEXT,
        uploaded_by INTEGER,
        uploaded_at TEXT
    )
''')
cursor.execute("CREATE INDEX IF NOT EXISTS idx_task_photos ON task_photos (task_id, week)")

# 🔐 BẢNG PHIÊN ĐĂNG NHẬP (cookie chỉ chứa token ngẫu nhiên, không chứa username)
cursor.execute('''
    CREATE TABLE IF NOT EXISTS sessions (
        token TEXT PRIMARY KEY,
        user_id INTEGER,
        expires_at TEXT
    )
''')


# ==========================================
# 🔑 HÀM BẢO MẬT MẬT KHẨU
# ==========================================
PBKDF2_ITERATIONS = 200_000

def hash_password(password):
    """Băm mật khẩu bằng PBKDF2-SHA256 có salt ngẫu nhiên."""
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERATIONS)
    return f"pbkdf2${PBKDF2_ITERATIONS}${salt.hex()}${digest.hex()}"

def verify_password(password, stored):
    """Kiểm tra mật khẩu. Hỗ trợ cả hash SHA-256 kiểu cũ để người dùng cũ vẫn đăng nhập được."""
    if not stored:
        return False
    if stored.startswith("pbkdf2$"):
        try:
            _, iters, salt_hex, digest_hex = stored.split("$")
            digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), int(iters))
            return hmac.compare_digest(digest.hex(), digest_hex)
        except Exception:
            return False
    legacy = hashlib.sha256(password.encode("utf-8")).hexdigest()
    return hmac.compare_digest(legacy, stored)

def is_legacy_hash(stored):
    return bool(stored) and not stored.startswith("pbkdf2$")

LEGACY_DEFAULT_ADMIN_HASH = hashlib.sha256("admin123".encode()).hexdigest()


# ==========================================
# 🎫 HÀM QUẢN LÝ PHIÊN ĐĂNG NHẬP
# ==========================================
def create_session(user_id):
    """Tạo phiên đăng nhập mới. Mỗi tài khoản chỉ được đăng nhập ở MỘT nơi:
    đăng nhập ở máy mới thì phiên ở máy cũ bị xóa (máy cũ sẽ tự đăng xuất)."""
    cursor.execute("DELETE FROM sessions WHERE user_id = ?", (user_id,))
    token = secrets.token_urlsafe(32)
    expires = (datetime.now() + timedelta(days=SESSION_DAYS)).isoformat()
    cursor.execute("INSERT INTO sessions (token, user_id, expires_at) VALUES (?, ?, ?)", (token, user_id, expires))
    conn.commit()
    return token

def session_is_valid(token):
    return cursor.execute("SELECT 1 FROM sessions WHERE token = ? AND expires_at > ?",
                          (token, datetime.now().isoformat())).fetchone() is not None

def get_user_by_session(token):
    row = cursor.execute("""
        SELECT u.id, u.username, u.fullname, u.role, u.workshop
        FROM sessions s JOIN users u ON u.id = s.user_id
        WHERE s.token = ? AND s.expires_at > ? AND u.is_deleted = 0
    """, (token, datetime.now().isoformat())).fetchone()
    return row

def delete_session(token):
    cursor.execute("DELETE FROM sessions WHERE token = ?", (token,))
    conn.commit()

def delete_user_sessions(user_id):
    cursor.execute("DELETE FROM sessions WHERE user_id = ?", (user_id,))
    conn.commit()

# Dọn các phiên đã hết hạn (việc dọn dẹp, lỡ database đang bận thì bỏ qua, lần sau dọn)
try:
    cursor.execute("DELETE FROM sessions WHERE expires_at <= ?", (datetime.now().isoformat(),))
except sqlite3.OperationalError:
    pass


# TẠO TÀI KHOẢN WOS MANAGER MẶC ĐỊNH (mật khẩu lấy từ Secrets, KHÔNG hardcode)
admin_exists = cursor.execute("SELECT id FROM users WHERE username = 'admin'").fetchone()
if not admin_exists:
    initial_admin_pw = get_secret("ADMIN_PASSWORD")
    if not initial_admin_pw:
        initial_admin_pw = secrets.token_urlsafe(12)
        print("=" * 60)
        print(f"[ShipControl] Mật khẩu admin tạm thời: {initial_admin_pw}")
        print("Hãy đăng nhập và đổi mật khẩu ngay, hoặc đặt ADMIN_PASSWORD trong Secrets.")
        print("=" * 60)
    cursor.execute("INSERT INTO users (username, password, fullname, role, is_deleted) VALUES (?, ?, ?, ?, 0)",
                   ('admin', hash_password(initial_admin_pw), 'System Admin', 'Admin'))

# Tài khoản 'admin' luôn là Admin hệ thống (chỉ cấp quyền WOS Manager + workshop)
if cursor.execute("SELECT 1 FROM users WHERE username = 'admin' AND (role IS NULL OR role != 'Admin')").fetchone():
    cursor.execute("UPDATE users SET role = 'Admin' WHERE username = 'admin'")

# 🆘 KHÔI PHỤC TÀI KHOẢN ADMIN: đặt RESET_ADMIN_PASSWORD trong Secrets để đặt lại mật khẩu admin.
# Sau khi đăng nhập được, hãy XÓA dòng RESET_ADMIN_PASSWORD khỏi Secrets.
reset_admin_pw = get_secret("RESET_ADMIN_PASSWORD")

@st.cache_resource(show_spinner=False)
def _apply_admin_reset(pw):
    """Chỉ đặt lại mật khẩu admin MỘT LẦN mỗi khi app khởi động (không phải mỗi lần bấm nút)."""
    c = sqlite3.connect("ship_control.db", timeout=30, isolation_level=None)
    c.execute("UPDATE users SET password = ?, role = 'Admin', is_deleted = 0 WHERE username = 'admin'",
              (hash_password(str(pw)),))
    c.close()
    return True

if reset_admin_pw:
    try:
        _apply_admin_reset(str(reset_admin_pw))
    except sqlite3.OperationalError:
        pass

conn.commit()


# 📧 HÀM GỬI EMAIL THÔNG BÁO CÓ NGƯỜI ĐĂNG KÝ MỚI
# Email chỉ thông báo. Việc cấp Role phải làm trong app bằng tài khoản WOS Manager,
# nên không ai có thể tự cấp quyền cho mình bằng cách sửa đường link.
def send_new_user_email(target_username, target_fullname):
    sender_email = get_secret("SMTP_EMAIL")
    sender_password = get_secret("SMTP_APP_PASSWORD")
    receiver_email = get_secret("ADMIN_EMAIL", sender_email)
    if not sender_email or not sender_password or not receiver_email:
        print("[ShipControl] Chưa cấu hình SMTP trong Secrets - bỏ qua gửi email.")
        return False

    safe_user = html.escape(target_username)
    safe_name = html.escape(target_fullname)
    subject = f"🔔 [ShipControl] Tài khoản mới chờ duyệt: {target_username}"
    body = f"""
    <div style="font-family: Arial, sans-serif; max-width: 600px; padding: 20px; border: 1px solid #e2e8f0; border-radius: 10px;">
        <h2 style="color: #0284c7;">🚢 CÓ TÀI KHOẢN ĐĂNG KÝ MỚI</h2>
        <p>Có người dùng vừa đăng ký trên hệ thống <b>ShipControl</b>:</p>
        <hr>
        <p><b>Họ và Tên:</b> {safe_name}</p>
        <p><b>Tên đăng nhập (Username):</b> {safe_user}</p>
        <hr>
        <p>Đăng nhập bằng tài khoản <b>Admin</b> hoặc <b>WOS Manager</b>, vào mục <b>👥 Quản Lý Phân Quyền</b> để cấp Role.</p>
        <a href="{APP_URL}" style="display: inline-block; background-color: #16a34a; color: white; padding: 10px 18px; text-decoration: none; border-radius: 6px; font-weight: bold;">👉 Mở ShipControl</a>
    </div>
    """

    msg = MIMEMultipart()
    msg['From'] = sender_email
    msg['To'] = receiver_email
    msg['Subject'] = subject
    msg.attach(MIMEText(body, 'html'))

    try:
        server = smtplib.SMTP('smtp.gmail.com', 587, timeout=15)
        server.starttls()
        server.login(sender_email, sender_password)
        server.sendmail(sender_email, receiver_email, msg.as_string())
        server.quit()
        return True
    except Exception as e:
        print("Lỗi gửi email:", e)
        return False

# --- 🔘 THANH MINI: tiêu đề + nút tròn nhỏ ➕ / ➖ (bấm để mở ô thao tác gọn, bấm ✖ để đóng) ---
def _mini_toggle(state_key, mode):
    st.session_state[state_key] = None if st.session_state.get(state_key) == mode else mode

def mini_bar(key, title, actions):
    """actions: [(mode, icon, help)] → trả về mode đang mở (hoặc None)."""
    sk = f"_mini_{key}"
    cur = st.session_state.get(sk)
    if cur not in [a[0] for a in actions]:
        cur = None
    with st.container(key=f"sc_mini_{key}"):
        cols = st.columns([8] + [1] * len(actions), vertical_alignment="center", gap="small")
        with cols[0]:
            if title:
                _t = tr("#### " + title)
                st.markdown(_t if _t != "#### " + title else "#### " + tr(title))
        for i, (mode, icon, hlp) in enumerate(actions):
            with cols[i + 1]:
                active = (cur == mode)
                st.button("✖" if active else icon, key=f"mini_{key}_{mode}", help=tr(hlp),
                          type="primary" if active else "secondary",
                          on_click=_mini_toggle, args=(sk, mode))
    return cur

def mini_select(key, label, options, actions, format_func=str):
    """Ô chọn (selectbox) có nút mini ➕ / ➖ ngay bên cạnh. Trả về (giá trị đang chọn, mode đang mở)."""
    sk = f"_mini_{key}"
    cur = st.session_state.get(sk)
    if cur not in [a[0] for a in actions]:
        cur = None
    pend = f"_pending_sel_{key}"
    if pend in st.session_state:
        st.session_state[f"sel_{key}"] = st.session_state.pop(pend)
    if st.session_state.get(f"sel_{key}") not in options:
        st.session_state.pop(f"sel_{key}", None)
    with st.container(key=f"sc_mini_{key}"):
        cols = st.columns([8] + [1] * len(actions), vertical_alignment="bottom", gap="small")
        with cols[0]:
            val = st.selectbox(label, options, format_func=format_func, key=f"sel_{key}")
        for i, (mode, icon, hlp) in enumerate(actions):
            with cols[i + 1]:
                active = (cur == mode)
                st.button("✖" if active else icon, key=f"mini_{key}_{mode}", help=tr(hlp),
                          type="primary" if active else "secondary",
                          on_click=_mini_toggle, args=(sk, mode))
    return val, cur

def mini_close(key):
    st.session_state[f"_mini_{key}"] = None

def mini_panel(key):
    return st.container(key=f"sc_minip_{key}", border=True)

# --- KHỞI TẠO STATE THEME, AUTH & SUB-TABS ---
if "theme_mode" not in st.session_state:
    st.session_state["theme_mode"] = "Light"

if "logged_in" not in st.session_state:
    st.session_state["logged_in"] = False

if "user_info" not in st.session_state:
    st.session_state["user_info"] = None

if "auth_tab" not in st.session_state:
    st.session_state["auth_tab"] = "login"

if "current_menu" not in st.session_state:
    st.session_state["current_menu"] = "🧰 Bảng Công Việc"

if "edit_sub_tab" not in st.session_state:
    st.session_state["edit_sub_tab"] = "task"

if "trash_sub_tab" not in st.session_state:
    st.session_state["trash_sub_tab"] = "task"

# --- ⚙️ NÚT CÀI ĐẶT (Chế độ Tối + Ngôn ngữ) ---
# Chỗ đặt nút được giữ sẵn ở đầu thanh bên; nút thật được vẽ sau khi đọc cookie (để nhớ ngôn ngữ đã chọn).
# 🧭 THANH MENU NGANG TRÊN CÙNG (thay cho thanh bên trái):
#    [🚢 SHIPCONTROL] [CÁC TRANG ...........................] [⚙️] [👤 tài khoản]
_topbar = st.container(key="sc_topbar", horizontal=True, vertical_alignment="center", gap="small", wrap=False)
with _topbar:
    st.markdown("<div class='sc-brand'>🚢 SHIPCONTROL</div>", unsafe_allow_html=True)
    _nav_slot = st.container(key="sc_nav", horizontal=True, vertical_alignment="center", gap=None, width="stretch", wrap=False)
    _settings_slot = st.container(key="sc_tools", horizontal=True, vertical_alignment="center", gap="small", width="content", wrap=False)
# Chỗ chứa các đoạn mã chạy ngầm (micro, menu tin nhắn): không chiếm chỗ trên trang
_js_slot = st.container(key="sc_js_slot")
dark_mode_on = st.session_state.get("dark_toggle", st.session_state["theme_mode"] == "Dark")
st.session_state["theme_mode"] = "Dark" if dark_mode_on else "Light"

ui_theme = st.session_state.get("ui_theme", "modern")
if ui_theme not in THEME_OPTIONS:
    ui_theme = "modern"
# Chủ đề Tương lai luôn dùng nền tối
is_dark = st.session_state["theme_mode"] == "Dark" or ui_theme == "futuristic"

main_bg = "#0f172a" if is_dark else "#f8f9fa"
text_color = "#ffffff" if is_dark else "#000000"
input_bg = "#1e293b" if is_dark else "#ffffff"
input_text = "#ffffff" if is_dark else "#000000"
border_color = "#334155" if is_dark else "#cbd5e1"
sec_bg = "#1e293b" if is_dark else "#ffffff"
sec_border = "#475569" if is_dark else "#cbd5e1"
sec_text = "#f1f5f9" if is_dark else "#0f172a"
danger_bg = "#1e293b" if is_dark else "#ffffff"

# ✨ VỆT SÁNG NGẪU NHIÊN: mỗi nút có nhịp riêng (2,4–4,4 giây) và bắt đầu lệch nhau,
#    nên các nút không bao giờ sáng cùng lúc. Giá trị ngẫu nhiên được giữ cố định trong 1 phiên
#    để vệt sáng không bị "giật" mỗi khi bấm nút.
import random as _random
if "shine_seed" not in st.session_state:
    st.session_state["shine_seed"] = _random.randint(1, 10**9)
_rng = _random.Random(st.session_state["shine_seed"])
_shine_rules = []
_CONTAINER = 'div[data-testid="stElementContainer"]'
_BTN_AFTER = [" div.stButton > button::after", " div.stDownloadButton > button::after",
              ' div[data-testid="stFormSubmitButton"] > button::after']
for _k in range(5):   # nhịp lặp: 5 giá trị khác nhau
    _sel = ", ".join(f"{_CONTAINER}:nth-child(5n+{_k}){b}" for b in _BTN_AFTER)
    _shine_rules.append(f"{_sel} {{ animation-duration: {_rng.uniform(2.4, 4.4):.2f}s; }}")
for _k in range(7):   # thời điểm bắt đầu: 7 giá trị khác nhau (5 x 7 = 35 tổ hợp)
    _sel = ", ".join(f"{_CONTAINER}:nth-child(7n+{_k}){b}" for b in _BTN_AFTER)
    _shine_rules.append(f"{_sel} {{ animation-delay: -{_rng.uniform(0, 4.4):.2f}s; }}")
shine_css = "\n    ".join(_shine_rules)

# 🚢 ẢNH NỀN TÀU: mỗi trang một ảnh (ảnh miễn phí từ Unsplash, giấy phép Unsplash License)
def _ship_photo(photo_id):
    return f"https://images.unsplash.com/{photo_id}?auto=format&fit=crop&w=1600&q=65"

SHIP_BACKGROUNDS = {
    "login":                   _ship_photo("photo-1605745341075-1b7460b99df8"),
    "🧰 Bảng Công Việc":       _ship_photo("photo-1605745341112-85968b19335b"),
    "➕ Thêm Công Việc":       _ship_photo("photo-1634638022845-1ab614a94128"),
    "📋 Giao Việc":            _ship_photo("photo-1585713181935-d5f622cc2415"),
    "👥 Team Của Tôi":         _ship_photo("photo-1604506522146-316c8bedd874"),
    "✏️ Chỉnh Sửa/Xóa":        _ship_photo("photo-1670121180530-cfcba4438038"),
    "👥 Quản Lý Phân Quyền":   _ship_photo("photo-1691591765923-3bd6f12f4209"),
    "🗑️ Thùng Rác":            _ship_photo("photo-1617952739760-1dcae19a1d93"),
    "📊 Báo Cáo & Khai Báo":   _ship_photo("photo-1724597500306-a4cbb7d1324e"),
    "🔑 Đổi Mật Khẩu":         _ship_photo("photo-1594110336951-5bc8c12d6b27"),
}
_bg_key = st.session_state.get("current_menu") if st.session_state.get("logged_in") else "login"
page_bg_url = SHIP_BACKGROUNDS.get(_bg_key, SHIP_BACKGROUNDS["🧰 Bảng Công Việc"])

# 🎨 GIAO DIỆN HIỆN ĐẠI: bảng màu theo chế độ Sáng / Tối
_danger_keys = ["btn_del_block", "btn_del_cc", "btn_role_del_ws", "chat_del_msg_btn", "chat_clear_btn", "btn_delete_team", "btn_leave_team", "btn_perm_del_task",
                "btn_perm_del_user", "btn_remove_team_member", "btn_soft_delete_task", "btn_soft_delete_user"]
_modern_tokens = {
    "__APP_BG__": (
        # Lớp phủ màu để chữ vẫn dễ đọc, ảnh tàu nhìn xuyên qua phía dưới
        ("linear-gradient(180deg, rgba(8,15,30,0.58) 0%, rgba(8,15,30,0.78) 45%, rgba(8,15,30,0.88) 100%), "
         if is_dark else
         "linear-gradient(180deg, rgba(244,247,251,0.38) 0%, rgba(244,247,251,0.70) 40%, rgba(244,247,251,0.86) 100%), ")
        + f"url('{page_bg_url}') center / cover no-repeat"
    ),
    "__BASE_BG__": "#0b1220" if is_dark else "#f4f7fb",
    "__BG_ANIM__": "bgFade_" + str(abs(hash(_bg_key)) % 100000),
    "__PRELOAD__": " ".join(f"url('{u}')" for u in SHIP_BACKGROUNDS.values()),
    "__SURFACE__": "#111a2e" if is_dark else "#ffffff",
    "__SURFACE_BORDER__": "#22304d" if is_dark else "#e3e8f0",
    "__INPUT_BG__": "#0a1222" if is_dark else "#fbfcfe",
    "__INPUT_BORDER__": "#33466b" if is_dark else "#d5dce8",
    "__INPUT_BORDER_HOVER__": "#4a6190" if is_dark else "#aab6c8",
    "__INPUT_TEXT__": "#f1f5f9" if is_dark else "#0f172a",
    "__PLACEHOLDER__": "#7c8aa5" if is_dark else "#94a3b8",
    "__ICON__": "#cbd5e1" if is_dark else "#475569",
    "__SCHEME__": "dark" if is_dark else "light",
    "__NAV_BG__": "linear-gradient(90deg, #1e3a8a, #1d4ed8)" if not is_dark else "linear-gradient(90deg, #0b1530, #1e3a8a)",
    "__BUBBLE_THEIRS__": "#1e293b" if is_dark else "#eef2f7",
    "__SELECTED_BG__": "rgba(14,165,233,0.18)" if is_dark else "#e0f2fe",
    "__TOGGLE_OFF__": "#475569" if is_dark else "#94a3b8",
    "__TOGGLE_OFF_BORDER__": "#64748b" if is_dark else "#64748b",
    "__HOVER_TINT__": "rgba(255,255,255,0.04)" if is_dark else "#f8fafc",
    "__DANGER_SOFT__": "rgba(239,68,68,0.10)" if is_dark else "#fef2f2",
    "__INFO_BG__": "rgba(14,165,233,0.12)" if is_dark else "#f0f9ff",
    "__INFO_TX__": "#7dd3fc" if is_dark else "#075985",
    "__OK_BG__": "rgba(34,197,94,0.12)" if is_dark else "#f0fdf4",
    "__OK_TX__": "#86efac" if is_dark else "#166534",
    "__WARN_BG__": "rgba(234,179,8,0.12)" if is_dark else "#fefce8",
    "__WARN_TX__": "#fde047" if is_dark else "#854d0e",
    "__ERR_BG__": "rgba(239,68,68,0.12)" if is_dark else "#fef2f2",
    "__ERR_TX__": "#fca5a5" if is_dark else "#991b1b",
    "__DANGER_SEL_HOVER__": ",\n".join(f"div.st-key-{k} div.stButton > button[kind]:hover:not(:disabled)" for k in _danger_keys),
    "__DANGER_SEL__": ",\n".join(f"div.st-key-{k} div.stButton > button[kind]" for k in _danger_keys),
}
modern_css = MODERN_CSS_TEMPLATE
for _tok, _val in _modern_tokens.items():
    modern_css = modern_css.replace(_tok, _val)

st.markdown(f"""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&display=swap');
    .stApp, .main {{
        background-color: {main_bg} !important;
        color: {text_color} !important;
    }}
    
    section[data-testid="stSidebar"] {{
        background-color: #0284c7 !important;
        width: 330px !important;
    }}
    
    .sidebar-header {{
        font-size: 2rem;
        font-weight: 800;
        color: #facc15;
        text-align: center;
        padding: 10px 0;
        margin-bottom: 10px;
    }}

    .made-by-minh {{
        background-color: #facc15;
        color: #0369a1;
        font-size: 1.5rem;
        font-weight: 900;
        font-style: italic;
        text-align: center;
        padding: 10px;
        border-radius: 12px;
        margin-top: 15px;
    }}

    .user-card {{
        background-color: rgba(255, 255, 255, 0.15);
        padding: 12px;
        border-radius: 12px;
        margin-top: 15px;
        text-align: center;
        color: white;
    }}

    /* ================= NÚT BẤM =================
       Nút to, dễ bấm trên điện thoại (kể cả khi đeo găng tay),
       có "đế" phía dưới giống nút bấm vật lý: nhấn vào thì lún xuống. */
    div.stButton > button,
    div.stDownloadButton > button,
    div[data-testid="stFormSubmitButton"] > button {{
        width: 100% !important;
        min-height: 52px !important;
        padding: 10px 18px !important;
        border-radius: 12px !important;
        margin-bottom: 8px !important;
        background-image: none !important;
        transition: transform 0.08s ease, box-shadow 0.08s ease, background-color 0.15s ease, border-color 0.15s ease !important;
    }}
    div.stButton > button *,
    div.stDownloadButton > button *,
    div[data-testid="stFormSubmitButton"] > button * {{
        font-size: 1.1rem !important;
        font-weight: 800 !important;
        letter-spacing: 0.01em !important;
        opacity: 1 !important;
    }}

    /* Nút chính (xanh lá): Lưu, Thêm, Đăng nhập... */
    div.stButton > button[kind="primary"],
    div.stButton > button[data-testid="baseButton-primary"],
    div.stButton > button[data-testid="stBaseButton-primary"],
    div[data-testid="stFormSubmitButton"] > button,
    div.stDownloadButton > button {{
        background-color: #16a34a !important;
        border: 1px solid #15803d !important;
        box-shadow: 0 4px 0 #166534, 0 6px 14px rgba(22, 101, 52, 0.25) !important;
    }}
    div.stButton > button[kind="primary"] *,
    div.stButton > button[data-testid="baseButton-primary"] *,
    div.stButton > button[data-testid="stBaseButton-primary"] *,
    div[data-testid="stFormSubmitButton"] > button *,
    div.stDownloadButton > button * {{
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
    }}
    div.stButton > button[kind="primary"]:hover:not(:disabled),
    div[data-testid="stFormSubmitButton"] > button:hover:not(:disabled),
    div.stDownloadButton > button:hover:not(:disabled) {{
        background-color: #15803d !important;
        transform: translateY(-1px) !important;
        box-shadow: 0 5px 0 #166534, 0 8px 18px rgba(22, 101, 52, 0.3) !important;
    }}

    /* Nút phụ: nền sáng, viền rõ */
    div.stButton > button[kind="secondary"],
    div.stButton > button[data-testid="baseButton-secondary"],
    div.stButton > button[data-testid="stBaseButton-secondary"] {{
        background-color: {sec_bg} !important;
        border: 1px solid {sec_border} !important;
        box-shadow: 0 4px 0 {sec_border} !important;
    }}
    div.stButton > button[kind="secondary"] *,
    div.stButton > button[data-testid="baseButton-secondary"] *,
    div.stButton > button[data-testid="stBaseButton-secondary"] * {{
        color: {sec_text} !important;
        -webkit-text-fill-color: {sec_text} !important;
    }}
    div.stButton > button[kind="secondary"]:hover:not(:disabled) {{
        border-color: #0284c7 !important;
        box-shadow: 0 4px 0 #0284c7 !important;
        transform: translateY(-1px) !important;
    }}

    /* Nút xóa / rời team: đỏ, để không bấm nhầm */
    div.st-key-btn_del_block div.stButton > button[kind],
    div.st-key-btn_del_cc div.stButton > button[kind],
    div.st-key-btn_delete_team div.stButton > button[kind],
    div.st-key-btn_leave_team div.stButton > button[kind],
    div.st-key-btn_perm_del_task div.stButton > button[kind],
    div.st-key-btn_perm_del_user div.stButton > button[kind],
    div.st-key-btn_remove_team_member div.stButton > button[kind],
    div.st-key-btn_soft_delete_task div.stButton > button[kind],
    div.st-key-btn_soft_delete_user div.stButton > button[kind] {{
        background-color: {danger_bg} !important;
        border: 1px solid #dc2626 !important;
        box-shadow: 0 4px 0 #b91c1c !important;
    }}
    div.st-key-btn_del_block div.stButton > button[kind] *,
    div.st-key-btn_del_cc div.stButton > button[kind] *,
    div.st-key-btn_delete_team div.stButton > button[kind] *,
    div.st-key-btn_leave_team div.stButton > button[kind] *,
    div.st-key-btn_perm_del_task div.stButton > button[kind] *,
    div.st-key-btn_perm_del_user div.stButton > button[kind] *,
    div.st-key-btn_remove_team_member div.stButton > button[kind] *,
    div.st-key-btn_soft_delete_task div.stButton > button[kind] *,
    div.st-key-btn_soft_delete_user div.stButton > button[kind] * {{
        color: #dc2626 !important;
        -webkit-text-fill-color: #dc2626 !important;
    }}
    div.st-key-btn_del_block div.stButton > button[kind]:hover:not(:disabled),
    div.st-key-btn_del_cc div.stButton > button[kind]:hover:not(:disabled),
    div.st-key-btn_delete_team div.stButton > button[kind]:hover:not(:disabled),
    div.st-key-btn_leave_team div.stButton > button[kind]:hover:not(:disabled),
    div.st-key-btn_perm_del_task div.stButton > button[kind]:hover:not(:disabled),
    div.st-key-btn_perm_del_user div.stButton > button[kind]:hover:not(:disabled),
    div.st-key-btn_remove_team_member div.stButton > button[kind]:hover:not(:disabled),
    div.st-key-btn_soft_delete_task div.stButton > button[kind]:hover:not(:disabled),
    div.st-key-btn_soft_delete_user div.stButton > button[kind]:hover:not(:disabled) {{
        background-color: #dc2626 !important;
        box-shadow: 0 4px 0 #991b1b !important;
    }}
    div.st-key-btn_del_block div.stButton > button[kind]:hover:not(:disabled) *,
    div.st-key-btn_del_cc div.stButton > button[kind]:hover:not(:disabled) *,
    div.st-key-btn_delete_team div.stButton > button[kind]:hover:not(:disabled) *,
    div.st-key-btn_leave_team div.stButton > button[kind]:hover:not(:disabled) *,
    div.st-key-btn_perm_del_task div.stButton > button[kind]:hover:not(:disabled) *,
    div.st-key-btn_perm_del_user div.stButton > button[kind]:hover:not(:disabled) *,
    div.st-key-btn_remove_team_member div.stButton > button[kind]:hover:not(:disabled) *,
    div.st-key-btn_soft_delete_task div.stButton > button[kind]:hover:not(:disabled) *,
    div.st-key-btn_soft_delete_user div.stButton > button[kind]:hover:not(:disabled) * {{
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
    }}

    /* Nhấn xuống: nút lún vào đế */
    div.stButton > button:active:not(:disabled),
    div.stDownloadButton > button:active:not(:disabled),
    div[data-testid="stFormSubmitButton"] > button:active:not(:disabled) {{
        transform: translateY(3px) !important;
        box-shadow: 0 1px 0 rgba(0, 0, 0, 0.25) !important;
    }}

    /* Nút bị khóa (chưa tick xác nhận...) */
    div.stButton > button:disabled {{
        opacity: 0.45 !important;
        box-shadow: none !important;
        cursor: not-allowed !important;
    }}

    /* Viền khi dùng bàn phím (Tab) */
    div.stButton > button:focus-visible,
    div.stDownloadButton > button:focus-visible,
    div[data-testid="stFormSubmitButton"] > button:focus-visible {{
        outline: 3px solid #facc15 !important;
        outline-offset: 2px !important;
    }}

    /* ================= MENU BÊN TRÁI =================
       Mục chưa chọn: trong suốt trên nền xanh. Mục đang chọn: vàng, giống nhãn "Made by Minh". */
    section[data-testid="stSidebar"] div.stButton > button {{
        min-height: 48px !important;
        justify-content: flex-start !important;
        text-align: left !important;
        margin-bottom: 4px !important;
    }}
    section[data-testid="stSidebar"] div.stButton > button > div,
    section[data-testid="stSidebar"] div.stButton > button p {{
        justify-content: flex-start !important;
        text-align: left !important;
        width: 100% !important;
    }}
    section[data-testid="stSidebar"] div.stButton > button[kind="secondary"],
    section[data-testid="stSidebar"] div.stButton > button[data-testid="baseButton-secondary"],
    section[data-testid="stSidebar"] div.stButton > button[data-testid="stBaseButton-secondary"] {{
        background-color: rgba(255, 255, 255, 0.10) !important;
        border: 1px solid rgba(255, 255, 255, 0.22) !important;
        box-shadow: none !important;
    }}
    section[data-testid="stSidebar"] div.stButton > button[kind="secondary"] *,
    section[data-testid="stSidebar"] div.stButton > button[data-testid="baseButton-secondary"] *,
    section[data-testid="stSidebar"] div.stButton > button[data-testid="stBaseButton-secondary"] * {{
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
        font-size: 1.05rem !important;
        font-weight: 700 !important;
    }}
    section[data-testid="stSidebar"] div.stButton > button[kind="secondary"]:hover:not(:disabled) {{
        background-color: rgba(255, 255, 255, 0.22) !important;
        border-color: rgba(255, 255, 255, 0.5) !important;
        box-shadow: none !important;
        transform: translateX(3px) !important;
    }}
    section[data-testid="stSidebar"] div.stButton > button[kind="primary"],
    section[data-testid="stSidebar"] div.stButton > button[data-testid="baseButton-primary"],
    section[data-testid="stSidebar"] div.stButton > button[data-testid="stBaseButton-primary"] {{
        background-color: #facc15 !important;
        border: 1px solid #eab308 !important;
        box-shadow: 0 4px 0 #a16207 !important;
    }}
    section[data-testid="stSidebar"] div.stButton > button[kind="primary"] *,
    section[data-testid="stSidebar"] div.stButton > button[data-testid="baseButton-primary"] *,
    section[data-testid="stSidebar"] div.stButton > button[data-testid="stBaseButton-primary"] * {{
        color: #0c4a6e !important;
        -webkit-text-fill-color: #0c4a6e !important;
        font-size: 1.05rem !important;
        font-weight: 900 !important;
    }}
    section[data-testid="stSidebar"] div.stButton > button[kind="primary"]:hover:not(:disabled) {{
        background-color: #fde047 !important;
        transform: none !important;
        box-shadow: 0 4px 0 #a16207 !important;
    }}
    /* Nút Đăng Xuất: viền trắng, căn giữa */
    section[data-testid="stSidebar"] div.st-key-btn_logout_bottom div.stButton > button[kind] {{
        justify-content: center !important;
        background-color: transparent !important;
        border: 2px solid rgba(255, 255, 255, 0.7) !important;
        margin-top: 12px !important;
    }}
    section[data-testid="stSidebar"] div.st-key-btn_logout_bottom div.stButton > button[kind] > div,
    section[data-testid="stSidebar"] div.st-key-btn_logout_bottom div.stButton > button[kind] p {{
        justify-content: center !important;
        text-align: center !important;
    }}

    /* ================= HIỆU ỨNG VỆT SÁNG TỰ ĐỘNG =================
       Vệt sáng tự lướt qua nút, mỗi nút một nhịp ngẫu nhiên (trung bình ~3 giây). */
    div.stButton > button,
    div.stDownloadButton > button,
    div[data-testid="stFormSubmitButton"] > button {{
        position: relative !important;
        overflow: hidden !important;
        isolation: isolate;
    }}
    div.stButton > button::after,
    div.stDownloadButton > button::after,
    div[data-testid="stFormSubmitButton"] > button::after {{
        content: "";
        position: absolute;
        top: 0;
        bottom: 0;
        left: 0;
        width: 60%;
        background: linear-gradient(100deg,
                    rgba(255, 255, 255, 0) 0%,
                    rgba(255, 255, 255, 0.55) 50%,
                    rgba(255, 255, 255, 0) 100%);
        transform: translateX(-130%) skewX(-20deg);
        will-change: transform;
        pointer-events: none;
        z-index: 1;
        animation: btnShine 3s cubic-bezier(0.45, 0, 0.35, 1) infinite;
    }}
    /* Lướt qua trong ~0,8 giây đầu, sau đó nghỉ đến hết 3 giây rồi lặp lại */
    @keyframes btnShine {{
        0%   {{ transform: translateX(-130%) skewX(-20deg); }}
        27%  {{ transform: translateX(260%) skewX(-20deg); }}
        100% {{ transform: translateX(260%) skewX(-20deg); }}
    }}
    /* Nhịp và thời điểm ngẫu nhiên cho từng nút */
    {shine_css}

    /* Nút đang bị khóa thì không có vệt sáng */
    div.stButton > button:disabled::after {{
        display: none;
    }}
    /* Nút trên nền trắng: vệt sáng màu xanh nhạt để nhìn thấy được */
    div.stButton > button[kind="secondary"]::after {{
        background: linear-gradient(100deg,
                    rgba(2, 132, 199, 0) 0%,
                    rgba(2, 132, 199, 0.22) 50%,
                    rgba(2, 132, 199, 0) 100%);
    }}
    section[data-testid="stSidebar"] div.stButton > button[kind="secondary"]::after {{
        background: linear-gradient(100deg,
                    rgba(255, 255, 255, 0) 0%,
                    rgba(255, 255, 255, 0.45) 50%,
                    rgba(255, 255, 255, 0) 100%);
    }}

    /* Mục menu vừa được chọn: nền vàng lướt vào từ trái */
    section[data-testid="stSidebar"] div.stButton > button[kind="primary"] {{
        animation: menuSlideIn 0.2s ease-out;
    }}
    @keyframes menuSlideIn {{
        from {{ background-position: 100% 0; transform: translateX(-6px); }}
        to   {{ background-position: 0 0;    transform: translateX(0); }}
    }}

    /* ================= CHUYỂN TRANG KIỂU LƯỚT =================
       Chọn mục bên dưới trong menu: trang mới lướt vào từ bên phải.
       Chọn mục bên trên: trang mới lướt vào từ bên trái. */
    section[data-testid="stMain"], .stMain, .main {{
        overflow-x: hidden !important;
    }}
    div[class*="st-key-pgR_"] {{
        animation: pageFromRight 0.26s cubic-bezier(0.2, 0.8, 0.2, 1) both;
        will-change: transform, opacity;
    }}
    div[class*="st-key-pgL_"] {{
        animation: pageFromLeft 0.26s cubic-bezier(0.2, 0.8, 0.2, 1) both;
        will-change: transform, opacity;
    }}
    @keyframes pageFromRight {{
        from {{ opacity: 0; transform: translate3d(32px, 0, 0); }}
        to   {{ opacity: 1; transform: translateX(0); }}
    }}
    @keyframes pageFromLeft {{
        from {{ opacity: 0; transform: translate3d(-32px, 0, 0); }}
        to   {{ opacity: 1; transform: translateX(0); }}
    }}

    @media (prefers-reduced-motion: reduce) {{
        div.stButton > button::after,
        div.stDownloadButton > button::after,
        div[data-testid="stFormSubmitButton"] > button::after {{
            display: none;
        }}
        section[data-testid="stSidebar"] div.stButton > button[kind="primary"],
        div[class*="st-key-pgR_"], div[class*="st-key-pgL_"] {{
            animation: none !important;
        }}
    }}

    @media (prefers-reduced-motion: reduce) {{
        div.stButton > button, div.stDownloadButton > button,
        div[data-testid="stFormSubmitButton"] > button {{
            transition: none !important;
            transform: none !important;
        }}
    }}

    .big-table-title {{
        font-size: 1.8rem !important;
        font-weight: 800 !important;
        color: {text_color} !important;
        margin-top: 10px !important;
        margin-bottom: 15px !important;
    }}

    div[data-testid="stDataFrame"] th, 
    div[data-testid="stTable"] th {{
        font-size: 1.15rem !important;
        font-weight: 800 !important;
        padding: 12px 8px !important;
    }}

    div[data-testid="stDataFrame"] td, 
    div[data-testid="stTable"] td,
    div[data-testid="stDataFrame"] [role="gridcell"] {{
        font-size: 1.1rem !important;
        font-weight: 600 !important;
        padding: 10px 8px !important;
    }}

    div[data-testid="stAlert"] {{
        background-color: #fef08a !important;
        border: 2px solid #eab308 !important;
        border-radius: 10px !important;
    }}
    
    div[data-testid="stAlert"] * {{
        color: #854d0e !important;
        -webkit-text-fill-color: #854d0e !important;
        font-size: 1.15rem !important;
        font-weight: 800 !important;
    }}


    /* Chữ của nút chọn Role (radio), nhãn các ô nhập và số liệu báo cáo: luôn cùng màu chữ chính */
    .stRadio label, .stRadio label p, .stRadio div[role="radiogroup"] *,
    .stNumberInput label, .stNumberInput label p,
    .stCheckbox label, .stCheckbox label p,
    .stMain [data-testid="stWidgetLabel"], .stMain [data-testid="stWidgetLabel"] p,
    .main [data-testid="stWidgetLabel"], .main [data-testid="stWidgetLabel"] p,
    [data-testid="stMetricLabel"], [data-testid="stMetricLabel"] *,
    [data-testid="stMetricValue"], [data-testid="stMetricValue"] * {{
        color: {text_color} !important;
        -webkit-text-fill-color: {text_color} !important;
        opacity: 1 !important;
    }}
    .stRadio div[role="radiogroup"] p {{
        font-size: 1.1rem !important;
        font-weight: 700 !important;
    }}

    .stTextInput label, .stTextArea label, .stSelectbox label, .stDateInput label, .stSlider label {{
        color: {text_color} !important;
        font-weight: 700 !important;
        font-size: 1.1rem !important;
    }}

    .main-title {{
        font-size: 2.3rem;
        color: {text_color} !important;
        font-weight: 800;
        text-align: center;
        padding: 10px 0;
        border-bottom: 3px solid #22c55e;
        margin-bottom: 25px;
    }}

    {modern_css}
    </style>
""", unsafe_allow_html=True)

_bg_css = BG_CSS_TEMPLATE
for _tok in ("__APP_BG__", "__BG_ANIM__"):
    _bg_css = _bg_css.replace(_tok, _modern_tokens[_tok])
st.markdown(f"<style>{_bg_css}</style>", unsafe_allow_html=True)

# 🎨 Lớp CSS của chủ đề đang chọn (Modern thì không cần thêm gì)
_theme_css = build_theme_css(ui_theme, is_dark, page_bg_url)
if _theme_css:
    st.markdown(f"<style>{_theme_css}</style>", unsafe_allow_html=True)

# --- TIÊU ĐỀ TRANG ---
st.markdown("<div class='main-title'>🚢 SHIPCONTROL - QUẢN LÝ CÔNG VIỆC TÀU</div>", unsafe_allow_html=True)

# --- KHỞI TẠO KHỔI PHỤC ĐĂNG NHẬP CHUẨN ĐỒNG BỘ COOKIE ---
all_cookies = cookie_manager.get_all()

if all_cookies is None:
    st.stop()

saved_token = all_cookies.get(SESSION_COOKIE)

# Tab vừa bị đăng xuất vì tài khoản được mở ở nơi khác thì KHÔNG tự đăng nhập lại
# (nếu không, 2 tab sẽ đá nhau qua lại).
if not st.session_state["logged_in"] and saved_token and not st.session_state.get("_kicked"):
    user_db = get_user_by_session(saved_token)
    if user_db and user_db[3] and user_db[3] != 'Pending':
        # Chỉ 1 tab được dùng: tab mới mở nhận phiên MỚI, phiên cũ bị xóa → tab cũ tự đăng xuất
        new_token = create_session(user_db[0])
        st.session_state["_pending_session_cookie"] = new_token
        st.session_state["logged_in"] = True
        st.session_state["session_token"] = new_token
        st.session_state["user_info"] = {"id": user_db[0], "username": user_db[1], "fullname": user_db[2], "role": user_db[3], "workshop": user_db[4]}
        st.rerun()

# Nếu đang đăng nhập, kiểm tra lại mỗi lần tải trang: tài khoản bị khóa hoặc đổi Role sẽ có hiệu lực ngay
if st.session_state["logged_in"]:
    fresh = cursor.execute("SELECT fullname, role, workshop FROM users WHERE id = ? AND is_deleted = 0",
                           (st.session_state["user_info"]["id"],)).fetchone()
    if not fresh or not fresh[1] or fresh[1] == 'Pending':
        tok = st.session_state.get("session_token")
        if tok:
            delete_session(tok)
        st.session_state["logged_in"] = False
        st.session_state["user_info"] = None
        st.rerun()
    st.session_state["user_info"]["fullname"] = fresh[0]
    st.session_state["user_info"]["role"] = fresh[1]
    st.session_state["user_info"]["workshop"] = fresh[2]

    # 🔒 Tài khoản đã được đăng nhập ở nơi khác → phiên này hết hiệu lực → đăng xuất
    def _kick_if_signed_in_elsewhere():
        tok = st.session_state.get("session_token")
        # Không có mã phiên (tab đăng nhập từ bản app cũ) cũng coi là hết hiệu lực → bắt đăng nhập lại
        if st.session_state.get("logged_in") and (not tok or not session_is_valid(tok)):
            st.session_state["logged_in"] = False
            st.session_state["user_info"] = None
            st.session_state["session_token"] = None
            st.session_state["_kicked"] = True
            st.rerun(scope="app")

    _kick_if_signed_in_elsewhere()

    # Kiểm tra lại mỗi 5 giây kể cả khi không bấm gì, để máy cũ tự đăng xuất ngay
    @st.fragment(run_every=5)
    def _session_watchdog():
        _kick_if_signed_in_elsewhere()

    _session_watchdog()

# 🌐 Nhớ ngôn ngữ đã chọn (lưu trong cookie của trình duyệt)
#    Cookie có thể đến chậm hơn lần chạy đầu, nên cứ khi nào cookie có mà chưa áp dụng thì áp dụng.
_lang_cookie = all_cookies.get(LANG_COOKIE)
if (_lang_cookie in LANG_OPTIONS and not st.session_state.get("_lang_restored")
        and not st.session_state.get("_lang_changed")):
    st.session_state["_lang_restored"] = True
    if _lang_cookie != ui_lang():
        st.session_state["ui_lang"] = _lang_cookie
        st.rerun()

def _on_lang_change():
    st.session_state["ui_lang"] = st.session_state["w_ui_lang"]
    st.session_state["_lang_changed"] = True

# 🎨 Nhớ chủ đề đã chọn (cookie), cùng cách với ngôn ngữ
THEME_COOKIE = "shipcontrol_theme"
_theme_cookie = all_cookies.get(THEME_COOKIE)
if (_theme_cookie in THEME_OPTIONS and not st.session_state.get("_theme_restored")
        and not st.session_state.get("_theme_changed")):
    st.session_state["_theme_restored"] = True
    if _theme_cookie != st.session_state.get("ui_theme", "modern"):
        st.session_state["ui_theme"] = _theme_cookie
        st.rerun()

def _on_theme_change():
    st.session_state["ui_theme"] = st.session_state["w_ui_theme"]
    st.session_state["_theme_changed"] = True

# 📱 Thiết bị (Máy tính / Điện thoại): web TỰ NHẬN BIẾT qua trình duyệt, không cần người dùng chọn
st.session_state["device_mode"] = _guess_device()

# Các nút chọn dùng khóa riêng (w_...), còn giá trị thật lưu ở ui_theme / ui_lang / device_mode.
# Mỗi lần chạy gán lại giá trị cho nút chọn, để nút luôn hiển thị đúng (không bị lệch hay bị reset).
st.session_state.setdefault("ui_lang", "vi")
st.session_state.setdefault("ui_theme", "modern")
st.session_state["w_ui_lang"] = st.session_state["ui_lang"]
st.session_state["w_ui_theme"] = st.session_state["ui_theme"]

with _settings_slot:
    if True:
        with st.popover("⚙️", key="pop_settings", help="Cài đặt"):
            # 🎨 Chủ đề: nút bấm xổ xuống (giống nút Ngôn ngữ)
            st.selectbox("🎨 Chủ đề", THEME_OPTIONS, key="w_ui_theme", on_change=_on_theme_change,
                         format_func=lambda t: THEME_LABELS[t])
            st.toggle("🌙 Chế độ Tối (Dark)", value=(st.session_state["theme_mode"] == "Dark"), key="dark_toggle",
                      disabled=(ui_theme == "futuristic"))
            if ui_theme == "futuristic":
                st.caption("Chủ đề Tương lai luôn dùng nền tối.")
            # 🌐 Ngôn ngữ: nút xanh có quả địa cầu + mũi tên, bấm để chọn
            st.selectbox("🌐 Ngôn ngữ", LANG_OPTIONS, key="w_ui_lang", on_change=_on_lang_change,
                         format_func=lambda c: LANG_LABELS[c])
            if st.session_state["device_mode"] == "mobile":
                st.caption("🎤 Bấm nút micro trong ô nhập chữ để nói thay vì gõ.")


# 🎤 Nút micro trong ô nhập chữ (chỉ ở chế độ Điện thoại)
render_voice_input(st.session_state["device_mode"] == "mobile")
# 🍪 Lưu cookie đăng nhập (để lần sau mở lại vẫn đăng nhập). Ghi lại cho tới khi trình duyệt xác nhận đã lưu.
_pending_cookie = st.session_state.get("_pending_session_cookie")
if _pending_cookie:
    if not st.session_state.get("logged_in") or all_cookies.get(SESSION_COOKIE) == _pending_cookie:
        st.session_state.pop("_pending_session_cookie", None)
    else:
        try:
            cookie_manager.set(SESSION_COOKIE, _pending_cookie, max_age=SESSION_DAYS * 24 * 3600, key="set_session_cookie")
        except Exception:
            pass

if st.session_state.get("_theme_changed") and _theme_cookie != st.session_state.get("ui_theme"):
    try:
        cookie_manager.set(THEME_COOKIE, st.session_state.get("ui_theme"), max_age=365 * 24 * 3600, key="set_theme_cookie")
    except Exception:
        pass

# Chỉ ghi cookie khi người dùng tự đổi ngôn ngữ
if st.session_state.get("_lang_changed") and _lang_cookie != ui_lang():
    try:
        cookie_manager.set(LANG_COOKIE, ui_lang(), max_age=365 * 24 * 3600, key="set_lang_cookie")
    except Exception:
        pass

# ==========================================
# 🔐 HỆ THỐNG XÁC THỰC
# ==========================================
if not st.session_state["logged_in"]:
    # 🏠 TRANG CHƯA ĐĂNG NHẬP: thanh menu có 2 nút  [VỀ CHÚNG TÔI]  [ĐĂNG NHẬP / ĐĂNG KÝ]
    st.session_state.setdefault("landing_view", "about")
    if st.session_state.get("_kicked"):
        st.session_state["landing_view"] = "auth"

    def _open_auth(tab="login"):
        st.session_state["landing_view"] = "auth"
        st.session_state["auth_tab"] = tab

    with _nav_slot:
        st.button("About us", key="btn_land_about",
                  type="primary" if st.session_state["landing_view"] == "about" else "secondary",
                  on_click=lambda: st.session_state.update({"landing_view": "about"}))
        st.button("Đăng nhập / Đăng ký", key="btn_land_auth",
                  type="primary" if st.session_state["landing_view"] == "auth" else "secondary",
                  on_click=_open_auth, args=("login",))

    if st.session_state["landing_view"] == "about":
        # 🎬 Video giới thiệu. Đổi video: đặt ABOUT_VIDEO_URL trong Secrets (link YouTube hoặc .mp4),
        #    hoặc thêm file about_video.mp4 vào cùng thư mục với app.py trên GitHub.
        _about_video = get_secret("ABOUT_VIDEO_URL")
        _local_mp4 = os.path.join(os.path.dirname(os.path.abspath(__file__)), "about_video.mp4")
        if not _about_video and os.path.exists(_local_mp4):
            _about_video = _local_mp4
        if not _about_video:
            _about_video = "https://www.youtube.com/watch?v=pitqMZO9TbA"

        # Phần đầu trang: chữ giới thiệu (trái) + VIDEO (phải) → thấy video ngay, không cần cuộn
        _h1, _h2 = st.columns([1, 1.15], gap="large", vertical_alignment="center")
        with _h1:
            st.markdown("""
<div class='sc-about-hero sc-about-hero-left'>
  <div class='sc-about-kicker'>🚢 SHIPCONTROL</div>
  <h1>Phần mềm quản lý công việc cho xưởng đóng tàu</h1>
  <p>ShipControl giúp đội ngũ xưởng tàu giao việc, theo dõi tiến độ và liên lạc với nhau ở cùng một nơi — trên máy tính lẫn điện thoại.</p>
</div>
""", unsafe_allow_html=True)
        with _h2:
            st.markdown("<div class='sc-about-video-title'><h3>🎬 Từ xưởng đến biển</h3>"
                        "<p>Xem một con tàu được đóng từ những khối thép đầu tiên đến khi hạ thủy.</p></div>",
                        unsafe_allow_html=True)
            with st.container(key="sc_about_video"):
                st.video(_about_video)
            if str(_about_video).startswith("http") and ("youtube.com" in _about_video or "youtu.be" in _about_video):
                st.caption("Không xem được video? " + f"[▶ Mở trên YouTube]({_about_video})")

        st.markdown("""
<div class='sc-about-grid'>
  <div class='sc-about-card'><h3>📋 Quản lý công việc</h3><p>Tạo công việc theo Block, Deck, Frame và theo dõi tiến độ từng phần trăm.</p></div>
  <div class='sc-about-card'><h3>👥 Team & phân quyền</h3><p>Admin, WOS Manager, Foreman, Team Leader và Worker — mỗi người thấy đúng phần việc của mình.</p></div>
  <div class='sc-about-card'><h3>💬 Tin nhắn & họp online</h3><p>Kênh chung, nhóm chat và phòng họp video có ID + mật khẩu.</p></div>
  <div class='sc-about-card'><h3>🌐 15 ngôn ngữ & giọng nói</h3><p>Tiếng Việt, English, 中文, 日本語, 한국어, ไทย, Indonesia, Español, Français, Deutsch, Русский… — và nói thay vì gõ trên điện thoại.</p></div>
</div>
""", unsafe_allow_html=True)

        _a1, _a2, _a3, _a4 = st.columns([1, 1.2, 1.2, 1])
        with _a2:
            st.button("🔑 Đăng Nhập", key="btn_about_login", type="primary", use_container_width=True,
                      on_click=_open_auth, args=("login",))
        with _a3:
            st.button("📝 Đăng Ký Tài Khoản", key="btn_about_register", use_container_width=True,
                      on_click=_open_auth, args=("register",))
        st.caption("Chưa có tài khoản? Đăng ký và chờ quản lý cấp quyền.")
        st.stop()

    if st.session_state.get("_kicked"):
        st.warning("🔒 Tài khoản của bạn vừa được mở ở một tab hoặc thiết bị khác, nên bạn đã bị đăng xuất khỏi trang này. "
                   "Nếu đó không phải bạn, hãy đăng nhập lại và đổi mật khẩu ngay.")

    col_space1, col_center, col_space2 = st.columns([1, 2, 1])
    
    with col_center:
        is_login = st.session_state["auth_tab"] == "login"
        t_col1, t_col2 = st.columns(2)
        
        with t_col1:
            st.button("🔑 Đăng Nhập", type="primary" if is_login else "secondary", use_container_width=True, key="btn_auth_tab_login",
                      on_click=lambda: st.session_state.update({"auth_tab": "login"}))
                
        with t_col2:
            st.button("📝 Đăng Ký Tài Khoản", type="primary" if not is_login else "secondary", use_container_width=True, key="btn_auth_tab_reg",
                      on_click=lambda: st.session_state.update({"auth_tab": "register"}))

        st.markdown("<div style='margin-bottom: 20px;'></div>", unsafe_allow_html=True)

        # 📲 Điện thoại được nhận biết tự động → có nút 🎤 nói để nhập chữ
        if st.session_state["device_mode"] == "mobile":
            st.caption("🎤 Chế độ điện thoại: bấm nút micro trong ô Tên tài khoản để nói thay vì gõ.")

        if st.session_state["auth_tab"] == "login":
            with st.form("form_login_system"):
                login_user = st.text_input("Tên tài khoản (Username):")
                login_pass = st.text_input("Mật khẩu (Password):", type="password")
                btn_login = st.form_submit_button("🚀 ĐĂNG NHẬP")

                if btn_login:
                    if not login_user or not login_pass:
                        st.error("Vui lòng nhập đầy đủ Username và Mật khẩu!")
                    else:
                        user = cursor.execute("SELECT id, username, password, fullname, role, workshop FROM users WHERE username = ? AND is_deleted = 0", (login_user.strip(),)).fetchone()
                        if user and not verify_password(login_pass, user[2]):
                            user = None
                        if user:
                            # Tự động nâng cấp hash cũ (SHA-256) lên PBKDF2
                            if is_legacy_hash(user[2]):
                                cursor.execute("UPDATE users SET password = ? WHERE id = ?", (hash_password(login_pass), user[0]))
                                conn.commit()
                            user_role = user[4] if len(user) > 4 and user[4] else 'Pending'
                            
                            if user_role == 'Pending':
                                st.warning("⏳ Tài khoản của bạn đang chờ Admin / WOS Manager cấp Role. Vui lòng quay lại sau!")
                            else:
                                st.session_state["logged_in"] = True
                                st.session_state["user_info"] = {"id": user[0], "username": user[1], "fullname": user[3], "role": user_role, "workshop": user[5]}
                                
                                token = create_session(user[0])
                                st.session_state["session_token"] = token
                                st.session_state["_kicked"] = False
                                # Cookie được ghi ở lần chạy sau (ghi ngay rồi st.rerun thì trình duyệt không kịp lưu)
                                st.session_state["_pending_session_cookie"] = token
                                st.success(f"Chào mừng {user[3]} ({user_role}) đã quay trở lại!")
                                st.rerun()
                        else:
                            st.error("Sai tên tài khoản, mật khẩu hoặc tài khoản đã bị khóa/xóa!")

        else:
            with st.form("form_register_system"):
                reg_fullname = st.text_input("Họ và Tên:")
                reg_user = st.text_input("Tên đăng nhập mới (Username):")
                reg_pass = st.text_input("Mật khẩu mới:", type="password")
                reg_confirm = st.text_input("Xác nhận lại mật khẩu:", type="password")
                
                btn_register = st.form_submit_button("✨ ĐĂNG KÝ NGAY")

                if btn_register:
                    reg_user = reg_user.strip()
                    reg_fullname = reg_fullname.strip()
                    if not reg_fullname or not reg_user or not reg_pass:
                        st.error("Vui lòng điền đầy đủ các thông tin!")
                    elif len(reg_pass) < 8:
                        st.error("Mật khẩu phải có ít nhất 8 ký tự!")
                    elif len(reg_user) > 50 or len(reg_fullname) > 100:
                        st.error("Username hoặc Họ tên quá dài!")
                    elif reg_pass != reg_confirm:
                        st.error("Mật khẩu xác nhận không trùng khớp!")
                    else:
                        try:
                            # LƯU VỚI ROLE LÀ Pending CHỜ ADMIN DUYỆT EMAIL
                            cursor.execute("INSERT INTO users (username, password, fullname, role, is_deleted) VALUES (?, ?, ?, 'Pending', 0)", 
                                           (reg_user, hash_password(reg_pass), reg_fullname))
                            conn.commit()
                            
                            # GỬI EMAIL THÔNG BÁO ĐẾN ADMIN (không chứa link cấp quyền)
                            send_new_user_email(reg_user, reg_fullname)
                            
                            st.success("🎉 Đăng ký thành công! Tài khoản đang chờ WOS Manager phê duyệt Role.")
                        except sqlite3.IntegrityError:
                            st.error("Tên đăng nhập này đã tồn tại, vui lòng chọn tên khác!")

# ==========================================
# 🚢 GIAO DIỆN CHÍNH
# ==========================================
else:
    user_data = st.session_state["user_info"]
    current_role = user_data.get("role", "Worker")


    if st.session_state.get("current_menu") in ("💬 Tin Nhắn & Họp", "🎥 Họp Online"):   # tên cũ
        st.session_state["current_menu"] = CONNECT_MENU
    if st.session_state.get("current_menu") in (CHAT_MENU, MEETING_MENU):
        st.session_state["connect_tab"] = "meet" if st.session_state["current_menu"] == MEETING_MENU else "chat"
        st.session_state["current_menu"] = CONNECT_MENU
    is_admin = current_role == "Admin"
    is_manager_up = current_role in ["Admin", "WOS Manager"]

    if is_manager_up:
        menu_options = [
            "🧰 Bảng Công Việc", 
            "💬 Tin Nhắn & Meeting",
            "➕ Thêm Công Việc", 
            "📋 Giao Việc",
            "✏️ Chỉnh Sửa/Xóa",
            "👥 Quản Lý Phân Quyền",
            "🗑️ Thùng Rác",
            "📊 Báo Cáo & Khai Báo",
            "🔑 Đổi Mật Khẩu"
        ]
    elif current_role == "Foreman":
        menu_options = [
            "🧰 Bảng Công Việc", 
            "💬 Tin Nhắn & Meeting",
            "➕ Thêm Công Việc", 
            "👥 Team Của Tôi",
            "📋 Giao Việc",
            "✏️ Chỉnh Sửa/Xóa",
            "📊 Báo Cáo & Khai Báo",
            "🔑 Đổi Mật Khẩu"
        ]
    elif current_role == "Team Leader":
        menu_options = [
            "🧰 Bảng Công Việc", 
            "💬 Tin Nhắn & Meeting",
            "➕ Thêm Công Việc",
            "👥 Team Của Tôi",
            "📋 Giao Việc",
            "📊 Báo Cáo & Khai Báo",
            "🔑 Đổi Mật Khẩu"
        ]
    else:
        menu_options = [
            "🧰 Bảng Công Việc",
            "💬 Tin Nhắn & Meeting",
            "👥 Team Của Tôi",
            "📊 Báo Cáo & Khai Báo",
            "🔑 Đổi Mật Khẩu"
        ]

    # 🛡️ Tài khoản Admin: chỉ Quản Lý Phân Quyền + Tin nhắn (Đổi mật khẩu nằm trong nút tài khoản)
    if is_admin:
        menu_options = ["👥 Quản Lý Phân Quyền", CONNECT_MENU]
        if st.session_state.get("current_menu") not in ("👥 Quản Lý Phân Quyền", CONNECT_MENU, "🔑 Đổi Mật Khẩu"):
            st.session_state["current_menu"] = "👥 Quản Lý Phân Quyền"

    def _go_to_page(target):
        st.session_state["current_menu"] = target

    try:
        _chat_unread_total = sum(chat_unread_by_conv(user_data["id"]).values())
    except sqlite3.OperationalError:
        _chat_unread_total = 0

    def _nav_text(item):
        """Chữ trên thanh menu: bỏ biểu tượng ở đầu cho gọn (giống thanh menu trang web)."""
        t = tr(item)
        first, _, rest = t.partition(" ")
        return rest if rest and not any(ch.isalnum() for ch in first) else t

    with _nav_slot:
        for item in menu_options:
            is_selected = (st.session_state["current_menu"] == item)
            _label = _nav_text(item) + (f" 🔴 {_chat_unread_total}" if item == CONNECT_MENU and _chat_unread_total else "")
            # on_click đổi trang TRƯỚC khi app chạy lại → chỉ chạy 1 lần thay vì 2 lần
            st.button(
                _label,
                type="primary" if is_selected else "secondary",
                key=f"btn_menu_{item}",
                on_click=_go_to_page,
                args=(item,),
            )

    menu = st.session_state["current_menu"]
    
    role_icons = {
        "Admin": "🛡️ Admin",
        "WOS Manager": "👑 WOS Manager",
        "Foreman": "👔 Foreman",
        "Team Leader": "🧢 Team Leader",
        "Worker": "👷 Worker"
    }
    role_badge = role_icons.get(current_role, f"👤 {current_role}")

    # 🪪 THẺ TÀI KHOẢN: ngay dưới menu, cho biết đang đăng nhập bằng tài khoản nào và vai trò gì
    _full = str(user_data.get('fullname') or user_data['username'])
    _words = [w for w in _full.split() if w]
    _initials = ((_words[0][0] + (_words[-1][0] if len(_words) > 1 else "")) if _words else "?").upper()
    _ws_code = user_data.get('workshop')
    _ws_row = ""
    if _ws_code:
        _ws_name = cursor.execute("SELECT name FROM custom_cost_codes WHERE code = ?", (_ws_code,)).fetchone()
        _ws_text = f"{_ws_code} · {_ws_name[0]}" if _ws_name and _ws_name[0] else _ws_code
        _ws_row = (f"<div class='pc-row'><span class='pc-label'>Workshop</span>"
                   f"<span class='pc-val'>{html.escape(_ws_text)}</span></div>")
    _acct_pop = _settings_slot.popover(f"{_initials}", key="pop_account", help=_full)
    _acct_pop.markdown(f"""<div class='profile-card'>
<div class='pc-top'><div class='pc-avatar'>{html.escape(_initials)}</div>
<div class='pc-id'><div class='pc-name'>{html.escape(_full)}</div><div class='pc-status'>● <span>Đang đăng nhập</span></div></div></div>
<div class='pc-row'><span class='pc-label'>Tài khoản</span><span class='pc-val'>@{html.escape(str(user_data['username']))}</span></div>
<div class='pc-row'><span class='pc-label'>Vai trò</span><span class='pc-role'>{html.escape(role_badge)}</span></div>{_ws_row}</div>""", unsafe_allow_html=True)


    my_pw = cursor.execute("SELECT password FROM users WHERE id = ?", (user_data['id'],)).fetchone()
    # Kiểm tra mật khẩu mặc định rất tốn thời gian (băm 200.000 vòng), nên chỉ kiểm tra lại khi mật khẩu thay đổi
    if my_pw and st.session_state.get("_pw_checked_hash") != my_pw[0]:
        st.session_state["_pw_checked_hash"] = my_pw[0]
        st.session_state["_pw_is_default"] = verify_password("admin123", my_pw[0])
    if my_pw and st.session_state.get("_pw_is_default"):
        st.error("⚠️ Tài khoản này vẫn dùng mật khẩu mặc định 'admin123'. Vào mục 🔑 Đổi Mật Khẩu và đổi NGAY!")

    if is_admin:
        _acct_pop.button("🔑 Đổi Mật Khẩu", key="btn_admin_change_pw", use_container_width=True,
                         on_click=_go_to_page, args=("🔑 Đổi Mật Khẩu",))
    if _acct_pop.button("🚪 Đăng Xuất", type="secondary", key="btn_logout_bottom", use_container_width=True):
        tok = st.session_state.get("session_token")
        if tok:
            delete_session(tok)
        try:
            cookie_manager.delete(SESSION_COOKIE)
        except Exception:
            pass
        st.session_state["logged_in"] = False
        st.session_state["user_info"] = None
        st.rerun()

    # 🎞️ KHUNG TRANG: mỗi lần đổi trang tạo khung mới để hiệu ứng lướt chạy lại.
    #    Bấm nút trong cùng một trang thì khung giữ nguyên, không bị lướt lại.
    page_idx = menu_options.index(menu) if menu in menu_options else 0
    prev_idx = st.session_state.get("page_prev_idx")
    if prev_idx != page_idx:
        st.session_state["page_dir"] = "L" if (prev_idx is not None and page_idx < prev_idx) else "R"
        st.session_state["page_prev_idx"] = page_idx
    _page_frame = st.container(key=f"pg{st.session_state.get('page_dir', 'R')}_{page_idx}")
    _page_frame.__enter__()

    # 💬 Tin Nhắn & Meeting: 1 nút menu, bên trong có 2 thẻ (Tin nhắn | Họp online)
    if menu == CONNECT_MENU:
        st.markdown("<div class='big-table-title'>💬 Tin Nhắn & Meeting</div>", unsafe_allow_html=True)
        st.session_state.setdefault("connect_tab", "chat")
        _ct1, _ct2, _ct_space = st.columns([1.3, 1.3, 2])
        with _ct1:
            st.button("💬 Tin Nhắn" + (f"  🔴 {_chat_unread_total}" if _chat_unread_total else ""),
                      key="btn_connect_chat", use_container_width=True,
                      type="primary" if st.session_state["connect_tab"] == "chat" else "secondary",
                      on_click=lambda: st.session_state.update({"connect_tab": "chat"}))
        with _ct2:
            st.button("🎥 Online Meeting", key="btn_connect_meet", use_container_width=True,
                      type="primary" if st.session_state["connect_tab"] == "meet" else "secondary",
                      on_click=lambda: st.session_state.update({"connect_tab": "meet"}))
        menu = CHAT_MENU if st.session_state["connect_tab"] == "chat" else MEETING_MENU

    # 1. BẢNG CÔNG VIỆC
    if menu == "🧰 Bảng Công Việc":
        st.markdown("<div class='big-table-title'>📋 Bảng Quản Lý Tiến Độ Công Việc</div>", unsafe_allow_html=True)
        
        df = pd.read_sql_query("""
            SELECT 
                t.id AS _id,
                t.assigned_leader_id AS _leader_id,
                t.assigned_worker_id AS _worker_id,
                t.task_id AS 'Task ID', 
                t.task_name AS 'Task Name', 
                t.task_cost_code AS 'WS Cost Code', 
                t.description AS 'Description',
                t.initial_by AS 'Initial By', 
                t.initial_date AS 'Initial Date', 
                t.block AS 'Block', 
                t.area AS 'Area', 
                t.deck AS 'Deck', 
                t.frame AS 'Frame', 
                t.in_charge_by AS 'In Charge By', 
                l.team_name AS 'Team',
                l.fullname AS 'Leader / Foreman',
                w.fullname AS 'Worker',
                t.plan_start_date AS 'Plan Start Date', 
                t.plan_finish_date AS 'Plan Finish Date', 
                t.progress AS 'Progress (%)', 
                t.remark AS 'Remark' 
            FROM tasks t
            LEFT JOIN users l ON l.id = t.assigned_leader_id
            LEFT JOIN users w ON w.id = t.assigned_worker_id
            WHERE t.is_deleted = 0
        """, conn)

        my_id = user_data['id']
        if current_role == "Worker":
            df = df[df['_worker_id'] == my_id]
            st.caption("👷 Đây là các công việc được giao cho bạn.")
        elif current_role in LEADER_ROLES:
            only_team = st.toggle("Chỉ xem công việc của team tôi", value=True)
            if only_team:
                df = df[df['_leader_id'] == my_id]
        
        if df.empty:
            st.info("Chưa có công việc nào để hiển thị.")
        else:
            st.dataframe(df.drop(columns=['_id', '_leader_id', '_worker_id']), use_container_width=True, height=400)
            
            st.markdown("---")
            st.markdown("### ⚡ Cập Nhật Tiến Độ Công Việc Nhanh")

            # Ai được cập nhật việc nào: Worker = việc của mình, Leader = việc của team, Manager = tất cả
            if current_role == "Worker":
                upd_df = df
            elif current_role in LEADER_ROLES:
                upd_df = df[df['_leader_id'] == my_id]
            else:
                upd_df = df

            if upd_df.empty:
                st.info("Bạn chưa có công việc nào để cập nhật tiến độ.")
            else:
                options_tasks = [f"{row['_id']} | {row['Task ID']} - {row['Task Name']} (Hiện tại: {row['Progress (%)']}%)" for _, row in upd_df.iterrows()]
                selected_update_task = st.selectbox("Chọn công việc cần cập nhật tiến độ:", options_tasks)
                selected_task_id = int(selected_update_task.split(" | ")[0])
                
                curr_task = cursor.execute("SELECT progress, remark FROM tasks WHERE id = ?", (selected_task_id,)).fetchone()
                
                with st.form("quick_update_progress_form"):
                    u_col1, u_col2 = st.columns([1, 2])
                    with u_col1:
                        new_progress = st.number_input("Mức Tiến Độ Mới (%)", min_value=0, max_value=100, value=int(curr_task[0] or 0), step=5)
                    with u_col2:
                        new_remark = st.text_input("Ghi Chú Thi Công (Remark):", value=curr_task[1] if curr_task[1] else "")
                    
                    btn_update_p = st.form_submit_button("🚀 CẬP NHẬT TIẾN ĐỘ")
                    if btn_update_p:
                        cursor.execute("UPDATE tasks SET progress = ?, remark = ? WHERE id = ?", (new_progress, new_remark.strip(), selected_task_id))
                        conn.commit()
                        st.success("Đã cập nhật tiến độ công việc thành công!")
                        st.rerun()

    # 3. THÊM CÔNG VIỆC MỚI
    elif menu == "➕ Thêm Công Việc" and current_role in ["Team Leader", "Foreman", "WOS Manager", "Admin"]:
        st.markdown("<div class='big-table-title'>➕ Thêm Công Việc Mới</div>", unsafe_allow_html=True)

        # 🏭 WS Cost Code + 🧱 Block nằm ngay trong form, nút mini ➕ / ➖ ngay cạnh danh sách (Foreman, WOS Manager, Admin)
        can_edit_lists = current_role in ["Foreman", "WOS Manager", "Admin"]
        list_actions = [("add", "➕", "Thêm"), ("del", "➖", "Xóa")] if can_edit_lists else []
        cost_codes_df = pd.read_sql_query("SELECT code, COALESCE(name, '') AS name FROM custom_cost_codes WHERE is_deleted = 0 ORDER BY code", conn)
        cc_label = {r['code']: (f"{r['code']} - {r['name']}" if r['name'] else r['code']) for _, r in cost_codes_df.iterrows()}
        NO_BLOCK = "— Không chọn —"
        block_options = [NO_BLOCK] + [r[0] for r in cursor.execute("SELECT name FROM blocks WHERE is_deleted = 0 ORDER BY name").fetchall()]

        # Lưu xong thì xóa trắng các ô (thay cho clear_on_submit của form)
        NT_KEYS = ["nt_id", "nt_name", "nt_desc", "nt_init_by", "nt_init_date", "nt_in_charge", "nt_area", "nt_deck",
                   "nt_frame", "nt_start", "nt_finish", "nt_progress", "nt_remark", "nt_leader"]
        if st.session_state.pop("_nt_reset", False):
            for k in NT_KEYS:
                st.session_state.pop(k, None)
            st.session_state["_pending_sel_block"] = NO_BLOCK
        if st.session_state.get("_nt_saved_msg"):
            st.success(st.session_state.pop("_nt_saved_msg"))

        with st.container(key="sc_add_task_box", border=True):
            c1, c2, c3 = st.columns([1, 1.6, 1.4])
            with c1:
                task_id = st.text_input("Task ID *", key="nt_id")
            with c2:
                task_name_input = st.text_input("Task Name *", key="nt_name")
            with c3:
                sel_cc, m = mini_select("cc", "WS Cost Code *", list(cc_label.keys()), list_actions,
                                        format_func=lambda c: cc_label.get(c, c))
                if m == "add":
                    with mini_panel("cc"):
                        nc_code = st.text_input("Mã (Ví dụ: WOS_07):", key="new_cc_code",
                                                placeholder="Mã (Ví dụ: WOS_07):", label_visibility="collapsed")
                        nc_name = st.text_input("Tên Workshop / Phòng ban:", key="new_cc_name",
                                                placeholder="Tên Workshop / Phòng ban:", label_visibility="collapsed")
                        if st.button("➕ THÊM COST CODE", key="btn_add_cc", type="primary", use_container_width=True):
                            code_clean = nc_code.strip().upper().replace(" ", "_")
                            name_clean = nc_name.strip()
                            if not code_clean or not name_clean:
                                st.error("Vui lòng nhập cả Mã và Tên!")
                            else:
                                ex = cursor.execute("SELECT is_deleted FROM custom_cost_codes WHERE code = ?", (code_clean,)).fetchone()
                                if ex and ex[0] == 0:
                                    st.error(f"Mã **{code_clean}** đã có trong danh sách!")
                                else:
                                    if ex:
                                        cursor.execute("UPDATE custom_cost_codes SET name = ?, is_deleted = 0 WHERE code = ?", (name_clean, code_clean))
                                    else:
                                        cursor.execute("INSERT INTO custom_cost_codes (code, name, description, is_deleted) VALUES (?, ?, '', 0)",
                                                       (code_clean, name_clean))
                                    conn.commit()
                                    st.toast(f"✅ {code_clean} - {name_clean}")
                                    st.session_state["_pending_sel_cc"] = code_clean
                                    mini_close("cc")
                                    st.rerun()
                elif m == "del":
                    with mini_panel("cc"):
                        if not sel_cc:
                            st.caption("Danh sách Cost Code đang trống.")
                        elif st.button(f"🗑️ {tr('Xóa')} {sel_cc}", key="btn_del_cc", use_container_width=True):
                            users_in_ws = cursor.execute("SELECT COUNT(*) FROM users WHERE workshop = ? AND is_deleted = 0", (sel_cc,)).fetchone()[0]
                            if users_in_ws > 0:
                                st.error(f"Không thể xóa: còn **{users_in_ws}** tài khoản thuộc workshop này. "
                                         "Hãy chuyển họ sang workshop khác ở 👥 Quản Lý Phân Quyền trước.")
                            else:
                                cursor.execute("UPDATE custom_cost_codes SET is_deleted = 1 WHERE code = ?", (sel_cc,))
                                conn.commit()
                                st.toast(f"🗑️ {cc_label[sel_cc]}")
                                mini_close("cc")
                                st.rerun()

            description = st.text_area("Description (Mô tả công việc):", height=80, key="nt_desc")

            c4, c5, c6 = st.columns(3)
            with c4:
                initial_by = st.text_input("Initial By (Người khởi tạo):", value=user_data['fullname'], key="nt_init_by")
            with c5:
                initial_date = st.date_input("Initial Date (Ngày tạo):", value=date.today(), key="nt_init_date")
            with c6:
                in_charge_by = st.text_input("In Charge By (Người phụ trách):", key="nt_in_charge")

            c7, c8, c9, c10 = st.columns([1.4, 1, 1, 1])
            with c7:
                sel_block, m = mini_select("block", "Block:", block_options, list_actions)
                if m == "add":
                    with mini_panel("block"):
                        nb = st.text_input("Block mới (Ví dụ: 170150):", key="new_block_name",
                                           placeholder="Block mới (Ví dụ: 170150):", label_visibility="collapsed")
                        if st.button("➕ THÊM BLOCK", key="btn_add_block", type="primary", use_container_width=True):
                            nb_clean = nb.strip().upper()
                            if not nb_clean:
                                st.error("Vui lòng nhập tên Block!")
                            else:
                                ex = cursor.execute("SELECT is_deleted FROM blocks WHERE name = ?", (nb_clean,)).fetchone()
                                if ex and ex[0] == 0:
                                    st.error(f"Block **{nb_clean}** đã có trong danh sách!")
                                else:
                                    if ex:
                                        cursor.execute("UPDATE blocks SET is_deleted = 0 WHERE name = ?", (nb_clean,))
                                    else:
                                        cursor.execute("INSERT INTO blocks (name, is_deleted) VALUES (?, 0)", (nb_clean,))
                                    conn.commit()
                                    st.toast(f"✅ Block {nb_clean}")
                                    st.session_state["_pending_sel_block"] = nb_clean
                                    mini_close("block")
                                    st.rerun()
                elif m == "del":
                    with mini_panel("block"):
                        if sel_block == NO_BLOCK:
                            st.caption("Chọn Block để xóa:")
                        elif st.button(f"🗑️ {tr('Xóa')} {sel_block}", key="btn_del_block", use_container_width=True):
                            cursor.execute("UPDATE blocks SET is_deleted = 1 WHERE name = ?", (sel_block,))
                            conn.commit()
                            st.toast(f"🗑️ Block {sel_block}")
                            st.session_state["_pending_sel_block"] = NO_BLOCK
                            mini_close("block")
                            st.rerun()
            block = "" if sel_block == NO_BLOCK else sel_block
            with c8:
                area = st.text_input("Area:", key="nt_area")
            with c9:
                deck = st.text_input("Deck:", key="nt_deck")
            with c10:
                frame = st.text_input("Frame:", key="nt_frame")

            c11, c12, c13 = st.columns(3)
            with c11:
                plan_start = st.date_input("Plan Start Date:", value=date.today(), key="nt_start")
            with c12:
                plan_finish = st.date_input("Plan Finish Date:", value=date.today(), key="nt_finish")
            with c13:
                progress_val = st.number_input("Progress (%)", min_value=0, max_value=100, value=0, step=5, key="nt_progress")

            remark = st.text_input("Remark (Ghi chú):", key="nt_remark")

            # Giao việc ngay khi tạo
            new_task_leader_id = None
            if current_role in LEADER_ROLES:
                new_task_leader_id = user_data['id']
                st.caption("📌 Công việc này sẽ thuộc team của bạn. Vào 📋 Giao Việc để giao cho Worker.")
            else:
                leaders_df_add = pd.read_sql_query(
                    "SELECT id, fullname, role, workshop FROM users WHERE is_deleted = 0 AND role IN ('Team Leader', 'Foreman') ORDER BY fullname", conn)
                leader_opts_add = {None: "— Chưa giao —"}
                for _, r in leaders_df_add.iterrows():
                    leader_opts_add[int(r['id'])] = f"{r['fullname']} ({r['role']}{', ' + r['workshop'] if r['workshop'] else ''})"
                new_task_leader_id = st.selectbox("Giao cho Team Leader / Foreman:", list(leader_opts_add.keys()),
                                                  format_func=lambda k: leader_opts_add[k], key="nt_leader")

            st.markdown("<br>", unsafe_allow_html=True)
            if st.button("💾 LƯU CÔNG VIỆC MỚI", key="btn_save_new_task", type="primary"):
                if not sel_cc:
                    st.error("⚠️ CHƯA CÓ WS COST CODE: Foreman / WOS Manager hãy thêm Cost Code ở mục ⚙️ phía trên trước khi thêm công việc!")
                elif not task_id.strip() or not task_name_input.strip():
                    st.error("Vui lòng điền đầy đủ thông tin bắt buộc: Task ID và Task Name!")
                elif plan_finish < plan_start:
                    st.error("Plan Finish Date không được trước Plan Start Date!")
                elif cursor.execute("SELECT 1 FROM tasks WHERE task_id = ?", (task_id.strip(),)).fetchone():
                    st.error(f"Task ID **{task_id.strip()}** đã tồn tại (có thể đang trong Thùng Rác). Vui lòng dùng ID khác!")
                else:
                    cursor.execute("""
                        INSERT INTO tasks (
                            task_id, task_name, task_cost_code, description,
                            initial_by, initial_date, block, area, deck, frame,
                            in_charge_by, plan_start_date, plan_finish_date,
                            progress, remark, assigned_leader_id, is_deleted
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
                    """, (
                        task_id.strip(), task_name_input.strip(), sel_cc,
                        description.strip(), initial_by.strip(), str(initial_date),
                        block.strip(), area.strip(), deck.strip(), frame.strip(),
                        in_charge_by.strip(), str(plan_start), str(plan_finish),
                        int(progress_val), remark.strip(), new_task_leader_id
                    ))
                    conn.commit()
                    st.session_state["_nt_saved_msg"] = f"Đã lưu thành công công việc **{task_id.strip()} - {task_name_input.strip()}**!"
                    st.session_state["_nt_reset"] = True
                    st.rerun()

    # 3B. GIAO VIỆC
    #    - WOS Manager / Admin: giao công việc cho Team Leader / Foreman
    #    - Team Leader / Foreman: giao công việc của team mình cho Worker trong team
    elif menu == "📋 Giao Việc" and current_role in ["Team Leader", "Foreman", "WOS Manager", "Admin"]:
        st.markdown("<div class='big-table-title'>📋 Giao Việc</div>", unsafe_allow_html=True)

        if is_manager_up:
            st.info("👑 Chọn công việc và giao cho một **Team Leader** hoặc **Foreman**. Họ sẽ giao tiếp cho Worker trong team.")
            tasks_df = pd.read_sql_query("""
                SELECT t.id, t.task_id, t.task_name, t.task_cost_code, t.assigned_leader_id, l.fullname AS leader_name
                FROM tasks t LEFT JOIN users l ON l.id = t.assigned_leader_id
                WHERE t.is_deleted = 0 ORDER BY t.id DESC
            """, conn)
            leaders_df = pd.read_sql_query(
                "SELECT id, fullname, role, workshop FROM users WHERE is_deleted = 0 AND role IN ('Team Leader', 'Foreman') ORDER BY fullname", conn)

            if tasks_df.empty:
                st.info("Chưa có công việc nào. Vào ➕ Thêm Công Việc để tạo.")
            elif leaders_df.empty:
                st.warning("Chưa có Team Leader / Foreman nào. Vào 👥 Quản Lý Phân Quyền để cấp role trước.")
            else:
                show_unassigned = st.toggle("Chỉ hiện công việc chưa giao", value=False)
                if show_unassigned:
                    tasks_df = tasks_df[tasks_df['assigned_leader_id'].isna()]
                if tasks_df.empty:
                    st.success("Tất cả công việc đã được giao!")
                else:
                    task_opts = {int(r['id']): f"{r['task_id']} - {r['task_name']} ({r['task_cost_code']}) → "
                                               f"{r['leader_name'] if r['leader_name'] else 'Chưa giao'}"
                                 for _, r in tasks_df.iterrows()}
                    sel_task = st.selectbox("Chọn công việc:", list(task_opts.keys()), format_func=lambda k: task_opts[k])

                    leader_opts = {None: "— Chưa giao —"}
                    for _, r in leaders_df.iterrows():
                        leader_opts[int(r['id'])] = f"{r['fullname']} ({r['role']}{', ' + r['workshop'] if r['workshop'] else ''})"
                    cur_leader = tasks_df.loc[tasks_df['id'] == sel_task, 'assigned_leader_id'].iloc[0]
                    cur_leader = None if pd.isna(cur_leader) else int(cur_leader)
                    keys = list(leader_opts.keys())
                    sel_leader = st.selectbox("Giao cho Team Leader / Foreman:", keys,
                                              index=keys.index(cur_leader) if cur_leader in keys else 0,
                                              format_func=lambda k: leader_opts[k])

                    if st.button("💾 LƯU GIAO VIỆC", type="primary", key="btn_assign_leader"):
                        if sel_leader != cur_leader:
                            # Đổi người phụ trách thì bỏ Worker cũ (vì Worker thuộc team cũ)
                            cursor.execute("UPDATE tasks SET assigned_leader_id = ?, assigned_worker_id = NULL WHERE id = ?",
                                           (sel_leader, sel_task))
                            conn.commit()
                        st.success(f"Đã giao công việc cho **{leader_opts[sel_leader]}**!")
                        st.rerun()

        else:
            my_id = user_data['id']
            team_df = pd.read_sql_query(
                "SELECT id, username, fullname FROM users WHERE is_deleted = 0 AND role = 'Worker' AND leader_id = ? ORDER BY fullname",
                conn, params=(my_id,))

            my_team_name = cursor.execute("SELECT team_name FROM users WHERE id = ?", (my_id,)).fetchone()[0]
            st.markdown(f"### 👷 Team: {html.escape(my_team_name) if my_team_name else '(chưa đặt tên)'} – {len(team_df)} Worker")
            if team_df.empty:
                st.info("Team của bạn chưa có Worker nào. Vào 👥 Team Của Tôi để tạo team và thêm Worker.")

            st.markdown("---")
            st.markdown("### 📌 Giao Việc Cho Worker")
            my_tasks = pd.read_sql_query("""
                SELECT t.id, t.task_id, t.task_name, t.progress, t.assigned_worker_id, w.fullname AS worker_name
                FROM tasks t LEFT JOIN users w ON w.id = t.assigned_worker_id
                WHERE t.is_deleted = 0 AND t.assigned_leader_id = ? ORDER BY t.id DESC
            """, conn, params=(my_id,))

            if my_tasks.empty:
                st.info("Bạn chưa được giao công việc nào.")
            elif team_df.empty:
                st.warning("Cần có Worker trong team trước khi giao việc.")
            else:
                task_opts = {int(r['id']): f"{r['task_id']} - {r['task_name']} ({r['progress']}%) → "
                                           f"{r['worker_name'] if r['worker_name'] else 'Chưa giao'}"
                             for _, r in my_tasks.iterrows()}
                sel_task = st.selectbox("Chọn công việc:", list(task_opts.keys()), format_func=lambda k: task_opts[k])

                worker_opts = {None: "— Chưa giao —"}
                for _, r in team_df.iterrows():
                    worker_opts[int(r['id'])] = f"{r['fullname']} (@{r['username']})"
                cur_worker = my_tasks.loc[my_tasks['id'] == sel_task, 'assigned_worker_id'].iloc[0]
                cur_worker = None if pd.isna(cur_worker) else int(cur_worker)
                keys = list(worker_opts.keys())
                sel_worker = st.selectbox("Giao cho Worker:", keys,
                                          index=keys.index(cur_worker) if cur_worker in keys else 0,
                                          format_func=lambda k: worker_opts[k])

                if st.button("💾 LƯU GIAO VIỆC", type="primary", key="btn_assign_worker"):
                    cursor.execute("UPDATE tasks SET assigned_worker_id = ? WHERE id = ? AND assigned_leader_id = ?",
                                   (sel_worker, sel_task, my_id))
                    conn.commit()
                    st.success(f"Đã giao công việc cho **{worker_opts[sel_worker]}**!")
                    st.rerun()

    # 3C. TEAM CỦA TÔI
    #    - Team Leader / Foreman: đặt tên team, thêm Worker CÙNG WORKSHOP (chưa thuộc team nào), xóa khỏi team
    #    - Worker: xem team của mình và có thể rời team
    elif menu == "👥 Team Của Tôi" and current_role in ["Team Leader", "Foreman", "Worker"]:
        st.markdown("<div class='big-table-title'>👥 Team Của Tôi</div>", unsafe_allow_html=True)
        my_id = user_data['id']
        my_ws = user_data.get("workshop")

        if current_role in LEADER_ROLES:
            my_team_name = cursor.execute("SELECT team_name FROM users WHERE id = ?", (my_id,)).fetchone()[0]

            if not my_ws:
                st.warning("⚠️ Tài khoản của bạn chưa được gán Workshop. Nhờ WOS Manager gán Workshop trước khi lập team.")
            elif not my_team_name:
                # Bước 1: tạo team (đặt tên)
                st.info("Bạn chưa có team. Hãy đặt tên để tạo team.")
                with st.form("create_team_form"):
                    new_team_name = st.text_input("Tên Team *", placeholder="Ví dụ: Team Hàn Block 170")
                    if st.form_submit_button("➕ TẠO TEAM"):
                        if not new_team_name.strip():
                            st.error("Vui lòng nhập tên team!")
                        elif len(new_team_name.strip()) > 60:
                            st.error("Tên team tối đa 60 ký tự!")
                        else:
                            cursor.execute("UPDATE users SET team_name = ? WHERE id = ?", (new_team_name.strip(), my_id))
                            conn.commit()
                            st.success(f"Đã tạo team **{new_team_name.strip()}**!")
                            st.rerun()
            else:
                st.markdown(f"### 🏷️ {html.escape(my_team_name)}")
                st.caption(f"🏭 Workshop: {my_ws}")

                with st.expander("✏️ Đổi tên team"):
                    with st.form("rename_team_form"):
                        renamed = st.text_input("Tên team mới:", value=my_team_name)
                        if st.form_submit_button("💾 LƯU TÊN"):
                            if not renamed.strip() or len(renamed.strip()) > 60:
                                st.error("Tên team không hợp lệ (1–60 ký tự)!")
                            else:
                                cursor.execute("UPDATE users SET team_name = ? WHERE id = ?", (renamed.strip(), my_id))
                                conn.commit()
                                st.success("Đã đổi tên team!")
                                st.rerun()

                with st.expander("🗑️ Xóa team"):
                    st.warning("Khi xóa team: tất cả Worker sẽ rời team, và các công việc đang giao cho họ sẽ trở về "
                               "trạng thái chưa giao. Công việc vẫn thuộc về bạn, bạn có thể tạo team mới sau.")
                    confirm_delete_team = st.checkbox("Tôi chắc chắn muốn xóa team này", key="chk_delete_team")
                    if st.button("🗑️ XÓA TEAM", type="secondary", key="btn_delete_team", disabled=not confirm_delete_team):
                        cursor.execute("UPDATE tasks SET assigned_worker_id = NULL WHERE assigned_leader_id = ?", (my_id,))
                        cursor.execute("UPDATE users SET leader_id = NULL WHERE leader_id = ?", (my_id,))
                        cursor.execute("UPDATE users SET team_name = NULL WHERE id = ?", (my_id,))
                        conn.commit()
                        st.success("Đã xóa team.")
                        st.rerun()

                team_df = pd.read_sql_query(
                    "SELECT id, username, fullname FROM users WHERE is_deleted = 0 AND role = 'Worker' AND leader_id = ? ORDER BY fullname",
                    conn, params=(my_id,))

                m = mini_bar("team", "👷 Thành Viên",
                             [("add", "➕", "Thêm Worker"), ("del", "➖", "Xóa Khỏi Team")])
                if m == "add":
                    with mini_panel("team"):
                        # Chỉ lấy Worker cùng workshop và chưa thuộc team nào
                        free_workers = pd.read_sql_query("""
                            SELECT id, username, fullname FROM users
                            WHERE is_deleted = 0 AND role = 'Worker' AND workshop = ? AND leader_id IS NULL
                            ORDER BY fullname
                        """, conn, params=(my_ws,))
                        if free_workers.empty:
                            st.caption(f"Không còn Worker nào trong workshop {my_ws} chưa có team.")
                        else:
                            fw_opts = {int(r['id']): f"{r['fullname']} (@{r['username']})" for _, r in free_workers.iterrows()}
                            to_add = st.multiselect("Chọn Worker (cùng workshop, chưa có team):", list(fw_opts.keys()),
                                                    format_func=lambda k: fw_opts[k], key="ms_team_add",
                                                    placeholder="Chọn Worker (cùng workshop, chưa có team):",
                                                    label_visibility="collapsed")
                            if st.button("➕ THÊM VÀO TEAM", type="primary", key="btn_add_team_members", use_container_width=True):
                                if not to_add:
                                    st.error("Hãy chọn ít nhất 1 Worker!")
                                else:
                                    for wid in to_add:
                                        # Kiểm tra lại phía server: đúng workshop và vẫn chưa có team
                                        cursor.execute("""UPDATE users SET leader_id = ?
                                                          WHERE id = ? AND role = 'Worker' AND workshop = ? AND leader_id IS NULL AND is_deleted = 0""",
                                                       (my_id, wid, my_ws))
                                    conn.commit()
                                    st.toast(f"✅ +{len(to_add)} Worker")
                                    mini_close("team")
                                    st.rerun()
                elif m == "del":
                    with mini_panel("team"):
                        if team_df.empty:
                            st.caption("Chưa có thành viên để xóa.")
                        else:
                            tm_opts = {int(r['id']): f"{r['fullname']} (@{r['username']})" for _, r in team_df.iterrows()}
                            to_remove = st.selectbox("Chọn Worker:", list(tm_opts.keys()), format_func=lambda k: tm_opts[k],
                                                     key="sb_team_rm", label_visibility="collapsed")
                            if st.button("➖ XÓA KHỎI TEAM", type="secondary", key="btn_remove_team_member", use_container_width=True):
                                remove_worker_from_team(to_remove, my_id)
                                st.toast(f"🗑️ {tm_opts[to_remove]}")
                                mini_close("team")
                                st.rerun()

                if team_df.empty:
                    st.info("Team chưa có Worker nào.")
                else:
                    st.dataframe(team_df.rename(columns={'username': 'Username', 'fullname': 'Họ và Tên'}).drop(columns=['id']),
                                 use_container_width=True)

        else:
            # WORKER
            my_leader = cursor.execute("""
                SELECT l.id, l.fullname, l.role, l.team_name FROM users me
                JOIN users l ON l.id = me.leader_id
                WHERE me.id = ? AND l.is_deleted = 0
            """, (my_id,)).fetchone()

            if not my_leader:
                st.info("Bạn chưa thuộc team nào. Team Leader / Foreman trong workshop của bạn có thể thêm bạn vào team.")
            else:
                leader_id, leader_name, leader_role, team_name = my_leader
                st.markdown(f"### 🏷️ {html.escape(team_name) if team_name else '(Team chưa đặt tên)'}")
                st.markdown(f"**Trưởng team:** {html.escape(leader_name)} ({leader_role})")

                mates = pd.read_sql_query(
                    "SELECT username, fullname FROM users WHERE is_deleted = 0 AND role = 'Worker' AND leader_id = ? ORDER BY fullname",
                    conn, params=(leader_id,))
                st.dataframe(mates.rename(columns={'username': 'Username', 'fullname': 'Họ và Tên'}), use_container_width=True)

                st.markdown("---")
                st.markdown("#### 🚪 Rời Team")
                st.caption("Khi rời team, các công việc của team đang giao cho bạn sẽ trở về trạng thái chưa giao.")
                confirm_leave = st.checkbox("Tôi chắc chắn muốn rời team này")
                if st.button("🚪 RỜI TEAM", type="secondary", key="btn_leave_team", disabled=not confirm_leave):
                    remove_worker_from_team(my_id, leader_id)
                    st.success("Bạn đã rời team.")
                    st.rerun()

    # 4. CHỈNH SỬA / XÓA TẠM
    elif menu == "✏️ Chỉnh Sửa/Xóa" and current_role in ["Foreman", "WOS Manager", "Admin"]:
        st.markdown("<div class='big-table-title'>✏️ Chỉnh Sửa & Xóa Quản Lý</div>", unsafe_allow_html=True)
        
        btn_col1, btn_col2, btn_col_space = st.columns([1.5, 2, 2.5])
        is_edit_task = st.session_state["edit_sub_tab"] == "task"
        
        with btn_col1:
            st.button("🧰 Xóa Tạm Công Việc", type="primary" if is_edit_task else "secondary", key="btn_sub_edit_task", on_click=lambda: st.session_state.update({"edit_sub_tab": "task"}))
                
        with btn_col2:
            st.button("👤 Xóa Tạm Tài Khoản Người Dùng", type="primary" if not is_edit_task else "secondary", key="btn_sub_edit_user", on_click=lambda: st.session_state.update({"edit_sub_tab": "user"}))

        st.markdown("---")

        if st.session_state["edit_sub_tab"] == "task":
            tasks_df = pd.read_sql_query("SELECT id, task_id, task_name FROM tasks WHERE is_deleted = 0", conn)
            if tasks_df.empty:
                st.info("Hiện không có công việc nào để chỉnh sửa hoặc xóa.")
            else:
                task_list = [f"{row['id']} | {row['task_id']} - {row['task_name']}" for _, row in tasks_df.iterrows()]
                selected_task_str = st.selectbox("Chọn công việc cần chuyển vào Thùng Rác:", task_list)
                selected_id = int(selected_task_str.split(" | ")[0])
                
                if st.button("🗑️ Chuyển Công Việc Vào Thùng Rác", type="secondary", key="btn_soft_delete_task"):
                    cursor.execute("UPDATE tasks SET is_deleted = 1 WHERE id = ?", (selected_id,))
                    conn.commit()
                    st.success("Đã chuyển công việc vào Thùng Rác thành công!")
                    st.rerun()

        else:
            users_df = pd.read_sql_query("SELECT id, username, fullname, role FROM users WHERE is_deleted = 0", conn)
            users_df = users_df[users_df['id'] != user_data['id']]
            users_df = users_df[users_df['role'] != "Admin"]
            if current_role == "WOS Manager":
                # WOS Manager không được khóa WOS Manager khác
                users_df = users_df[users_df['role'] != "WOS Manager"]
            elif current_role != "Admin":
                # Foreman chỉ được khóa Worker, Team Leader và tài khoản đang chờ duyệt
                users_df = users_df[users_df['role'].isin(["Worker", "Team Leader", "Pending"])]
            
            if users_df.empty:
                st.info("Không có tài khoản khác khả dụng để xóa.")
            else:
                user_list = [f"{row['id']} | @{row['username']} - {row['fullname']} ({row['role']})" for _, row in users_df.iterrows()]
                selected_user_str = st.selectbox("Chọn tài khoản muốn khóa/xóa tạm:", user_list)
                selected_user_id = int(selected_user_str.split(" | ")[0])
                
                if st.button("🗑️ Khóa/Xóa Tạm Tài Khoản Này", type="secondary", key="btn_soft_delete_user"):
                    cursor.execute("UPDATE users SET is_deleted = 1 WHERE id = ?", (selected_user_id,))
                    conn.commit()
                    delete_user_sessions(selected_user_id)
                    st.success("Đã khóa/chuyển tài khoản vào Thùng Rác thành công!")
                    st.rerun()

    # 5. QUẢN LÝ PHÂN QUYỀN
    #    - Admin: CHỈ cấp role WOS Manager và chọn workshop cho Manager
    #    - WOS Manager: cấp Worker / Team Leader / Foreman trong workshop của mình
    elif menu == "👥 Quản Lý Phân Quyền" and is_manager_up:
        st.markdown("<div class='big-table-title'>👥 Quản Lý & Cấp Quyền Tài Khoản (Role List)</div>", unsafe_allow_html=True)

        ws_df = pd.read_sql_query("SELECT code, COALESCE(name, '') AS name FROM custom_cost_codes WHERE is_deleted = 0 ORDER BY code", conn)
        ws_label = {row['code']: f"{row['code']} - {row['name']}" for _, row in ws_df.iterrows()}

        all_users = pd.read_sql_query("""
            SELECT u.id, u.username, u.fullname, u.role, u.workshop, u.leader_id, u.team_name,
                   COALESCE(l.team_name || ' (' || l.fullname || ')', l.fullname) AS leader_name
            FROM users u LEFT JOIN users l ON l.id = u.leader_id
            WHERE u.is_deleted = 0
        """, conn)
        all_users = all_users[(all_users['id'] != user_data['id']) & (all_users['role'] != "Admin")]

        if is_admin:
            visible_users = all_users
            st.info("🛡️ Admin chỉ cấp quyền **WOS Manager** và chọn **Workshop** mà Manager đó phụ trách. "
                    "Các role Worker / Team Leader / Foreman do WOS Manager cấp.")
        else:
            st.info("👑 Bạn có thể cấp **Worker / Team Leader / Foreman**, chọn **Workshop**, "
                    "và xếp Worker vào **team** của một Team Leader / Foreman.")
            visible_users = all_users[all_users['role'] != "WOS Manager"]

        show_df = visible_users.copy()
        show_df['workshop'] = show_df['workshop'].map(lambda c: ws_label.get(c, c) if c else "—")
        show_df['leader_name'] = show_df['leader_name'].fillna("—")
        show_df = show_df.rename(columns={'username': 'Username', 'fullname': 'Họ và Tên',
                                          'role': 'Vai Trò (Role)', 'workshop': 'Workshop',
                                          'leader_name': 'Team của'})
        st.dataframe(show_df.drop(columns=['id', 'leader_id', 'team_name']), use_container_width=True)

        # 🏭 QUẢN LÝ WORKSHOP ngay tại trang phân quyền (thêm workshop mới rồi chọn luôn cho tài khoản)
        ws_stats = pd.read_sql_query("""
            SELECT c.code AS 'Mã', COALESCE(c.name, '') AS 'Tên Workshop',
                   (SELECT COUNT(*) FROM users u WHERE u.workshop = c.code AND u.is_deleted = 0) AS 'Số tài khoản',
                   (SELECT COUNT(*) FROM tasks t WHERE t.task_cost_code = c.code AND t.is_deleted = 0) AS 'Số công việc'
            FROM custom_cost_codes c WHERE c.is_deleted = 0 ORDER BY c.code
        """, conn)
        m = mini_bar("ws", "🏭 Workshop", [("add", "➕", "Thêm Workshop"), ("del", "➖", "Xóa Workshop")])
        if m == "add":
            with mini_panel("ws"):
                wa1, wa2 = st.columns([1, 2])
                with wa1:
                    new_ws_code = st.text_input("Mã Workshop (Ví dụ: WOS_07):", key="role_new_ws_code",
                                                placeholder="Mã Workshop (Ví dụ: WOS_07):", label_visibility="collapsed")
                with wa2:
                    new_ws_name = st.text_input("Tên Workshop / Phòng ban:", key="role_new_ws_name",
                                                placeholder="Tên Workshop / Phòng ban:", label_visibility="collapsed")
                if st.button("➕ THÊM WORKSHOP", type="primary", key="btn_role_add_ws", use_container_width=True):
                    code_clean = new_ws_code.strip().upper().replace(" ", "_")
                    name_clean = new_ws_name.strip()
                    if not code_clean or not name_clean:
                        st.error("Vui lòng nhập cả Mã và Tên!")
                    elif len(code_clean) > 20 or len(name_clean) > 80:
                        st.error("Mã tối đa 20 ký tự, tên tối đa 80 ký tự!")
                    else:
                        ex = cursor.execute("SELECT is_deleted FROM custom_cost_codes WHERE code = ?", (code_clean,)).fetchone()
                        if ex and ex[0] == 0:
                            st.error(f"Mã **{code_clean}** đã có trong danh sách!")
                        else:
                            if ex:
                                cursor.execute("UPDATE custom_cost_codes SET name = ?, is_deleted = 0 WHERE code = ?", (name_clean, code_clean))
                            else:
                                cursor.execute("INSERT INTO custom_cost_codes (code, name, description, is_deleted) VALUES (?, ?, '', 0)",
                                               (code_clean, name_clean))
                            conn.commit()
                            st.toast(f"✅ {code_clean} - {name_clean}")
                            mini_close("ws")
                            st.rerun()
        elif m == "del":
            with mini_panel("ws"):
                if ws_stats.empty:
                    st.caption("Danh sách Cost Code đang trống.")
                else:
                    del_opts = {r['Mã']: f"{r['Mã']} - {r['Tên Workshop']}" for _, r in ws_stats.iterrows()}
                    del_ws = st.selectbox("Chọn Workshop để xóa:", list(del_opts.keys()),
                                          format_func=lambda c: del_opts[c], key="role_del_ws", label_visibility="collapsed")
                    if st.button("🗑️ XÓA WORKSHOP", key="btn_role_del_ws", use_container_width=True):
                        users_in_ws = cursor.execute("SELECT COUNT(*) FROM users WHERE workshop = ? AND is_deleted = 0", (del_ws,)).fetchone()[0]
                        if users_in_ws > 0:
                            st.error(f"Không thể xóa: còn **{users_in_ws}** tài khoản thuộc workshop này. "
                                     "Hãy chuyển họ sang workshop khác ở 👥 Quản Lý Phân Quyền trước.")
                        else:
                            cursor.execute("UPDATE custom_cost_codes SET is_deleted = 1 WHERE code = ?", (del_ws,))
                            conn.commit()
                            st.toast(f"🗑️ {del_opts[del_ws]}")
                            mini_close("ws")
                            st.rerun()
        if ws_stats.empty:
            st.info("Chưa có Workshop nào. Hãy thêm Workshop đầu tiên bên dưới.")
        else:
            st.dataframe(ws_stats, use_container_width=True, hide_index=True)

        st.markdown("---")
        st.markdown("### 🔄 Thay Đổi Quyền Hạn Cho Tài Khoản")

        if visible_users.empty:
            st.info("Không có tài khoản nào để cấp quyền.")
        else:
            user_roles_list = [f"{row['id']} | @{row['username']} - {row['fullname']} (Hiện tại: {row['role']})"
                               for _, row in visible_users.iterrows()]
            target_role_user = st.selectbox("Chọn tài khoản cần chuyển đổi Role:", user_roles_list)
            target_user_id = int(target_role_user.split(" | ")[0])

            if is_admin:
                if ws_df.empty:
                    st.warning("⚠️ Chưa có Workshop nào. Vào ➕ Thêm Công Việc → ⚙️ Thêm / Xóa trong danh sách để thêm.")
                else:
                    action = st.radio("Chọn thao tác:", ["👑 Cấp quyền WOS Manager", "⛔ Thu hồi quyền (về Pending)"], horizontal=True)
                    chosen_ws = None
                    if action.startswith("👑"):
                        chosen_ws = st.selectbox("Workshop mà Manager này phụ trách:", list(ws_label.keys()),
                                                 format_func=lambda c: ws_label[c])
                    if st.button("💾 LƯU THAY ĐỔI", type="primary", key="btn_save_role_admin"):
                        if chosen_ws:
                            cursor.execute("UPDATE users SET role = 'WOS Manager', workshop = ? WHERE id = ?", (chosen_ws, target_user_id))
                            msg = f"Đã cấp **WOS Manager** cho workshop **{ws_label[chosen_ws]}**!"
                        else:
                            cursor.execute("UPDATE users SET role = 'Pending' WHERE id = ?", (target_user_id,))
                            msg = "Đã thu hồi quyền, tài khoản trở về trạng thái Pending."
                        conn.commit()
                        delete_user_sessions(target_user_id)
                        st.success(msg)
                        st.rerun()
            else:
                target_row = visible_users[visible_users['id'] == target_user_id].iloc[0]
                role_choices = ["Worker", "Team Leader", "Foreman", "Pending"]
                cur_role = target_row['role'] if target_row['role'] in role_choices else "Worker"
                new_role = st.radio("Chọn Role Mới:", role_choices, index=role_choices.index(cur_role), horizontal=True)

                ws_keys = list(ws_label.keys())
                default_ws = target_row['workshop'] if target_row['workshop'] in ws_keys else user_data.get("workshop")
                new_ws = st.selectbox("Workshop:", ws_keys,
                                      index=ws_keys.index(default_ws) if default_ws in ws_keys else 0,
                                      format_func=lambda c: ws_label[c]) if ws_keys else None

                new_leader = None
                if new_role == "Worker":
                    leaders_in_ws = all_users[(all_users['role'].isin(LEADER_ROLES)) & (all_users['workshop'] == new_ws)]
                    leader_opts = {None: "— Chưa xếp team —"}
                    for _, r in leaders_in_ws.iterrows():
                        leader_opts[int(r['id'])] = f"{r['team_name'] + ' – ' if r['team_name'] else ''}{r['fullname']} ({r['role']})"
                    cur_ld = None if pd.isna(target_row['leader_id']) else int(target_row['leader_id'])
                    ld_keys = list(leader_opts.keys())
                    new_leader = st.selectbox("Thuộc team của (Team Leader / Foreman):", ld_keys,
                                              index=ld_keys.index(cur_ld) if cur_ld in ld_keys else 0,
                                              format_func=lambda k: leader_opts[k])
                    if len(leader_opts) == 1:
                        st.caption("Workshop này chưa có Team Leader / Foreman nào.")

                if st.button("💾 LƯU THAY ĐỔI ROLE", type="primary", key="btn_save_role"):
                    cursor.execute("UPDATE users SET role = ?, workshop = ?, leader_id = ? WHERE id = ?",
                                   (new_role, new_ws, new_leader, target_user_id))
                    if new_role not in LEADER_ROLES or new_ws != target_row['workshop']:
                        # Không còn là Leader (hoặc đổi workshop): giải tán team và trả công việc về trạng thái chưa giao
                        cursor.execute("UPDATE users SET team_name = NULL WHERE id = ?", (target_user_id,))
                        cursor.execute("UPDATE users SET leader_id = NULL WHERE leader_id = ?", (target_user_id,))
                        cursor.execute("UPDATE tasks SET assigned_leader_id = NULL, assigned_worker_id = NULL WHERE assigned_leader_id = ?",
                                       (target_user_id,))
                    if new_role != "Worker":
                        cursor.execute("UPDATE tasks SET assigned_worker_id = NULL WHERE assigned_worker_id = ?", (target_user_id,))
                    else:
                        # Worker đổi team: bỏ các việc không thuộc team mới
                        cursor.execute("""UPDATE tasks SET assigned_worker_id = NULL
                                          WHERE assigned_worker_id = ? AND (assigned_leader_id IS NULL OR assigned_leader_id != ?)""",
                                       (target_user_id, new_leader if new_leader else -1))
                    conn.commit()
                    st.success(f"Đã cập nhật: **{new_role}** – workshop **{ws_label.get(new_ws, '')}**!")
                    st.rerun()

    # 6. THÙNG RÁC TỔNG HỢP (CHỈ WOS MANAGER)
    elif menu == "🗑️ Thùng Rác" and is_manager_up:
        st.markdown("<div class='big-table-title'>🗑️ Thùng Rác & Khôi Phục Tổng Hợp</div>", unsafe_allow_html=True)
        
        t_col1, t_col3, t_space = st.columns([1.5, 1.5, 3])
        
        is_t_task = st.session_state["trash_sub_tab"] == "task"
        is_t_ws = st.session_state["trash_sub_tab"] == "ws"
        is_t_user = st.session_state["trash_sub_tab"] == "user"
        
        with t_col1:
            st.button("🧰 Thùng Rác Công Việc", type="primary" if is_t_task else "secondary", key="btn_t_task", on_click=lambda: st.session_state.update({"trash_sub_tab": "task"}))
                
        with t_col3:
            st.button("👤 Thùng Rác Tài Khoản", type="primary" if is_t_user else "secondary", key="btn_t_user", on_click=lambda: st.session_state.update({"trash_sub_tab": "user"}))

        st.markdown("---")

        if st.session_state["trash_sub_tab"] == "task":
            deleted_tasks = pd.read_sql_query("SELECT id, task_id, task_name, task_cost_code, initial_by FROM tasks WHERE is_deleted = 1", conn)
            if deleted_tasks.empty:
                st.info("Thùng rác công việc đang trống.")
            else:
                st.dataframe(deleted_tasks, use_container_width=True)
                task_del_options = [f"{row['id']} | {row['task_id']} - {row['task_name']}" for _, row in deleted_tasks.iterrows()]
                restore_task_target = st.selectbox("Chọn công việc để xử lý:", task_del_options, key="sb_res_task")
                res_task_id = int(restore_task_target.split(" | ")[0])
                
                col_r1, col_r2 = st.columns(2)
                with col_r1:
                    if st.button("♻️ KHÔI PHỤC CÔNG VIỆC", type="primary", key="btn_res_task"):
                        cursor.execute("UPDATE tasks SET is_deleted = 0 WHERE id = ?", (res_task_id,))
                        conn.commit()
                        st.success("Đã khôi phục công việc!")
                        st.rerun()
                with col_r2:
                    if st.button("💥 XÓA VĨNH VIỄN CÔNG VIỆC", type="secondary", key="btn_perm_del_task"):
                        cursor.execute("DELETE FROM tasks WHERE id = ?", (res_task_id,))
                        conn.commit()
                        st.warning("Đã xóa vĩnh viễn công việc!")
                        st.rerun()

        else:
            deleted_users = pd.read_sql_query("SELECT id, username, fullname, role FROM users WHERE is_deleted = 1", conn)
            if deleted_users.empty:
                st.info("Thùng rác tài khoản đang trống.")
            else:
                st.dataframe(deleted_users, use_container_width=True)
                user_res_options = [f"{row['id']} | @{row['username']} - {row['fullname']} ({row['role']})" for _, row in deleted_users.iterrows()]
                restore_user_target = st.selectbox("Chọn tài khoản để xử lý:", user_res_options, key="sb_res_user")
                res_user_id = int(restore_user_target.split(" | ")[0])
                
                col_u1, col_u2 = st.columns(2)
                with col_u1:
                    if st.button("♻️ MỞ KHÓA / KHÔI PHỤC TÀI KHOẢN", type="primary", key="btn_res_user"):
                        cursor.execute("UPDATE users SET is_deleted = 0 WHERE id = ?", (res_user_id,))
                        conn.commit()
                        st.success("Đã khôi phục tài khoản người dùng!")
                        st.rerun()
                with col_u2:
                    if st.button("💥 XÓA VĨNH VIỄN TÀI KHOẢN", type="secondary", key="btn_perm_del_user"):
                        cursor.execute("DELETE FROM users WHERE id = ?", (res_user_id,))
                        conn.commit()
                        delete_user_sessions(res_user_id)
                        st.warning("Đã xóa vĩnh viễn tài khoản khỏi cơ sở dữ liệu!")
                        st.rerun()

    # 7. BÁO CÁO & THỐNG KÊ
    elif menu == "📊 Báo Cáo & Khai Báo":
        st.markdown("<div class='big-table-title'>📊 Báo Cáo & Thống Kê Tiến Độ</div>", unsafe_allow_html=True)
        _rep_tab = st.session_state.get("report_tab", "photos")
        _rt1, _rt2, _ = st.columns([1, 1, 1])
        with _rt1:
            st.button("📸 Báo cáo hình ảnh theo tuần", key="btn_rep_photos", use_container_width=True,
                      type="primary" if _rep_tab == "photos" else "secondary",
                      on_click=lambda: st.session_state.update({"report_tab": "photos"}))
        with _rt2:
            st.button("📈 Thống kê", key="btn_rep_stats", use_container_width=True,
                      type="primary" if _rep_tab == "stats" else "secondary",
                      on_click=lambda: st.session_state.update({"report_tab": "stats"}))

        if _rep_tab == "photos":
            render_weekly_photo_report(user_data, current_role)
        else:
        
            df_all = pd.read_sql_query("SELECT task_id, task_name, task_cost_code, block, in_charge_by, progress FROM tasks WHERE is_deleted = 0", conn)
        
            if df_all.empty:
                st.info("Chưa có dữ liệu công việc để tạo báo cáo. Vui lòng thêm công việc trước!")
            else:
                df_all['progress'] = pd.to_numeric(df_all['progress'], errors='coerce').fillna(0)
            
                total_tasks = len(df_all)
                completed_tasks = len(df_all[df_all['progress'] == 100])
                in_progress_tasks = len(df_all[(df_all['progress'] > 0) & (df_all['progress'] < 100)])
                avg_progress = round(df_all['progress'].mean(), 1)
            
                m1, m2, m3, m4 = st.columns(4)
                m1.metric("Tổng Số Công Việc", f"{total_tasks}")
                m2.metric("Đã Hoàn Thành (100%)", f"{completed_tasks}", f"{(completed_tasks/total_tasks*100):.1f}%")
                m3.metric("Đang Thực Hiện", f"{in_progress_tasks}")
                m4.metric("Tiến Độ Trung Bình", f"{avg_progress}%")
            
                st.markdown("---")
            
                col_chart1, col_chart2 = st.columns(2)
            
                with col_chart1:
                    st.markdown("### 📈 Phân Bố Tiến Độ Công Việc")
                    progress_bins = pd.cut(df_all['progress'], bins=[-1, 0, 50, 99, 100], labels=[tr('Chưa bắt đầu (0%)'), tr('Đang làm (1-50%)'), tr('Sắp xong (51-99%)'), tr('Hoàn thành (100%)')])
                    progress_dist = progress_bins.value_counts().reset_index()
                    _c_status, _c_count = tr('Trạng Thái'), tr('Số Lượng')
                    progress_dist.columns = [_c_status, _c_count]
                
                    chart = alt.Chart(progress_dist).mark_bar(color='#3b82f6').encode(
                        x=alt.X(f'{_c_status}:N', axis=alt.Axis(labelAngle=0, labelFontSize=12, title=_c_status)),
                        y=alt.Y(f'{_c_count}:Q', axis=alt.Axis(title=tr("Số Lượng Công Việc"), tickMinStep=1, format='d')),
                        tooltip=[_c_status, _c_count]
                    ).properties(height=350)
                
                    st.altair_chart(chart, use_container_width=True)

                with col_chart2:
                    st.markdown("### 🏭 Thống Kê Theo Workshop")
                    ws_stats = df_all.groupby('task_cost_code').agg(
                        Tổng_CV=('task_id', 'count'),
                        Tiến_Độ_TB=('progress', 'mean')
                    ).reset_index()
                    ws_stats['Tiến_Độ_TB'] = ws_stats['Tiến_Độ_TB'].round(1)
                    ws_stats.columns = ['WS Cost Code', 'Tổng Công Việc', 'Tiến Độ Trung Bình (%)']
                    st.dataframe(ws_stats, use_container_width=True)

                st.markdown("---")
                st.markdown("### 📥 Tải Báo Cáo Dữ Liệu")
            
                csv_data = df_all.to_csv(index=False).encode('utf-8-sig')
                st.download_button(
                    label="📥 Tải Báo Cáo Bảng Công Việc (File CSV/Excel)",
                    data=csv_data,
                    file_name=f"ShipControl_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
                    mime="text/csv",
                    type="primary"
                )

    # 8. ĐỔI MẬT KHẨU (MỌI ROLE)
    elif menu == "🔑 Đổi Mật Khẩu":
        st.markdown("<div class='big-table-title'>🔑 Đổi Mật Khẩu</div>", unsafe_allow_html=True)
        with st.form("change_password_form", clear_on_submit=True):
            old_pw = st.text_input("Mật khẩu hiện tại:", type="password")
            new_pw = st.text_input("Mật khẩu mới (ít nhất 8 ký tự):", type="password")
            new_pw2 = st.text_input("Nhập lại mật khẩu mới:", type="password")
            btn_change_pw = st.form_submit_button("💾 ĐỔI MẬT KHẨU")
            if btn_change_pw:
                row = cursor.execute("SELECT password FROM users WHERE id = ?", (user_data['id'],)).fetchone()
                if not row or not verify_password(old_pw, row[0]):
                    st.error("Mật khẩu hiện tại không đúng!")
                elif len(new_pw) < 8:
                    st.error("Mật khẩu mới phải có ít nhất 8 ký tự!")
                elif new_pw != new_pw2:
                    st.error("Mật khẩu xác nhận không trùng khớp!")
                elif new_pw == "admin123":
                    st.error("Không được dùng lại mật khẩu mặc định!")
                else:
                    cursor.execute("UPDATE users SET password = ? WHERE id = ?", (hash_password(new_pw), user_data['id']))
                    conn.commit()
                    # Đăng xuất mọi thiết bị khác
                    cursor.execute("DELETE FROM sessions WHERE user_id = ? AND token != ?",
                                   (user_data['id'], st.session_state.get("session_token", "")))
                    conn.commit()
                    st.success("✅ Đã đổi mật khẩu thành công! Các thiết bị khác đã bị đăng xuất.")

    # 9. 💬 TIN NHẮN: kênh chung + nhóm chat tự tạo + nhắn riêng (bấm vào tên tài khoản)
    elif menu == CHAT_MENU:
        st.session_state.setdefault("chat_conv", "general")
        if str(st.session_state["chat_conv"]).startswith("dm:"):
            st.session_state["chat_conv"] = "general"

        def _open_conv(conv):
            st.session_state["chat_conv"] = conv

        @st.fragment(run_every=3)
        def _chat_room():
            me = user_data["id"]
            conv = st.session_state.get("chat_conv", "general")
            # 🖱️ Lệnh từ menu chuột phải (ghim, đánh dấu, thu hồi, xóa phía tôi...)
            _raw_action = st.session_state.get("chat_action")
            if _raw_action:
                handle_chat_action(_raw_action, conv, me, current_role)
                st.session_state["chat_action"] = ""
            all_people = cursor.execute("""
                SELECT id, username, fullname, role FROM users
                WHERE is_deleted = 0 AND id != ? AND role IS NOT NULL AND role != 'Pending'
                ORDER BY fullname COLLATE NOCASE
            """, (me,)).fetchall()
            people_label = {p[0]: f"{CHAT_ROLE_ICONS.get(p[3], '👤')} {p[2] or p[1]} (@{p[1]})" for p in all_people}

            col_list, col_chat = st.columns([1, 2.2], gap="medium")

            # ---------- Cột trái: Kênh chung + Nhóm chat (không còn nhắn riêng từng người) ----------
            with col_list:
                unread = chat_unread_by_conv(me)
                unread.pop(conv, None)        # cuộc đang mở thì coi như đã đọc
                g_badge = f"  🔴 {unread['general']}" if unread.get("general") else ""
                st.button("📢 Kênh chung" + g_badge, key="chat_open_general", use_container_width=True,
                          type="primary" if conv == "general" else "secondary",
                          on_click=_open_conv, args=("general",))

                # 💬 Tiêu đề "Nhóm chat" + nút ➕ nhỏ (bấm để mở/đóng ô tạo nhóm gọn)
                if st.session_state.pop("_chat_grp_reset", False):
                    st.session_state["chat_new_group_name"] = ""
                    st.session_state["chat_new_group_members"] = []
                show_new = mini_bar("grpnew", "💬 Nhóm chat", [("add", "➕", "Tạo nhóm chat")]) == "add"
                if show_new:
                    with mini_panel("grpnew"):
                        g_name = st.text_input("Tên nhóm", key="chat_new_group_name", placeholder="Ví dụ: Tổ hàn Block 170",
                                               label_visibility="collapsed")
                        g_members = st.multiselect("Thêm người vào nhóm", list(people_label.keys()), placeholder="Chọn người...",
                                                   format_func=lambda i: people_label[i], key="chat_new_group_members",
                                                   label_visibility="collapsed")
                        if st.button("✨ TẠO NHÓM", key="chat_create_group", type="primary", use_container_width=True):
                            if not g_name.strip():
                                st.error("Vui lòng đặt tên nhóm!")
                            elif not g_members:
                                st.error("Hãy chọn ít nhất 1 người!")
                            else:
                                new_gid = create_chat_group(g_name.strip()[:60], me, g_members)
                                st.session_state["chat_conv"] = f"group:{new_gid}"
                                mini_close("grpnew")
                                st.session_state["_chat_grp_reset"] = True
                                st.rerun(scope="fragment")

                my_groups = user_chat_groups(me)
                if not my_groups:
                    st.caption("Chưa có nhóm nào — bấm ➕ để tạo.")
                else:
                    for gid, gname, gowner, gcount in my_groups:
                        c_id = f"group:{gid}"
                        badge = f"  🔴 {unread[c_id]}" if unread.get(c_id) else ""
                        st.button(f"💬 {gname} · {gcount}{badge}", key=f"chat_open_g{gid}", use_container_width=True,
                                  type="primary" if conv == c_id else "secondary",
                                  on_click=_open_conv, args=(c_id,))


            # ---------- Cột phải: khung trò chuyện ----------
            with col_chat:
                if conv.startswith("group:"):
                    gid = int(conv.split(":")[1])
                    group = cursor.execute("SELECT id, name, created_by FROM chat_groups WHERE id = ?", (gid,)).fetchone()
                    if not group or not is_group_member(gid, me):
                        conv = "general"
                        st.session_state["chat_conv"] = "general"
                if conv.startswith("group:"):
                    is_owner = (group[2] == me)
                    members = group_members(gid)
                    st.markdown(f"### 💬 {html.escape(group[1])}")
                    owner_row = next((m for m in members if m[0] == group[2]), None)
                    if owner_row:
                        st.caption("👑 Quản trị nhóm" + f": {owner_row[2] or owner_row[1]}")
                    st.caption(" · ".join(f"{CHAT_ROLE_ICONS.get(r, '👤')} {n or u}" for _, u, n, r in members))
                    g_actions = ([("add", "➕", "Thêm người vào nhóm"), ("del", "➖", "Xóa khỏi nhóm"),
                                  ("kill", "🗑️", "Xóa nhóm")] if is_owner else []) + [("leave", "🚪", "Rời nhóm")]
                    m = mini_bar(f"grp{gid}", "⚙️ Thành viên nhóm", g_actions)
                    if m:
                        with mini_panel(f"grp{gid}"):
                            if m == "add":
                                outsiders = [p for p in all_people if p[0] not in {x[0] for x in members}]
                                add_ids = st.multiselect("Thêm người vào nhóm", [p[0] for p in outsiders], placeholder="Chọn người...",
                                                         format_func=lambda i: people_label[i], key=f"chat_add_{gid}",
                                                         label_visibility="collapsed")
                                if st.button("➕ THÊM VÀO NHÓM", key=f"chat_add_btn_{gid}", type="primary", use_container_width=True):
                                    if add_ids:
                                        for uid in add_ids:
                                            cursor.execute("INSERT OR IGNORE INTO chat_group_members (group_id, user_id) VALUES (?, ?)", (gid, uid))
                                        st.toast(f"✅ +{len(add_ids)}")
                                        mini_close(f"grp{gid}")
                                        st.rerun(scope="fragment")
                            elif m == "del":
                                others = [x for x in members if x[0] != me]
                                if others:
                                    rm_id = st.selectbox("Xóa khỏi nhóm", [x[0] for x in others],
                                                         format_func=lambda i: next(f"{n or u} (@{u})" for x, u, n, r in others if x == i),
                                                         key=f"chat_rm_{gid}", label_visibility="collapsed")
                                    if st.button("➖ XÓA KHỎI NHÓM", key=f"chat_rm_btn_{gid}", use_container_width=True):
                                        cursor.execute("DELETE FROM chat_group_members WHERE group_id = ? AND user_id = ?", (gid, rm_id))
                                        mini_close(f"grp{gid}")
                                        st.rerun(scope="fragment")
                                else:
                                    st.caption("Chưa có thành viên để xóa.")
                            elif m == "kill":
                                if st.button("🗑️ XÓA NHÓM", key=f"chat_del_group_{gid}", use_container_width=True):
                                    delete_chat_group(gid)
                                    mini_close(f"grp{gid}")
                                    st.session_state["chat_conv"] = "general"
                                    st.rerun(scope="fragment")
                            elif m == "leave":
                                if st.button("🚪 RỜI NHÓM", key=f"chat_leave_{gid}", use_container_width=True):
                                    leave_chat_group(gid, me)
                                    mini_close(f"grp{gid}")
                                    st.session_state["chat_conv"] = "general"
                                    st.rerun(scope="fragment")
                    if not conv.startswith("group:"):
                        st.markdown("### 📢 Kênh chung")
                elif conv == "general":
                    st.markdown("### 📢 Kênh chung")
                    st.caption("Mọi người trong app đều thấy kênh này.")
                else:                                  # nhắn riêng đã bỏ → quay về kênh chung
                    conv = "general"
                    st.session_state["chat_conv"] = "general"
                    st.markdown("### 📢 Kênh chung")
                    st.caption("Mọi người trong app đều thấy kênh này.")

                # Chỗ hiện tin nhắn được giữ trước; ô nhập nằm dưới. Gửi xong thì tin mới hiện ngay.
                msg_box = st.container()
                new_msg = st.chat_input("Nhập tin nhắn...", key=f"chat_input_{conv}")
                if new_msg and new_msg.strip():
                    post_chat_message(conv, me, new_msg)

                msgs = cursor.execute("""
                    SELECT m.id, m.sender_id, m.body, m.created_at, u.fullname, u.username, u.role,
                           COALESCE(m.recalled, 0),
                           EXISTS (SELECT 1 FROM chat_stars s WHERE s.user_id = ? AND s.message_id = m.id),
                           EXISTS (SELECT 1 FROM chat_pins p WHERE p.message_id = m.id)
                    FROM chat_messages m LEFT JOIN users u ON u.id = m.sender_id
                    WHERE m.conv = ? AND m.id NOT IN (SELECT message_id FROM chat_hidden WHERE user_id = ?)
                    ORDER BY m.id DESC LIMIT 200
                """, (me, conv, me)).fetchall()
                pinned = cursor.execute("""
                    SELECT p.message_id, m.body, COALESCE(u.fullname, u.username)
                    FROM chat_pins p JOIN chat_messages m ON m.id = p.message_id
                    LEFT JOIN users u ON u.id = m.sender_id
                    WHERE p.conv = ? AND COALESCE(m.recalled, 0) = 0
                    ORDER BY p.pinned_at DESC LIMIT 3
                """, (conv,)).fetchall()
                with msg_box:
                    render_chat_messages(msgs, me, is_admin=is_chat_admin(conv, me, current_role), pinned=pinned)
                if msgs:
                    mark_chat_read(me, conv, msgs[0][0])

                # Ô ẩn nhận lệnh từ menu chuột phải + đoạn mã tạo menu (giống Zalo)
                st.text_input("chat_action", key="chat_action", label_visibility="collapsed")
                render_chat_context_menu()

        _chat_room()

    # 10. 🎥 HỌP ONLINE: Foreman / Team Leader tạo phòng họp có ID + mật khẩu; ai cũng vào được nếu có mã
    elif menu == MEETING_MENU:
        me = user_data["id"]
        can_host = current_role in MEETING_HOST_ROLES
        my_name = user_data.get("fullname") or user_data["username"]

        # ---------- 📨 Lời mời họp (vào thẳng, không cần gõ mã) ----------
        invites = cursor.execute("""
            SELECT m.id, m.room_id, m.password, m.title, m.video_room, m.start_at, u.fullname, u.username
            FROM meeting_invites i JOIN meetings m ON m.id = i.meeting_id
            LEFT JOIN users u ON u.id = m.host_id
            WHERE i.user_id = ? AND m.is_active = 1 ORDER BY m.start_at
        """, (me,)).fetchall()
        if invites:
            st.markdown("### 📨 Lời mời họp của bạn")
            for mid, rid, pw, ttl, vroom, start_at, hname, huser in invites:
                with st.container(border=True):
                    ic1, ic2 = st.columns([2, 1], vertical_alignment="center")
                    with ic1:
                        when = datetime.fromisoformat(start_at).strftime("%d/%m/%Y %H:%M") if start_at else ""
                        st.markdown(f"**🎥 {html.escape(ttl)}**  \n🕒 {when} · " + "Chủ phòng" + f": {html.escape(hname or huser or '')}")
                        st.markdown(meeting_code_card(rid, pw, compact=True), unsafe_allow_html=True)
                    with ic2:
                        st.link_button("🎥 VÀO HỌP", meeting_url(vroom, my_name), type="primary", use_container_width=True)
            st.markdown("---")

        # ---------- 🔑 Vào cuộc họp ----------
        st.markdown("### 🔑 Vào cuộc họp")
        jc1, jc2, jc3 = st.columns([1.2, 1, 0.9], vertical_alignment="bottom")
        with jc1:
            j_room = st.text_input("ID phòng họp", key="meet_join_room", placeholder="123 456 789")
        with jc2:
            j_pass = st.text_input("Mật khẩu phòng", key="meet_join_pass", type="password")
        with jc3:
            do_join = st.button("🎥 VÀO PHÒNG", key="meet_join_btn", type="primary", use_container_width=True)
        if do_join:
            st.session_state["meet_attempts"] = st.session_state.get("meet_attempts", 0) + 1
            if st.session_state["meet_attempts"] > 15:
                st.error("Bạn đã nhập sai quá nhiều lần. Hãy tải lại trang và thử lại sau.")
            else:
                mt = find_meeting(j_room, j_pass)
                if mt:
                    st.session_state["meet_joined"] = mt[0]
                    st.session_state["meet_attempts"] = 0
                else:
                    st.session_state.pop("meet_joined", None)
                    st.error("Sai ID phòng hoặc mật khẩu, hoặc cuộc họp đã kết thúc!")
        joined = st.session_state.get("meet_joined")
        if joined:
            mt = cursor.execute("""SELECT m.title, m.video_room, m.start_at, u.fullname, u.username
                                   FROM meetings m LEFT JOIN users u ON u.id = m.host_id
                                   WHERE m.id = ? AND m.is_active = 1""", (joined,)).fetchone()
            if mt:
                st.success(f"✅ Đúng mã phòng: **{mt[0]}** · " + "Chủ phòng" + f": {mt[3] or mt[4]}")
                st.link_button("🎥 MỞ PHÒNG HỌP (camera + micro)", meeting_url(mt[1], my_name), type="primary")
                st.caption("Phòng họp mở trong tab mới. Lần đầu, trình duyệt sẽ hỏi quyền dùng camera và micro — hãy bấm Cho phép.")

        # ---------- ➕ Tạo cuộc họp (Foreman / Team Leader) ----------
        if can_host:
            st.markdown("---")
            st.markdown("### ➕ Tạo cuộc họp mới")
            people = cursor.execute("""
                SELECT id, username, fullname, role FROM users
                WHERE is_deleted = 0 AND id != ? AND role IS NOT NULL AND role != 'Pending'
                ORDER BY fullname COLLATE NOCASE
            """, (me,)).fetchall()
            p_label = {p[0]: f"{CHAT_ROLE_ICONS.get(p[3], '👤')} {p[2] or p[1]} (@{p[1]})" for p in people}
            with st.form("meet_create_form", clear_on_submit=True):
                m_title = st.text_input("Tên cuộc họp *", placeholder="Ví dụ: Họp giao ca sáng")
                mc1, mc2 = st.columns(2)
                with mc1:
                    m_date = st.date_input("Ngày họp", value=date.today())
                with mc2:
                    m_time = st.time_input("Giờ họp", value=(datetime.now() + timedelta(minutes=5)).time().replace(second=0, microsecond=0))
                m_invite = st.multiselect("Mời người (họ sẽ thấy lời mời trong 🎥 Online Meeting)", list(p_label.keys()), placeholder="Chọn người...",
                                          format_func=lambda i: p_label[i])
                m_post_general = st.checkbox("📢 Thông báo vào Kênh chung", value=False)
                if st.form_submit_button("🎥 TẠO CUỘC HỌP"):
                    if not m_title.strip():
                        st.error("Vui lòng nhập tên cuộc họp!")
                    else:
                        start_at = datetime.combine(m_date, m_time).isoformat(timespec="minutes")
                        room_id, password = create_meeting(m_title.strip()[:80], me, start_at)
                        invite = (tr("🎥 Mời họp") + f": {m_title.strip()[:80]}\n"
                                  + tr("🕒 Thời gian") + f": {datetime.fromisoformat(start_at).strftime('%d/%m/%Y %H:%M')}\n"
                                  + tr("🔢 ID phòng") + f": {format_room_id(room_id)}\n"
                                  + tr("🔑 Mật khẩu") + f": {password}\n"
                                  + tr("👉 Vào mục 🎥 Online Meeting để tham gia."))
                        new_mid = cursor.execute("SELECT id FROM meetings WHERE room_id = ?", (room_id,)).fetchone()[0]
                        for uid in m_invite:
                            cursor.execute("INSERT OR IGNORE INTO meeting_invites (meeting_id, user_id) VALUES (?, ?)", (new_mid, uid))
                        if m_post_general:
                            post_chat_message("general", me, invite)
                        st.session_state["meet_just_created"] = (room_id, password, m_title.strip()[:80])
            if st.session_state.get("meet_just_created"):
                rid, pw, ttl = st.session_state.pop("meet_just_created")
                st.success(f"🎉 Đã tạo cuộc họp **{ttl}**!")
                st.markdown(meeting_code_card(rid, pw), unsafe_allow_html=True)

            # ---------- 📋 Cuộc họp của tôi ----------
            st.markdown("---")
            st.markdown("### 📋 Cuộc họp của tôi")
            mine = cursor.execute("""SELECT id, room_id, password, title, video_room, start_at FROM meetings
                                     WHERE host_id = ? AND is_active = 1 ORDER BY start_at DESC""", (me,)).fetchall()
            if not mine:
                st.info("Bạn chưa có cuộc họp nào.")
            for mid, rid, pw, ttl, vroom, start_at in mine:
                with st.container(border=True):
                    cc1, cc2 = st.columns([2, 1], vertical_alignment="center")
                    with cc1:
                        when = datetime.fromisoformat(start_at).strftime("%d/%m/%Y %H:%M") if start_at else ""
                        st.markdown(f"**🎥 {html.escape(ttl)}**  \n🕒 {when}")
                        st.markdown(meeting_code_card(rid, pw, compact=True), unsafe_allow_html=True)
                    with cc2:
                        st.link_button("▶️ BẮT ĐẦU", meeting_url(vroom, my_name), type="primary", use_container_width=True)
                        if st.button("⛔ KẾT THÚC", key=f"meet_end_{mid}", use_container_width=True):
                            cursor.execute("UPDATE meetings SET is_active = 0 WHERE id = ? AND host_id = ?", (mid, me))
                            st.rerun()
        else:
            st.caption("Chỉ Foreman và Team Leader mới tạo được cuộc họp. Bạn có thể vào họp khi có ID phòng và mật khẩu.")